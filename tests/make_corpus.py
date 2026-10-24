"""Generate the test corpus: every deck the suite reads, built from code.

Each deck carries made-up filler text, solid-colour pictures, a 3x3 table,
speaker notes and, for the template, the .potx content type. Nothing in it is
copied from any real document, and no test reads a deck it did not build
here. tests/corpus/ is a build output:

- a deck this generator wrote is kept while the generator is unchanged
  (GENERATOR_VERSION and each deck's digest are recorded in
  tests/corpus/.generated.json) and rebuilt when it changes;
- a deck under one of the corpus names that this generator did not write is
  never read: ensure_corpus() raises ForeignDeck and the test run stops, so
  the file can be moved out of the folder first.

Run directly:  python -X utf8 tests/make_corpus.py
Or automatically: tests/conftest.py calls ensure_corpus() at session start.
"""

from __future__ import annotations

import hashlib
import json
import os
import random
import re
import struct
import sys
import tempfile
import time
import zipfile
import zlib
from contextlib import contextmanager
from pathlib import Path

CORPUS = Path(__file__).parent / "corpus"
MANIFEST = CORPUS / ".generated.json"
LOCK = CORPUS / ".generate.lock"

#: Bump whenever build_deck or GENERATORS change, so decks written by an
#: older generator are rebuilt instead of reused.
GENERATOR_VERSION = 2

#: Slide 1 of every deck build_deck writes carries this subtitle. A deck
#: without it was not written here.
SIGNATURE = "Structural stand-in, zero real content"

_WORDS = (
    "the framework extends across cases while compliance rises and the "
    "institution adapts its mandate toward renewed cooperation under "
    "conditions of uncertainty where actors weigh obligations against "
    "expected consequences and prior commitments shape present choices"
).split()

#: Content type stamped on ppt/presentation.xml for a .potx stand-in.
_CT_TEMPLATE_MAIN = (
    "application/vnd.openxmlformats-officedocument.presentationml"
    ".template.main+xml"
)


def _sentence(rng: random.Random, n: int = 8) -> str:
    ws = [rng.choice(_WORDS) for _ in range(n)]
    return (" ".join(ws)).capitalize() + "."


def _png(path: Path, w: int = 60, h: int = 30, rgb=(120, 60, 40)) -> Path:
    """Minimal hand-rolled truecolor PNG; no imaging library needed."""

    def chunk(tag, data):
        c = tag + data
        return struct.pack(">I", len(data)) + c + struct.pack(
            ">I", zlib.crc32(c) & 0xFFFFFFFF
        )

    raw = b"".join(b"\x00" + bytes(rgb) * w for _ in range(h))
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )
    return path


def build_deck(
    path: Path,
    *,
    seed: int = 0,
    extra_slides: int = 2,
    potx: bool = False,
) -> Path:
    """One synthetic deck: title slide, bullet slide, table slide, picture
    slide, speaker notes on two slides, plus `extra_slides` filler slides.
    Structure only, zero real content."""
    from pptx import Presentation
    from pptx.util import Inches

    rng = random.Random(seed)
    prs = Presentation()

    s = prs.slides.add_slide(prs.slide_layouts[0])  # title
    s.shapes.title.text = "Synthetic Deck"
    s.placeholders[1].text = "Structural stand-in, zero real content"
    s.notes_slide.notes_text_frame.text = "Synthetic speaker notes, slide one."

    s = prs.slides.add_slide(prs.slide_layouts[1])  # title + content bullets
    s.shapes.title.text = "Bullets"
    tf = s.placeholders[1].text_frame
    tf.text = _sentence(rng, 5)
    for level in (0, 1, 1, 2):
        para = tf.add_paragraph()
        para.text = _sentence(rng, 4)
        para.level = level

    s = prs.slides.add_slide(prs.slide_layouts[5])  # title only + table
    s.shapes.title.text = "Table"
    rows, cols = 3, 3
    table = s.shapes.add_table(
        rows, cols, Inches(1), Inches(1.8), Inches(8), Inches(3)
    ).table
    for r in range(rows):
        for c in range(cols):
            table.cell(r, c).text = f"r{r}c{c}"

    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank + picture
    with tempfile.TemporaryDirectory() as td:
        png = _png(Path(td) / "syn.png", rgb=(60, 90, 140))
        s.shapes.add_picture(str(png), Inches(1), Inches(1))
    s.notes_slide.notes_text_frame.text = "Synthetic speaker notes, picture slide."

    for _ in range(extra_slides):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = _sentence(rng, 3)[:-1]
        s.placeholders[1].text_frame.text = _sentence(rng, 6)

    prs.save(str(path))

    if potx:
        # Restamp the presentation part's content type so the stand-in is a
        # real template package, not a renamed .pptx.
        sys.path.insert(0, str(Path(__file__).parents[1] / "src"))
        from kitchensink4ppt.core.package import PptxPackage

        pkg = PptxPackage(path)
        pkg.add_content_type_override("ppt/presentation.xml", _CT_TEMPLATE_MAIN)
        pkg.save(do_backup=False)
    return path


GENERATORS = {
    "proposal_deck.pptx": lambda p: build_deck(p, seed=26, extra_slides=22),
    "conference_talk.pptx": lambda p: build_deck(p, seed=13, extra_slides=6),
    "image_deck.pptx": lambda p: build_deck(p, seed=3, extra_slides=8),
    "conference_template.potx": lambda p: build_deck(p, seed=7, potx=True),
    "heavy_deck.pptx": lambda p: build_deck(p, seed=18, extra_slides=30),
    "table_deck.pptx": lambda p: build_deck(p, seed=23, extra_slides=4),
}


class ForeignDeck(RuntimeError):
    """A file under a corpus name that this generator did not write."""


def is_generated(path: Path) -> bool:
    """True when `path` is a deck build_deck wrote (slide 1 carries SIGNATURE)."""
    try:
        with zipfile.ZipFile(path) as zf:
            xml = zf.read("ppt/slides/slide1.xml").decode("utf-8", "replace")
    except (OSError, KeyError, zipfile.BadZipFile):
        return False
    return SIGNATURE in xml


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _read_manifest() -> dict:
    try:
        data = json.loads(MANIFEST.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return data if isinstance(data, dict) else {}


def _write_manifest(files: dict[str, str]) -> None:
    tmp = MANIFEST.with_name(f".tmp-{os.getpid()}-{MANIFEST.name}")
    tmp.write_text(
        json.dumps({"generator": GENERATOR_VERSION, "files": files},
                   indent=1, sort_keys=True),
        encoding="utf-8",
    )
    os.replace(tmp, MANIFEST)


@contextmanager
def _lock(timeout: float = 600.0):
    """One generator at a time per corpus folder (two runs in one tree)."""
    CORPUS.mkdir(exist_ok=True)
    fh = open(LOCK, "a+b")
    try:
        deadline = time.monotonic() + timeout
        while True:
            try:
                if os.name == "nt":
                    import msvcrt

                    fh.seek(0)
                    msvcrt.locking(fh.fileno(), msvcrt.LK_NBLCK, 1)
                else:
                    import fcntl

                    fcntl.flock(fh.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
                break
            except OSError:
                if time.monotonic() > deadline:
                    raise
                time.sleep(0.2)
        try:
            yield
        finally:
            if os.name == "nt":
                import msvcrt

                fh.seek(0)
                msvcrt.locking(fh.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                import fcntl

                fcntl.flock(fh.fileno(), fcntl.LOCK_UN)
    finally:
        fh.close()


def _build_atomically(target: Path, gen) -> None:
    # The temporary name keeps the extension: the .potx build reopens it.
    tmp = target.with_name(f".tmp-{os.getpid()}-{target.name}")
    try:
        gen(tmp)
        os.replace(tmp, target)
    finally:
        if tmp.exists():
            tmp.unlink()


def ensure_corpus(verbose: bool = False) -> list[str]:
    """Make tests/corpus/ hold the generated decks the suite reads.

    Returns the names built by this call. Raises ForeignDeck, before writing
    anything, when a file under a corpus name was not written here."""
    CORPUS.mkdir(exist_ok=True)
    with _lock():
        foreign = sorted(
            name for name in GENERATORS
            if (CORPUS / name).exists() and not is_generated(CORPUS / name)
        )
        if foreign:
            raise ForeignDeck(
                f"tests/corpus/ holds file(s) this generator did not write: "
                f"{foreign}. The suite runs on generated decks only and never "
                f"reads one it did not build; move these files out of "
                f"tests/corpus/ and run again."
            )
        manifest = _read_manifest()
        current = manifest.get("generator") == GENERATOR_VERSION
        known = manifest.get("files", {}) if current else {}
        made = []
        for name, gen in GENERATORS.items():
            target = CORPUS / name
            if target.exists() and known.get(name) == _sha256(target):
                continue
            if verbose:
                print(f"generating {name} ...")
            _build_atomically(target, gen)
            made.append(name)
        if made or not current:
            _write_manifest({name: _sha256(CORPUS / name) for name in GENERATORS})
    return made


#: Older name of ensure_corpus, kept for callers outside this file.
generate_missing = ensure_corpus


# ---------------------------------------------- feature parts for the
# preservation tests: made-up payloads stored the way Office stores them.

_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _rewrite_package(path: Path, edit) -> Path:
    """Rewrite the package at `path` through edit(parts) -> parts, where
    parts is the ordered {member name: bytes}; every entry is deflated."""
    with zipfile.ZipFile(path) as zf:
        parts = {info.filename: zf.read(info) for info in zf.infolist()}
    parts = edit(parts)
    tmp = path.with_name(f".tmp-{os.getpid()}-{path.name}")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in parts.items():
            zf.writestr(name, data)
    os.replace(tmp, path)
    return path


def _add_type(types: bytes, element: str) -> bytes:
    return types.replace(b"</Types>", element.encode("utf-8") + b"</Types>", 1)


def _add_rel(rels: bytes, rel_type: str, target: str) -> tuple[bytes, str]:
    taken = set(re.findall(rb'Id="([^"]+)"', rels))
    n = 1
    while f"rId{n}".encode() in taken:
        n += 1
    rid = f"rId{n}"
    rel = f'<Relationship Id="{rid}" Type="{_R}/{rel_type}" Target="{target}"/>'
    return rels.replace(b"</Relationships>", rel.encode("utf-8") + b"</Relationships>", 1), rid


def add_custom_xml(path: Path) -> Path:
    """Add a custom XML data part and its properties part to the deck, linked
    from the presentation the way Office links document-level data."""

    def edit(parts):
        parts["customXml/item1.xml"] = (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            b'<record xmlns="urn:example:fixture"><project>Sample project</project>'
            b"<status>Draft</status></record>"
        )
        parts["customXml/itemProps1.xml"] = (
            b'<?xml version="1.0" encoding="UTF-8" standalone="no"?>\n'
            b'<ds:datastoreItem ds:itemID="{3F2A6C10-5B7D-4E21-9C3A-0D1E2F3A4B5C}" '
            b'xmlns:ds="http://schemas.openxmlformats.org/officeDocument/2006/customXml">'
            b'<ds:schemaRefs><ds:schemaRef ds:uri="urn:example:fixture"/></ds:schemaRefs>'
            b"</ds:datastoreItem>"
        )
        parts["customXml/_rels/item1.xml.rels"] = (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            + f'<Relationship Id="rId1" Type="{_R}/customXmlProps" Target="itemProps1.xml"/>'.encode()
            + b"</Relationships>"
        )
        parts["[Content_Types].xml"] = _add_type(
            parts["[Content_Types].xml"],
            '<Override PartName="/customXml/itemProps1.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.customXmlProperties+xml"/>',
        )
        parts["ppt/_rels/presentation.xml.rels"], _rid = _add_rel(
            parts["ppt/_rels/presentation.xml.rels"], "customXml", "../customXml/item1.xml"
        )
        return parts

    return _rewrite_package(path, edit)


def add_embedded_font(path: Path) -> Path:
    """Add an embedded font part and its presentation.xml entry. The payload
    is made-up bytes; nothing in the suite renders it."""

    def edit(parts):
        parts["ppt/fonts/font1.fntdata"] = b"KS4P-SYNTHETIC-FONT-PAYLOAD\x00" * 32
        types = parts["[Content_Types].xml"]
        if b'Extension="fntdata"' not in types:
            parts["[Content_Types].xml"] = _add_type(
                types, '<Default Extension="fntdata" ContentType="application/x-fontdata"/>'
            )
        parts["ppt/_rels/presentation.xml.rels"], rid = _add_rel(
            parts["ppt/_rels/presentation.xml.rels"], "font", "fonts/font1.fntdata"
        )
        pres = parts["ppt/presentation.xml"]
        notes_size = re.search(rb"<p:notesSz[^>]*/>", pres)
        entry = (
            '<p:embeddedFontLst><p:embeddedFont><p:font typeface="Fixture Sans"/>'
            f'<p:regular r:id="{rid}"/></p:embeddedFont></p:embeddedFontLst>'
        ).encode("utf-8")
        parts["ppt/presentation.xml"] = pres[:notes_size.end()] + entry + pres[notes_size.end():]
        return parts

    return _rewrite_package(path, edit)


if __name__ == "__main__":
    made = ensure_corpus(verbose=True)
    print(f"generated {len(made)} corpus deck(s): {made or 'none needed'}")
