"""Recovery from a stopped save does not depend on how the deck's path is spelled.

The same deck can be named in many ways: with another case, with forward
slashes, with the trailing dots and spaces Windows drops, or by its 8.3 short
name (C:\\Users\\JOHNSM~1\\...), which some programs and settings hand out. A
save stopped in its last instants under one spelling must be finished, cleanly
and with nothing left beside the deck, by a change or a read that names the
deck in another.

Each case is a real kill of a saving process (see test_recovery_marked_decks.py
for the instants). The 8.3 cases skip where the volume makes no short names.
"""

from __future__ import annotations

import asyncio
import ctypes
import os
import subprocess
import sys
import time
from pathlib import Path

import pytest
from pptx import Presentation

from kitchensink4ppt import server
from kitchensink4ppt.core import safesave, swapwait
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.plan import build

pytestmark = pytest.mark.skipif(sys.platform != "win32", reason="ReplaceFileW")

NO_WINDOW = 0x08000000
LONG_FOLDER = "Folder with a long name"
LONG_NAME = "Deck with a long name for short names.pptx"

VICTIM = r'''
import asyncio, os, sys, time
from kitchensink4ppt import server  # noqa: F401
from kitchensink4ppt.core import safesave
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.plan import build
deck, marker, mode = sys.argv[1], sys.argv[2], sys.argv[3]
fn, move, last_error, text = safesave._replace_file()

def stop():
    open(marker, "w").close()
    time.sleep(600)

def replace(ldst, lsrc, lbak, flags, a, b):
    if mode == "mid":
        os.replace(ldst, lbak)
        stop()
    ok = fn(ldst, lsrc, lbak, flags, a, b)
    if mode == "late":
        stop()
    return ok

safesave._REPLACE_FILE = (replace, move, last_error, text)
plan = build(Settings())
asyncio.run(dispatch(plan, "ppt_edit", {"op": "set_notes", "args": {
    "file_path": deck, "slide": 0, "text": "v2", "live": "off"}}))
'''


@pytest.fixture(autouse=True)
def own_journal(tmp_path, monkeypatch):
    monkeypatch.setenv("KS4P_SAVE_JOURNAL_DIR", str(tmp_path / "journal"))
    safesave.take_notices()
    yield
    safesave.take_notices()


def _short(path: Path) -> str:
    buf = ctypes.create_unicode_buffer(1024)
    n = ctypes.windll.kernel32.GetShortPathNameW(str(path), buf, 1024)
    return buf.value if n else str(path)


def _deck(path: Path, notes: str) -> Path:
    prs = Presentation()
    for title in ("Alpha", "Bravo"):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.notes_slide.notes_text_frame.text = notes
    prs.save(path)
    return path


def _notes(path) -> str | None:
    if not os.path.lexists(path):
        return None
    return Presentation(str(path)).slides[0].notes_slide.notes_text_frame.text


def _call(op, args):
    plan = build(Settings())
    return asyncio.run(dispatch(plan, plan.specs[op].group, {"op": op, "args": args}))


def _long_deck(tmp_path: Path) -> tuple[Path, str]:
    """v0, then v1 through the product: a deck with a long name in a folder
    with a long name, and its 8.3 short spelling."""
    folder = tmp_path / LONG_FOLDER
    folder.mkdir()
    deck = _deck(folder / LONG_NAME, "v0")
    short = _short(deck)
    if os.path.basename(short).lower() == LONG_NAME.lower():
        pytest.skip("this volume makes no 8.3 short names")
    _payload, refused = _call("set_notes", {"file_path": str(deck), "slide": 0, "text": "v1",
                                            "live": "off"})
    assert not refused
    return deck, short


def _stop_a_save(tmp_path: Path, spelling: str, mode: str) -> None:
    script = tmp_path / "victim.py"
    script.write_text(VICTIM, encoding="utf-8")
    marker = tmp_path / "stopped"
    victim = subprocess.Popen([sys.executable, "-X", "utf8", str(script), spelling, str(marker),
                               mode], stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                              stderr=subprocess.DEVNULL, creationflags=NO_WINDOW)
    try:
        deadline = time.monotonic() + 120
        while not marker.exists():
            assert victim.poll() is None, "the saving process ended before the stop"
            assert time.monotonic() < deadline, "the saving process never reached the stop"
            time.sleep(0.02)
    finally:
        victim.kill()  # TerminateProcess on this test's own child
        victim.wait(timeout=60)
    time.sleep(0.3)


def _clean(deck: Path) -> dict:
    folder = safesave.slot_dir(deck)
    return {
        "beside_the_deck": sorted(n for n in os.listdir(deck.parent)
                                  if n != deck.name and n != safesave.BACKUP_DIR_NAME),
        "staged_links": sorted(n for n in os.listdir(folder) if n.startswith(".slot-")),
        "journal": sorted(os.listdir(safesave._journal_dir())),
    }


CLEAN = {"beside_the_deck": [], "staged_links": [], "journal": []}


@pytest.mark.parametrize("saved_as, changed_as", [("long", "short"), ("short", "long")])
def test_a_change_under_another_spelling_finishes_a_stopped_save(tmp_path, saved_as, changed_as):
    deck, short = _long_deck(tmp_path)
    spell = {"long": str(deck), "short": short}
    _stop_a_save(tmp_path, spell[saved_as], "late")
    assert _notes(deck) == "v2" and len(swapwait.asides(deck)) + len(swapwait.asides(short)) == 1
    payload, refused = _call("set_notes", {"file_path": spell[changed_as], "slide": 0,
                                           "text": "v3", "live": "off"})
    assert not refused, payload
    assert _notes(deck) == "v3"
    assert _notes(safesave.slot_dir(deck) / safesave.PREV_SLOT) == "v2"
    assert _clean(deck) == CLEAN, "the stopped save's old copy or record was left behind"
    assert not payload.get("notices"), payload.get("notices")


def test_a_deck_left_aside_is_found_by_its_short_name(tmp_path):
    """A save stopped between its two moves left the deck aside. A read that
    names the deck by its short name says where it is; a change that names it
    so puts it back first and then applies."""
    deck, short = _long_deck(tmp_path)
    _stop_a_save(tmp_path, str(deck), "mid")
    assert not deck.exists()
    aside = swapwait.asides(deck)
    assert len(aside) == 1 and _notes(aside[0]) == "v1"
    payload, is_error = _call("get_text", {"file_path": short, "live": "off"})
    assert is_error is True and payload["error"]["code"] == "NOT_FOUND", payload
    assert aside[0] in payload["error"]["message"], payload["error"]["message"]
    payload, refused = _call("set_notes", {"file_path": short, "slide": 0, "text": "v3",
                                           "live": "off"})
    assert not refused, payload
    assert _notes(deck) == "v3"
    assert _notes(safesave.slot_dir(deck) / safesave.PREV_SLOT) == "v1"
    assert _clean(deck) == CLEAN


def test_case_separators_and_trailing_dots_name_the_same_deck(tmp_path):
    deck = _deck(tmp_path / "Deck.pptx", "v0")
    assert not _call("set_notes", {"file_path": str(deck), "slide": 0, "text": "v1",
                                   "live": "off"})[1]
    _stop_a_save(tmp_path, str(deck), "late")
    odd = str(deck).upper().replace("\\", "/") + ". "
    payload, refused = _call("set_notes", {"file_path": odd, "slide": 0, "text": "v3",
                                           "live": "off"})
    assert not refused, payload
    assert _notes(deck) == "v3"
    assert _notes(safesave.slot_dir(deck) / safesave.PREV_SLOT) == "v2"
    leftover = _clean(deck)
    leftover["beside_the_deck"] = [n for n in leftover["beside_the_deck"]
                                   if n not in ("victim.py", "stopped", "journal")]
    assert leftover == CLEAN


def test_start_up_finishes_a_save_made_under_the_short_name(tmp_path):
    deck, short = _long_deck(tmp_path)
    _stop_a_save(tmp_path, short, "late")
    server._recover_at_start(Settings())
    assert _notes(deck) == "v2"
    assert _notes(safesave.slot_dir(deck) / safesave.PREV_SLOT) == "v1"
    assert _clean(deck) == CLEAN
    assert safesave.take_notices() == []
