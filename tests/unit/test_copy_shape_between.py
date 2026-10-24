"""copy_shape_between: one shape or group from a deck onto a slide.

A canonical figure group was reused on twelve slides by deep-copying its XML
by hand: new shape ids, the group frame rewritten, the legend and frame
removed by name. The operation does that. The parts a figure carries
(pictures, charts, tags, links) travel with it, theme colours and theme
fonts are written against the SOURCE deck's theme, inherited text sizes are
written out so the text keeps its size in the other deck, connector glue
follows the copied shapes, and what cannot travel is refused with the
reason. The source deck is read from disk and never changed, so a copy
within one deck works too.
"""

from __future__ import annotations

import asyncio
import hashlib
import struct
import zlib
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation

from kitchensink4ppt import server  # noqa: F401  (the import registers ops)
from kitchensink4ppt.core.package import PptxPackage, qn, rels_name, resolve_target
from kitchensink4ppt.ops import shapes as shp
from kitchensink4ppt.registry import public_ops
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.plan import build

EMU = 914400
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
RT_IMAGE = R_NS + "/image"
RT_SLIDE = R_NS + "/slide"
RT_CHART = R_NS + "/chart"


def _call(tool: str, op: str, args: dict):
    plan = build(Settings(), public_ops())
    return asyncio.run(dispatch(plan, tool, {"op": op, "args": args}))


def _result(payload: dict) -> dict:
    inner = payload.get("changed")
    return inner if isinstance(inner, dict) else payload


def _ok(tool: str, op: str, args: dict) -> dict:
    payload, is_error = _call(tool, op, args)
    assert is_error is False, payload
    return _result(payload)


def _copy(args: dict) -> tuple[dict, dict]:
    """(full payload, the copy's answer)."""
    payload, is_error = _call("ppt_graphics", "copy_shape_between", args)
    assert is_error is False, payload
    return payload, _result(payload)


def _refused(args: dict) -> str:
    payload, is_error = _call("ppt_graphics", "copy_shape_between", args)
    assert is_error is True, payload
    return payload["error"]["message"]


def _md5(path) -> str:
    return hashlib.md5(Path(path).read_bytes()).hexdigest()


def _png(path: Path, rgb=(200, 30, 30)) -> str:
    width = height = 4
    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (struct.pack(">I", len(data)) + tag + data
                + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF))

    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))
    return str(path)


def _deck(path: Path, slides: int = 1, layout: int = 6) -> str:
    prs = Presentation()
    for _ in range(slides):
        prs.slides.add_slide(prs.slide_layouts[layout])
    prs.save(str(path))
    return str(path)


def _insert(deck, name, x, y, w=1.0, h=0.5, slide=0, shape_type="rect", **extra):
    return _ok("ppt_graphics", "insert_shape", {
        "file_path": deck, "slide": slide, "shape_type": shape_type, "name": name,
        "x": x, "y": y, "w": w, "h": h, **extra})["shape_id"]


@pytest.fixture
def source(tmp_path) -> dict:
    """Slide 0 holds the group "Figure": "Band A" (2 pt outline), "Label"
    (20 pt text), the text box "Caption" (no size of its own), the picture
    "Logo", and the members a reuse leaves out, "Legend" and "Frame". The
    deck's default text size is 32 pt, so Caption is drawn at 32 pt."""
    deck = _deck(tmp_path / "source.pptx", slides=2)
    band = _insert(deck, "Band A", 1, 1, 2, 1, line={"color": "#203040", "width": 2})
    label = _insert(deck, "Label", 1, 2.5, 2, 0.5, text="Model",
                    text_style={"size": 20})
    caption = _ok("ppt_edit", "insert_textbox", {
        "file_path": deck, "slide": 0, "text": "Plain", "x": 1, "y": 3.2,
        "w": 3, "h": 0.5})
    caption_id = caption.get("shape_id") or caption.get("id")
    _ok("ppt_graphics", "set_shape", {
        "file_path": deck, "slide": 0, "shape": caption_id, "name": "Caption"})
    logo = _ok("ppt_graphics", "insert_image", {
        "file_path": deck, "slide": 0, "image": _png(tmp_path / "logo.png"),
        "x": 3.5, "y": 1, "w": 1, "h": 1})
    logo_id = logo.get("shape_id") or logo.get("id")
    _ok("ppt_graphics", "set_shape", {
        "file_path": deck, "slide": 0, "shape": logo_id, "name": "Logo"})
    legend = _insert(deck, "Legend", 6, 1, 1, 1)
    frame = _insert(deck, "Frame", 0.5, 0.5, 7, 3.5, fill="none")
    figure = _ok("ppt_graphics", "group_shapes", {
        "file_path": deck, "slide": 0,
        "ids": [band, label, caption_id, logo_id, legend, frame],
        "name": "Figure"})["group_id"]
    pkg = PptxPackage(deck)
    pres = pkg.presentation()
    lvl1 = pres.find(f"{qn('p:defaultTextStyle')}/{qn('a:lvl1pPr')}/{qn('a:defRPr')}")
    lvl1.set("sz", "3200")
    pkg.mark_dirty()
    pkg.save(do_backup=False)
    return {"deck": deck, "figure": figure, "band": band, "label": label,
            "caption": caption_id, "logo": logo_id, "legend": legend, "frame": frame,
            "png": str(tmp_path / "logo.png"), "tmp": tmp_path}


@pytest.fixture
def dest(tmp_path) -> str:
    """Two blank slides; slide 1 already holds shapes, so ids collide."""
    deck = _deck(tmp_path / "dest.pptx", slides=2)
    for i in range(4):
        _insert(deck, f"Busy {i}", 1 + i, 5, 0.5, 0.5, slide=1)
    return deck


def _tree(deck: str, slide: int = 1):
    pkg = PptxPackage(deck)
    part = pkg.slide_parts()[slide]
    return pkg, part, shp._sp_tree(pkg, part)


def _named(deck: str, name: str, slide: int = 1):
    pkg, part, _tree_ = _tree(deck, slide)
    return pkg, part, shp._find_shape(pkg, part, name)


def _box_in(elem, chain) -> tuple[float, ...]:
    return tuple(round(v / EMU, 3) for v in shp._slide_box(elem, chain))


def _ids(tree) -> list[int]:
    return [int(c.get("id")) for c in tree.iter(qn("p:cNvPr"))]


# ------------------------------------------------------------------ copying


def test_copies_a_group_onto_another_deck_with_fresh_ids(source, dest):
    before = _md5(source["deck"])
    payload, res = _copy({
        "file_path": dest, "source_path": source["deck"], "source_slide": 0,
        "shape": "Figure", "slide": 1})
    assert _md5(source["deck"]) == before, "the source deck changed"
    pkg, part, tree = _tree(dest)
    ids = _ids(tree)
    assert len(ids) == len(set(ids)), ids
    elem, chain = shp._find_shape(pkg, part, "Figure")
    assert elem.tag == qn("p:grpSp") and not chain
    assert res["shape_id"] == shp._shape_id(elem)
    members = {shp._shape_name(e) for e, _c in shp._all_shapes(elem)}
    assert {"Band A", "Label", "Caption", "Logo", "Legend", "Frame"} <= members
    known = {str(source[key]) for key in (
        "figure", "band", "label", "caption", "logo", "legend", "frame")}
    assert set(res["ids"]) == known
    new_ids = list(res["ids"].values())
    assert len(set(new_ids)) == 7 and set(new_ids) <= set(ids)
    # The answer leads with the DESTINATION slide and names the source.
    assert payload["slide_number"] == 2
    assert res["source"]["slide_number"] == 1
    assert res["source"]["shape_id"] == source["figure"]
    Presentation(dest)  # opens


def test_default_place_is_the_source_place(source, dest):
    _copy({"file_path": dest, "source_path": source["deck"], "source_slide": 0,
           "shape": "Figure", "slide": 1})
    _p, _part, (elem, chain) = _named(dest, "Figure")
    src_pkg = PptxPackage(source["deck"])
    src_elem, src_chain = shp._find_shape(src_pkg, src_pkg.slide_parts()[0], "Figure")
    assert _box_in(elem, chain) == _box_in(src_elem, src_chain)


def test_one_side_keeps_the_aspect_and_the_members_follow(source, dest):
    _p, _r = _copy({
        "file_path": dest, "source_path": source["deck"], "source_slide": 0,
        "shape": "Figure", "slide": 1, "x": 2, "y": 3, "w": 3.5})
    _pkg, _part, (elem, chain) = _named(dest, "Figure")
    x, y, w, h = _box_in(elem, chain)
    assert (x, y, w) == (2.0, 3.0, 3.5)
    assert h == pytest.approx(1.75, abs=0.002)  # 7 x 3.5 in, halved
    _pkg, _part, (band, band_chain) = _named(dest, "Band A")
    bx, by, bw, bh = _box_in(band, band_chain)
    assert (bx, by) == (pytest.approx(2.25, abs=0.002), pytest.approx(3.25, abs=0.002))
    assert (bw, bh) == (pytest.approx(1.0, abs=0.002), pytest.approx(0.5, abs=0.002))


def test_scale_text_and_lines_follow_the_new_size(source, dest):
    _p, res = _copy({
        "file_path": dest, "source_path": source["deck"], "source_slide": 0,
        "shape": "Figure", "slide": 1, "w": 3.5, "scale_text": True,
        "scale_lines": True, "no_wrap": True})
    assert res["text_scaled"]["factor"] == 0.5
    _pkg, _part, (label, _c) = _named(dest, "Label")
    sizes = {r.get("sz") for r in label.iter(qn("a:rPr"))}
    assert sizes == {"1000"}, sizes
    _pkg, _part, (band, _c) = _named(dest, "Band A")
    assert band.find(f"{qn('p:spPr')}/{qn('a:ln')}").get("w") == "12700"
    _pkg, _part, (caption, _c) = _named(dest, "Caption")
    assert caption.find(f"{qn('p:txBody')}/{qn('a:bodyPr')}").get("wrap") == "none"


def test_inherited_text_size_is_written_from_the_source_deck(source, dest):
    _copy({"file_path": dest, "source_path": source["deck"], "source_slide": 0,
           "shape": "Caption", "slide": 1})
    _pkg, _part, (caption, _c) = _named(dest, "Caption")
    sizes = {r.get("sz") for r in caption.iter(qn("a:rPr"))}
    assert sizes == {"3200"}, sizes  # the destination's own default is 18 pt


def test_leave_out_drops_members_and_refits_the_group(source, dest):
    _p, res = _copy({
        "file_path": dest, "source_path": source["deck"], "source_slide": 0,
        "shape": "Figure", "slide": 1, "leave_out": ["Legend", "Frame"]})
    _pkg, _part, (elem, chain) = _named(dest, "Figure")
    members = {shp._shape_name(e) for e, _c in shp._all_shapes(elem)}
    assert "Legend" not in members and "Frame" not in members
    assert sorted(res["left_out"]) == sorted([source["legend"], source["frame"]])
    x, y, w, h = _box_in(elem, chain)
    assert (x, y) == (1.0, 1.0)
    assert w == pytest.approx(3.5, abs=0.002) and h == pytest.approx(2.7, abs=0.002)
    _pkg, _part, (band, band_chain) = _named(dest, "Band A")
    assert _box_in(band, band_chain)[:2] == (1.0, 1.0)  # nothing moved


def test_picture_travels_and_media_is_shared(source, dest):
    _copy({"file_path": dest, "source_path": source["deck"], "source_slide": 0,
           "shape": "Figure", "slide": 1})
    _p, res2 = _copy({"file_path": dest, "source_path": source["deck"],
                      "source_slide": 0, "shape": "Logo", "slide": 0})
    pkg = PptxPackage(dest)
    for index in (0, 1):
        part = pkg.slide_parts()[index]
        logo, _c = shp._find_shape(pkg, part, "Logo")
        rid = logo.find(f".//{qn('a:blip')}").get(qn("r:embed"))
        target = None
        for rel in pkg.root(rels_name(part)):
            if rel.get("Id") == rid:
                assert rel.get("Type") == RT_IMAGE
                target = resolve_target(part, rel.get("Target"))
        assert target and pkg.part_bytes(target) == Path(source["png"]).read_bytes()
    assert res2["media_reused"], res2


def test_connector_glue_follows_the_copy_and_is_dropped_outside_it(tmp_path, dest):
    deck = _deck(tmp_path / "flow.pptx")
    a = _insert(deck, "A", 1, 1)
    b = _insert(deck, "B", 4, 1)
    conn = _ok("ppt_graphics", "insert_connector", {
        "file_path": deck, "slide": 0, "start_shape": a, "end_shape": b})
    conn_id = conn.get("shape_id") or conn.get("id")
    _ok("ppt_graphics", "group_shapes", {
        "file_path": deck, "slide": 0, "ids": [a, b, conn_id], "name": "Flow"})
    _p, res = _copy({"file_path": dest, "source_path": deck, "source_slide": 0,
                     "shape": "Flow", "slide": 1})
    pkg, part, (flow, _c) = _named(dest, "Flow")
    new_a, new_b = res["ids"][str(a)], res["ids"][str(b)]
    cxn = flow.find(f".//{qn('p:cxnSp')}")
    st = cxn.find(f".//{qn('a:stCxn')}")
    end = cxn.find(f".//{qn('a:endCxn')}")
    assert (int(st.get("id")), int(end.get("id"))) == (new_a, new_b)
    _p, alone = _copy({"file_path": dest, "source_path": deck, "source_slide": 0,
                       "shape": conn_id, "slide": 0})
    pkg = PptxPackage(dest)
    part = pkg.slide_parts()[0]
    copied, _c = shp._find_shape(pkg, part, alone["shape_id"])
    assert copied.find(f".//{qn('a:stCxn')}") is None
    assert copied.find(f".//{qn('a:endCxn')}") is None
    assert any("glued" in w for w in _p.get("warnings", [])), _p


def test_theme_colours_are_written_from_the_source_theme(source, dest):
    _ok("ppt_design", "set_theme_colors", {
        "file_path": dest, "colors": {"accent1": "#00FF00"}})
    _p, res = _copy({"file_path": dest, "source_path": source["deck"],
                     "source_slide": 0, "shape": "Band A", "slide": 1})
    assert res["theme"] == "written"
    src = PptxPackage(source["deck"])
    theme = next(n for n in src.part_names() if n.startswith("ppt/theme/"))
    accent1 = src.root(theme).find(
        f".//{qn('a:clrScheme')}/{qn('a:accent1')}/{qn('a:srgbClr')}").get("val")
    _pkg, _part, (band, _c) = _named(dest, "Band A")
    assert band.find(f".//{qn('a:schemeClr')}") is None
    assert accent1 in {el.get("val") for el in band.iter(qn("a:srgbClr"))}


def test_a_copy_within_one_deck_keeps_its_theme_links(source):
    deck = source["deck"]
    _p, res = _copy({"file_path": deck, "source_path": deck, "source_slide": 0,
                     "shape": "Band A", "slide": 1})
    assert res["theme"] == "kept"
    pkg = PptxPackage(deck)
    part = pkg.slide_parts()[1]
    band, _c = shp._find_shape(pkg, part, "Band A")
    assert band.find(f".//{qn('a:schemeClr')}") is not None
    # The original is still on slide 0, unchanged.
    shp._find_shape(pkg, pkg.slide_parts()[0], "Figure")


def test_a_jump_link_keeps_its_slide_within_one_deck(source):
    deck = source["deck"]
    _ok("ppt_edit", "set_hyperlink", {
        "file_path": deck, "slide": 0, "target": source["band"], "to_slide": 1})
    _copy({"file_path": deck, "source_path": deck, "source_slide": 0,
           "shape": "Band A", "slide": 1})
    pkg = PptxPackage(deck)
    part = pkg.slide_parts()[1]
    band, _c = shp._find_shape(pkg, part, "Band A")
    rid = band.find(f".//{qn('a:hlinkClick')}").get(qn("r:id"))
    rels = {rel.get("Id"): rel for rel in pkg.root(rels_name(part))}
    assert rels[rid].get("Type") == RT_SLIDE
    assert resolve_target(part, rels[rid].get("Target")) == part


def test_a_chart_travels_with_its_parts(tmp_path, dest):
    deck = _deck(tmp_path / "chart.pptx")
    chart = _ok("ppt_tables", "create_chart", {
        "file_path": deck, "slide": 0, "chart_type": "bar",
        "series": [{"name": "Qty", "values": [1, 2]}], "categories": ["A", "B"],
        "x": 1, "y": 1, "w": 5, "h": 3})
    chart_id = chart.get("shape_id") or chart.get("id")
    _p, res = _copy({"file_path": dest, "source_path": deck, "source_slide": 0,
                     "shape": chart_id, "slide": 1, "w": 2.5})
    pkg = PptxPackage(dest)
    part = pkg.slide_parts()[1]
    frame, _c = shp._find_shape(pkg, part, res["shape_id"])
    rid = frame.find(f".//{{{'http://schemas.openxmlformats.org/drawingml/2006/chart'}}}chart").get(qn("r:id"))
    rels = {rel.get("Id"): rel for rel in pkg.root(rels_name(part))}
    assert rels[rid].get("Type") == RT_CHART
    assert pkg.has_part(resolve_target(part, rels[rid].get("Target")))
    Presentation(dest)


# ----------------------------------------------------------------- refusals


def _inject(deck: str, xml: str, slide: int = 0) -> None:
    pkg = PptxPackage(deck)
    part = pkg.slide_parts()[slide]
    tree = shp._sp_tree(pkg, part)
    tree.append(etree.fromstring(xml))
    pkg.mark_dirty(part)
    pkg.save(do_backup=False)


_NS = ('xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
       'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
       f'xmlns:r="{R_NS}"')


def _frame(uri: str, sid: int, name: str, inner: str = "") -> str:
    return (
        f'<p:graphicFrame {_NS}><p:nvGraphicFramePr><p:cNvPr id="{sid}" name="{name}"/>'
        '<p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
        '<p:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="914400"/></p:xfrm>'
        f'<a:graphic><a:graphicData uri="{uri}">{inner}</a:graphicData></a:graphic>'
        '</p:graphicFrame>')


def test_refuses_a_placeholder(tmp_path, dest):
    deck = _deck(tmp_path / "titled.pptx", layout=0)
    pkg = PptxPackage(deck)
    part = pkg.slide_parts()[0]
    title = next(e for e, _c in shp._all_shapes(shp._sp_tree(pkg, part))
                 if e.find(f".//{qn('p:ph')}") is not None)
    message = _refused({"file_path": dest, "source_path": deck, "source_slide": 0,
                        "shape": shp._shape_id(title), "slide": 1})
    assert "placeholder" in message


@pytest.mark.parametrize("uri, word", [
    ("http://schemas.openxmlformats.org/drawingml/2006/diagram", "SmartArt"),
    ("http://schemas.openxmlformats.org/presentationml/2006/ole", "embedded object"),
])
def test_refuses_smartart_and_embedded_objects(tmp_path, dest, uri, word):
    deck = _deck(tmp_path / "frames.pptx")
    _inject(deck, _frame(uri, 40, "Thing"))
    message = _refused({"file_path": dest, "source_path": deck, "source_slide": 0,
                        "shape": "Thing", "slide": 1})
    assert word in message


def test_refuses_ink_inside_a_group(tmp_path, dest):
    deck = _deck(tmp_path / "ink.pptx")
    a = _insert(deck, "A", 1, 1)
    b = _insert(deck, "B", 3, 1)
    _ok("ppt_graphics", "group_shapes", {
        "file_path": deck, "slide": 0, "ids": [a, b], "name": "Inked"})
    pkg = PptxPackage(deck)
    part = pkg.slide_parts()[0]
    pkg.add_part_with_content_type(
        "ppt/ink/ink1.xml", b'<inkml:ink xmlns:inkml="http://www.w3.org/2003/InkML"/>',
        "application/inkml+xml")
    rid = pkg.add_relationship(
        part, "http://schemas.microsoft.com/office/2011/relationships/ink",
        "../ink/ink1.xml")
    grp, _c = shp._find_shape(pkg, part, "Inked")
    grp.append(etree.fromstring(f'<p:contentPart {_NS} r:id="{rid}"/>'))
    pkg.mark_dirty(part)
    pkg.save(do_backup=False)
    message = _refused({"file_path": dest, "source_path": deck, "source_slide": 0,
                        "shape": "Inked", "slide": 1})
    assert "ink" in message


def test_refuses_a_jump_link_into_another_deck(source, dest):
    _ok("ppt_edit", "set_hyperlink", {
        "file_path": source["deck"], "slide": 0, "target": source["band"],
        "to_slide": 1})
    message = _refused({"file_path": dest, "source_path": source["deck"],
                        "source_slide": 0, "shape": "Figure", "slide": 1})
    assert "jumps to a slide" in message


def test_refuses_an_unknown_relationship(source, dest):
    pkg = PptxPackage(source["deck"])
    part = pkg.slide_parts()[0]
    rid = pkg.add_relationship(part, "http://example.com/relationships/odd",
                               "https://example.com/odd", external=True)
    band, _c = shp._find_shape(pkg, part, "Band A")
    band.find(qn("p:nvSpPr")).find(qn("p:cNvPr")).append(etree.fromstring(
        f'<a:hlinkHover {_NS} r:id="{rid}"/>'))
    pkg.mark_dirty(part)
    pkg.save(do_backup=False)
    message = _refused({"file_path": dest, "source_path": source["deck"],
                        "source_slide": 0, "shape": "Band A", "slide": 1})
    assert "http://example.com/relationships/odd" in message


def test_refuses_a_member_of_a_rotated_group(source, dest):
    _ok("ppt_graphics", "set_shape", {
        "file_path": source["deck"], "slide": 0, "shape": "Figure", "rotation": 30})
    message = _refused({"file_path": dest, "source_path": source["deck"],
                        "source_slide": 0, "shape": "Figure/Label", "slide": 1})
    assert "rotated or flipped" in message
    # The rotated group itself copies whole.
    _copy({"file_path": dest, "source_path": source["deck"], "source_slide": 0,
           "shape": "Figure", "slide": 1})


def test_refuses_a_new_size_for_a_table(tmp_path, dest):
    deck = _deck(tmp_path / "table.pptx")
    table = _ok("ppt_tables", "create_table", {
        "file_path": deck, "slide": 0, "rows": 2, "cols": 2, "x": 1, "y": 1,
        "w": 4, "h": 1.5, "data": [["a", "b"], ["c", "d"]]})
    table_id = table.get("shape_id") or table.get("id")
    message = _refused({"file_path": dest, "source_path": deck, "source_slide": 0,
                        "shape": table_id, "slide": 1, "w": 2})
    assert "table" in message
    _copy({"file_path": dest, "source_path": deck, "source_slide": 0,
           "shape": table_id, "slide": 1, "x": 2})


def test_refuses_leave_out_without_a_group_or_with_nothing_left(source, dest):
    message = _refused({"file_path": dest, "source_path": source["deck"],
                        "source_slide": 0, "shape": "Band A", "slide": 1,
                        "leave_out": ["Label"]})
    assert "group" in message
    every = ["Band A", "Label", "Caption", "Logo", "Legend", "Frame"]
    message = _refused({"file_path": dest, "source_path": source["deck"],
                        "source_slide": 0, "shape": "Figure", "slide": 1,
                        "leave_out": every})
    assert "nothing" in message
    message = _refused({"file_path": dest, "source_path": source["deck"],
                        "source_slide": 0, "shape": "Figure", "slide": 1,
                        "leave_out": ["No such member"]})
    assert "No such member" in message


def test_refuses_a_source_slide_out_of_range_naming_the_source(source, dest):
    message = _refused({"file_path": dest, "source_path": source["deck"],
                        "source_slide": 5, "shape": "Figure", "slide": 1})
    assert "source_slide" in message and "2 slides" in message
    unchanged = _md5(dest)
    _refused({"file_path": dest, "source_path": source["deck"],
              "source_slide": 0, "shape": "Nope", "slide": 1})
    assert _md5(dest) == unchanged
