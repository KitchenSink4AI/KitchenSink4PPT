"""A refused save or restore never costs the undo of the last change.

The save rotated prev from the current content BEFORE the swap. When the
swap was then refused (another program holding the deck; "The original file
was NOT modified"), prev already held the current content: the version
before the last change was gone, and a later "restore prev" answered ok and
changed nothing. A refused restore did the same to the very backup it was
restoring. Now the rotation is staged first and lands on the slots only
after the swap went through (a two-phase rotation), and a
restore that would change nothing says so.
"""

from __future__ import annotations

import asyncio
import json
import os
import sys
import time
from pathlib import Path

import pytest
from pptx import Presentation

from kitchensink4ppt import server  # noqa: F401  (registers the ops)
from kitchensink4ppt.core import safesave
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.plan import build

windows_only = pytest.mark.skipif(sys.platform != "win32", reason="Windows sharing modes")


@pytest.fixture(autouse=True)
def own_journal(tmp_path, monkeypatch):
    monkeypatch.setenv("KS4P_SAVE_JOURNAL_DIR", str(tmp_path / "journal"))


def _deck(path: Path, notes: str) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "T"
    slide.notes_slide.notes_text_frame.text = notes
    prs.save(path)
    return path


def _notes(path) -> str:
    return Presentation(str(path)).slides[0].notes_slide.notes_text_frame.text


def _call(op, args):
    plan = build(Settings())
    return asyncio.run(dispatch(plan, plan.specs[op].group, {"op": op, "args": args}))


def _set_notes(deck, text):
    return _call("set_notes", {"file_path": str(deck), "slide": 0, "text": text, "live": "off"})


def _three_versions(tmp_path) -> tuple[Path, Path]:
    deck = _deck(tmp_path / "deck.pptx", "v1")
    for version in ("v2", "v3"):
        payload, is_error = _set_notes(deck, version)
        assert not is_error, payload
    prev = safesave.slot_dir(deck) / safesave.PREV_SLOT
    assert _notes(deck) == "v3" and _notes(prev) == "v2"
    return deck, prev


def _slot_temps(deck) -> list[str]:
    return [n for n in os.listdir(safesave.slot_dir(deck)) if n.startswith(".slot-")]


@windows_only
def test_a_refused_save_keeps_the_undo_of_the_last_change(tmp_path):
    deck, prev = _three_versions(tmp_path)
    with open(deck, "rb"):  # shares read and write, not delete: the swap is refused
        payload, is_error = _set_notes(deck, "v4")
    assert is_error and payload["error"]["code"] == "DOCUMENT_LOCKED", payload
    assert _notes(deck) == "v3" and _notes(prev) == "v2"
    assert _slot_temps(deck) == []
    payload, is_error = _call("manage_backups", {"action": "restore", "file_path": str(deck),
                                                 "source": "prev"})
    assert not is_error, payload
    assert _notes(deck) == "v2", "undo still goes back to the version before the last change"


@windows_only
def test_a_refused_restore_keeps_the_backup_it_was_restoring(tmp_path):
    deck, prev = _three_versions(tmp_path)
    with open(deck, "rb"):
        payload, is_error = _call("manage_backups", {"action": "restore",
                                                     "file_path": str(deck), "source": "prev"})
    assert is_error and payload["error"]["code"] == "DOCUMENT_LOCKED", payload
    assert _notes(deck) == "v3" and _notes(prev) == "v2"
    assert _slot_temps(deck) == []


@windows_only
def test_a_refused_copy_overwrite_keeps_the_destinations_undo(tmp_path):
    deck, prev = _three_versions(tmp_path)
    other = _deck(tmp_path / "other.pptx", "other")
    with open(deck, "rb"):
        payload, is_error = _call("copy_presentation", {"file_path": str(other),
                                                        "dest_path": str(deck),
                                                        "overwrite": True})
    assert is_error and payload["error"]["code"] == "DOCUMENT_LOCKED", payload
    assert _notes(deck) == "v3" and _notes(prev) == "v2"


def test_a_save_that_goes_through_still_rotates_prev_and_anchor(tmp_path):
    deck = _deck(tmp_path / "deck.pptx", "v1")
    payload, is_error = _set_notes(deck, "v2")
    assert not is_error, payload
    folder = safesave.slot_dir(deck)
    assert _notes(folder / safesave.PREV_SLOT) == "v1"
    assert _notes(folder / safesave.ANCHOR_SLOT) == "v1"
    assert _slot_temps(deck) == []


def test_a_restore_that_would_change_nothing_says_so(tmp_path):
    deck = _deck(tmp_path / "deck.pptx", "v1")
    assert not _set_notes(deck, "v2")[1]
    assert not _call("manage_backups", {"action": "restore", "file_path": str(deck),
                                        "source": "prev"})[1]
    assert _notes(deck) == "v1"
    before = deck.read_bytes()
    prev = safesave.slot_dir(deck) / safesave.PREV_SLOT
    prev_before = prev.read_bytes()
    payload, is_error = _call("manage_backups", {"action": "restore", "file_path": str(deck),
                                                 "source": "anchor"})
    assert is_error is True and payload["error"]["code"] == "CONFLICT", payload
    assert "would change nothing" in payload["error"]["message"]
    assert deck.read_bytes() == before and prev.read_bytes() == prev_before


def test_a_failed_slot_update_after_a_good_save_is_a_warning(tmp_path, monkeypatch):
    deck = _deck(tmp_path / "deck.pptx", "v1")
    real_commit = safesave.RotationTicket.commit

    def refuse(self):
        raise PermissionError(13, "the slot is held", "prev.pptx")

    monkeypatch.setattr(safesave.RotationTicket, "commit", refuse)
    payload, is_error = _set_notes(deck, "v2")
    monkeypatch.setattr(safesave.RotationTicket, "commit", real_commit)
    assert is_error is False, payload
    assert _notes(deck) == "v2"
    assert any("backup could not be updated" in w for w in payload["warnings"]), payload
    assert _slot_temps(deck) == []


# ------------------------------- a server stopped between the swap and the slots


def _entry(deck: Path, ticket, new: Path) -> str:
    folder = safesave._journal_dir()
    folder.mkdir(parents=True, exist_ok=True)
    entry = folder / "stopped.json"
    entry.write_text(json.dumps({
        "deck": str(deck), "aside": str(deck.parent / "nothing-aside.tmp"), "new": str(new),
        "staged": ticket.staged(), "pid": os.getpid(), "token": safesave._OWNER_TOKEN,
        "pid_created": safesave._OWNER_CREATED, "host": safesave._local_host(),
        "time": time.time()}), encoding="utf-8")
    return str(entry)


def test_a_stop_after_the_swap_lands_the_staged_backups_at_the_next_write(tmp_path):
    """The swap went through and the server was stopped before the staged
    rotation landed: the next write lock lands it, so undo goes back to the
    version from before that save."""
    deck = _deck(tmp_path / "deck.pptx", "v1")
    assert not _set_notes(deck, "v2")[1]  # prev = v1
    ticket = safesave.prepare_rotation(deck)  # stages v2
    new = _deck(tmp_path / "deck.pptx.ppt-mcp-tmp", "v3")
    os.replace(new, deck)  # the swap went through
    entry = _entry(deck, ticket, new)
    with safesave.write_lock(deck):
        pass
    prev = safesave.slot_dir(deck) / safesave.PREV_SLOT
    assert _notes(deck) == "v3" and _notes(prev) == "v2"
    assert _slot_temps(deck) == [] and not os.path.exists(entry)


def test_a_stop_before_the_swap_drops_the_staged_backups_at_the_next_write(tmp_path):
    deck = _deck(tmp_path / "deck.pptx", "v1")
    assert not _set_notes(deck, "v2")[1]  # prev = v1
    ticket = safesave.prepare_rotation(deck)
    new = _deck(tmp_path / "deck.pptx.ppt-mcp-tmp", "v3")  # never swapped in
    entry = _entry(deck, ticket, new)
    with safesave.write_lock(deck):
        pass
    prev = safesave.slot_dir(deck) / safesave.PREV_SLOT
    assert _notes(deck) == "v2" and _notes(prev) == "v1"
    assert _slot_temps(deck) == [] and not new.exists() and not os.path.exists(entry)
