"""list_hyperlinks lists the links on a shape PowerPoint keeps in a wrapper.

When PowerPoint saves a text box, a table or a notes body that holds an
equation, it keeps it in a compatibility wrapper (mc:AlternateContent): the
branch PowerPoint shows, and a Fallback copy for older readers. The other
reads list such a shape and read its words (get_slide_info, find_text), but
list_hyperlinks walked past the wrapper, so a link on the box, its text, a
cell of the table or the notes was missing, while its op line says "every
link". It now reads the branch PowerPoint shows, so each link is listed
once.

Every deck here is made up by code.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.package import PptxPackage
from kitchensink4ppt.ops import equations, links

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A14 = "http://schemas.microsoft.com/office/drawing/2010/main"


def _wrap(element: etree._Element) -> None:
    """Put `element` in an mc:AlternateContent where it stands, with the
    element in the Choice (Requires a14) and a copy of it in the Fallback,
    as PowerPoint keeps a box, table or notes body that holds an equation."""
    wrapper = etree.Element(f"{{{MC}}}AlternateContent", nsmap={"mc": MC})
    choice = etree.SubElement(wrapper, f"{{{MC}}}Choice", nsmap={"a14": A14})
    choice.set("Requires", "a14")
    fallback = etree.SubElement(wrapper, f"{{{MC}}}Fallback")
    element.addprevious(wrapper)
    fallback.append(etree.fromstring(etree.tostring(element)))
    choice.append(element)


def _deck(path: Path, *, wrapped: bool) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.name = "Link box"
    box.text_frame.text = "Link box words"
    box.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://example.com/box-text"
    box.click_action.hyperlink.address = "https://example.com/box-click"
    plain = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(4), Inches(1))
    plain.name = "Plain box"
    plain.text_frame.text = "Plain words"
    plain.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://example.com/plain"
    frame = slide.shapes.add_table(1, 2, Inches(1), Inches(4), Inches(4), Inches(1))
    frame.name = "Link table"
    cell = frame.table.cell(0, 1)
    cell.text = "Cell words"
    cell.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://example.com/cell"
    notes = slide.notes_slide.notes_text_frame
    notes.text = "Notes words"
    notes.paragraphs[0].runs[0].hyperlink.address = "https://example.com/notes"
    prs.save(str(path))
    pkg = PptxPackage(path)
    equations.add_equation_to_shape(pkg, 0, "Link box", "x+1=2")
    pkg.save(do_backup=False)
    if not wrapped:
        return path

    def fix(name, data):
        root = etree.fromstring(data)
        if name == "ppt/slides/slide1.xml":
            for el in list(root.iter(f"{{{P}}}sp", f"{{{P}}}graphicFrame")):
                cnvpr = next(el.iter(f"{{{P}}}cNvPr"))
                if cnvpr.get("name") in ("Link box", "Link table"):
                    _wrap(el)
        else:
            body = next(sp for sp in root.iter(f"{{{P}}}sp")
                        if sp.find(f".//{{{P}}}ph") is not None
                        and sp.find(f".//{{{P}}}ph").get("type") == "body")
            _wrap(body)
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    tmp = path.with_name(path.name + ".rw")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        notes_parts = [n for n in zin.namelist()
                       if n.startswith("ppt/notesSlides/") and n.endswith(".xml")]
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "ppt/slides/slide1.xml" or info.filename in notes_parts:
                data = fix(info.filename, data)
            zout.writestr(info, data)
    tmp.replace(path)
    return path


EXPECTED = {
    ("https://example.com/box-text", "text", "Link box"),
    ("https://example.com/box-click", "shape", "Link box"),
    ("https://example.com/plain", "text", "Plain box"),
    ("https://example.com/cell", "table", "Link table"),
    ("https://example.com/notes", "notes", None),
}


def _listed(path: Path) -> list[tuple]:
    out = links.list_hyperlinks(PptxPackage(path))
    assert out["count"] == len(out["hyperlinks"])
    rows = []
    for h in out["hyperlinks"]:
        name = h.get("shape_name") if h["where"] != "notes" else None
        rows.append((h.get("url"), h["where"], name))
    return rows


def test_every_link_is_listed_on_a_bare_deck(tmp_path):
    rows = _listed(_deck(tmp_path / "made_up_bare.pptx", wrapped=False))
    assert set(rows) == EXPECTED and len(rows) == len(EXPECTED), rows


def test_links_in_powerpoint_wrappers_are_listed_once(tmp_path):
    rows = _listed(_deck(tmp_path / "made_up_wrapped.pptx", wrapped=True))
    assert set(rows) == EXPECTED, rows
    assert len(rows) == len(EXPECTED), f"a link was listed from both branches: {rows}"


def test_the_help_names_what_is_left_out():
    """The op line says "every link and click action"; a click that only
    plays a sound is listed (test_list_hyperlinks_sounds.py), and the help
    says the play controls of a video or audio shape are not links."""
    from kitchensink4ppt import server  # noqa: F401
    from kitchensink4ppt.settings import Settings
    from kitchensink4ppt.surface.help import help_payload
    from kitchensink4ppt.surface.plan import build

    payload = help_payload(build(Settings()), "list_hyperlinks")
    assert "every link and click action" in payload["summary"], payload["summary"]
    scope = payload["args_schema"]["properties"]["scope"]["description"]
    assert "only plays a sound is listed" in scope and "play controls" in scope, scope
