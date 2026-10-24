"""Safety tests for the bounded server-exit PowerPoint reaper."""

from __future__ import annotations

import gc
import os
import sys
import time
import types

import pytest

from kitchensink4ppt.com import bridge
from kitchensink4ppt.core import safesave


class _Presentation:
    def __init__(self, full_name: str, *, saved: bool):
        self.FullName = full_name
        self.Saved = -1 if saved else 0
        self.close_calls = 0

    def Close(self):  # noqa: N802 - COM spelling
        self.close_calls += 1


class _Presentations:
    def __init__(self, *items: _Presentation):
        self._items = items
        self.Count = len(items)

    def Item(self, index: int):  # noqa: N802 - COM spelling
        return self._items[index - 1]


class _App:
    def __init__(self, *presentations: _Presentation, visible: bool = False):
        self.Visible = -1 if visible else 0
        self.Presentations = _Presentations(*presentations)


def _install_fake_com(monkeypatch, app: _App, pid: int = 4242):
    pythoncom = types.ModuleType("pythoncom")
    client = types.ModuleType("win32com.client")
    client.GetActiveObject = lambda _progid: app
    win32com = types.ModuleType("win32com")
    win32com.__path__ = []
    win32com.client = client
    monkeypatch.setitem(sys.modules, "pythoncom", pythoncom)
    monkeypatch.setitem(sys.modules, "win32com", win32com)
    monkeypatch.setitem(sys.modules, "win32com.client", client)
    monkeypatch.setattr(bridge, "_ensure_apartment", lambda _module: "initialized")
    monkeypatch.setattr(bridge, "_powerpoint_app_pid", lambda _app: pid)


def test_exit_inspection_closes_only_saved_owned_presentations(monkeypatch, tmp_path):
    owned = _Presentation(str(tmp_path / "owned.pptx"), saved=True)
    foreign = _Presentation(str(tmp_path / "foreign.pptx"), saved=False)
    _install_fake_com(monkeypatch, _App(owned, foreign))

    safe, why = bridge._inspect_exit_ownership(
        4242, {safesave.canonical_key(owned.FullName)}
    )

    assert safe is False
    assert "did not open" in why
    assert owned.close_calls == 1
    assert foreign.close_calls == 0


def test_exit_inspection_preserves_unsaved_owned_presentation(monkeypatch, tmp_path):
    owned = _Presentation(str(tmp_path / "owned.pptx"), saved=False)
    _install_fake_com(monkeypatch, _App(owned))

    safe, why = bridge._inspect_exit_ownership(
        4242, {safesave.canonical_key(owned.FullName)}
    )

    assert safe is False
    assert "unsaved changes" in why
    assert owned.close_calls == 0


def test_exit_inspection_preserves_new_untitled_presentation(monkeypatch):
    untitled = _Presentation("", saved=False)
    _install_fake_com(monkeypatch, _App(untitled))

    safe, why = bridge._inspect_exit_ownership(4242, set())

    assert safe is False
    assert "did not open" in why
    assert untitled.close_calls == 0


def test_exit_inspection_allows_only_saved_owned_presentations(monkeypatch, tmp_path):
    owned = _Presentation(str(tmp_path / "owned.pptx"), saved=True)
    _install_fake_com(monkeypatch, _App(owned))

    safe, why = bridge._inspect_exit_ownership(
        4242, {safesave.canonical_key(owned.FullName)}
    )

    assert safe is True
    assert "no foreign or unsaved" in why
    assert owned.close_calls == 1


def test_hidden_proxy_pid_check_uses_only_an_unambiguous_process_table(monkeypatch):
    app = object()
    monkeypatch.setattr(bridge, "powerpnt_pids", lambda: {4242})
    assert bridge._powerpoint_app_pid(app) == 4242
    monkeypatch.setattr(bridge, "powerpnt_pids", lambda: {4242, 4343})
    assert bridge._powerpoint_app_pid(app) is None
    monkeypatch.setattr(bridge, "powerpnt_pids", lambda: None)
    assert bridge._powerpoint_app_pid(app) is None


def test_identity_only_cleanup_accepts_owned_modal_when_com_cannot_answer(
    monkeypatch, tmp_path
):
    from kitchensink4ppt.com import dialogs

    monkeypatch.setattr(
        dialogs,
        "window_classes",
        lambda pids: [{"title": "Password", "class": "NUIDialog"}],
    )

    safe, why = bridge._identity_only_modal_cleanup_safe(
        4242,
        {safesave.canonical_key(tmp_path / "password.pptx")},
        "ownership inspection failed: AttributeError",
    )

    assert safe is True
    assert "only the owned pid's modal dialog" in why


def test_identity_only_cleanup_refuses_a_visible_application_frame(
    monkeypatch, tmp_path
):
    from kitchensink4ppt.com import dialogs

    monkeypatch.setattr(
        dialogs,
        "window_classes",
        lambda pids: [
            {"title": "Password", "class": "NUIDialog"},
            {"title": "User deck", "class": "PPTFrameClass"},
        ],
    )

    safe, why = bridge._identity_only_modal_cleanup_safe(
        4242,
        {safesave.canonical_key(tmp_path / "password.pptx")},
        "ownership inspection timed out",
    )

    assert safe is False
    assert "visible non-dialog window" in why


def test_identity_only_cleanup_never_overrides_a_com_safety_refusal(
    monkeypatch, tmp_path
):
    from kitchensink4ppt.com import dialogs

    monkeypatch.setattr(
        dialogs,
        "window_classes",
        lambda pids: [{"title": "Password", "class": "NUIDialog"}],
    )

    safe, why = bridge._identity_only_modal_cleanup_safe(
        4242,
        {safesave.canonical_key(tmp_path / "password.pptx")},
        "it holds a presentation this server did not open: user.pptx",
    )

    assert safe is False
    assert "answered and refused" in why


def test_identity_only_cleanup_requires_a_journalled_read_only_open(monkeypatch):
    from kitchensink4ppt.com import dialogs

    monkeypatch.setattr(
        dialogs,
        "window_classes",
        lambda pids: [{"title": "Password", "class": "NUIDialog"}],
    )

    safe, why = bridge._identity_only_modal_cleanup_safe(
        4242,
        set(),
        "ownership inspection failed: AttributeError",
    )

    assert safe is False
    assert "no journalled read-only" in why


@pytest.mark.skipif(sys.platform != "win32", reason="Windows process API shape")
def test_reaper_uses_identity_only_route_when_owned_modal_blocks_com(
    monkeypatch, tmp_path
):
    class _Function:
        def __init__(self, fn):
            self.fn = fn
            self.argtypes = None

        def __call__(self, *args):
            return self.fn(*args)

    terminated = []

    class _Kernel:
        def __init__(self):
            self.OpenProcess = _Function(lambda *_args: 99)
            self.CloseHandle = _Function(lambda *_args: True)
            self.GetProcessTimes = _Function(lambda *_args: True)
            self.QueryFullProcessImageNameW = _Function(self._image)
            self.TerminateProcess = _Function(self._terminate)

        @staticmethod
        def _image(_process, _flags, image, _size):
            image.value = r"C:\Program Files\Microsoft Office\POWERPNT.EXE"
            return True

        @staticmethod
        def _terminate(_process, _exit_code):
            terminated.append(4242)
            return True

    monkeypatch.setattr(bridge, "_kernel32", lambda: _Kernel())
    monkeypatch.setattr(bridge, "_creation_time", lambda *_a, **_k: 123456)
    monkeypatch.setattr(bridge, "_is_automation_instance", lambda _pid: True)
    monkeypatch.setattr(
        bridge,
        "_inspect_exit_ownership",
        lambda *_a, **_k: (
            False,
            "ownership inspection failed: AttributeError",
        ),
    )
    monkeypatch.setattr(
        bridge,
        "_identity_only_modal_cleanup_safe",
        lambda *_a, **_k: (True, "owned modal only"),
    )

    result = bridge._end_owned_powerpoint_at_exit(
        {4242: 123456},
        {4242: {safesave.canonical_key(tmp_path / "password.pptx")}},
    )

    assert result == {"ended": [4242], "left": []}
    assert terminated == [4242]


def _wait_for_powerpoint_exit(pid: int, timeout: float = 20.0):
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        if pid not in (bridge.powerpnt_pids() or set()):
            return
        time.sleep(0.25)
    raise AssertionError(f"PowerPoint pid {pid} did not exit after Quit")


@pytest.mark.live
@pytest.mark.skipif(
    os.name != "nt" or os.environ.get("KS4_VERIFY_PPT_COM") != "1",
    reason="needs PowerPoint in a clean window: set KS4_VERIFY_PPT_COM=1",
)
def test_live_exit_inspection_preserves_new_untitled_presentation():
    import pythoncom
    import win32com.client

    if bridge.powerpnt_pids():
        pytest.skip("a PowerPoint is already running")
    pythoncom.CoInitialize()
    app = win32com.client.DispatchEx("PowerPoint.Application")
    started = bridge.powerpnt_pids() or set()
    assert len(started) == 1
    pid = next(iter(started))
    pres = app.Presentations.Add(False)
    try:
        safe, why = bridge._inspect_exit_ownership(pid, set())
        assert safe is False
        assert "did not open" in why
        assert pid in (bridge.powerpnt_pids() or set())
        assert int(app.Presentations.Count) == 1
    finally:
        pres.Saved = -1
        pres.Close()
        pres = None
        app.Quit()
        app = None
        gc.collect()
        _wait_for_powerpoint_exit(pid)


@pytest.mark.live
@pytest.mark.skipif(
    os.name != "nt" or os.environ.get("KS4_VERIFY_PPT_COM") != "1",
    reason="needs PowerPoint in a clean window: set KS4_VERIFY_PPT_COM=1",
)
def test_live_exit_inspection_preserves_unsaved_owned_presentation(tmp_path):
    import pythoncom
    import win32com.client
    from pptx import Presentation

    if bridge.powerpnt_pids():
        pytest.skip("a PowerPoint is already running")
    deck = tmp_path / "owned-unsaved.pptx"
    Presentation().save(deck)
    pythoncom.CoInitialize()
    app = win32com.client.DispatchEx("PowerPoint.Application")
    started = bridge.powerpnt_pids() or set()
    assert len(started) == 1
    pid = next(iter(started))
    pres = app.Presentations.Open(
        str(deck), ReadOnly=False, Untitled=False, WithWindow=False
    )
    pres.Slides.Add(int(pres.Slides.Count) + 1, 12)
    assert not bool(pres.Saved)
    try:
        safe, why = bridge._inspect_exit_ownership(
            pid, {safesave.canonical_key(deck)}
        )
        assert safe is False
        assert "unsaved changes" in why
        assert pid in (bridge.powerpnt_pids() or set())
        assert int(app.Presentations.Count) == 1
    finally:
        pres.Saved = -1
        pres.Close()
        pres = None
        app.Quit()
        app = None
        gc.collect()
        _wait_for_powerpoint_exit(pid)
