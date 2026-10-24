"""Damage inside a deck is reported as damage, not as the caller's mistake.

A slide part cut short, or a shape that lost its id, is damage PowerPoint
repairs when it opens the deck. A truncated part surfaced as a bare parser
message under BAD_PARAMS, naming no part; a shape without an id made
get_slide_info, list_elements, find_text, the view, the accessibility audit
and set_placeholder_text fail with Python's own TypeError text, also as
BAD_PARAMS. Now a damaged part is refused by name with the next step, and a
shape without an id reads with id null.

Every deck here is made up by code.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.errors import DocumentCorrupt
from kitchensink4ppt.core.package import PptxPackage, qn
from kitchensink4ppt.ops import access, read, text, view
from kitchensink4ppt.surface import envelope


def _rewrite(path, part, fn):
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            z.writestr(info, fn(data) if info.filename == part else data)
    path.write_bytes(out.getvalue())


def _deck(tmp_path, name):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Made-up title"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.name = "Made-up box"
    box.text_frame.text = "Made-up words"
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "Second made-up slide"
    path = tmp_path / name
    prs.save(str(path))
    return path


def test_a_truncated_slide_is_refused_by_name_as_damage(tmp_path):
    path = _deck(tmp_path, "truncated.pptx")
    _rewrite(path, "ppt/slides/slide1.xml", lambda data: data[: len(data) // 2])
    with pytest.raises(DocumentCorrupt) as refused:
        read.get_text(PptxPackage(path))
    message = str(refused.value)
    assert "ppt/slides/slide1.xml" in message and "PowerPoint" in message
    assert envelope.refusal(refused.value)["error"]["code"] == "UNSUPPORTED_CONTENT"
    # Code that reads around a part it cannot parse still can.
    assert isinstance(refused.value, etree.XMLSyntaxError)
    assert read.get_text(PptxPackage(path), scope=1)["slides"][0]["text"] == "Second made-up slide"


def _without_title_id(data):
    root = etree.fromstring(data)
    for cnv in root.iter(qn("p:cNvPr")):
        if cnv.get("name", "").startswith("Title"):
            del cnv.attrib["id"]
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def test_a_shape_without_an_id_reads_with_id_null(tmp_path):
    path = _deck(tmp_path, "no_id.pptx")
    _rewrite(path, "ppt/slides/slide1.xml", _without_title_id)
    pkg = PptxPackage(path)
    info = read.get_slide_info(pkg, 0)
    ids = {s["name"]: s["id"] for s in info["shapes"]}
    assert ids["Made-up box"] is not None
    assert [v for k, v in ids.items() if k.startswith("Title")] == [None]
    assert read.list_elements(pkg, "shapes", scope=0)["count"] == 2
    assert read.find_text(pkg, "Made-up title")["count"] == 1
    shown = view.get_presentation_view(pkg)["view"]
    assert "[no id] title: Made-up title" in shown and "Made-up words" in shown
    access.audit_accessibility(pkg)


def test_placeholder_text_goes_into_a_shape_without_an_id(tmp_path):
    path = _deck(tmp_path, "no_id_write.pptx")
    _rewrite(path, "ppt/slides/slide1.xml", _without_title_id)
    pkg = PptxPackage(path)
    text.set_placeholder_text(pkg, 0, "title", "New made-up title")
    pkg.save(do_backup=False)
    assert "New made-up title" in read.get_text(PptxPackage(path), scope=0)["slides"][0]["text"]
