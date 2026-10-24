"""search_and_replace's case_note leaves out what match_case=false cannot replace.

With match_case true (the default), case_note counts the matches in another
case that were left unchanged and says that match_case=false replaces them
too. A match inside an equation (or other content PowerPoint keeps in a
compatibility wrapper) is never rewritten, in any case: a replace with
match_case=false leaves it and counts it in wrapped_matches_skipped. The note
counted such a match anyway, so its promise was false for it. It now counts
only matches a replace with match_case=false would change. find_text's note
still counts them, because find_text with match_case=false does list them.

Every deck here is made up by code.
"""

from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.package import PptxPackage, qn
from kitchensink4ppt.ops import read
from kitchensink4ppt.ops import text as tx

MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
A14 = "http://schemas.microsoft.com/office/drawing/2010/main"
OMML = "http://schemas.openxmlformats.org/officeDocument/2006/math"

#: An inline equation whose math reads "Budget=income-costs", as a
#: paragraph-level wrapper with a plain run for older readers.
EQUATION = (
    f'<mc:AlternateContent xmlns:mc="{MC}" xmlns:a="{A}"><mc:Choice xmlns:a14="{A14}" '
    f'Requires="a14"><a14:m><m:oMath xmlns:m="{OMML}"><m:r><m:t>Budget=income-costs</m:t>'
    "</m:r></m:oMath></a14:m></mc:Choice><mc:Fallback><a:r><a:t>Budget=income-costs</a:t>"
    "</a:r></mc:Fallback></mc:AlternateContent>"
)


def _deck(tmp_path):
    """Slide 1: "The budget line: " then the equation. Slide 2: "Budget
    risks are listed here." in plain text."""
    prs = Presentation()
    for words in ("The budget line: ", "Budget risks are listed here."):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1)).text_frame.text = words
    path = tmp_path / "case_wrapped.pptx"
    prs.save(str(path))
    pkg = PptxPackage(path)
    part = pkg.slide_parts()[0]
    p = next(pkg.root(part).iter(qn("a:p")))
    p.append(etree.fromstring(EQUATION))
    pkg.mark_dirty(part)
    pkg.save(do_backup=False)
    return path


def test_the_premise_a_replace_in_any_case_leaves_the_match_in_the_equation(tmp_path):
    pkg = PptxPackage(_deck(tmp_path))
    out = tx.search_and_replace(pkg, "budget", "plan", match_case=False)
    assert out["total"] == 2  # slide 1's plain "budget" and slide 2's "Budget"
    assert out["wrapped_matches_skipped"] == 1  # the equation's "Budget"


def test_the_replace_case_note_does_not_count_a_match_in_the_equation(tmp_path):
    pkg = PptxPackage(_deck(tmp_path))
    out = tx.search_and_replace(pkg, "budget", "plan")
    assert out["total"] == 1
    assert out["case_note"] == (
        '1 match in another case ("Budget") on slide_number 2 was left unchanged; '
        "send match_case=false to replace it too."
    )


def test_no_case_note_when_the_only_other_case_match_is_in_the_equation(tmp_path):
    path = _deck(tmp_path)
    pkg = PptxPackage(path)
    out = tx.search_and_replace(pkg, "budget", "plan", scope=0)
    assert out["total"] == 1
    assert "case_note" not in out


def test_find_text_still_counts_the_match_in_the_equation(tmp_path):
    """find_text with match_case=false lists the equation's match, so its
    note still counts it."""
    pkg = PptxPackage(_deck(tmp_path))
    out = read.find_text(pkg, "budget")
    assert out["count"] == 1
    assert out["case_note"].startswith('2 more matches in another case ("Budget")')
    assert "slide_numbers 1 and 2" in out["case_note"]
    assert read.find_text(pkg, "budget", match_case=False)["count"] == 3
