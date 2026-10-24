"""Copying a slide whose picture points at a missing image.

A picture names its image through a relationship; when the image part is
not in the deck, PowerPoint opens the deck and shows a missing picture.
copy_slide_between and merge_decks failed on such a slide with an internal
error. Now the copy shows the picture as missing, as PowerPoint does, and the
answer says so; a missing part that something other than a picture needs is
refused by name, with nothing changed.

Every deck here is made up by code.
"""

from __future__ import annotations

import io
import struct
import zipfile
import zlib

import pytest
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.errors import UnsupportedStructure
from kitchensink4ppt.core.package import PptxPackage, qn, rels_name
from kitchensink4ppt.ops import assembly, interdeck

R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _png_bytes():
    raw = b"".join(b"\x00" + b"\x40\x80\xc0" * 2 for _ in range(2))

    def chunk(kind, data):
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def _drop_part(path, name):
    """Remove one part from the zip, as a damaged deck has lost it."""
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist() if i.filename != name]
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            z.writestr(info, data)
    path.write_bytes(out.getvalue())


def _source(tmp_path, *, chart=False):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "Made-up words"
    png = tmp_path / "band.png"
    png.write_bytes(_png_bytes())
    pic = slide.shapes.add_picture(str(png), Inches(5), Inches(3), Inches(1), Inches(1))
    pic.name = "Made-up band"
    if chart:
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE

        data = CategoryChartData()
        data.categories = ["A", "B"]
        data.add_series("Made-up", (1, 2))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3),
                               Inches(3), Inches(2), data)
    path = tmp_path / ("src_chart.pptx" if chart else "src.pptx")
    prs.save(str(path))
    _drop_part(path, "ppt/charts/chart1.xml" if chart else "ppt/media/image1.png")
    return str(path)


def _dest(tmp_path, name="dst.pptx"):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / name
    prs.save(str(path))
    return path


def _last_slide(path):
    pkg = PptxPackage(path)
    part = pkg.slide_parts()[-1]
    return pkg, part


def _check_missing_picture_copy(path):
    pkg, part = _last_slide(path)
    blip = next(pkg.root(part).iter(qn("a:blip")))
    assert blip.get(f"{{{R}}}embed") is None
    targets = [r.get("Target") for r in pkg.root(rels_name(part))]
    assert not any("image1.png" in t for t in targets)
    texts = [t.text for t in pkg.root(part).iter(qn("a:t"))]
    assert "Made-up words" in texts


def test_copy_slide_between_shows_the_picture_as_missing(tmp_path):
    source = _source(tmp_path)
    dest = _dest(tmp_path)
    pkg = PptxPackage(dest)
    out = interdeck.copy_slide_between(pkg, source, 0)
    pkg.save(do_backup=False)
    assert any("Made-up band" in w and "missing" in w for w in out["warnings"])
    _check_missing_picture_copy(dest)


def test_merge_decks_shows_the_picture_as_missing(tmp_path):
    source = _source(tmp_path)
    dest = _dest(tmp_path)
    pkg = PptxPackage(dest)
    out = assembly.merge_decks(pkg, [source])
    pkg.save(do_backup=False)
    assert any("Made-up band" in w and "missing" in w for w in out.get("warnings", []))
    _check_missing_picture_copy(dest)


def test_a_missing_part_a_chart_needs_is_refused_by_name(tmp_path):
    source = _source(tmp_path, chart=True)
    dest = _dest(tmp_path)
    before = dest.read_bytes()
    pkg = PptxPackage(dest)
    with pytest.raises(UnsupportedStructure) as refused:
        interdeck.copy_slide_between(pkg, source, 0)
    message = str(refused.value)
    assert "ppt/charts/chart1.xml" in message and "nothing was changed" in message
    assert dest.read_bytes() == before
