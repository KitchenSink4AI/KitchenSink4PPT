"""Paragraph edits in apply_edits, and what set_text does to formatting.

Field reports: a shape held two paragraphs with different formatting and
only the first had to go, but apply_edits had no paragraph edit and its
help said nothing about how set_text treats the paragraphs it replaces, so
the slide was rebuilt elsewhere and copied in. And search_and_replace that
emptied a paragraph left the empty paragraph behind as a blank line.

apply_edits now takes set_paragraph, insert_paragraph and delete_paragraph,
each addressed by a shape and a paragraph counted from 0, as find_text and
format_text count them, or a paragraph_number counted from 1 (not both).
Every paragraph address in one batch is read as the shape was before the
batch, so several edits to one shape need no renumbering.
search_and_replace takes remove_empty_paragraphs. The help states exactly
what set_text does to a shape whose paragraphs differ, and the tests below
hold it to that statement.
"""

from __future__ import annotations

import asyncio
import hashlib
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

from kitchensink4ppt import server  # noqa: F401  (the import registers ops)
from kitchensink4ppt.core.package import PptxPackage, qn
from kitchensink4ppt.ops import batch
from kitchensink4ppt.ops.read import paragraph_text, slide_table
from kitchensink4ppt.registry import public_ops
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.help import help_payload
from kitchensink4ppt.surface.plan import build


def _plan():
    return build(Settings(), public_ops())


def _call(tool: str, op: str, args: dict):
    return asyncio.run(dispatch(_plan(), tool, {"op": op, "args": args}))


def _ok(tool: str, op: str, args: dict) -> dict:
    payload, is_error = _call(tool, op, args)
    assert is_error is False, payload
    return payload


def _refused(tool: str, op: str, args: dict) -> dict:
    payload, is_error = _call(tool, op, args)
    assert is_error is True, payload
    return payload["error"]


def _md5(path) -> str:
    return hashlib.md5(Path(path).read_bytes()).hexdigest()


def _batch(deck, edits) -> dict:
    return _ok("ppt_edit", "apply_edits", {"file_path": str(deck), "edits": edits})


def _batch_refused(deck, edits) -> dict:
    return _refused("ppt_edit", "apply_edits", {"file_path": str(deck), "edits": edits})


def _deck(tmp_path: Path, name: str = "paras.pptx") -> tuple[Path, int]:
    """One slide with a text box of three differently formatted paragraphs:
    a large bold line, an italic line and a plain line. Returns (deck,
    the box's shape id)."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3))
    tf = box.text_frame
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = "Big headline"
    r0.font.size, r0.font.bold = Pt(40), True
    p1 = tf.add_paragraph()
    r1 = p1.add_run()
    r1.text = "An italic aside"
    r1.font.size, r1.font.italic = Pt(18), True
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "A plain closing line"
    r2.font.size = Pt(16)
    other = prs.slides.add_slide(prs.slide_layouts[6])
    table = other.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "cell"
    path = tmp_path / name
    prs.save(str(path))
    return path, box.shape_id


def _paras(deck, shape_id: int, slide: int = 0) -> list[dict]:
    """Each paragraph of the box: text, size, bold, italic of its first run."""
    pkg = PptxPackage(deck)
    part = slide_table(pkg)[slide]["part"]
    for sp in pkg.root(part).iter(qn("p:sp")):
        cnv = sp.find(f"{qn('p:nvSpPr')}/{qn('p:cNvPr')}")
        if cnv is not None and int(cnv.get("id")) == shape_id:
            out = []
            for p in sp.find(qn("p:txBody")).findall(qn("a:p")):
                rpr = p.find(f"{qn('a:r')}/{qn('a:rPr')}")
                if rpr is None:
                    rpr = p.find(qn("a:endParaRPr"))
                get = (lambda k: rpr.get(k)) if rpr is not None else (lambda k: None)
                out.append({"text": paragraph_text(p), "sz": get("sz"),
                            "b": get("b"), "i": get("i")})
            return out
    raise AssertionError(f"no shape {shape_id}")


def _anchor(deck) -> str:
    """The text box's view anchor: the only shape on slide 1."""
    view = _ok("ppt_read", "get_presentation_view", {"file_path": str(deck), "scope": 0})
    section = view["view"].split("## Slide 1", 1)[1]
    return "a:" + section.split("[a:", 1)[1].split("]", 1)[0]


# ------------------------------------------------------------ delete_paragraph


def test_delete_one_paragraph_keeps_the_others_formatting(tmp_path):
    deck, sid = _deck(tmp_path)
    out = _batch(deck, [{"op": "delete_paragraph", "slide_number": 1, "shape": sid,
                         "paragraph": 0}])
    assert out["changed"]["0"]["slide_number"] == 1
    assert out["changed"]["0"]["shape_id"] == sid
    paras = _paras(deck, sid)
    assert [p["text"] for p in paras] == ["An italic aside", "A plain closing line"]
    assert paras[0]["i"] == "1" and paras[0]["sz"] == "1800"
    assert paras[1]["sz"] == "1600" and paras[1]["i"] is None


def test_paragraph_number_counts_from_one(tmp_path):
    deck, sid = _deck(tmp_path)
    _batch(deck, [{"op": "delete_paragraph", "slide_number": 1, "shape": sid,
                   "paragraph_number": 2}])
    assert [p["text"] for p in _paras(deck, sid)] == ["Big headline", "A plain closing line"]


def test_paragraph_and_paragraph_number_together_are_refused(tmp_path):
    deck, sid = _deck(tmp_path)
    before = _md5(deck)
    error = _batch_refused(deck, [{"op": "delete_paragraph", "slide_number": 1, "shape": sid,
                                   "paragraph": 1, "paragraph_number": 2}])
    assert "paragraph_number" in error["message"] and "not both" in error["message"]
    assert _md5(deck) == before


def test_paragraph_number_zero_is_refused(tmp_path):
    deck, sid = _deck(tmp_path)
    error = _batch_refused(deck, [{"op": "delete_paragraph", "slide_number": 1, "shape": sid,
                                   "paragraph_number": 0}])
    assert "counts from 1" in error["message"]


def test_indexes_are_read_before_the_batch(tmp_path):
    """Deleting paragraphs 0 and 2 in one batch deletes those two, not
    paragraph 0 and then whatever moved into place 2."""
    deck, sid = _deck(tmp_path)
    _batch(deck, [
        {"op": "delete_paragraph", "slide_number": 1, "shape": sid, "paragraph": 0},
        {"op": "delete_paragraph", "slide_number": 1, "shape": sid, "paragraph": 2},
    ])
    assert [p["text"] for p in _paras(deck, sid)] == ["An italic aside"]


def test_deleting_every_paragraph_is_refused(tmp_path):
    deck, sid = _deck(tmp_path)
    before = _md5(deck)
    error = _batch_refused(deck, [
        {"op": "delete_paragraph", "slide_number": 1, "shape": sid, "paragraph": i}
        for i in range(3)
    ])
    assert "every paragraph" in error["message"]
    assert _md5(deck) == before


def test_an_out_of_range_paragraph_names_the_count(tmp_path):
    deck, sid = _deck(tmp_path)
    error = _batch_refused(deck, [{"op": "delete_paragraph", "slide_number": 1, "shape": sid,
                                   "paragraph": 3}])
    assert "3 paragraphs" in error["message"]


def test_a_table_cell_anchor_is_refused(tmp_path):
    deck, _sid = _deck(tmp_path)
    view = _ok("ppt_read", "get_presentation_view", {"file_path": str(deck), "scope": 1})
    cell = "t:" + view["view"].split("cell anchors t:", 1)[1].split(":", 1)[0] + ":r1c1"
    error = _batch_refused(deck, [{"op": "delete_paragraph", "anchor": cell, "paragraph": 0}])
    assert "cannot target a table CELL anchor" in error["message"]


# ------------------------------------------------------------- set_paragraph


def test_set_paragraph_keeps_that_paragraphs_formatting(tmp_path):
    deck, sid = _deck(tmp_path)
    _batch(deck, [{"op": "set_paragraph", "slide_number": 1, "shape": sid,
                   "paragraph": 1, "text": "A new aside\von two lines"}])
    paras = _paras(deck, sid)
    assert paras[1]["text"] == "A new aside\von two lines"
    assert paras[1]["i"] == "1" and paras[1]["sz"] == "1800"
    assert paras[0]["text"] == "Big headline" and paras[0]["b"] == "1"
    assert paras[2]["text"] == "A plain closing line"


def test_set_paragraph_refuses_a_paragraph_break(tmp_path):
    deck, sid = _deck(tmp_path)
    error = _batch_refused(deck, [{"op": "set_paragraph", "slide_number": 1, "shape": sid,
                                   "paragraph": 1, "text": "one\ntwo"}])
    assert "insert_paragraph" in error["message"]


def test_set_and_delete_of_one_paragraph_are_refused(tmp_path):
    deck, sid = _deck(tmp_path)
    before = _md5(deck)
    error = _batch_refused(deck, [
        {"op": "set_paragraph", "slide_number": 1, "shape": sid, "paragraph": 1, "text": "x"},
        {"op": "delete_paragraph", "slide_number": 1, "shape": sid, "paragraph": 1},
    ])
    assert "both change paragraph 1" in error["message"]
    assert _md5(deck) == before


# ---------------------------------------------------------- insert_paragraph


def test_insert_takes_the_formatting_of_the_paragraph_it_goes_before(tmp_path):
    deck, sid = _deck(tmp_path)
    _batch(deck, [{"op": "insert_paragraph", "slide_number": 1, "shape": sid,
                   "paragraph": 1, "text": "Another aside"}])
    paras = _paras(deck, sid)
    assert [p["text"] for p in paras] == [
        "Big headline", "Another aside", "An italic aside", "A plain closing line"]
    assert paras[1]["i"] == "1" and paras[1]["sz"] == "1800"


def test_insert_with_no_paragraph_adds_it_last_like_the_last(tmp_path):
    deck, sid = _deck(tmp_path)
    out = _batch(deck, [{"op": "insert_paragraph", "slide_number": 1, "shape": sid,
                         "text": "Closing words"}])
    paras = _paras(deck, sid)
    assert paras[-1]["text"] == "Closing words" and paras[-1]["sz"] == "1600"
    assert out["changed"]["0"]["paragraphs"] == 4


def test_insert_then_delete_by_start_indexes(tmp_path):
    deck, sid = _deck(tmp_path)
    _batch(deck, [
        {"op": "insert_paragraph", "slide_number": 1, "shape": sid, "paragraph_number": 1,
         "text": "New first line"},
        {"op": "delete_paragraph", "slide_number": 1, "shape": sid, "paragraph_number": 1},
    ])
    paras = _paras(deck, sid)
    assert [p["text"] for p in paras] == [
        "New first line", "An italic aside", "A plain closing line"]
    assert paras[0]["b"] == "1" and paras[0]["sz"] == "4000"


def test_paragraph_edits_by_view_anchor(tmp_path):
    deck, sid = _deck(tmp_path)
    anchor = _anchor(deck)
    _batch(deck, [{"op": "delete_paragraph", "anchor": anchor, "paragraph": 0}])
    assert _paras(deck, sid)[0]["text"] == "An italic aside"


# --------------------------------------------- conflicts within one batch


def test_set_text_and_paragraph_edits_on_one_shape_are_refused(tmp_path):
    deck, sid = _deck(tmp_path)
    before = _md5(deck)
    error = _batch_refused(deck, [
        {"op": "set_text", "slide_number": 1, "shape": sid, "text": "All new"},
        {"op": "delete_paragraph", "slide_number": 1, "shape": sid, "paragraph": 0},
    ])
    assert "set_text" in error["message"] and "second batch" in error["message"]
    assert _md5(deck) == before


def test_format_text_by_paragraph_beside_an_insert_is_refused(tmp_path):
    deck, sid = _deck(tmp_path)
    error = _batch_refused(deck, [
        {"op": "insert_paragraph", "slide_number": 1, "shape": sid, "paragraph": 0, "text": "x"},
        {"op": "format_text", "slide_number": 1, "shape": sid, "paragraph": 1, "bold": True},
    ])
    assert "format_text" in error["message"] and "second batch" in error["message"]


def test_format_text_takes_paragraph_number_in_a_batch(tmp_path):
    deck, sid = _deck(tmp_path)
    _batch(deck, [{"op": "format_text", "slide_number": 1, "shape": sid,
                   "paragraph_number": 3, "bold": True}])
    assert _paras(deck, sid)[2]["b"] == "1"


# ------------------------------------------- search_and_replace and blanks


def test_remove_empty_paragraphs_drops_the_line_it_emptied(tmp_path):
    deck, sid = _deck(tmp_path)
    _ok("ppt_edit", "search_and_replace", {
        "file_path": str(deck), "find": "An italic aside", "replace": "",
        "remove_empty_paragraphs": True, "live": "off"})
    assert [p["text"] for p in _paras(deck, sid)] == ["Big headline", "A plain closing line"]


def test_without_the_option_the_empty_paragraph_stays(tmp_path):
    deck, sid = _deck(tmp_path)
    _ok("ppt_edit", "search_and_replace", {
        "file_path": str(deck), "find": "An italic aside", "replace": "", "live": "off"})
    assert [p["text"] for p in _paras(deck, sid)] == [
        "Big headline", "", "A plain closing line"]


def test_remove_empty_paragraphs_leaves_blank_lines_it_did_not_empty(tmp_path):
    deck, sid = _deck(tmp_path)
    _batch(deck, [{"op": "insert_paragraph", "slide_number": 1, "shape": sid,
                   "paragraph": 2, "text": ""}])
    assert [p["text"] for p in _paras(deck, sid)][2] == ""
    out = _batch(deck, [{"op": "search_and_replace", "find": "Big headline", "replace": "",
                         "remove_empty_paragraphs": True}])
    assert [p["text"] for p in _paras(deck, sid)] == [
        "An italic aside", "", "A plain closing line"]
    assert out["changed"]["0"]["paragraphs_removed"] == 1


def test_a_shape_keeps_one_paragraph(tmp_path):
    deck, sid = _deck(tmp_path)
    _ok("ppt_edit", "search_and_replace", {
        "file_path": str(deck), "find": r"^.+$", "replace": "", "regex": True,
        "remove_empty_paragraphs": True, "live": "off"})
    assert len(_paras(deck, sid)) == 1


def test_remove_empty_paragraphs_has_no_live_route(tmp_path, monkeypatch):
    from kitchensink4ppt.com import live_ops

    def boom(*_a, **_k):  # pragma: no cover - reaching PowerPoint is the failure
        raise AssertionError("reached the live layer")

    monkeypatch.setattr(live_ops, "live_search_and_replace", boom)
    deck, _sid = _deck(tmp_path)
    error = _refused("ppt_edit", "search_and_replace", {
        "file_path": str(deck), "find": "x", "replace": "", "live": "force",
        "remove_empty_paragraphs": True})
    assert "remove_empty_paragraphs" in error["message"]


# ------------------------------------------------- what set_text does, stated


SET_TEXT_CLAIMS = (
    "paragraph properties",
    "first run",
    "last",
    "dropped",
    "runs_collapsed_to_first",
    "set_paragraph",
)


def test_the_help_states_what_set_text_does_to_formatting():
    rows = {row["op"]: row for row in batch.batchable_ops()}
    note = rows["set_text"]["note"]
    for words in SET_TEXT_CLAIMS:
        assert words in note, (words, note)
    text_help = help_payload(_plan(), "set_shape")["args_schema"]["properties"]["text"]
    assert "first run" in text_help["description"]
    for name in ("set_paragraph", "insert_paragraph", "delete_paragraph"):
        row = rows[name]
        assert "counted from 0" in row["address"] and "paragraph_number" in row["address"]
        assert "before the batch" in row["address"]


def test_set_text_does_what_the_help_says(tmp_path):
    """Paragraph N of the new text takes old paragraph N's look; extra new
    paragraphs take the last old one's; old paragraphs past the new count are
    dropped with their look."""
    deck, sid = _deck(tmp_path)
    _batch(deck, [{"op": "set_text", "slide_number": 1, "shape": sid,
                   "text": "One\nTwo\nThree\nFour"}])
    paras = _paras(deck, sid)
    assert [p["text"] for p in paras] == ["One", "Two", "Three", "Four"]
    assert (paras[0]["b"], paras[0]["sz"]) == ("1", "4000")
    assert (paras[1]["i"], paras[1]["sz"]) == ("1", "1800")
    assert paras[2]["sz"] == "1600" and paras[3]["sz"] == "1600"

    deck2, sid2 = _deck(tmp_path, "paras2.pptx")
    _batch(deck2, [{"op": "set_text", "slide_number": 1, "shape": sid2,
                    "text": "Only the aside is left"}])
    paras = _paras(deck2, sid2)
    assert len(paras) == 1
    # It takes the FIRST paragraph's look, not the italic aside's.
    assert (paras[0]["b"], paras[0]["sz"], paras[0]["i"]) == ("1", "4000", None)


def test_a_paragraph_with_mixed_runs_collapses_to_its_first_run(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    p = box.text_frame.paragraphs[0]
    a = p.add_run()
    a.text, a.font.bold = "Bold start ", True
    b = p.add_run()
    b.text, b.font.italic = "italic end", True
    path = tmp_path / "mixed.pptx"
    prs.save(str(path))
    out = _batch(path, [{"op": "set_text", "slide_number": 1, "shape": box.shape_id,
                         "text": "Replaced"}])
    assert out["changed"]["0"]["runs_collapsed_to_first"] is True
    paras = _paras(path, box.shape_id)
    assert (paras[0]["b"], paras[0]["i"]) == ("1", None)


def test_help_lists_the_paragraph_edits_and_the_new_option():
    payload = help_payload(_plan(), "set_paragraph")
    assert payload["batch_only"] is True and payload["use"] == "apply_edits"
    assert "paragraph_number" in payload["fields"]
    options = help_payload(_plan(), "search_and_replace")["args_schema"]["properties"]
    assert "remove_empty_paragraphs" in options
    assert "empty" in options["remove_empty_paragraphs"]["description"]
