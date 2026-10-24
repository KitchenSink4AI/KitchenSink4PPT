"""merge_cells moves a covered cell's equation or newer run, and counts it.

A covered cell's text moves into the merged cell, as PowerPoint does. A cell
whose paragraph held its words in a compatibility wrapper (an inline
equation, a run for a newer version) read as empty, so merge_cells emptied
it without moving anything and answered text_moved_from_cells: 0. The
wrapper now moves with its paragraph and the cell is counted; a wrapper is
never dropped, even one that shows no text.

Every deck here is made up by code.
"""

from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.package import PptxPackage, qn
from kitchensink4ppt.ops import tables

MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
A14 = "http://schemas.microsoft.com/office/drawing/2010/main"
OMML = "http://schemas.openxmlformats.org/officeDocument/2006/math"
FUTURE = "http://example.com/made-up/future/2033"


def _future(fallback):
    return (
        f'<mc:AlternateContent xmlns:mc="{MC}" xmlns:x33="{FUTURE}"><mc:Choice Requires="x33">'
        f'<a:r xmlns:a="{A}"><a:t>choice words</a:t></a:r></mc:Choice><mc:Fallback>'
        f'<a:r xmlns:a="{A}"><a:t>{fallback}</a:t></a:r></mc:Fallback></mc:AlternateContent>'
    )


def _math(expr):
    return (
        f'<mc:AlternateContent xmlns:mc="{MC}"><mc:Choice xmlns:a14="{A14}" Requires="a14">'
        f'<a14:m><m:oMathPara xmlns:m="{OMML}"><m:oMath>{expr}</m:oMath></m:oMathPara></a14:m>'
        f'</mc:Choice><mc:Fallback><a:r xmlns:a="{A}"><a:t>[made-up equation]</a:t></a:r>'
        "</mc:Fallback></mc:AlternateContent>"
    )


def _deck(tmp_path, covered_xml):
    """A 1x2 table: (0, 0) says "plain"; (0, 1) holds only `covered_xml`."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    frame = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(6), Inches(0.8))
    frame.table.cell(0, 0).text = "plain"
    path = tmp_path / "merge.pptx"
    prs.save(str(path))
    pkg = PptxPackage(path)
    part = pkg.slide_parts()[0]
    p = list(pkg.root(part).iter(qn("a:tc")))[1].find(f".//{qn('a:p')}")
    for r in p.findall(qn("a:r")):
        p.remove(r)
    p.insert(0, etree.fromstring(covered_xml))
    pkg.mark_dirty(part)
    pkg.save(do_backup=False)
    return path


def _merge(path):
    pkg = PptxPackage(path)
    out = tables.merge_cells(pkg, 0, {"index": 0}, 0, 0, 0, 1)
    pkg.save(do_backup=False)
    return out


def _origin_xml(path):
    pkg = PptxPackage(path)
    return etree.tostring(list(pkg.root(pkg.slide_parts()[0]).iter(qn("a:tc")))[0])


def test_a_newer_run_moves_into_the_merged_cell_and_is_counted(tmp_path):
    deck = _deck(tmp_path, _future("fallback words in a cell"))
    out = _merge(deck)
    assert out["text_moved_from_cells"] == 1
    grid = tables.get_table(PptxPackage(deck), 0, {"index": 0})
    assert grid["cells"][0]["text"] == "plain\nfallback words in a cell"
    assert b"AlternateContent" in _origin_xml(deck)


def test_an_equation_moves_into_the_merged_cell_and_is_counted(tmp_path):
    deck = _deck(tmp_path, _math("<m:r><m:t>z=1</m:t></m:r>"))
    out = _merge(deck)
    assert out["text_moved_from_cells"] == 1
    assert b"<m:t>z=1</m:t>" in _origin_xml(deck)


def test_a_wrapper_that_shows_no_text_still_moves(tmp_path):
    deck = _deck(tmp_path, _math(""))
    out = _merge(deck)
    assert out["text_moved_from_cells"] == 1
    assert _origin_xml(deck).count(b"AlternateContent") == 2
