"""A backup, a restore source or a file given to validate is checked for a
decompression bomb before any part of it is inflated.

Opening a deck checks the sizes its package declares before it inflates
anything (core/zipsafe.py). The payload check that restore_backup runs on a
backup, that known_problem_keys runs on the deck being replaced, and that
validate runs on its file read every XML part in full without that check, so
a part built to exhaust memory, placed in the hidden backup folder or handed
to validate, was inflated whole. The payload check now runs the same
directory check first and refuses such a file before reading any part.

The bomb here is made up: a deck with one XML part of 300 MiB of zeros,
about 300 KB in the file. No test inflates it: reading it is refused in the
test itself, and the tests check that nothing tried to.
"""

from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from kitchensink4ppt import server
from kitchensink4ppt.core.errors import DocumentCorrupt
from kitchensink4ppt.core.package import known_problem_keys
from kitchensink4ppt.ops import backups

MiB = 1024 * 1024
BOMB_PART = "ppt/made_up_zeros.xml"


class Inflated(Exception):
    """Raised instead of inflating the bomb part."""


@pytest.fixture()
def opened(monkeypatch):
    """Record every part a ZIP reader opens; opening the bomb part raises
    Inflated instead of inflating it."""
    names: list[str] = []
    original = zipfile.ZipFile.open

    def watched(self, name, *args, **kwargs):
        part = name.filename if isinstance(name, zipfile.ZipInfo) else name
        mode = args[0] if args else kwargs.get("mode", "r")
        if mode == "r":
            names.append(part)
            if part == BOMB_PART:
                raise Inflated(part)
        return original(self, name, *args, **kwargs)

    monkeypatch.setattr(zipfile.ZipFile, "open", watched)
    return names


def _deck(path: Path, words: str = "Made-up deck") -> Path:
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = words
    prs.save(str(path))
    return path


def _bomb(path: Path, words: str = "Made-up bomb") -> Path:
    """A deck plus one XML part of 300 MiB of zeros, deflated."""
    _deck(path, words)
    with zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED, allowZip64=True) as zf:
        info = zipfile.ZipInfo(BOMB_PART, (2026, 1, 1, 0, 0, 0))
        info.compress_type = zipfile.ZIP_DEFLATED
        chunk = bytes(MiB)
        with zf.open(info, "w", force_zip64=True) as fh:
            for _ in range(300):
                fh.write(chunk)
    assert path.stat().st_size < 2 * MiB
    return path


def test_a_bomb_in_a_backup_slot_is_refused_before_it_is_inflated(tmp_path, opened):
    deck = _deck(tmp_path / "deck.pptx")
    server.set_slide_hidden(str(deck), 0, True)  # a change, so the prev slot exists
    prev = tmp_path / ".ks4p-backups" / "deck.pptx" / "prev.pptx"
    assert prev.exists()
    _bomb(tmp_path / "made_up_bomb.pptx")
    shutil.copy2(tmp_path / "made_up_bomb.pptx", prev)
    before = deck.read_bytes()
    opened.clear()
    with pytest.raises(DocumentCorrupt) as refused:
        backups.restore_backup(str(deck), "prev")
    assert BOMB_PART not in opened
    message = str(refused.value)
    assert "exhaust memory" in message and "Nothing was changed" in message
    assert deck.read_bytes() == before


def test_a_bomb_given_as_the_restore_source_is_refused_before_it_is_inflated(tmp_path, opened):
    deck = _deck(tmp_path / "deck.pptx")
    source = _bomb(tmp_path / "made_up_source.pptx")
    before = deck.read_bytes()
    opened.clear()
    with pytest.raises(DocumentCorrupt, match="exhaust memory"):
        backups.restore_backup(str(deck), str(source))
    assert BOMB_PART not in opened
    assert deck.read_bytes() == before


def test_known_problem_keys_never_inflates_a_bomb_on_disk(tmp_path, opened):
    deck = _bomb(tmp_path / "made_up_bomb.pptx")
    opened.clear()
    assert known_problem_keys(deck) == frozenset()  # unreadable: nothing tolerated
    assert BOMB_PART not in opened


def test_validate_refuses_a_bomb_before_it_is_inflated(tmp_path, opened, monkeypatch):
    from kitchensink4ppt.com import bridge

    def no_powerpoint(*_a, **_k):
        raise RuntimeError("made-up: the opens-clean check is not run in this test")

    monkeypatch.setattr(bridge, "com_validate_opens_clean", no_powerpoint)
    deck = _bomb(tmp_path / "made_up_bomb.pptx")
    opened.clear()
    out = server.validate(str(deck))
    assert BOMB_PART not in opened
    assert out["payload_valid"] is False
    assert "exhaust memory" in str(out.get("payload_error"))
