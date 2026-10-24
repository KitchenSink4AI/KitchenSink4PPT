"""Two shapes that share an id on one slide.

A damaged deck can give two shapes on a slide the same id; PowerPoint opens
it and gives the second shape a new id. Here an id then named two shapes: a
shape edit by id silently changed the first one, an animation by id was
refused with "Address the shape by id instead", and the view printed both
with the same anchor, refused with "use a longer prefix". Now every such
address is refused with the shapes it could mean and a next step that works:
the shape's name for an edit, PowerPoint's own repair for an animation,
which names its shape by id.

Every deck here is made up by code.
"""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.errors import AmbiguousTarget
from kitchensink4ppt.core.package import PptxPackage, qn
from kitchensink4ppt.ops import animations, shapes, view


@pytest.fixture()
def deck(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    first = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    first.name = "First box"
    first.text_frame.text = "Made-up first"
    second = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1))
    second.name = "Second box"
    second.text_frame.text = "Made-up second"
    path = tmp_path / "shared_id.pptx"
    prs.save(str(path))
    pkg = PptxPackage(path)
    part = pkg.slide_parts()[0]
    ids = [c for c in pkg.root(part).iter(qn("p:cNvPr")) if c.get("name") == "Second box"]
    shared = [c for c in pkg.root(part).iter(qn("p:cNvPr")) if c.get("name") == "First box"][0].get("id")
    ids[0].set("id", shared)
    pkg.mark_dirty(part)
    pkg.save(do_backup=False)
    return path, int(shared)


def _offsets(path):
    pkg = PptxPackage(path)
    out = {}
    for sp in pkg.root(pkg.slide_parts()[0]).iter(qn("p:sp")):
        name = sp.find(f".//{qn('p:cNvPr')}").get("name")
        out[name] = int(sp.find(f".//{qn('a:off')}").get("x"))
    return out


def test_an_edit_by_a_shared_id_is_refused_with_the_names(deck):
    path, sid = deck
    before = path.read_bytes()
    pkg = PptxPackage(path)
    with pytest.raises(AmbiguousTarget) as refused:
        shapes.set_shape(pkg, 0, sid, x=2.0)
    message = str(refused.value)
    assert "'First box'" in message and "'Second box'" in message
    assert "name" in message and "PowerPoint" in message
    assert path.read_bytes() == before


def test_the_name_reaches_the_second_shape(deck):
    path, _sid = deck
    pkg = PptxPackage(path)
    shapes.set_shape(pkg, 0, "Second box", x=2.0)
    pkg.save(do_backup=False)
    assert _offsets(path) == {"First box": 914400, "Second box": 1828800}


def test_an_animation_on_a_shared_id_names_powerpoints_repair(deck):
    path, sid = deck
    pkg = PptxPackage(path)
    with pytest.raises(AmbiguousTarget) as refused:
        animations.add_entrance_animation(pkg, 0, sid, "fade")
    message = str(refused.value)
    assert "Address the shape by id instead" not in message
    assert "'Second box'" in message and "PowerPoint" in message


def test_the_shared_anchor_is_refused_without_the_longer_prefix_advice(deck):
    path, sid = deck
    pkg = PptxPackage(path)
    slide_id = int(pkg.presentation().find(f"{qn('p:sldIdLst')}/{qn('p:sldId')}").get("id"))
    anchor = "a:" + view.anchor_digest(slide_id, sid)
    with pytest.raises(AmbiguousTarget) as refused:
        view.resolve_anchor(pkg, anchor)
    assert "longer" not in str(refused.value)
    assert "'Second box'" in str(refused.value)
    with pytest.raises(AmbiguousTarget) as refused_edit:
        shapes.set_shape(pkg, 0, anchor, x=2.0)
    assert "whole anchor" not in str(refused_edit.value)
    assert "'Second box'" in str(refused_edit.value)
