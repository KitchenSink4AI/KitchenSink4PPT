"""Live search and replace lands every edit on the matched text, whatever
characters come before it.

PowerPoint's TextRange counts characters in UTF-16 code units: an emoji, a
CJK Extension B character or any other character outside the Basic
Multilingual Plane is TWO units, while Python counts it as one code point.
The live route found matches in code points and wrote through
TextRange.Characters(start, length) with those numbers, so one emoji before
a match moved the write one unit to the left: the edit landed on the wrong
characters and the read-back check then blamed protection. A case-insensitive
search also matched on text.lower(), which is longer than the text when it
holds a dotted capital I (U+0130), so offsets drifted the other way.

Two layers of proof:
- in-process, a stand-in TextRange that counts in UTF-16 units exactly as
  PowerPoint does drives the whole live_search_and_replace body;
- a real PowerPoint round (marked live; own instance, temp copy of a
  synthetic deck, closed without saving) confirms the stand-in's model and
  the fix on the real application.

All text here is made up.
"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest

from kitchensink4ppt.com import live, live_ops

REPO = Path(__file__).resolve().parents[2]
EMOJI = "\U0001F600"  # one code point, two UTF-16 units
CJK_B = "\U00020000"  # CJK Extension B: one code point, two UTF-16 units
FAMILY = "\U0001F468\u200d\U0001F469\u200d\U0001F467"  # ZWJ sequence: 5 cp, 8 units


# ------------------------------------------- a TextRange that counts in UTF-16


class _Units:
    """The text of one frame as UTF-16 code units, like PowerPoint keeps it."""

    def __init__(self, text: str):
        self.units = list(text.encode("utf-16-le"))

    def get(self, start: int, length: int) -> str:
        return bytes(self.units[2 * start: 2 * (start + length)]).decode(
            "utf-16-le", errors="surrogatepass")

    def put(self, start: int, length: int, text: str) -> None:
        self.units[2 * start: 2 * (start + length)] = list(text.encode("utf-16-le"))

    def length(self) -> int:
        return len(self.units) // 2


class _Range:
    def __init__(self, store: _Units, start: int = 0, length: int | None = None):
        self._s, self._start, self._len = store, start, length

    def _length(self):
        return self._s.length() - self._start if self._len is None else self._len

    @property
    def Text(self):  # noqa: N802 - COM spelling
        return self._s.get(self._start, self._length())

    @Text.setter
    def Text(self, value):  # noqa: N802
        self._s.put(self._start, self._length(), value)

    def Characters(self, start, length):  # noqa: N802
        # PowerPoint: 1-based start, both in UTF-16 units, clipped to the text.
        begin = min(self._start + start - 1, self._s.length())
        return _Range(self._s, begin, max(0, min(length, self._s.length() - begin)))


class _Frame:
    def __init__(self, text):
        self.store = _Units(text)

    @property
    def HasText(self):  # noqa: N802
        return self.store.length() > 0

    @property
    def TextRange(self):  # noqa: N802
        return _Range(self.store)


class _Shape:
    HasTable = False
    HasTextFrame = True
    Type = 1

    def __init__(self, sid, text):
        self.Id = sid
        self.TextFrame = _Frame(text)


class _Coll:
    def __init__(self, items):
        self._items = items
        self.Count = len(items)

    def Item(self, i):  # noqa: N802
        return self._items[i - 1]


class _Slide:
    SlideIndex = 1
    SlideID = 256

    def __init__(self, shapes):
        self.Shapes = _Coll(shapes)


class _Pres:
    def __init__(self, slides):
        self.Slides = _Coll(slides)


class _Session:
    def __init__(self, pres):
        self.pres = pres

    def result(self, out):
        return out


@pytest.fixture()
def fake_live(monkeypatch):
    """live_search_and_replace runs its body against stand-in shapes."""
    holder = {}

    def run_live(path, tool, body, *, mutating=True):
        return body(_Session(holder["pres"]))

    monkeypatch.setattr(live, "run_live", run_live)
    monkeypatch.setattr(live_ops.safesave, "claim_publish", lambda *a, **k: None)
    monkeypatch.setattr(live_ops, "check_path", lambda p, what: p)

    def make(*texts):
        shapes = [_Shape(10 + n, t) for n, t in enumerate(texts)]
        holder["pres"] = _Pres([_Slide(shapes)])
        return shapes

    return make


def _text(shape):
    return shape.TextFrame.TextRange.Text


def test_the_stand_in_counts_like_powerpoint():
    """The model the stand-in encodes (confirmed on PowerPoint by the live
    round below): an emoji is two units long."""
    rng = _Frame(f"{EMOJI}ab").TextRange
    assert rng.Characters(3, 1).Text == "a"


@pytest.mark.parametrize(
    "before",
    [EMOJI, CJK_B, FAMILY, f"{EMOJI}{EMOJI} x {CJK_B}", "plain "],
    ids=["emoji", "cjk-ext-b", "zwj-family", "several", "bmp-only"],
)
def test_replace_after_astral_characters_lands_on_the_match(fake_live, before):
    (shp,) = fake_live(f"{before} alpha beta gamma beta")
    out = live_ops.live_search_and_replace("deck.pptx", "beta", "delta")
    assert out["total"] == 2
    assert _text(shp) == f"{before} alpha delta gamma delta"


def test_find_text_holding_an_astral_character(fake_live):
    (shp,) = fake_live(f"x {EMOJI}y {EMOJI}y z")
    out = live_ops.live_search_and_replace("deck.pptx", f"{EMOJI}y", "Q")
    assert out["total"] == 2
    assert _text(shp) == "x Q Q z"


def test_case_insensitive_after_a_dotted_capital_i(fake_live):
    """'\u0130'.lower() is two code points; matching on the lowered copy
    shifted every later offset by one."""
    (shp,) = fake_live("\u0130stanbul Beta and BETA")
    out = live_ops.live_search_and_replace(
        "deck.pptx", "beta", "Delta", match_case=False)
    assert out["total"] == 2
    assert _text(shp) == "\u0130stanbul Delta and Delta"


def test_case_insensitive_matches_what_the_file_route_matches(fake_live):
    """Parity with the closed-file search_and_replace, which matches a
    literal case-insensitively with re.IGNORECASE on the text itself."""
    from kitchensink4ppt.ops.text import _compile_matcher

    text = f"{EMOJI} STRASSE stra\u00dfe \u0130 i \u03a3\u03b9\u03c3 \u03c3"
    for find in ("strasse", "i", "\u03c3"):
        expected = [(s, e) for s, e, _r in
                    _compile_matcher(find, "_", False, False)(text)]
        got = [(cp_s, cp_e) for cp_s, cp_e, _u, _n in
               live_ops._live_match_spans(text, find, match_case=False)]
        assert got == expected, find


def test_spans_report_utf16_positions():
    spans = live_ops._live_match_spans(f"{EMOJI}ab{FAMILY}ab", "ab", match_case=True)
    # code points: emoji=1, family=5; UTF-16: emoji=2, family=8
    assert spans == [(1, 3, 2, 2), (8, 10, 12, 2)]


# ------------------------------------------------- the real PowerPoint round

_SCENARIO = r"""
import contextlib, gc, json, sys, time
from pathlib import Path

import pythoncom
import win32com.client

from kitchensink4ppt.com import bridge, live, live_ops

deck = str(Path(sys.argv[1]).resolve())
EMOJI = "\U0001F600"
CJK_B = "\U00020000"
out = {}


def pids():
    import subprocess
    r = subprocess.run(["tasklist", "/FI", "IMAGENAME eq POWERPNT.EXE", "/FO", "CSV",
                        "/NH"], capture_output=True, text=True, timeout=30)
    found = set()
    for ln in r.stdout.splitlines():
        if "POWERPNT" in ln.upper():
            with contextlib.suppress(Exception):
                found.add(int(ln.split('","')[1].strip('"')))
    return found


pre = pids()
out["pre_powerpnt"] = len(pre)
if pre:
    out["skipped"] = "a PowerPoint was already running; refusing to share it"
    print("RESULT " + json.dumps(out))
    sys.exit(0)

pythoncom.CoInitialize()
app = win32com.client.DispatchEx("PowerPoint.Application")
ours = pids() - pre
app.DisplayAlerts = 1
pres = app.Presentations.Open(deck, ReadOnly=False, Untitled=False, WithWindow=True)
with contextlib.suppress(Exception):
    pres.Windows(1).WindowState = 2


def frame_texts():
    sl = pres.Slides.Item(1)
    texts = {}
    for i in range(1, int(sl.Shapes.Count) + 1):
        shp = sl.Shapes.Item(i)
        if shp.HasTable:
            tbl = shp.Table
            texts["cell"] = str(tbl.Cell(1, 1).Shape.TextFrame.TextRange.Text)
        elif shp.HasTextFrame:
            texts[str(shp.Name)] = str(shp.TextFrame.TextRange.Text)
    phs = sl.NotesPage.Shapes.Placeholders
    for i in range(1, int(phs.Count) + 1):
        ph = phs.Item(i)
        if int(ph.PlaceholderFormat.Type) == 2 and ph.HasTextFrame:
            texts["notes"] = str(ph.TextFrame.TextRange.Text)
    return texts


try:
    # The fact the fix rests on, measured on PowerPoint itself.
    tr = pres.Slides.Item(1).Shapes.Item(1).TextFrame.TextRange
    out["model_length_units"] = int(tr.Length)
    out["model_text_len_cp"] = len(str(tr.Text))
    out["model_char_3_1"] = str(tr.Characters(3, 1).Text)
    tr = None
    out["before"] = frame_texts()
    try:
        rep = live_ops.live_search_and_replace(deck, "beta", "delta", include_notes=True)
        out["case_total"] = rep["total"]
    except Exception as exc:
        out["case_error"] = f"{type(exc).__name__}: {exc}"
    out["after_case"] = frame_texts()
    try:
        rep = live_ops.live_search_and_replace(deck, "GAMMA", "omega", match_case=False,
                                               include_notes=True)
        out["nocase_total"] = rep["total"]
    except Exception as exc:
        out["nocase_error"] = f"{type(exc).__name__}: {exc}"
    out["after_nocase"] = frame_texts()
except Exception:
    import traceback
    out["scenario_error"] = traceback.format_exc()
finally:
    with contextlib.suppress(Exception):
        pres.Saved = True
    with contextlib.suppress(Exception):
        pres.Close()
    pres = None
    with contextlib.suppress(Exception):
        app.Quit()
    app = None
    gc.collect()
    deadline = time.monotonic() + 45.0
    while time.monotonic() < deadline and (pids() & ours):
        time.sleep(1.0)
out["our_zombies"] = sorted(pids() & ours)
print("RESULT " + json.dumps(out))
"""


def _deck(path: Path) -> Path:
    """One slide: two text boxes and a 1x1 table whose text starts with
    characters outside the Basic Multilingual Plane, and speaker notes."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    a = s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    a.name = "first"
    a.text_frame.text = f"{EMOJI}ab alpha beta gamma beta"
    b = s.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    b.name = "second"
    b.text_frame.text = f"{CJK_B}{EMOJI} \u0130 Gamma beta"
    t = s.shapes.add_table(1, 1, Inches(1), Inches(3.5), Inches(4), Inches(1)).table
    t.cell(0, 0).text = f"{EMOJI} cell beta"
    s.notes_slide.notes_text_frame.text = f"{EMOJI}{EMOJI} notes beta GAMMA"
    prs.save(str(path))
    return path


def _gate():
    if sys.platform != "win32":
        pytest.skip("PowerPoint is Windows-only")
    try:
        import win32com.client  # noqa: F401
        from kitchensink4ppt.com import bridge
    except ImportError:
        pytest.skip("pywin32 not installed")
    if not bridge.powerpoint_installed():
        pytest.skip("PowerPoint is not installed on this machine")
    import com_test_gate

    reason = com_test_gate.powerpoint_blocks_com_tests(bridge.powerpnt_count)
    if reason:
        pytest.skip(reason)


@pytest.mark.live
@pytest.mark.timeout(600)
def test_live_replace_on_powerpoint_after_astral_characters(tmp_path):
    _gate()
    deck = _deck(tmp_path / "astral_live.pptx")
    script = tmp_path / "astral_scenario.py"
    script.write_text(_SCENARIO, encoding="utf-8")
    proc = subprocess.run(
        [sys.executable, "-X", "utf8", str(script), str(deck)],
        capture_output=True, text=True, encoding="utf-8", timeout=480,
        cwd=str(REPO), stdin=subprocess.DEVNULL,
    )
    line = next((ln for ln in reversed((proc.stdout or "").splitlines())
                 if ln.startswith("RESULT ")), None)
    assert proc.returncode == 0 and line, (proc.stdout, proc.stderr)
    out = json.loads(line[len("RESULT "):])
    if "skipped" in out:
        pytest.skip(out["skipped"])
    assert "scenario_error" not in out, out["scenario_error"]
    # PowerPoint counts UTF-16 units: the emoji is two of them.
    assert out["model_length_units"] == out["model_text_len_cp"] + 1
    assert out["model_char_3_1"] == "a"
    assert "case_error" not in out, out["case_error"]
    assert "nocase_error" not in out, out["nocase_error"]
    after = out["after_nocase"]
    assert after["first"] == f"{EMOJI}ab alpha delta omega delta"
    assert after["second"] == f"{CJK_B}{EMOJI} \u0130 omega delta"
    assert after["cell"] == f"{EMOJI} cell delta"
    assert after["notes"] == f"{EMOJI}{EMOJI} notes delta omega"
    assert out["case_total"] == 5
    assert out["nocase_total"] == 3
    assert out["our_zombies"] == []
