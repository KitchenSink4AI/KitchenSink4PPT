"""Three refusals that were worded wrongly.

- An argument with no live route, refused on a deck open in PowerPoint, got
  the generic hint that the deck "holds something this operation does not
  handle, or is protected or damaged". The deck is fine; the argument needs a
  closed deck.
- A slide sent as slide_number 9 on a 5-slide deck was refused as "slide index
  8 out of range", a number the caller never sent.
- PowerPoint opens nothing under a folder whose name Windows keeps for a
  device (con, nul, com1 ...), and export and validate then called a healthy
  deck damaged.
The decks are made up.
"""

from __future__ import annotations

import asyncio
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt import server
from kitchensink4ppt.com import bridge
from kitchensink4ppt.registry import public_ops
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface import envelope
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.plan import build


def _five(tmp_path):
    path = tmp_path / "five.pptx"
    prs = Presentation()
    for n in range(5):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = f"Slide {n + 1}"
    prs.save(path)
    return path


def _call(tool, op, **args):
    plan = build(Settings(), public_ops())
    return asyncio.run(dispatch(plan, tool, {"op": op, "args": args}))


def test_a_file_only_argument_on_the_live_route_has_its_own_hint():
    try:
        server._live_refuse(notes=[{"slide_number": 1, "text": "x"}])
    except Exception as exc:  # noqa: BLE001 - the refusal is what is checked
        error = envelope.refusal(exc)["error"]
    else:
        raise AssertionError("no refusal")
    assert "damaged" not in error["hint"] and "protected" not in error["hint"], error
    assert "closed deck" in error["hint"], error


def test_a_batch_slide_number_out_of_range_is_named_as_sent(tmp_path):
    path = _five(tmp_path)
    payload, is_error = _call("ppt_edit", "apply_edits", file_path=str(path),
                              edits=[{"op": "set_notes", "slide_number": 9, "text": "x"}])
    assert is_error is True
    text = str(payload)
    assert "slide_number 9" in text and "slide index 8" not in text, payload
    assert "1 to 5" in text, payload


def test_a_call_slide_number_out_of_range_is_named_as_sent(tmp_path):
    path = _five(tmp_path)
    payload, is_error = _call("ppt_edit", "insert_textbox", file_path=str(path), slide_number=9,
                              text="x", x=1, y=1, w=2, h=1, live="off")
    assert is_error is True
    message = payload["error"]["message"]
    assert "slide_number 9" in message and "slide index 8" not in message, payload
    # A 0-based slide keeps the 0-based wording.
    payload, _ = _call("ppt_edit", "insert_textbox", file_path=str(path), slide=8,
                       text="x", x=1, y=1, w=2, h=1, live="off")
    assert "slide index 8" in payload["error"]["message"], payload


@pytest.mark.skipif(sys.platform != "win32", reason="Windows device names")
def test_a_reserved_device_folder_is_named_instead_of_damage():
    exc = bridge._open_failure(Path(r"C:\work\com8\plan.pptx"), RuntimeError("E_FAIL"))
    message = str(exc)
    assert "com8" in message and "damaged" not in message, message
    assert "another name" in exc.hint, exc.hint
    plain = str(bridge._open_failure(Path(r"C:\work\decks\plan.pptx"), RuntimeError("E_FAIL")))
    assert "could not open plan.pptx" in plain
