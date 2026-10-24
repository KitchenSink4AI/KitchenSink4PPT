"""Every read answer that lists slides counts them both ways.

A slide record in an answer carries its 0-based position (index or
slide_index, as `slide` takes it) and, beside it, slide_number: the same slide
counted from 1, the way people count, as export_slide_image and
manage_custom_show already do. An assistant can then tell the user "slide 3"
without doing the arithmetic.
"""

from __future__ import annotations

import asyncio

import pytest
from pptx import Presentation

from kitchensink4ppt import server  # noqa: F401  (registers the operations)
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.plan import build

SLIDE_MARKS = ("slide_id", "title", "layout", "hidden", "notes", "text", "transition",
               "comments", "shapes", "words_body", "check")


@pytest.fixture
def deck(tmp_path, monkeypatch):
    monkeypatch.setenv("KS4P_SAVE_JOURNAL_DIR", str(tmp_path / "journal"))
    path = tmp_path / "Deck.pptx"
    prs = Presentation()
    for title in ("Alpha", "Bravo", "Charlie"):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.notes_slide.notes_text_frame.text = f"notes for {title}"
    prs.save(path)
    return path


def _call(op: str, args: dict) -> dict:
    plan = build(Settings())
    payload, is_error = asyncio.run(
        dispatch(plan, plan.specs[op].group, {"op": op, "args": args}))
    assert is_error is False, payload
    return payload


def _slide_records(answer, found=None) -> list[dict]:
    """Every dict in an answer that names a slide by its 0-based position."""
    found = [] if found is None else found
    if isinstance(answer, dict):
        for key in ("index", "slide_index"):
            value = answer.get(key)
            if (isinstance(value, int) and not isinstance(value, bool)
                    and (key == "slide_index" or any(m in answer for m in SLIDE_MARKS))):
                found.append(answer)
                break
        for value in answer.values():
            _slide_records(value, found)
    elif isinstance(answer, list):
        for value in answer:
            _slide_records(value, found)
    return found


CASES = [
    ("get_text", {}),
    ("extract_text", {}),
    ("list_elements", {"kind": "slides"}),
    ("list_elements", {"kind": "shapes"}),
    ("list_comments", {}),
    ("get_transitions", {}),
    ("get_autofit_state", {}),
    ("deck_statistics", {}),
    ("check_layout", {}),
]


@pytest.mark.parametrize("op, extra", CASES, ids=[f"{op}-{'-'.join(map(str, e.values())) or 'all'}"
                                                  for op, e in CASES])
def test_each_slide_in_a_read_answer_carries_its_slide_number(deck, op, extra):
    args = {"file_path": str(deck), **extra}
    plan = build(Settings())
    if "live" in plan.specs[op].args.model_fields:
        args["live"] = "off"
    answer = _call(op, args)
    records = _slide_records(answer)
    assert records, f"{op} listed no slide records to check: {answer}"
    for record in records:
        position = record.get("index", record.get("slide_index"))
        assert record.get("slide_number") == position + 1, (op, record)
