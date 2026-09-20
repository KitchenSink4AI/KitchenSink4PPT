"""Regressions for the 2026-09-20 defense-deck field report.

One test (or more) per punchlist issue in
`Developer Feedback/General Feedback/Beta Test Report - 2026-09-20 -
Final Defense Deck Build.md`:

#869  extract_text paged a 43-slide deck but put the `page` block last,
      behind the whole joined "text", where a client-side truncation ate
      it; and the tool took no offset, so the block's own continuation
      hint named a call that could not be made.
#870  check_layout's tiny_text check skipped every run with no explicit
      sz, which is every ordinary body run on a templated deck.
#871  no contrast audit against the backdrop beneath a shape, and the
      overlap audit saw neither in-group siblings nor placeholders whose
      geometry is inherited.
#872  create_presentation with no template built a 4:3 deck while its
      docstring promised 16:9, and never stated the size it built.
#873  line={"type":"none"} left a themed outline painted, and
      insert_connector dropped end_arrow/start_arrow in silence.
#874  set_shape refused to seed geometry on a layout-inheriting
      placeholder, and vertical anchor had no route that did not replace
      the shape's whole text.
#875  a text replacement dropped explicit run colour, reset paragraph
      alignment and bullets, and wrote only <a:latin> for a font.
#876  delete_table_cols documented a widths value it rejects, and
      get_workflows had no recipe for inheriting a deck and fixing it.
"""

from __future__ import annotations

import json

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core import budget
from kitchensink4ppt.core.package import PptxPackage
from kitchensink4ppt.ops import read as rd


# --------------------------------------------------------------- #869


def _wordy_deck(path, slides: int = 12):
    """A deck with enough text per slide to overflow a small budget."""
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for i in range(slides):
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.5), Inches(8), Inches(5)
        )
        frame = box.text_frame
        frame.text = f"Slide {i} heading"
        for j in range(6):
            para = frame.add_paragraph()
            para.text = f"Slide {i} paragraph {j}: " + ("lorem ipsum " * 12)
    prs.save(str(path))
    return path


def test_paged_text_leads_with_the_page_block_not_trails_it(
    tmp_path, monkeypatch
):
    """#869: the honesty signal has to survive a truncated payload.

    All the weight of this answer is in "text". A `page` block written
    after it is the first thing any downstream cut removes, which is how
    a 43-slide deck came back as 23 slides that looked complete.
    """
    monkeypatch.setenv(budget.ENV_MAX_CHARS, "6000")
    pkg = PptxPackage(str(_wordy_deck(tmp_path / "wordy.pptx")))
    result = rd.get_text(pkg, include_notes=True)

    keys = list(result)
    assert keys[0] == "page", f"page block is not first: {keys}"
    assert "returned" in keys
    assert keys.index("returned") < keys.index("text")
    assert result["returned"] == len(result["slides"])
    assert result["returned"] < result["slide_count"]
    assert result["page"]["truncated"] is True

    # And it survives a cut that keeps only the first fifth of the wire form.
    wire = json.dumps(result, ensure_ascii=False, default=str)
    assert "next_offset" in wire[: len(wire) // 5]


def test_extract_text_can_act_on_its_own_continuation_hint(
    tmp_path, monkeypatch
):
    """#869: the page block said "call again with offset=N" at a tool
    that had no offset parameter."""
    from kitchensink4ppt import server

    monkeypatch.setenv(budget.ENV_MAX_CHARS, "6000")
    deck = str(_wordy_deck(tmp_path / "wordy2.pptx"))

    seen: list[int] = []
    offset = 0
    for _ in range(50):
        result = server.extract_text(deck, offset=offset)
        seen.extend(s["index"] for s in result["slides"])
        page = result.get("page")
        if page is None or page["next_offset"] is None:
            break
        offset = page["next_offset"]
    assert seen == list(range(12)), seen


def test_a_deck_that_fits_keeps_its_historical_get_text_shape(tmp_path):
    """#869 guard: the reorder is invisible below the ceiling."""
    pkg = PptxPackage(str(_wordy_deck(tmp_path / "small.pptx", slides=2)))
    result = rd.get_text(pkg)
    assert list(result) == ["slide_count", "slides", "text"]
