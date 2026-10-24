"""list_hyperlinks reaches every place a deck keeps a link or a click action.

It used to read only slide shapes' cNvPr and the a:r runs of a p:txBody, so a link in
a table cell, on a field (a:fld), in the speaker notes, in a SmartArt node or
on a layout or master was not listed, and a click action with no relationship
(run a custom show, run a macro, an OLE verb) was dropped without a word. The
deck here is made up and built by code.
"""

from __future__ import annotations

import asyncio

from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt import server  # noqa: F401  (the import registers ops)
from kitchensink4ppt.core.package import PptxPackage, qn
from kitchensink4ppt.ops import links, read
from kitchensink4ppt.registry import public_ops
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.help import help_payload
from kitchensink4ppt.surface.plan import build

from test_diagrams_write import add_smartart

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_HL = _R + "/hyperlink"


def _hlink(rid: str, action: str | None = None) -> etree._Element:
    el = etree.Element(qn("a:hlinkClick"), nsmap={"a": _A, "r": _R})
    el.set(qn("r:id"), rid)
    if action:
        el.set("action", action)
    return el


def _link_first_run(container: etree._Element, rid: str, action: str | None = None,
                    tag: str = "a:r") -> None:
    run = next(container.iter(qn(tag)))
    rpr = run.find(qn("a:rPr"))
    if rpr is None:
        rpr = etree.Element(qn("a:rPr"))
        run.insert(0, rpr)
    rpr.append(_hlink(rid, action))


def _deck(tmp_path):
    prs = Presentation()
    blank, title_only = prs.slide_layouts[6], prs.slide_layouts[5]
    s1 = prs.slides.add_slide(blank)
    frame = s1.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1))
    frame.table.cell(1, 0).text = "Cell link"
    s1.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1)).text_frame.text = "x"
    s2 = prs.slides.add_slide(title_only)
    s2.shapes.title.text = "Actions"
    box = s2.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    box.text_frame.text = "Show"
    for word in ("Macro", "Verb", "Nothing"):
        box.text_frame.add_paragraph().text = word
    s2.notes_slide.notes_text_frame.text = "See the notes link"
    s3 = prs.slides.add_slide(blank)
    s3.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "y"
    # A plain shape on the blank layout and one on the master.
    blank.shapes._spTree.append(etree.fromstring(
        f'<p:sp xmlns:p="{_P}" xmlns:a="{_A}"><p:nvSpPr><p:cNvPr id="90" name="Home"/>'
        '<p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:p>'
        '<a:r><a:rPr lang="en-US"/><a:t>Home</a:t></a:r></a:p></p:txBody></p:sp>'))
    prs.slide_master.shapes._spTree.append(etree.fromstring(
        f'<p:sp xmlns:p="{_P}" xmlns:a="{_A}"><p:nvSpPr><p:cNvPr id="91" name="Logo link"/>'
        '<p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr/></p:sp>'))
    path = tmp_path / "links_everywhere.pptx"
    prs.save(path)

    pkg = PptxPackage(path)
    table = read.slide_table(pkg)
    p1, p2, p3 = (r["part"] for r in table)
    root1 = pkg.root(p1)
    # Table cell (row 1, col 0) and a slide-number field.
    tc = root1.findall(f".//{qn('a:tc')}")[2]
    _link_first_run(tc, pkg.add_relationship(p1, _HL, "https://example.com/cell", external=True))
    para = [b for b in root1.iter(qn("p:txBody"))][-1].find(qn("a:p"))
    para.append(etree.fromstring(
        f'<a:fld xmlns:a="{_A}" id="{{AAAAAAAA-0000-0000-0000-000000000001}}" type="slidenum">'
        '<a:rPr lang="en-US"/><a:t>1</a:t></a:fld>'))
    _link_first_run(para, pkg.add_relationship(p1, _HL, "https://example.com/field", external=True),
                    tag="a:fld")
    # A program action on a whole shape (it keeps its external relationship).
    shape = [sp for sp in root1.iter(qn("p:sp"))][-1]
    cnvpr = next(shape.iter(qn("p:cNvPr")))
    cnvpr.insert(0, _hlink(pkg.add_relationship(p1, _HL, "file:///C:/Tools/run.exe", external=True),
                           "ppaction://program"))
    pkg.mark_dirty(p1)
    # Slide 2: actions with no relationship, one per paragraph.
    root2 = pkg.root(p2)
    paras = [b for b in root2.iter(qn("p:txBody"))][-1].findall(qn("a:p"))
    _link_first_run(paras[0], "", "ppaction://customshow?id=0&return=true")
    _link_first_run(paras[1], "", "ppaction://macro?name=Module1.Tally")
    _link_first_run(paras[2], "", "ppaction://ole?verb=0")
    _link_first_run(paras[3], "", "ppaction://noaction")
    pkg.mark_dirty(p2)
    pres = pkg.presentation()
    shows = etree.fromstring(
        f'<p:custShowLst xmlns:p="{_P}" xmlns:r="{_R}"><p:custShow name="Short version" id="0">'
        '<p:sldLst/></p:custShow></p:custShowLst>')
    pres.insert(list(pres).index(pres.find(qn("p:notesSz"))) + 1, shows)
    pkg.mark_dirty("ppt/presentation.xml")
    # Slide 2's notes.
    npart = read.notes_part_for(pkg, p2)
    _link_first_run(pkg.root(npart), pkg.add_relationship(npart, _HL, "https://example.com/notes", external=True))
    pkg.mark_dirty(npart)
    # Slide 3: SmartArt whose second node carries a link (in the data part).
    add_smartart(pkg, p3, ["Plan", "Build"])
    data = read._diagram_data_part(pkg, p3, next(pkg.root(p3).iter(qn("p:graphicFrame"))))
    pt = [p for p in pkg.root(data).iter("{http://schemas.openxmlformats.org/drawingml/2006/diagram}pt")
          if p.get("type") is None][1]
    _link_first_run(pt, pkg.add_relationship(data, _HL, "https://example.com/node", external=True))
    pkg.mark_dirty(data)
    # The blank layout's Home button (a show jump) and the master's logo link.
    layout = read._layout_info(pkg, p1)[0]
    home = [sp for sp in pkg.root(layout).iter(qn("p:sp"))][-1]
    _link_first_run(home, "", "ppaction://hlinkshowjump?jump=firstslide")
    pkg.mark_dirty(layout)
    master = "ppt/slideMasters/slideMaster1.xml"
    logo = [sp for sp in pkg.root(master).iter(qn("p:sp"))][-1]
    next(logo.iter(qn("p:cNvPr"))).insert(
        0, _hlink(pkg.add_relationship(master, _HL, "https://example.com/logo", external=True)))
    pkg.mark_dirty(master)
    pkg.save(do_backup=False)
    return path


def _by_where(out):
    rows: dict[str, list[dict]] = {}
    for rec in out["hyperlinks"]:
        rows.setdefault(rec["where"], []).append(rec)
    return rows


def test_table_cells_fields_and_notes_are_listed(tmp_path):
    out = links.list_hyperlinks(PptxPackage(_deck(tmp_path)))
    rows = _by_where(out)
    cell = rows["table"][0]
    assert (cell["row"], cell["col"], cell["paragraph"]) == (1, 0, 0)
    assert cell["url"] == "https://example.com/cell" and cell["text"] == "Cell link"
    field = [r for r in rows["text"] if r.get("field")][0]
    assert field["field"] == "slidenum" and field["url"] == "https://example.com/field"
    note = rows["notes"][0]
    assert note["slide_index"] == 1 and note["url"] == "https://example.com/notes"
    assert note["paragraph"] == 0 and note["text"] == "See the notes link"


def test_actions_without_a_relationship_are_reported(tmp_path):
    out = links.list_hyperlinks(PptxPackage(_deck(tmp_path)), scope=[1])
    kinds = {r["kind"]: r for r in out["hyperlinks"] if r["where"] == "text"}
    show = kinds["custom_show"]
    assert show["custom_show_id"] == 0 and show["custom_show_name"] == "Short version"
    assert show["broken"] is False and show["text"] == "Show"
    assert kinds["macro"]["macro"] == "Module1.Tally"
    assert kinds["ole_verb"]["verb"] == 0
    # "no action" is still not a link.
    assert all(r.get("text") != "Nothing" for r in out["hyperlinks"])


def test_custom_show_missing_from_the_deck_is_broken(tmp_path):
    path = _deck(tmp_path)
    pkg = PptxPackage(path)
    pres = pkg.presentation()
    pres.remove(pres.find(qn("p:custShowLst")))
    pkg.mark_dirty("ppt/presentation.xml")
    show = [r for r in links.list_hyperlinks(pkg)["hyperlinks"] if r["kind"] == "custom_show"][0]
    assert show["broken"] is True and "custom show 0" in show["problem"]


def test_program_action_is_named(tmp_path):
    out = links.list_hyperlinks(PptxPackage(_deck(tmp_path)), scope=[0])
    prog = [r for r in out["hyperlinks"] if r["where"] == "shape"][0]
    assert prog["kind"] == "external" and prog["action"] == "ppaction://program"


def test_smartart_layout_and_master_links_are_listed_once(tmp_path):
    out = links.list_hyperlinks(PptxPackage(_deck(tmp_path)))
    rows = _by_where(out)
    node = rows["diagram"][0]
    assert node["slide_index"] == 2 and node["node"] == 1 and node["text"] == "Build"
    assert node["url"] == "https://example.com/node"
    layout = rows["layout"]
    assert len(layout) == 1 and layout[0]["kind"] == "show_jump"
    assert layout[0]["layout"] == "Blank" and "slide_index" not in layout[0]
    master = rows["master"]
    assert len(master) == 1 and master[0]["url"] == "https://example.com/logo"
    # Slide 2 uses another layout, so the Blank layout's button is not in its scope.
    scoped = _by_where(links.list_hyperlinks(PptxPackage(_deck(tmp_path)), scope=[1]))
    assert "layout" not in scoped and len(scoped["master"]) == 1


def test_served_answer_and_help(tmp_path):
    plan = build(Settings(), public_ops())
    args = {"file_path": str(_deck(tmp_path))}
    payload, is_error = asyncio.run(dispatch(plan, "ppt_read", {"op": "list_hyperlinks", "args": args}))
    assert is_error is False, payload
    wheres = {r["where"] for r in payload["hyperlinks"]}
    assert {"table", "notes", "diagram", "layout", "master", "shape", "text"} <= wheres
    helped = help_payload(plan, "list_hyperlinks")
    for word in ("table", "notes", "SmartArt", "layout", "master"):
        assert word in helped["summary"], (word, helped["summary"])
    scope = helped["args_schema"]["properties"]["scope"]["description"]
    assert "layouts and masters they use" in scope and "listed for reading" in scope
