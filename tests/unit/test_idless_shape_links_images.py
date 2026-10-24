"""A shape without an id reads with id null in list_hyperlinks and in
replace_image_everywhere's record, as every other read now answers for it.

PowerPoint gives such a shape a new id when it opens the deck; a damaged or
hand-edited deck can carry one. These two operations still read the id with
a bare int() and failed the whole call with Python's TypeError text.

Every deck and picture here is made up by code.
"""

from __future__ import annotations

import io
import zipfile

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.package import PptxPackage, qn
from kitchensink4ppt.ops import links, sweeps


def _rewrite(path, part, fn):
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            z.writestr(info, fn(data) if info.filename == part else data)
    path.write_bytes(out.getvalue())


def _drop_id(name):
    def fn(data):
        root = etree.fromstring(data)
        for cnv in root.iter(qn("p:cNvPr")):
            if cnv.get("name") == name:
                del cnv.attrib["id"]
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return fn


def _png(path, colour):
    Image.new("RGB", (8, 8), colour).save(path, format="PNG")
    return path


def test_list_hyperlinks_lists_a_link_on_a_shape_without_an_id(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.name = "Made-up link box"
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Made-up link"
    run.hyperlink.address = "https://example.com/made-up"
    path = tmp_path / "link_no_id.pptx"
    prs.save(str(path))
    _rewrite(path, "ppt/slides/slide1.xml", _drop_id("Made-up link box"))

    found = links.list_hyperlinks(PptxPackage(path))
    rows = [h for h in found["hyperlinks"] if h.get("shape_name") == "Made-up link box"]
    assert rows and all(h["shape_id"] is None for h in rows)
    assert any("example.com/made-up" in str(h.get("target", h)) for h in rows)


def test_replace_image_everywhere_records_a_picture_without_an_id(tmp_path):
    old = _png(tmp_path / "old.png", (200, 30, 30))
    new = _png(tmp_path / "new.png", (30, 30, 200))
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pic = slide.shapes.add_picture(str(old), Inches(1), Inches(1), Inches(1), Inches(1))
    pic.name = "Made-up picture"
    path = tmp_path / "pic_no_id.pptx"
    prs.save(str(path))
    _rewrite(path, "ppt/slides/slide1.xml", _drop_id("Made-up picture"))

    pkg = PptxPackage(path)
    done = sweeps.replace_image_everywhere(pkg, str(old), str(new))
    records = [r for r in done["instances"] if r.get("name") == "Made-up picture"]
    assert records and records[0]["shape_id"] is None
