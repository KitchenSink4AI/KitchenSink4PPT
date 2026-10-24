"""Stacking and reading order with a shape in a compatibility wrapper.

PowerPoint stores an equation text box (and ink) inside a slide-level
mc:AlternateContent, and its stacking and reading order count that wrapper
as one shape. set_reading_order moved only plain shapes, re-appending them
after the wrapper (which so went to the back and was read first, without a
word) and after a trailing p:extLst, where the schema allows no shape; it
could not name the wrapped shape at all. set_z_order's forward and backward
stepped over a wrapper as if it were not there. Both now treat the wrapper
as one member of the order, named by the wrapped shape's id.

Every deck here is made up by code.
"""

from __future__ import annotations

import copy

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.errors import TargetNotFound
from kitchensink4ppt.core.package import PptxPackage, qn
from kitchensink4ppt.ops import access
from kitchensink4ppt.ops import shapes as sh

MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
AC = f"{{{MC}}}AlternateContent"
A14 = "http://schemas.microsoft.com/office/drawing/2010/main"


@pytest.fixture()
def deck(tmp_path):
    """Three text boxes (ids 2, 3, 4); id 3 wrapped as PowerPoint wraps an
    equation box: mc:Choice Requires="a14" with the shape, mc:Fallback with a
    copy."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for n in range(3):
        box = slide.shapes.add_textbox(Inches(1), Inches(1 + n), Inches(4), Inches(1))
        box.text_frame.text = f"Made-up box {n + 1}"
    path = tmp_path / "wrapped.pptx"
    prs.save(str(path))
    pkg = PptxPackage(path)
    part = pkg.slide_parts()[0]
    tree = pkg.root(part).find(f"{qn('p:cSld')}/{qn('p:spTree')}")
    middle = [c for c in tree if c.tag == qn("p:sp")][1]
    wrapper = etree.Element(AC, nsmap={"mc": MC})
    choice = etree.SubElement(wrapper, f"{{{MC}}}Choice", nsmap={"a14": A14})
    choice.set("Requires", "a14")
    fallback = etree.SubElement(wrapper, f"{{{MC}}}Fallback")
    middle.addprevious(wrapper)
    fallback.append(copy.deepcopy(middle))
    choice.append(middle)
    pkg.mark_dirty(part)
    pkg.save(do_backup=False)
    return path


def _order(path):
    pkg = PptxPackage(path)
    tree = pkg.root(pkg.slide_parts()[0]).find(f"{qn('p:cSld')}/{qn('p:spTree')}")
    out = []
    for c in tree:
        if c.tag == qn("p:sp"):
            out.append(int(c.find(f"{qn('p:nvSpPr')}/{qn('p:cNvPr')}").get("id")))
        elif c.tag == AC:
            out.append(("wrapped", int(next(c.iter(qn("p:cNvPr"))).get("id"))))
        elif c.tag == qn("p:extLst"):
            out.append("extLst")
    return out


def _run(path, fn, *args):
    pkg = PptxPackage(path)
    out = fn(pkg, 0, *args)
    pkg.save(do_backup=False)
    return out


def test_reading_order_places_the_wrapped_shape_where_listed(deck):
    out = _run(deck, access.set_reading_order, [4, 3, 2])
    assert out["changed"] is True
    assert _order(deck) == [4, ("wrapped", 3), 2]


def test_reading_order_without_the_wrapped_shape_is_refused_by_name(deck):
    before = deck.read_bytes()
    with pytest.raises(TargetNotFound) as refused:
        _run(deck, access.set_reading_order, [4, 2])
    message = str(refused.value)
    assert "missing ids [3]" in message and "compatibility wrapper" in message
    assert deck.read_bytes() == before


def test_reading_order_keeps_a_trailing_extlst_last(deck):
    pkg = PptxPackage(deck)
    part = pkg.slide_parts()[0]
    tree = pkg.root(part).find(f"{qn('p:cSld')}/{qn('p:spTree')}")
    etree.SubElement(tree, qn("p:extLst"))
    pkg.mark_dirty(part)
    pkg.save(do_backup=False)
    _run(deck, access.set_reading_order, [2, 4, 3])
    assert _order(deck) == [2, 4, ("wrapped", 3), "extLst"]


def test_z_order_forward_passes_the_wrapper_one_step(deck):
    out = _run(deck, sh.set_z_order, 2, "forward")
    assert _order(deck) == [("wrapped", 3), 2, 4]
    assert out["z"] == 1 and out["of"] == 3


def test_z_order_backward_passes_the_wrapper_one_step(deck):
    _run(deck, sh.set_z_order, 4, "backward")
    assert _order(deck) == [2, 4, ("wrapped", 3)]
