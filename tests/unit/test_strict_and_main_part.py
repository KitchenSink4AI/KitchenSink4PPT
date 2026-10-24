"""A deck this server cannot read is refused by name, never read as empty.

The server used to find the presentation by one fixed part name and one set of
namespaces. A deck saved as Strict Open XML (the same part names, other
namespaces) therefore read as a deck with no slides, answered ok, and an
edit said "slide index 0 out of range". A package whose main part has another
name (the root relationship says where it is) was refused as "not a
PowerPoint presentation", and one that also held a ppt/presentation.xml the
root relationship does not name would be read from the wrong part. The decks
are made up and converted by code.
"""

from __future__ import annotations

import asyncio
import zipfile

from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt import server  # noqa: F401  (the import registers ops)
from kitchensink4ppt.registry import public_ops
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.plan import build

_TO_STRICT = [
    (b"http://schemas.openxmlformats.org/presentationml/2006/main",
     b"http://purl.oclc.org/ooxml/presentationml/main"),
    (b"http://schemas.openxmlformats.org/drawingml/2006/main",
     b"http://purl.oclc.org/ooxml/drawingml/main"),
    (b"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
     b"http://purl.oclc.org/ooxml/officeDocument/relationships"),
]


def _plain(tmp_path):
    path = tmp_path / "plain.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Quarterly plan"
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text_frame.text = "Made up"
    prs.save(path)
    return path


def _rewrite(src, dst, fn) -> None:
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            name, data = fn(info.filename, zin.read(info.filename))
            if name is not None:
                zout.writestr(name, data)


def _strict(tmp_path):
    dst = tmp_path / "strict.pptx"

    def fn(name, data):
        if name.endswith((".xml", ".rels")) and name != "[Content_Types].xml":
            for a, b in _TO_STRICT:
                data = data.replace(a, b)
        return name, data

    _rewrite(_plain(tmp_path), dst, fn)
    return dst


def _moved_main(tmp_path, keep_decoy: bool):
    dst = tmp_path / ("decoy.pptx" if keep_decoy else "moved.pptx")
    src = _plain(tmp_path)
    with zipfile.ZipFile(src) as zf:
        pres = zf.read("ppt/presentation.xml")
        pres_rels = zf.read("ppt/_rels/presentation.xml.rels")

    def fn(name, data):
        if name == "_rels/.rels":
            return name, data.replace(b"ppt/presentation.xml", b"ppt/main.xml")
        if name == "[Content_Types].xml":
            return name, data.replace(b"/ppt/presentation.xml", b"/ppt/main.xml")
        if name in ("ppt/presentation.xml", "ppt/_rels/presentation.xml.rels") and not keep_decoy:
            return None, None
        return name, data

    _rewrite(src, dst, fn)
    with zipfile.ZipFile(dst, "a") as zf:
        zf.writestr("ppt/main.xml", pres)
        zf.writestr("ppt/_rels/main.xml.rels", pres_rels)
    return dst


def _call(tool, op, path, **extra):
    plan = build(Settings(), public_ops())
    args = {"file_path": str(path), **extra}
    return asyncio.run(dispatch(plan, tool, {"op": op, "args": args}))


def test_strict_deck_is_refused_by_reads_and_edits(tmp_path):
    path = _strict(tmp_path)
    before = path.read_bytes()
    for tool, op, extra in (
        ("ppt_read", "get_text", {}),
        ("ppt_read", "get_presentation_view", {}),
        ("ppt_edit", "insert_textbox", {"slide": 0, "text": "x", "x": 1, "y": 1, "w": 2, "h": 1}),
    ):
        payload, is_error = _call(tool, op, path, **extra)
        assert is_error is True, (op, payload)
        error = payload["error"]
        assert error["code"] == "UNSUPPORTED_CONTENT", (op, error)
        assert "Strict Open XML" in error["message"], (op, error)
        assert "PowerPoint Presentation (.pptx)" in error["hint"], (op, error)
    assert path.read_bytes() == before


def test_main_part_under_another_name_is_refused_by_name(tmp_path):
    payload, is_error = _call("ppt_read", "get_text", _moved_main(tmp_path, keep_decoy=False))
    assert is_error is True
    assert "ppt/main.xml" in payload["error"]["message"], payload


def test_a_decoy_presentation_part_is_never_read(tmp_path):
    payload, is_error = _call("ppt_read", "get_text", _moved_main(tmp_path, keep_decoy=True))
    assert is_error is True, payload
    assert "ppt/main.xml" in payload["error"]["message"], payload
