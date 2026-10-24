"""Text inserted into an empty table cell keeps its insertion style."""

from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.package import PptxPackage, qn
from kitchensink4ppt.ops import tables as tb


def _one_cell_deck(path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_table(
        1, 1, Inches(1), Inches(1), Inches(4), Inches(1)
    )
    prs.save(path)
    pkg = PptxPackage(path)
    tc = tb.resolve_table(pkg, 0, 0)["tbl"].find(
        f"{qn('a:tr')}/{qn('a:tc')}"
    )
    return pkg, tc, path


def _set_empty_cell_style(tc, *, source: str) -> None:
    body = tc.find(qn("a:txBody"))
    paragraph = body.find(qn("a:p"))
    for child in list(paragraph):
        paragraph.remove(child)
    ppr = etree.SubElement(paragraph, qn("a:pPr"))
    ppr.set("algn", "r")
    if source == "endParaRPr":
        rpr = etree.SubElement(paragraph, qn("a:endParaRPr"))
    else:
        rpr = etree.SubElement(ppr, qn("a:defRPr"))
    rpr.set("lang", "en-US")
    rpr.set("sz", "900")
    rpr.set("b", "1")
    rpr.set("i", "1")
    fill = etree.SubElement(rpr, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr")).set("val", "0000C0")
    for tag, face in (
        ("a:latin", "Aptos"),
        ("a:ea", "Malgun Gothic"),
        ("a:cs", "Arial"),
    ):
        etree.SubElement(rpr, qn(tag)).set("typeface", face)


def _inserted_run(tc):
    paragraph = tc.find(f"{qn('a:txBody')}/{qn('a:p')}")
    return paragraph, paragraph.find(f"{qn('a:r')}/{qn('a:rPr')}")


def test_filling_empty_cell_preserves_end_paragraph_formatting(tmp_path):
    pkg, tc, path = _one_cell_deck(tmp_path / "end-paragraph.pptx")
    _set_empty_cell_style(tc, source="endParaRPr")

    out = tb.set_table_cells(
        pkg, 0, 0, [{"row": 0, "col": 0, "text": "77"}]
    )

    pkg.save()
    reopened = PptxPackage(path)
    reopened_tc = tb.resolve_table(reopened, 0, 0)["tbl"].find(
        f"{qn('a:tr')}/{qn('a:tc')}"
    )
    paragraph, rpr = _inserted_run(reopened_tc)
    assert paragraph.find(qn("a:pPr")).get("algn") == "r"
    assert rpr.get("sz") == "900"
    assert rpr.get("b") == "1"
    assert rpr.get("i") == "1"
    assert rpr.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}").get(
        "val"
    ) == "0000C0"
    assert rpr.find(qn("a:latin")).get("typeface") == "Aptos"
    assert rpr.find(qn("a:ea")).get("typeface") == "Malgun Gothic"
    assert rpr.find(qn("a:cs")).get("typeface") == "Arial"
    assert set(out["preserved"]) == {
        "alignment", "bold", "font", "italic", "run_color", "size",
    }
    assert out["preserved_by_cell"] == [{
        "row": 0,
        "col": 0,
        "preserved": [
            "alignment", "bold", "font", "italic", "run_color", "size",
        ],
    }]


def test_filling_empty_cell_preserves_default_run_formatting(tmp_path):
    pkg, tc, _path = _one_cell_deck(tmp_path / "default-run.pptx")
    _set_empty_cell_style(tc, source="defRPr")

    out = tb.set_table_cells(
        pkg, 0, 0, [{"row": 0, "col": 0, "text": "77"}]
    )

    paragraph, rpr = _inserted_run(tc)
    assert paragraph.find(qn("a:pPr")).get("algn") == "r"
    assert rpr.get("sz") == "900"
    assert rpr.get("b") == "1"
    assert rpr.get("i") == "1"
    assert rpr.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}").get(
        "val"
    ) == "0000C0"
    assert rpr.find(qn("a:latin")).get("typeface") == "Aptos"
    assert set(out["preserved"]) == {
        "alignment", "bold", "font", "italic", "run_color", "size",
    }


def test_empty_cell_explicit_overrides_win_and_others_survive(tmp_path):
    pkg, tc, _path = _one_cell_deck(tmp_path / "overrides.pptx")
    _set_empty_cell_style(tc, source="endParaRPr")

    out = tb.set_table_cells(
        pkg,
        0,
        0,
        [{
            "row": 0,
            "col": 0,
            "text": "77",
            "align": "left",
            "bold": False,
            "size_pt": 18,
            "color": "C00000",
            "font": "Calibri",
        }],
    )

    paragraph, rpr = _inserted_run(tc)
    assert paragraph.find(qn("a:pPr")).get("algn") == "l"
    assert rpr.get("sz") == "1800"
    assert rpr.get("b") == "0"
    assert rpr.get("i") == "1"
    assert rpr.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}").get(
        "val"
    ) == "C00000"
    assert rpr.find(qn("a:latin")).get("typeface") == "Calibri"
    assert set(out["preserved"]) == {"italic"}


def test_nonempty_cell_preservation_result_stays_compatible(tmp_path):
    pkg, tc, _path = _one_cell_deck(tmp_path / "nonempty.pptx")
    tb.set_table_cells(
        pkg,
        0,
        0,
        [{
            "row": 0,
            "col": 0,
            "text": "before",
            "align": "right",
            "bold": True,
            "italic": True,
            "size_pt": 9,
            "color": "0000C0",
            "font": "Aptos",
        }],
    )

    out = tb.set_table_cells(
        pkg, 0, 0, [{"row": 0, "col": 0, "text": "after"}]
    )

    paragraph, rpr = _inserted_run(tc)
    assert paragraph.find(qn("a:pPr")).get("algn") == "r"
    assert rpr.get("b") == "1"
    assert rpr.get("i") == "1"
    assert rpr.get("sz") == "900"
    assert set(out["preserved"]) == {
        "alignment", "font", "run_color", "size",
    }
