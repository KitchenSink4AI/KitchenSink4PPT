"""A save stopped in its last instants, on a deck that carries a Mark of the Web.

Every deck that came by email, Teams or a download carries a Mark of the Web
(a Zone.Identifier stream). When ReplaceFileW puts a new version in place of
such a deck, it copies the stream onto the new file, and that moves the new
file's modification time to the instant of the swap. Recovery must not read
that as "saved again since": the deck is still the version the stopped save
put in place, and the backups must land exactly as they would have had the
save not been stopped.

Each case is a real kill of a saving process at one instant:
  late        ReplaceFileW went through; the old file is still aside; the
              backups have not landed.
  postremove  the old file was removed too; the backups have not landed.
  midland     the anchor backup landed, the prev backup not yet.
and is then finished by the next start of a server or by the next change.
The same kills on a deck without the stream are the controls.
"""

from __future__ import annotations

import asyncio
import os
import subprocess
import sys
import time
from pathlib import Path

import pytest
from pptx import Presentation

from kitchensink4ppt import server
from kitchensink4ppt.core import safesave, swapwait
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.plan import build

pytestmark = pytest.mark.skipif(sys.platform != "win32", reason="ReplaceFileW")

NO_WINDOW = 0x08000000
ZONE = b"[ZoneTransfer]\r\nZoneId=3\r\nReferrerUrl=https://example.com/deck\r\n"

VICTIM = r'''
import asyncio, os, sys, time
from kitchensink4ppt import server  # noqa: F401
from kitchensink4ppt.core import safesave
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.plan import build
deck, marker, mode = sys.argv[1], sys.argv[2], sys.argv[3]
fn, move, last_error, text = safesave._replace_file()

def stop():
    open(marker, "w").close()
    time.sleep(600)

def replace(ldst, lsrc, lbak, flags, a, b):
    ok = fn(ldst, lsrc, lbak, flags, a, b)
    if mode == "late":
        stop()
    return ok

safesave._REPLACE_FILE = (replace, move, last_error, text)
if mode == "postremove":
    _land = safesave.RotationTicket.land
    def land(self, name):
        stop()
        return _land(self, name)
    safesave.RotationTicket.land = land
if mode == "midland":
    _replace = safesave.replace_with_retry
    def replace_with_retry(src, dst, **kw):
        if (os.path.basename(str(src)).startswith(".slot-")
                and os.path.basename(str(dst)) == safesave.PREV_SLOT):
            stop()
        return _replace(src, dst, **kw)
    safesave.replace_with_retry = replace_with_retry
plan = build(Settings())
asyncio.run(dispatch(plan, "ppt_edit", {"op": "set_notes", "args": {
    "file_path": deck, "slide": 0, "text": "v2", "live": "off"}}))
'''


@pytest.fixture(autouse=True)
def own_journal(tmp_path, monkeypatch):
    monkeypatch.setenv("KS4P_SAVE_JOURNAL_DIR", str(tmp_path / "journal"))
    safesave.take_notices()
    yield
    safesave.take_notices()


def _deck(path: Path, notes: str, marked: bool) -> Path:
    prs = Presentation()
    for title in ("Alpha", "Bravo"):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.notes_slide.notes_text_frame.text = notes
    prs.save(path)
    if marked:
        with open(str(path) + ":Zone.Identifier", "wb") as fh:
            fh.write(ZONE)
    return path


def _notes(path) -> str | None:
    if not os.path.lexists(path):
        return None
    return Presentation(str(path)).slides[0].notes_slide.notes_text_frame.text


def _call(op, args):
    plan = build(Settings())
    return asyncio.run(dispatch(plan, plan.specs[op].group, {"op": op, "args": args}))


def _stopped(tmp_path: Path, mode: str, marked: bool) -> Path:
    """v0, then v1 through the product (prev = anchor = v0), then a save of v2
    stopped at `mode`. For midland the prev backup is two hours old first, so
    the stopped save rotates the anchor too."""
    deck = _deck(tmp_path / "Deck.pptx", "v0", marked)
    _payload, refused = _call("set_notes", {"file_path": str(deck), "slide": 0, "text": "v1",
                                            "live": "off"})
    assert not refused
    if mode == "midland":
        prev = safesave.slot_dir(deck) / safesave.PREV_SLOT
        aged = time.time() - 7200
        os.utime(prev, (aged, aged))
    script = tmp_path / "victim.py"
    script.write_text(VICTIM, encoding="utf-8")
    marker = tmp_path / "stopped"
    victim = subprocess.Popen([sys.executable, "-X", "utf8", str(script), str(deck), str(marker),
                               mode], stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                              stderr=subprocess.DEVNULL, creationflags=NO_WINDOW)
    try:
        deadline = time.monotonic() + 120
        while not marker.exists():
            assert victim.poll() is None, "the saving process ended before the stop"
            assert time.monotonic() < deadline, "the saving process never reached the stop"
            time.sleep(0.02)
    finally:
        victim.kill()  # TerminateProcess on this test's own child
        victim.wait(timeout=60)
    time.sleep(0.3)
    assert _notes(deck) == "v2"
    return deck


def _left_over(deck: Path) -> dict:
    folder = safesave.slot_dir(deck)
    key = safesave.canonical_key(deck)
    return {
        "asides": swapwait.asides(deck),
        "staged_links": sorted(n for n in os.listdir(folder) if n.startswith(".slot-")),
        "entries": [entry for entry, record in safesave._journal_entries()
                    if isinstance(record.get("deck"), str)
                    and safesave.canonical_key(record["deck"]) == key],
        "save_tmp": os.path.lexists(str(deck) + ".ppt-mcp-tmp"),
    }


CLEAN = {"asides": [], "staged_links": [], "entries": [], "save_tmp": False}


@pytest.mark.parametrize("marked", [False, True], ids=["plain", "marked"])
@pytest.mark.parametrize("mode", ["late", "postremove", "midland"])
def test_the_next_start_finishes_a_stopped_save_as_the_save_would_have(tmp_path, mode, marked):
    deck = _stopped(tmp_path, mode, marked)
    server._recover_at_start(Settings())
    notices = safesave.take_notices()
    folder = safesave.slot_dir(deck)
    assert _notes(deck) == "v2"
    assert _notes(folder / safesave.PREV_SLOT) == "v1", "the undo of the last change did not land"
    assert _notes(folder / safesave.ANCHOR_SLOT) == ("v1" if mode == "midland" else "v0")
    assert _left_over(deck) == CLEAN
    assert notices == [], "nothing saved over the deck after the stop"


@pytest.mark.parametrize("marked", [False, True], ids=["plain", "marked"])
@pytest.mark.parametrize("mode", ["late", "postremove", "midland"])
def test_the_next_change_finishes_a_stopped_save_first(tmp_path, mode, marked):
    deck = _stopped(tmp_path, mode, marked)
    payload, refused = _call("set_notes", {"file_path": str(deck), "slide": 0, "text": "v3",
                                           "live": "off"})
    assert not refused, payload
    folder = safesave.slot_dir(deck)
    assert _notes(deck) == "v3"
    assert _notes(folder / safesave.PREV_SLOT) == "v2"
    assert _notes(folder / safesave.ANCHOR_SLOT) == ("v1" if mode == "midland" else "v0"), (
        "the session anchor was replaced")
    assert _left_over(deck) == CLEAN
    assert not payload.get("notices"), payload.get("notices")
    if marked:
        with open(str(deck) + ":Zone.Identifier", "rb") as fh:
            assert fh.read() == ZONE, "the Mark of the Web stays on the deck"
