"""Deck-wide sweeps reach parts by what they are, not by what they are called.

A sweep used to put a part in a bucket (layouts, notes, charts, ...) only when its
name matched PowerPoint's usual pattern, so a layout, a notes page or a chart
that another program saved under another name was skipped by replace_fonts,
set_language, replace_colors and font_inventory, and so was a chartex chart
(waterfall, treemap, sunburst: content type application/vnd.ms-office.chartex+xml,
part name chartExN.xml). The [Content_Types].xml entry says what a part is,
so that decides now. The deck is made up and built by code.
"""

from __future__ import annotations

import zipfile

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from kitchensink4ppt.core.package import PptxPackage
from kitchensink4ppt.ops import _traverse as tv, sweeps

_RENAMES = {
    "notesSlide1.xml": "speaker1.xml",
    "chart1.xml": "sales.xml",
    "slideLayout6.xml": "titleOnly.xml",
}
_CX = "http://schemas.microsoft.com/office/drawing/2014/chartex"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_CHARTEX = (
    f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    f'<cx:chartSpace xmlns:cx="{_CX}" xmlns:a="{_A}"><cx:chart><cx:title><cx:tx><cx:rich>'
    '<a:bodyPr/><a:p><a:r><a:rPr lang="en-US"><a:solidFill><a:srgbClr val="112233"/></a:solidFill>'
    '<a:latin typeface="Arial"/></a:rPr><a:t>Flow</a:t></a:r></a:p></cx:rich></cx:tx></cx:title>'
    '<cx:plotArea><cx:plotAreaRegion><cx:series layoutId="waterfall"><cx:spPr><a:solidFill>'
    '<a:srgbClr val="112233"/></a:solidFill></cx:spPr></cx:series></cx:plotAreaRegion></cx:plotArea>'
    '</cx:chart></cx:chartSpace>'
).encode()


def _deck(tmp_path):
    src = tmp_path / "canonical.pptx"
    prs = Presentation()
    layout = prs.slide_layouts[5]
    layout.placeholders[0].text_frame.paragraphs[0].runs[0].font.name = "Arial"
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Plan"
    notes = slide.notes_slide.notes_text_frame
    notes.text = "Speaker text"
    notes.paragraphs[0].runs[0].font.name = "Arial"
    data = CategoryChartData()
    data.categories = ["North", "South"]
    data.add_series("Units", (3, 5))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(5), Inches(3), data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Sales"
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.name = "Arial"
    prs.save(src)

    dst = tmp_path / "renamed.pptx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            name, data_bytes = info.filename, zin.read(info.filename)
            for old, new in _RENAMES.items():
                name = name.replace(old, new)
                if info.filename.endswith(".rels") or info.filename == "[Content_Types].xml":
                    data_bytes = data_bytes.replace(old.encode(), new.encode())
            if info.filename == "[Content_Types].xml":
                data_bytes = data_bytes.replace(
                    b"</Types>",
                    b'<Override PartName="/ppt/charts/chartEx1.xml" '
                    b'ContentType="application/vnd.ms-office.chartex+xml"/></Types>')
            zout.writestr(name, data_bytes)
        zout.writestr("ppt/charts/chartEx1.xml", _CHARTEX)
    return dst


_PARTS = ("ppt/notesSlides/speaker1.xml", "ppt/charts/sales.xml",
          "ppt/slideLayouts/titleOnly.xml", "ppt/charts/chartEx1.xml")


def test_buckets_follow_the_content_type(tmp_path):
    pkg = PptxPackage(_deck(tmp_path))
    bucket_of = dict(tv.parts_in_scope(pkg, "all"))
    assert bucket_of.get("ppt/notesSlides/speaker1.xml") == "notes"
    assert bucket_of.get("ppt/charts/sales.xml") == "charts"
    assert bucket_of.get("ppt/slideLayouts/titleOnly.xml") == "layouts"
    assert bucket_of.get("ppt/charts/chartEx1.xml") == "charts"


def test_font_and_language_sweeps_reach_renamed_and_chartex_parts(tmp_path):
    path = _deck(tmp_path)
    pkg = PptxPackage(path)
    assert sweeps.font_inventory(pkg)["fonts"], "the inventory saw nothing"
    sweeps.replace_fonts(pkg, {"Arial": "Georgia"})
    sweeps.set_language(pkg, "fr-FR")
    sweeps.replace_colors(pkg, {"112233": "445566"})
    pkg.save(do_backup=False)
    with zipfile.ZipFile(path) as zf:
        for part in _PARTS:
            xml = zf.read(part)
            assert b'typeface="Arial"' not in xml, part
            assert b'lang="fr-FR"' in xml, part
        cx = zf.read("ppt/charts/chartEx1.xml")
        assert b"112233" not in cx and cx.count(b"445566") == 2


def test_colour_roles_in_chartex(tmp_path):
    pkg = PptxPackage(_deck(tmp_path))
    roles = {ctx.role for ctx in tv.iter_colors(pkg, "charts")
             if ctx.part == "ppt/charts/chartEx1.xml"}
    assert roles == {"text", "fill"}, roles
