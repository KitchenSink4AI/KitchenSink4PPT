"""export_shape_image: one shape or group as a PNG at a set size.

A figure was rendered by exporting the whole slide and cropping it by hand.
export_shape_image renders the slide in PowerPoint at the scale that gives
the shape the width asked for and keeps only the area the shape is drawn
in: its box through its groups, turned when it is rotated, grown by half
its outline, and clipped to the slide (said so when clipped). The answer
gives that area in inches and in pixels. Like every export it refuses an
existing file unless overwrite=true, and it says plainly what it needs:
PowerPoint to render, Pillow to crop.
"""

from __future__ import annotations

import asyncio
import sys
from pathlib import Path

import pytest
from pptx import Presentation

from kitchensink4ppt import server  # noqa: F401  (the import registers ops)
from kitchensink4ppt.com import bridge
from kitchensink4ppt.ops import export as ex
from kitchensink4ppt.registry import public_ops
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.plan import build

PIL = pytest.importorskip("PIL")
from PIL import Image  # noqa: E402

SLIDE_W, SLIDE_H = 10.0, 7.5  # python-pptx's default 4:3 slide, inches


def _call(op: str, args: dict):
    plan = build(Settings(), public_ops())
    return asyncio.run(dispatch(plan, plan.specs[op].group, {"op": op, "args": args}))


def _ok(op: str, args: dict) -> dict:
    payload, is_error = _call(op, args)
    assert is_error is False, payload
    return payload


def _deck(path: Path) -> str:
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(path))
    return str(path)


def _shape(deck, name, x, y, w, h, **extra) -> int:
    payload = _ok("insert_shape", {
        "file_path": deck, "slide": 0, "shape_type": "rect", "name": name,
        "x": x, "y": y, "w": w, "h": h, **extra})
    return payload["changed"]["shape_id"]


@pytest.fixture
def renders(monkeypatch):
    """PowerPoint present but never started: a slide 'render' is a white
    image of the size asked for with the slide's inch grid drawn in, so a
    crop can be checked against where the shape is."""
    calls: list[dict] = []

    def images(path, output_dir=None, slides=None, width=1280, height=None, **kw):
        calls.append({"slides": slides, "width": width, "height": height})
        out = Path(output_dir)
        out.mkdir(parents=True, exist_ok=True)
        image = Image.new("RGB", (width, height), "white")
        image.save(out / f"slide{slides[0] + 1}.png")
        return {"images": [{"slide": slides[0], "file": str(out / f"slide{slides[0] + 1}.png")}],
                "width": width, "height": height, "engine": "powerpoint-com"}

    monkeypatch.setattr(ex, "_com_available", lambda: True)
    monkeypatch.setattr(bridge, "com_export_slide_images", images)
    return calls


def _bounds(payload) -> tuple[float, ...]:
    box = payload["bounds"]["inches"]
    return tuple(round(box[k], 3) for k in ("x", "y", "w", "h"))


def test_renders_the_shape_area_at_the_width_asked(tmp_path, renders):
    deck = _deck(tmp_path / "deck.pptx")
    _shape(deck, "Box", 2, 1, 4, 2, line="none")
    out = tmp_path / "box.png"
    payload = _ok("export_shape_image", {
        "file_path": deck, "slide": 0, "shape": "Box", "output": str(out), "width": 400})
    assert _bounds(payload) == (2.0, 1.0, 4.0, 2.0)
    with Image.open(out) as image:
        assert image.size == (400, 200)
    # The slide was rendered at 100 px per inch, so the box is 400 px wide.
    assert renders[-1]["width"] == 1000 and renders[-1]["height"] == 750
    assert payload["bounds"]["pixels"] == {"x": 200, "y": 100, "w": 400, "h": 200}
    assert payload["slide_number"] == 1 and payload["image"]["width"] == 400
    assert payload["clipped"] is False


def test_height_alone_keeps_the_proportions(tmp_path, renders):
    deck = _deck(tmp_path / "deck.pptx")
    _shape(deck, "Box", 2, 1, 4, 2, line="none")
    out = tmp_path / "box.png"
    _ok("export_shape_image", {
        "file_path": deck, "slide": 0, "shape": "Box", "output": str(out), "height": 100})
    with Image.open(out) as image:
        assert image.size == (200, 100)


def test_the_outline_and_rotation_count_in_the_area(tmp_path, renders):
    deck = _deck(tmp_path / "deck.pptx")
    _shape(deck, "Lined", 2, 1, 4, 2, line={"color": "#000000", "width": 7.2})
    payload = _ok("export_shape_image", {
        "file_path": deck, "slide": 0, "shape": "Lined", "output": str(tmp_path / "a.png")})
    assert _bounds(payload) == (1.95, 0.95, 4.1, 2.1)  # half of 0.1 in each side
    _shape(deck, "Turned", 4, 4, 2, 1, line="none", rotation=90)
    payload = _ok("export_shape_image", {
        "file_path": deck, "slide": 0, "shape": "Turned", "output": str(tmp_path / "b.png")})
    assert _bounds(payload) == (4.5, 3.5, 1.0, 2.0)  # same centre, turned box


def test_a_member_of_a_group_is_found_where_it_is_drawn(tmp_path, renders):
    deck = _deck(tmp_path / "deck.pptx")
    a = _shape(deck, "A", 1, 1, 2, 1, line="none")
    b = _shape(deck, "B", 4, 1, 2, 1, line="none")
    _ok("group_shapes", {"file_path": deck, "slide": 0, "ids": [a, b], "name": "Pair"})
    _ok("set_shape", {"file_path": deck, "slide": 0, "shape": "Pair", "w": 2.5})
    payload = _ok("export_shape_image", {
        "file_path": deck, "slide": 0, "shape": "Pair/B", "output": str(tmp_path / "b.png")})
    assert _bounds(payload) == (2.5, 1.0, 1.0, 1.0)  # only the width was halved


def test_an_area_off_the_slide_is_clipped_and_said_so(tmp_path, renders):
    deck = _deck(tmp_path / "deck.pptx")
    _shape(deck, "Edge", 9, 1, 2, 1, line="none")
    payload = _ok("export_shape_image", {
        "file_path": deck, "slide": 0, "shape": "Edge", "output": str(tmp_path / "e.png")})
    assert _bounds(payload) == (9.0, 1.0, 1.0, 1.0) and payload["clipped"] is True
    _shape(deck, "Gone", 11, 1, 1, 1, line="none")
    payload, is_error = _call("export_shape_image", {
        "file_path": deck, "slide": 0, "shape": "Gone", "output": str(tmp_path / "g.png")})
    assert is_error and "off the slide" in payload["error"]["message"]


def test_an_existing_png_needs_overwrite(tmp_path, renders):
    deck = _deck(tmp_path / "deck.pptx")
    _shape(deck, "Box", 2, 1, 4, 2, line="none")
    out = tmp_path / "box.png"
    out.write_bytes(b"old")
    payload, is_error = _call("export_shape_image", {
        "file_path": deck, "slide": 0, "shape": "Box", "output": str(out)})
    assert is_error and "overwrite=true" in payload["error"]["message"]
    assert not renders and out.read_bytes() == b"old"
    _ok("export_shape_image", {
        "file_path": deck, "slide": 0, "shape": "Box", "output": str(out), "overwrite": True})
    with Image.open(out) as image:
        assert image.size[0] == 1280


def test_the_default_name_sits_beside_the_deck(tmp_path, renders):
    deck = _deck(tmp_path / "deck.pptx")
    _shape(deck, "Box", 2, 1, 4, 2, line="none")
    payload = _ok("export_shape_image", {"file_path": deck, "slide_number": 1, "shape": "Box"})
    assert Path(payload["file"]).parent == tmp_path
    assert Path(payload["file"]).suffix == ".png"


def test_a_very_small_shape_is_rendered_at_the_largest_slide_render(tmp_path, renders):
    deck = _deck(tmp_path / "deck.pptx")
    _shape(deck, "Dot", 1, 1, 0.1, 0.1, line="none")
    payload = _ok("export_shape_image", {
        "file_path": deck, "slide": 0, "shape": "Dot", "output": str(tmp_path / "d.png"),
        "width": 2000})
    assert max(renders[-1]["width"], renders[-1]["height"]) <= ex.MAX_RENDER_PX
    assert payload["image"]["width"] == 2000
    assert any("enlarged" in w for w in payload.get("warnings", []))


def test_it_says_what_it_needs(tmp_path, monkeypatch):
    deck = _deck(tmp_path / "deck.pptx")
    _shape(deck, "Box", 2, 1, 4, 2)
    monkeypatch.setattr(ex, "_com_available", lambda: False)
    payload, is_error = _call("export_shape_image", {
        "file_path": deck, "slide": 0, "shape": "Box", "output": str(tmp_path / "x.png")})
    assert is_error and "needs PowerPoint" in payload["error"]["message"]
    monkeypatch.setattr(ex, "_com_available", lambda: True)
    monkeypatch.setitem(sys.modules, "PIL", None)
    payload, is_error = _call("export_shape_image", {
        "file_path": deck, "slide": 0, "shape": "Box", "output": str(tmp_path / "x.png")})
    assert is_error and "Pillow" in payload["error"]["message"]


# --------------------------------------------- a sandboxed server (KS4P_ALLOWED_ROOTS)


class _FakeSlide:
    def __init__(self, seen: list):
        self.seen = seen

    def Export(self, path, fmt, width, height):  # noqa: N802 (COM name)
        self.seen.append(path)
        Image.new("RGB", (width, height), "white").save(path)


class _FakeSlides:
    Count = 1

    def __init__(self, seen: list):
        self.seen = seen

    def Item(self, index):  # noqa: N802 (COM name)
        return _FakeSlide(self.seen)


class _FakePresentation:
    def __init__(self, seen: list):
        self.Slides = _FakeSlides(seen)


class _FakePresentations(list):
    def __init__(self, seen: list):
        super().__init__()
        self.seen = seen

    def Open(self, target, ReadOnly, Untitled, WithWindow):  # noqa: N802,N803
        return _FakePresentation(self.seen)


@pytest.fixture
def fake_powerpoint(monkeypatch):
    """The real bridge (its own path checks included) over a PowerPoint that
    is never started: Slide.Export writes a white PNG where it is told to and
    records that path."""
    import contextlib
    import types

    seen: list[str] = []

    @contextlib.contextmanager
    def session():
        app = types.SimpleNamespace(Presentations=_FakePresentations(seen))
        yield bridge._PowerPointSession(app, launched=False)

    monkeypatch.setattr(ex, "_com_available", lambda: True)
    monkeypatch.setattr(bridge, "_powerpoint", session)
    return seen


def test_a_sandboxed_server_renders_in_its_own_scratch_folder(tmp_path, monkeypatch,
                                                              fake_powerpoint):
    """With KS4P_ALLOWED_ROOTS set, the slide render the shape is cut
    from used to go into a temporary folder that the bridge then refused as
    outside the allowed roots, so a sandboxed server could never export a
    shape. The render goes into the server's own scratch folder, which no
    caller names; the deck and the output still have to be inside the roots,
    and the scratch folder is gone afterwards."""
    inside = tmp_path / "inside"
    inside.mkdir()
    deck = _deck(inside / "deck.pptx")
    _shape(deck, "Box", 2, 1, 4, 2, line="none")
    monkeypatch.setenv("KS4P_ALLOWED_ROOTS", str(inside))
    out = inside / "box.png"
    payload = _ok("export_shape_image", {
        "file_path": deck, "slide": 0, "shape": "Box", "output": str(out), "width": 400})
    with Image.open(out) as image:
        assert image.size == (400, 200)
    assert payload["file"] == str(out)
    (rendered,) = fake_powerpoint
    scratch = Path(rendered).parent
    assert not scratch.exists(), "the scratch folder outlived the call"
    assert inside not in scratch.parents, "the render went into the caller's folder"
    assert sorted(p.name for p in inside.iterdir() if not p.name.startswith(".")) == [
        "box.png", "deck.pptx"]


def test_a_sandboxed_server_still_refuses_paths_a_caller_names(tmp_path, monkeypatch,
                                                               fake_powerpoint):
    """The scratch folder is admitted as a ServerScratch object only, which a
    call's JSON arguments cannot produce: a folder a caller names outside the
    roots is still refused, by export_slide_image and by the bridge itself,
    and so is a shape image written outside them."""
    inside, outside = tmp_path / "inside", tmp_path / "outside"
    inside.mkdir()
    outside.mkdir()
    deck = _deck(inside / "deck.pptx")
    _shape(deck, "Box", 2, 1, 4, 2, line="none")
    monkeypatch.setenv("KS4P_ALLOWED_ROOTS", str(inside))
    payload, is_error = _call("export_slide_image", {
        "file_path": deck, "output_dir": str(outside), "engine": "com"})
    assert is_error and "outside the allowed roots" in payload["error"]["message"]
    payload, is_error = _call("export_shape_image", {
        "file_path": deck, "slide": 0, "shape": "Box", "output": str(outside / "x.png")})
    assert is_error and "outside the allowed roots" in payload["error"]["message"]
    from kitchensink4ppt.core.sandbox import SandboxViolation
    with pytest.raises(SandboxViolation):
        bridge.com_export_slide_images(deck, output_dir=str(outside), slides=[0])
    assert not fake_powerpoint and not list(outside.iterdir())


@pytest.mark.live
def test_live_render_of_a_shape(tmp_path):
    """Real PowerPoint, only when none is running (another session's
    PowerPoint is never touched): the PNG has the width asked for."""
    if sys.platform != "win32" or not ex._com_available():
        pytest.skip("needs PowerPoint on Windows")
    running = bridge.powerpnt_pids()
    if running is None or running:
        pytest.skip("PowerPoint is running or the process table cannot be read")
    deck = _deck(tmp_path / "deck.pptx")
    _shape(deck, "Box", 2, 1, 4, 2, fill="#C00000", line="none")
    out = tmp_path / "live.png"
    payload = _ok("export_shape_image", {
        "file_path": deck, "slide": 0, "shape": "Box", "output": str(out), "width": 300})
    with Image.open(out) as image:
        assert image.size == (300, 150)
        red = image.convert("RGB").getpixel((150, 75))
        assert red[0] > 150 and red[1] < 80 and red[2] < 80, red
    assert payload["engine"] == "powerpoint-com"
