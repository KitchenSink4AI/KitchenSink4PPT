"""Reads of a deck that is open in PowerPoint.

Field report: get_presentation_view and get_slide_info on a deck open in
PowerPoint were refused DOCUMENT_LOCKED, the same refusal an edit gets,
although a read changes nothing. The cause: every load probed the deck by
opening it for writing, and PowerPoint denies other programs write access
to a deck it has open while it shares read access (measured in real
PowerPoint with a made-up deck: read open allowed, write open denied, the
file holding what was last saved).

Now a read reads the file as last saved and says so in open_deck_note; the
two reads with a live route read the open copy under live='auto' and fall
back to the saved file when it cannot be reached; an edit is refused as
before; and a hold that is not PowerPoint's (another program's save) is
still looked at again and reported as busy.

The hold is simulated the way PowerPoint holds a deck: its owner file
~$name beside the deck, and the deck open with write sharing denied and
read sharing allowed. No PowerPoint is started.
"""

from __future__ import annotations

import asyncio
import ctypes
import hashlib
import sys
from contextlib import contextmanager
from ctypes import wintypes
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt import server
from kitchensink4ppt.core import errors as _err
from kitchensink4ppt.registry import public_ops
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.plan import build

pytestmark = pytest.mark.skipif(sys.platform != "win32", reason="Windows share modes")

GENERIC_READ, GENERIC_WRITE = 0x80000000, 0x40000000
FILE_SHARE_READ, OPEN_EXISTING, NORMAL = 0x1, 3, 0x80


def _call(tool: str, op: str, args: dict):
    plan = build(Settings(), public_ops())
    return asyncio.run(dispatch(plan, tool, {"op": op, "args": args}))


def _md5(path) -> str:
    return hashlib.md5(Path(path).read_bytes()).hexdigest()


@contextmanager
def powerpoint_style_hold(deck: Path, share: int = FILE_SHARE_READ):
    """The deck held as PowerPoint holds an open deck: the owner file beside
    it, and a handle that denies others write access (share=0 denies read
    access too)."""
    k32 = ctypes.WinDLL("kernel32", use_last_error=True)
    k32.CreateFileW.restype = wintypes.HANDLE
    k32.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD,
                                ctypes.c_void_p, wintypes.DWORD, wintypes.DWORD,
                                wintypes.HANDLE]
    owner = deck.with_name("~$" + deck.name)
    owner.write_bytes(bytes(165))
    handle = k32.CreateFileW(str(deck), GENERIC_READ | GENERIC_WRITE, share, None,
                             OPEN_EXISTING, NORMAL, None)
    assert handle and handle != wintypes.HANDLE(-1).value
    try:
        yield
    finally:
        k32.CloseHandle(handle)
        owner.unlink()


@pytest.fixture
def deck(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Quarter review"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Region"
    slide.notes_slide.notes_text_frame.text = "Say the totals first."
    other = prs.slides.add_slide(prs.slide_layouts[6])
    other.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = (
        "Second slide text")
    path = tmp_path / "held.pptx"
    prs.save(str(path))
    return path


@pytest.mark.parametrize("tool,op,args", [
    ("ppt_read", "get_presentation_view", {}),
    ("ppt_read", "find_text", {"query": "Second"}),
    ("ppt_read", "get_presentation_info", {}),
    ("ppt_read", "get_notes", {"slide_number": 1}),
    ("ppt_read", "list_elements", {"kind": "shapes"}),
    ("ppt_read", "get_text", {"live": "off"}),
    ("ppt_read", "get_slide_info", {"slide_number": 2, "live": "off"}),
    ("ppt_inspect", "deck_statistics", {}),
    ("ppt_inspect", "check_layout", {}),
])
def test_a_read_reads_the_saved_file_and_says_so(deck, tool, op, args):
    with powerpoint_style_hold(deck):
        payload, is_error = _call(tool, op, {"file_path": str(deck), **args})
    assert is_error is False, payload
    note = payload["open_deck_note"]
    assert "held.pptx is open in PowerPoint" in note
    assert "as last saved" in note and "not included" in note


def test_a_mutating_operations_read_path_says_so_too(deck):
    """export_table writes no deck: it reads the table and answers inline."""
    with powerpoint_style_hold(deck):
        payload, is_error = _call("ppt_deck", "export_table", {
            "file_path": str(deck), "slide_number": 1, "table": {"index": 0},
            "format": "json"})
    assert is_error is False, payload
    assert "open in PowerPoint" in payload["open_deck_note"]


def test_live_auto_reads_the_open_copy_when_it_can(deck, monkeypatch):
    called = []
    monkeypatch.setattr(server._lo, "live_get_text",
                        lambda *a, **k: called.append(a) or {"slides": [], "live": True})
    with powerpoint_style_hold(deck):
        payload, is_error = _call("ppt_read", "get_text",
                                  {"file_path": str(deck), "live": "auto"})
    assert is_error is False, payload
    assert len(called) == 1
    assert "open_deck_note" not in payload


def test_live_auto_falls_back_to_the_saved_file(deck, monkeypatch):
    def unreachable(*_a, **_k):
        raise _err.DocumentNotOpenInPowerPoint("held.pptx is not open in this PowerPoint")

    monkeypatch.setattr(server._lo, "live_get_text", unreachable)
    with powerpoint_style_hold(deck):
        payload, is_error = _call("ppt_read", "get_text",
                                  {"file_path": str(deck), "live": "auto"})
    assert is_error is False, payload
    assert "Second slide text" in str(payload)
    note = payload["open_deck_note"]
    assert "could not be read live" in note and "not open in this PowerPoint" in note


def test_an_edit_is_still_refused_while_powerpoint_holds_the_deck(deck):
    before = _md5(deck)
    with powerpoint_style_hold(deck):
        payload, is_error = _call("ppt_edit", "set_notes", {
            "file_path": str(deck), "slide": 0, "text": "Changed", "live": "off"})
    assert is_error is True
    assert payload["error"]["code"] == "DOCUMENT_LOCKED"
    assert _md5(deck) == before


def test_a_powerpoint_that_refuses_reads_too_is_reported_busy(deck):
    """PowerPoint shares read access, but a hold that refuses it (as during
    its own save) is looked at again and then reported as busy, naming
    PowerPoint, never as a reason to close the deck."""
    with powerpoint_style_hold(deck, share=0):
        payload, is_error = _call("ppt_read", "get_presentation_info",
                                  {"file_path": str(deck)})
    assert is_error is True
    message = payload["error"]["message"]
    assert "open in PowerPoint" in message and "right now" in message
    assert "Close the file" not in message


def test_a_stale_owner_file_changes_nothing(deck):
    deck.with_name("~$" + deck.name).write_bytes(bytes(165))
    payload, is_error = _call("ppt_read", "get_presentation_view", {"file_path": str(deck)})
    assert is_error is False, payload
    assert "open_deck_note" not in payload


def test_a_deck_nobody_holds_has_no_note(deck):
    payload, is_error = _call("ppt_read", "find_text", {"file_path": str(deck),
                                                        "query": "Second"})
    assert is_error is False, payload
    assert "open_deck_note" not in payload


def test_validate_checks_the_saved_file_of_a_held_deck(deck, monkeypatch):
    """validate never writes the deck. Its payload check read a held deck
    with edit intent and reported PowerPoint's hold as an invalid payload;
    it now reads the file as last saved, as the reads do, and says so."""
    from kitchensink4ppt.com import bridge

    monkeypatch.setattr(bridge, "com_validate_opens_clean",
                        lambda fp, **_kw: {"opens_clean": True})
    with powerpoint_style_hold(deck):
        payload, is_error = _call(public_ops()["validate"].group, "validate",
                                  {"file_path": str(deck)})
    assert is_error is False, payload
    assert payload["payload_valid"] is True, payload
    assert "open in PowerPoint" in payload["open_deck_note"]
