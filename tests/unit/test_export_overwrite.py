"""Exports never replace a file unless the call says so.

export_slide_image wrote over slideN.png in silence and export_table over
its CSV or JSON, so ppt_help marked every export destructive, which made a
strict client ask before a render that changes nothing in the deck. Now
every export refuses a file that already exists, before PowerPoint starts,
listing what is there and naming overwrite=true or another output; with
overwrite=true the file is replaced. The exports are therefore not
destructive: a file is replaced only on the caller's explicit
overwrite=true. Output paths that name the deck itself or its backups are
refused even then.
"""

from __future__ import annotations

import asyncio
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt import server  # noqa: F401  (the import registers ops)
from kitchensink4ppt.com import bridge
from kitchensink4ppt.core import safesave
from kitchensink4ppt.core import tool_annotations as ann
from kitchensink4ppt.ops import export as ex
from kitchensink4ppt.registry import public_ops
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.help import help_payload
from kitchensink4ppt.surface.plan import build

EXPORTS = ("export_pdf", "export_handout", "export_slide_image", "export_table")


def _call(op: str, args: dict):
    plan = build(Settings(), public_ops())
    return asyncio.run(dispatch(plan, plan.specs[op].group, {"op": op, "args": args}))


def _deck(path: Path, slides: int = 2) -> str:
    prs = Presentation()
    for i in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = f"Slide {i + 1}"
    table = prs.slides[0].shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1))
    for r in range(2):
        for c in range(2):
            table.table.cell(r, c).text = f"{r}{c}"
    prs.save(str(path))
    return str(path)


@pytest.fixture
def fake_com(monkeypatch):
    """PowerPoint present but never started: each bridge export records its
    call and writes what PowerPoint would."""
    calls: list[tuple[str, dict]] = []

    def images(path, output_dir=None, slides=None, width=1280, height=None, **kw):
        calls.append(("images", {"slides": slides, **kw}))
        out = Path(output_dir)
        out.mkdir(parents=True, exist_ok=True)
        wanted = slides if slides is not None else range(2)
        for index in wanted:
            (out / f"slide{index + 1}.png").write_bytes(b"new")
        return {"images": [{"slide": i, "file": str(out / f"slide{i + 1}.png")}
                           for i in wanted], "engine": "powerpoint-com"}

    def pdf(path, output=None, overwrite=False, **kw):
        calls.append(("pdf", {"overwrite": overwrite}))
        Path(output).write_bytes(b"%PDF-new")
        return {"pdf": output, "engine": "powerpoint-com"}

    def handout(path, output=None, slides_per_page=3, include_notes=False,
                overwrite=False, **kw):
        calls.append(("handout", {"overwrite": overwrite}))
        Path(output).write_bytes(b"%PDF-new")
        return {"pdf": output, "engine": "powerpoint-com"}

    monkeypatch.setattr(ex, "_com_available", lambda: True)
    monkeypatch.setattr(bridge, "com_export_slide_images", images)
    monkeypatch.setattr(bridge, "com_export_pdf", pdf)
    monkeypatch.setattr(bridge, "com_export_handout", handout)
    return calls


def test_a_rerender_into_the_same_folder_needs_overwrite(tmp_path, fake_com):
    deck = _deck(tmp_path / "deck.pptx")
    out = tmp_path / "renders"
    payload, is_error = _call("export_slide_image", {
        "file_path": deck, "output_dir": str(out), "engine": "com"})
    assert is_error is False, payload
    (out / "slide1.png").write_bytes(b"old")
    payload, is_error = _call("export_slide_image", {
        "file_path": deck, "output_dir": str(out), "engine": "com"})
    assert is_error is True
    message = payload["error"]["message"]
    assert "slide1.png" in message and "overwrite=true" in message
    assert len(fake_com) == 1, "PowerPoint was asked to render after the refusal"
    assert (out / "slide1.png").read_bytes() == b"old"
    payload, is_error = _call("export_slide_image", {
        "file_path": deck, "output_dir": str(out), "engine": "com", "overwrite": True})
    assert is_error is False, payload
    assert (out / "slide1.png").read_bytes() == b"new"


def test_only_the_slides_asked_for_count(tmp_path, fake_com):
    deck = _deck(tmp_path / "deck.pptx")
    out = tmp_path / "renders"
    out.mkdir()
    (out / "slide1.png").write_bytes(b"old")
    (out / "slide9.png").write_bytes(b"other deck")
    payload, is_error = _call("export_slide_image", {
        "file_path": deck, "output_dir": str(out), "engine": "com", "slide_numbers": [2]})
    assert is_error is False, payload
    assert (out / "slide1.png").read_bytes() == b"old"
    payload, is_error = _call("export_slide_image", {
        "file_path": deck, "output_dir": str(out), "engine": "com"})
    assert is_error is True and "slide1.png" in payload["error"]["message"]
    assert "slide9.png" not in payload["error"]["message"]  # the deck has 2 slides


@pytest.mark.parametrize("op, key, name", [
    ("export_pdf", "output", "out.pdf"),
    ("export_handout", "output", "handout.pdf"),
])
def test_a_pdf_is_replaced_only_with_overwrite(tmp_path, fake_com, op, key, name):
    deck = _deck(tmp_path / "deck.pptx")
    target = tmp_path / name
    target.write_bytes(b"%PDF-old")
    args = {"file_path": deck, key: str(target)}
    if op == "export_pdf":
        args["engine"] = "com"
    payload, is_error = _call(op, args)
    assert is_error is True
    assert "already exists" in payload["error"]["message"]
    assert "overwrite=true" in payload["error"]["message"]
    assert not fake_com and target.read_bytes() == b"%PDF-old"
    payload, is_error = _call(op, {**args, "overwrite": True})
    assert is_error is False, payload
    assert fake_com[-1][1]["overwrite"] is True
    assert target.read_bytes() == b"%PDF-new"


def test_the_powerpoint_pdf_path_settles_with_overwrite(tmp_path):
    source = Path(_deck(tmp_path / "deck.pptx"))
    out = tmp_path / "out.pdf"
    out.write_bytes(b"%PDF-old")
    with pytest.raises(Exception, match="already exists"):
        bridge._settled_pdf_output(source, str(out), source.with_suffix(".pdf"))
    assert bridge._settled_pdf_output(
        source, str(out), source.with_suffix(".pdf"), overwrite=True) == out
    with pytest.raises(Exception, match="already exists"):
        ex._pdf_output(source, str(out), purpose="PDF export")
    assert ex._pdf_output(source, str(out), purpose="PDF export", overwrite=True) == out


def test_a_table_export_is_replaced_only_with_overwrite(tmp_path):
    deck = _deck(tmp_path / "deck.pptx")
    table = next(s["id"] for s in server.get_slide_info(deck, 0, live="off")["shapes"]
                 if "table" in (s.get("type") or "").lower())
    target = tmp_path / "table.csv"
    target.write_text("old", encoding="utf-8")
    args = {"file_path": deck, "slide": 0, "table": table, "path": str(target)}
    payload, is_error = _call("export_table", args)
    assert is_error is True and "overwrite=true" in payload["error"]["message"]
    assert target.read_text(encoding="utf-8") == "old"
    payload, is_error = _call("export_table", {**args, "overwrite": True})
    assert is_error is False, payload
    assert "00" in target.read_text(encoding="utf-8-sig")


def test_overwrite_never_reaches_the_deck_or_its_backups(tmp_path):
    deck = Path(_deck(tmp_path / "guarded.pptx"))
    server.set_slide_hidden(str(deck), 0, True)
    slot = safesave.slot_dir(deck) / safesave.PREV_SLOT
    before = (deck.read_bytes(), slot.read_bytes())
    table = next(s["id"] for s in server.get_slide_info(str(deck), 0, live="off")["shapes"]
                 if "table" in (s.get("type") or "").lower())
    for op, key, extra in (
        ("export_pdf", "output", {"engine": "libreoffice"}),
        ("export_handout", "output", {}),
        ("export_table", "path", {"slide": 0, "table": table, "format": "csv"}),
    ):
        for target in (deck, slot):
            payload, is_error = _call(op, {
                "file_path": str(deck), key: str(target), "overwrite": True, **extra})
            assert is_error, (op, target, payload)
    payload, is_error = _call("export_slide_image", {
        "file_path": str(deck), "output_dir": str(slot.parent), "overwrite": True,
        "engine": "libreoffice"})
    assert is_error
    assert (deck.read_bytes(), slot.read_bytes()) == before


def test_exports_are_not_destructive_and_help_says_so():
    plan = build(Settings(), public_ops())
    for op in EXPORTS:
        assert ann.destructive_hint(op) is False, op
        detail = help_payload(plan, op)
        assert "overwrite" in detail["args_schema"]["properties"], op
        assert detail["args_schema"]["properties"]["overwrite"]["default"] is False
        flags = {key: detail.get(key) for key in ("writes", "destructive")}
        assert flags.get("destructive") in (False, None), (op, flags)


def test_the_render_workflow_says_a_rerender_needs_overwrite():
    text = str(server.get_workflows("render-and-review"))
    assert "overwrite=true" in text
