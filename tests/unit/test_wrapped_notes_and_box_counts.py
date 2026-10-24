"""Counts agree with what PowerPoint shows when it keeps text in a wrapper.

PowerPoint keeps a notes body that holds an equation in a compatibility
wrapper on the notes page: the body it shows, and a Fallback copy for older
readers. find_text and search_and_replace walked every body placeholder in
the notes part, so a word in such notes was found twice and a replace
counted it twice (notes_count 2 for one match). Both now count the body
PowerPoint shows; the replace still writes both copies, so they stay in
step.

A text box PowerPoint saved with an equation sits in a slide-level wrapper.
find_text finds its words, but search_and_replace walked past it and
answered matches_found 0 with no word about it. It now counts those matches
in wrapped_matches_skipped and its warning, as it does for a match inside an
equation, and leaves the box as it is (edits do not rewrite a wrapper).

The wrapped forms follow PowerPoint's own save of a made-up deck. Every deck
here is made up by code.
"""

from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.package import PptxPackage
from kitchensink4ppt.ops import read
from kitchensink4ppt.ops import text as tx

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A14 = "http://schemas.microsoft.com/office/drawing/2010/main"
OMML = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _wrap(element: etree._Element, choice_extra=None) -> None:
    """Put `element` in an mc:AlternateContent where it stands: the element
    (given its equation by `choice_extra`) in the Choice, which Requires a14,
    and a copy without one in the Fallback."""
    wrapper = etree.Element(f"{{{MC}}}AlternateContent", nsmap={"mc": MC})
    choice = etree.SubElement(wrapper, f"{{{MC}}}Choice", nsmap={"a14": A14})
    choice.set("Requires", "a14")
    fallback = etree.SubElement(wrapper, f"{{{MC}}}Fallback")
    element.addprevious(wrapper)
    fallback.append(etree.fromstring(etree.tostring(element)))
    if choice_extra is not None:
        choice_extra(element)
    choice.append(element)


def _math(chars: str) -> etree._Element:
    return etree.fromstring(
        f'<a:p xmlns:a="{A}" xmlns:a14="{A14}"><a14:m><m:oMathPara xmlns:m="{OMML}"><m:oMath>'
        f"<m:r><m:t>{chars}</m:t></m:r></m:oMath></m:oMathPara></a14:m><a:endParaRPr/></a:p>")


def _deck(path: Path) -> Path:
    """Slide 1: a plain box "Budget words", and a box "Budget box" with an
    equation that PowerPoint keeps in a slide-level wrapper. Its notes:
    "Budget notes" and an equation, the body in a notes-level wrapper."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    plain = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    plain.name = "Plain box"
    plain.text_frame.text = "Budget words"
    boxed = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1))
    boxed.name = "Budget box"
    boxed.text_frame.text = "Budget box words"
    slide.notes_slide.notes_text_frame.text = "Budget notes"
    prs.save(str(path))

    def fix(name, data):
        root = etree.fromstring(data)
        if name == "ppt/slides/slide1.xml":
            sp = next(e for e in root.iter(f"{{{P}}}sp")
                      if next(e.iter(f"{{{P}}}cNvPr")).get("name") == "Budget box")
            _wrap(sp, lambda el: el.find(f"{{{P}}}txBody").append(_math("x=1")))
        else:
            body = next(sp for sp in root.iter(f"{{{P}}}sp")
                        if sp.find(f".//{{{P}}}ph") is not None
                        and sp.find(f".//{{{P}}}ph").get("type") == "body")
            _wrap(body, lambda el: el.find(f"{{{P}}}txBody").append(_math("y=2")))
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    tmp = path.with_name(path.name + ".rw")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        notes_parts = [n for n in zin.namelist()
                       if n.startswith("ppt/notesSlides/") and n.endswith(".xml")]
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "ppt/slides/slide1.xml" or info.filename in notes_parts:
                data = fix(info.filename, data)
            zout.writestr(info, data)
    tmp.replace(path)
    return path


def _part_text(path: Path, prefix: str) -> str:
    with zipfile.ZipFile(path) as z:
        name = next(n for n in z.namelist() if n.startswith(prefix) and n.endswith(".xml"))
        return z.read(name).decode("utf-8")


def test_find_text_counts_a_word_in_wrapped_notes_once(tmp_path):
    path = _deck(tmp_path / "made_up.pptx")
    out = read.find_text(PptxPackage(path), "notes")
    assert out["count"] == 1, out["matches"]
    assert out["matches"][0]["where"] == "notes"


def test_get_text_reads_the_notes_body_powerpoint_shows(tmp_path):
    path = _deck(tmp_path / "made_up.pptx")
    pkg = PptxPackage(path)
    assert read.notes_text(pkg, pkg.slide_parts()[0]) == "Budget notes\ny=2"


def test_search_and_replace_counts_a_wrapped_notes_match_once_and_keeps_both_copies(tmp_path):
    path = _deck(tmp_path / "made_up.pptx")
    pkg = PptxPackage(path)
    out = tx.search_and_replace(pkg, "Budget notes", "Plan notes", include_notes=True)
    pkg.save(do_backup=False)
    assert out["total"] == 1, out
    assert out["slides"][0]["notes_count"] == 1, out
    notes = _part_text(path, "ppt/notesSlides/")
    assert notes.count("Plan notes") == 2 and "Budget notes" not in notes


def test_search_and_replace_counts_matches_in_a_wrapped_box_as_skipped(tmp_path):
    """find_text finds the word in the box PowerPoint saved with an
    equation; the replace leaves the box as it is and says so."""
    path = _deck(tmp_path / "made_up.pptx")
    found = read.find_text(PptxPackage(path), "Budget", include_notes=False)
    assert found["count"] == 2
    before = hashlib.sha256(_part_text(path, "ppt/slides/slide1").encode()).hexdigest()
    pkg = PptxPackage(path)
    out = tx.search_and_replace(pkg, "Budget", "Plan")
    pkg.save(do_backup=False)
    assert out["total"] == 1, out
    assert out["matches_found"] == found["count"], out
    assert out["wrapped_matches_skipped"] == 1, out
    assert out["slides"][0]["wrapped_matches_skipped"] == 1, out
    assert any("compatibility wrapper" in w and "slide_number 1" in w for w in out["warnings"]), out
    slide = _part_text(path, "ppt/slides/slide1")
    assert "Plan words" in slide and slide.count("Budget box words") == 2
    assert hashlib.sha256(slide.encode()).hexdigest() != before
