"""Shapes inside a group are edited where they are, by name or by id.

A figure built as one group was recoloured, recaptioned and extended by hand
because nothing reached its members by name: set_shape took only an id, and
nothing put a new shape inside an existing group. set_shape now takes a name
(inside groups too, with "Group/Shape" when a name repeats), format_text
takes the same path, insert_shape takes the group to insert into, and a group
whose members moved or grew is refitted the way PowerPoint refits it: the
group's frame (off/ext with chOff/chExt) encloses its members again and
nothing moves on the slide.
"""

from __future__ import annotations

import asyncio
from pathlib import Path

import pytest
from pptx import Presentation

from kitchensink4ppt import server  # noqa: F401  (the import registers ops)
from kitchensink4ppt.core.package import PptxPackage, qn
from kitchensink4ppt.ops import shapes as shp
from kitchensink4ppt.registry import public_ops
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.help import help_payload
from kitchensink4ppt.surface.plan import build

EMU = 914400


def _call(tool: str, op: str, args: dict):
    plan = build(Settings(), public_ops())
    return asyncio.run(dispatch(plan, tool, {"op": op, "args": args}))


def _result(payload: dict) -> dict:
    inner = payload.get("changed")
    return inner if isinstance(inner, dict) else payload


def _ok(tool: str, op: str, args: dict) -> dict:
    payload, is_error = _call(tool, op, args)
    assert is_error is False, payload
    return _result(payload)


def _insert(deck, name, x, y, w=1.0, h=0.5, shape_type="rect", **extra):
    return _ok("ppt_graphics", "insert_shape", {
        "file_path": deck, "slide": 0, "shape_type": shape_type, "name": name,
        "x": x, "y": y, "w": w, "h": h, **extra})["shape_id"]


@pytest.fixture
def figure(tmp_path) -> dict:
    """A slide with two groups, "Figure" and "Legend", each holding a shape
    named "Label"; Figure also holds "Band A" and a text box "Caption"."""
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    deck = str(tmp_path / "figure.pptx")
    prs.save(deck)
    band = _insert(deck, "Band A", 1, 1, 2, 1)
    label = _insert(deck, "Label", 1, 2.5, 2, 0.5, text="Model")
    caption = _ok("ppt_edit", "insert_textbox", {
        "file_path": deck, "slide": 0, "text": "Plain and bold",
        "x": 1, "y": 3.2, "w": 3, "h": 0.5})
    caption_id = caption.get("shape_id") or caption.get("id")
    _ok("ppt_graphics", "set_shape", {
        "file_path": deck, "slide": 0, "shape": caption_id, "name": "Caption"})
    fig = _ok("ppt_graphics", "group_shapes", {
        "file_path": deck, "slide": 0, "ids": [band, label, caption_id],
        "name": "Figure"})["group_id"]
    key = _insert(deck, "Key", 6, 1, 1, 1)
    key_label = _insert(deck, "Label", 6, 2.5, 1, 0.5, text="Key")
    legend = _ok("ppt_graphics", "group_shapes", {
        "file_path": deck, "slide": 0, "ids": [key, key_label],
        "name": "Legend"})["group_id"]
    return {"deck": deck, "band": band, "label": label, "caption": caption_id,
            "figure": fig, "key_label": key_label, "legend": legend}


def _elem(deck: str, ref):
    pkg = PptxPackage(deck)
    part = pkg.slide_parts()[0]
    return shp._find_shape(pkg, part, ref)


def _box(deck: str, ref) -> tuple[float, ...]:
    elem, chain = _elem(deck, ref)
    return tuple(round(v / EMU, 4) for v in shp._slide_box(elem, chain))


def _frame(deck: str, ref) -> dict:
    elem, _chain = _elem(deck, ref)
    xfrm = elem.find(f"{qn('p:grpSpPr')}/{qn('a:xfrm')}")
    out = {}
    for tag in ("a:off", "a:ext", "a:chOff", "a:chExt"):
        node = xfrm.find(qn(tag))
        out[tag[2:]] = tuple(int(v) for v in node.attrib.values())
    return out


def test_a_member_is_recoloured_by_its_name(figure):
    deck = figure["deck"]
    result = _ok("ppt_graphics", "set_shape", {
        "file_path": deck, "slide": 0, "shape": "Band A",
        "fill": "#C00000", "line": {"color": "#002060", "width": 2}})
    assert result["shape_id"] == figure["band"]
    elem, chain = _elem(deck, figure["band"])
    assert [shp._shape_id(g) for g in chain] == [figure["figure"]]
    sppr = elem.find(qn("p:spPr"))
    assert sppr.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}").get("val") == "C00000"
    assert sppr.find(qn("a:ln")).get("w") == str(2 * 12700)


def test_a_repeated_name_is_refused_with_every_candidate(figure):
    payload, is_error = _call("ppt_graphics", "set_shape", {
        "file_path": figure["deck"], "slide": 0, "shape": "Label",
        "fill": "#FF0000"})
    assert is_error is True
    assert payload["error"]["code"] == "AMBIGUOUS_LOCATION"
    message = payload["error"]["message"]
    assert str(figure["label"]) in message and str(figure["key_label"]) in message
    assert "Group/Shape" in message


def test_group_slash_name_picks_the_member_of_that_group(figure):
    result = _ok("ppt_graphics", "set_shape", {
        "file_path": figure["deck"], "slide": 0, "shape": "Legend/Label",
        "fill": "#00B050"})
    assert result["shape_id"] == figure["key_label"]


def test_an_unknown_name_lists_the_shapes_with_their_groups(figure):
    payload, is_error = _call("ppt_graphics", "set_shape", {
        "file_path": figure["deck"], "slide": 0, "shape": "Band Z",
        "fill": "#FF0000"})
    assert is_error is True
    assert payload["error"]["code"] == "NOT_FOUND"
    assert "'Band A' in 'Figure'" in payload["error"]["message"]


def test_runs_in_a_member_text_box_are_formatted_through_the_group_path(figure):
    deck = figure["deck"]
    result = _ok("ppt_edit", "format_text", {
        "file_path": deck, "slide": 0, "shape": "Figure/Caption",
        "paragraph": 0, "start": 10, "end": 14, "bold": True, "color": "#C00000", "live": "off"})
    assert result
    elem, _chain = _elem(deck, figure["caption"])
    runs = [(r.findtext(qn("a:t")), r.find(qn("a:rPr"))) for r in elem.iter(qn("a:r"))]
    bold = [text for text, rpr in runs if rpr is not None and rpr.get("b") == "1"]
    assert bold == ["bold"]


def test_a_new_shape_lands_inside_the_group_where_it_was_asked(figure):
    deck = figure["deck"]
    before = {ref: _box(deck, figure[ref]) for ref in ("band", "label", "caption")}
    new = _insert(deck, "Badge", 1.5, 1.2, 0.8, 0.4, group="Figure", text="New")
    elem, chain = _elem(deck, new)
    assert [shp._shape_id(g) for g in chain] == [figure["figure"]]
    assert _box(deck, new) == (1.5, 1.2, 0.8, 0.4)
    for ref, box in before.items():
        assert _box(deck, figure[ref]) == box


def test_a_shape_inserted_outside_the_frame_grows_the_group(figure):
    deck = figure["deck"]
    before = {ref: _box(deck, figure[ref]) for ref in ("band", "label", "caption")}
    new = _insert(deck, "Note", 0.2, 4.5, 1.5, 0.6, group=figure["figure"])
    assert _box(deck, new) == (0.2, 4.5, 1.5, 0.6)
    for ref, box in before.items():
        assert _box(deck, figure[ref]) == box
    frame = _frame(deck, figure["figure"])
    left, top = frame["off"]
    width, height = frame["ext"]
    assert left == round(0.2 * EMU) and top == round(1 * EMU)
    assert left + width == round(4 * EMU) and top + height == round(5.1 * EMU)


def test_inserting_into_a_scaled_group_maps_through_its_scale(figure):
    deck = figure["deck"]
    _ok("ppt_graphics", "set_shape", {
        "file_path": deck, "slide": 0, "shape": "Figure", "w": 1.5, "h": 1.15})
    band = _box(deck, figure["band"])
    new = _insert(deck, "Dot", 1.1, 1.1, 0.2, 0.2, shape_type="ellipse",
                  group="Figure")
    assert _box(deck, new) == (1.1, 1.1, 0.2, 0.2)
    assert _box(deck, figure["band"]) == band


def test_moving_a_member_out_of_the_frame_refits_the_group(figure):
    deck = figure["deck"]
    label = _box(deck, figure["label"])
    _ok("ppt_graphics", "set_shape", {
        "file_path": deck, "slide": 0, "shape": "Band A", "x": 0.5, "y": 0.4})
    assert _box(deck, figure["band"])[:2] == (0.5, 0.4)
    assert _box(deck, figure["label"]) == label
    frame = _frame(deck, figure["figure"])
    assert frame["off"] == (round(0.5 * EMU), round(0.4 * EMU))


def test_inserting_into_a_shape_that_is_not_a_group_is_refused(figure):
    payload, is_error = _call("ppt_graphics", "insert_shape", {
        "file_path": figure["deck"], "slide": 0, "shape_type": "rect",
        "x": 1, "y": 1, "w": 1, "h": 1, "group": "Band A"})
    assert is_error is True
    assert "not a group" in payload["error"]["message"]


def test_inserting_into_a_rotated_group_is_refused(figure):
    deck = figure["deck"]
    _ok("ppt_graphics", "set_shape", {
        "file_path": deck, "slide": 0, "shape": "Legend", "rotation": 30})
    payload, is_error = _call("ppt_graphics", "insert_shape", {
        "file_path": deck, "slide": 0, "shape_type": "rect",
        "x": 6, "y": 1, "w": 1, "h": 1, "group": "Legend"})
    assert is_error is True
    assert "rotated" in payload["error"]["message"]


def test_help_names_the_new_arguments():
    plan = build(Settings(), public_ops())
    shape_arg = help_payload(plan, "set_shape")["args_schema"]["properties"]["shape"]
    assert "Group/Shape" in shape_arg["description"]
    group_arg = help_payload(plan, "insert_shape")["args_schema"]["properties"]["group"]
    assert "group" in group_arg["description"].lower()
