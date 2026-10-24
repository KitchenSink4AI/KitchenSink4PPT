"""remove_empty_paragraphs: the help says where paragraphs are removed, and the
behaviour is held to it.

The help said emptied paragraphs are removed "in slide text, table cells and
notes", but notes are searched only when include_notes is true (its default
is false), so a notes paragraph stayed until include_notes was sent; and
SmartArt text, which the replacement does reach, keeps its paragraphs. The
help now says both, and each claim below is checked against a made-up deck.
"""

from __future__ import annotations

import asyncio

from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt import server  # noqa: F401  (the import registers ops)
from kitchensink4ppt.core.package import PptxPackage
from kitchensink4ppt.ops import read, slides
from kitchensink4ppt.registry import public_ops
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.help import help_payload
from kitchensink4ppt.surface.plan import build

from test_diagrams_write import add_smartart


def _plan():
    return build(Settings(), public_ops())


def _replace(path, **extra) -> dict:
    args = {"file_path": str(path), "find": "Draft only", "replace": "",
            "remove_empty_paragraphs": True, "live": "off", **extra}
    payload, is_error = asyncio.run(dispatch(
        _plan(), "ppt_edit", {"op": "search_and_replace", "args": args}))
    assert is_error is False, payload
    return payload["changed"]


def _help() -> str:
    props = help_payload(_plan(), "search_and_replace")["args_schema"]["properties"]
    return props["remove_empty_paragraphs"]["description"]


def _deck(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    box.text_frame.text = "Keep this line\nDraft only"
    slide.notes_slide.notes_text_frame.text = "Keep this note\nDraft only"
    path = tmp_path / "blanks.pptx"
    prs.save(str(path))
    return path


def _paragraphs(path):
    prs = Presentation(str(path))
    slide = prs.slides[0]
    box = next(sh for sh in slide.shapes if sh.has_text_frame)
    body = [p.text for p in box.text_frame.paragraphs]
    notes = [p.text for p in slide.notes_slide.notes_text_frame.paragraphs]
    return body, notes


def test_the_help_names_include_notes_for_notes_and_smartart():
    desc = _help()
    assert "notes when include_notes is true" in desc
    assert "SmartArt keeps its paragraphs" in desc
    assert "table cells and notes" not in desc


def test_notes_are_left_alone_without_include_notes(tmp_path):
    path = _deck(tmp_path)
    out = _replace(path)
    assert _paragraphs(path) == (["Keep this line"], ["Keep this note", "Draft only"])
    assert out["paragraphs_removed"] == 1


def test_notes_paragraph_goes_with_include_notes(tmp_path):
    path = _deck(tmp_path)
    out = _replace(path, include_notes=True)
    assert _paragraphs(path) == (["Keep this line"], ["Keep this note"])
    assert out["paragraphs_removed"] == 2


def test_smartart_keeps_the_paragraph_a_replacement_emptied(tmp_path):
    path = tmp_path / "smartart.pptx"
    slides.create_presentation(path)
    pkg = PptxPackage(path)
    slides.insert_slide(pkg, 6)
    add_smartart(pkg, pkg.slide_parts()[0], ["Alpha", "Draft only", "Gamma"])
    pkg.save(do_backup=False)
    out = _replace(path)
    assert out["total"] >= 1  # the data model, and the cached drawing when there is one
    assert out["paragraphs_removed"] == 0
    nodes = [n["text"] for n in read.diagram_text(PptxPackage(path))["items"][0]["nodes"]]
    assert nodes == ["Alpha", "", "Gamma"]
