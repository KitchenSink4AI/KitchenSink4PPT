"""A list of slides can be counted from 1, and every slide in the answer says
its number the way people count.

A caller asked export_slide_image for "slide_numbers": [1, 2, 4] and was
refused: only the 0-based ``slides`` existed. Its answer then listed
"slide": 0 beside a file named slide1.png. Wherever an operation takes a
``slides`` list, ``slide_numbers`` (each counted from 1) is accepted in its
place, sending both is refused with a hint, and each slide in the answer
carries its slide_number beside the 0-based index. ``slides`` itself is
unchanged.
"""

from __future__ import annotations

import asyncio
from pathlib import Path

import pytest
from pptx import Presentation

from kitchensink4ppt import server  # noqa: F401  (the import registers ops)
from kitchensink4ppt.com import bridge
from kitchensink4ppt.ops import export as export_ops
from kitchensink4ppt.registry import public_ops
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface import copy
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.help import help_payload
from kitchensink4ppt.surface.plan import build

#: Every operation that takes a list of slides.
SLIDE_LIST_OPS = ("export_slide_image", "manage_custom_show")


def _deck(path: Path) -> str:
    prs = Presentation()
    for title in ("Opening", "Revenue", "Budget", "Close"):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
    prs.save(path)
    return str(path)


def _call(tool: str, op: str, args: dict):
    plan = build(Settings(), public_ops())
    return asyncio.run(dispatch(plan, tool, {"op": op, "args": args}))


@pytest.fixture
def rendered(monkeypatch):
    """PowerPoint's slide export stood in for: it writes slideN.png for each
    0-based index it is given and reports them as the bridge does."""
    asked: dict = {}

    def export(path, output_dir=None, slides=None, width=1280, height=None, **_):
        asked["slides"] = slides
        out = Path(output_dir)
        out.mkdir(parents=True, exist_ok=True)
        images = []
        for idx in (range(4) if slides is None else slides):
            png = out / f"slide{idx + 1}.png"
            png.write_bytes(b"\x89PNG\r\n\x1a\n")
            images.append({"slide": idx, "file": str(png)})
        return {"images": images, "width": width, "height": 720, "engine": "powerpoint-com"}

    monkeypatch.setattr(export_ops, "_com_available", lambda: True)
    monkeypatch.setattr(bridge, "com_export_slide_images", export)
    return asked


def test_export_slide_image_takes_slide_numbers_counted_from_1(tmp_path, rendered):
    deck = _deck(tmp_path / "deck.pptx")
    payload, is_error = _call("ppt_deck", "export_slide_image", {
        "file_path": deck, "output_dir": str(tmp_path / "png"),
        "slide_numbers": [1, 2, 4], "engine": "com"})
    assert is_error is False, payload
    assert rendered["slides"] == [0, 1, 3]
    assert [(i["slide"], i["slide_number"]) for i in payload["images"]] == [(0, 1), (1, 2), (3, 4)]
    for image in payload["images"]:
        assert Path(image["file"]).name == f"slide{image['slide_number']}.png"


def test_slides_still_count_from_0_and_the_answer_names_each_slide_number(tmp_path, rendered):
    deck = _deck(tmp_path / "deck.pptx")
    payload, is_error = _call("ppt_deck", "export_slide_image", {
        "file_path": deck, "output_dir": str(tmp_path / "png"), "slides": [0, 2],
        "engine": "com"})
    assert is_error is False, payload
    assert rendered["slides"] == [0, 2]
    assert [(i["slide"], i["slide_number"]) for i in payload["images"]] == [(0, 1), (2, 3)]


def test_both_forms_of_a_slide_list_are_refused_with_a_hint(tmp_path, rendered):
    deck = _deck(tmp_path / "deck.pptx")
    payload, is_error = _call("ppt_deck", "export_slide_image", {
        "file_path": deck, "slides": [0], "slide_numbers": [1], "engine": "com"})
    assert is_error is True and payload["error"]["code"] == "BAD_ARGS", payload
    assert "slide_numbers" in payload["error"]["hint"] and "slides" in payload["error"]["hint"]
    assert "slides" not in rendered, "nothing was exported"


@pytest.mark.parametrize("bad", [[0], [2, 0], [1.5], "1, 2", ["two"]])
def test_slide_numbers_must_be_whole_numbers_counted_from_1(tmp_path, rendered, bad):
    deck = _deck(tmp_path / "deck.pptx")
    payload, is_error = _call("ppt_deck", "export_slide_image", {
        "file_path": deck, "slide_numbers": bad, "engine": "com"})
    assert is_error is True and payload["error"]["code"] == "BAD_ARGS", payload
    assert "counted from 1" in payload["error"]["message"] + payload["error"]["hint"]
    assert "slides" not in rendered


def test_a_null_slide_numbers_is_read_as_not_sent(tmp_path, rendered):
    deck = _deck(tmp_path / "deck.pptx")
    payload, is_error = _call("ppt_deck", "export_slide_image", {
        "file_path": deck, "output_dir": str(tmp_path / "png"), "slides": [1],
        "slide_numbers": None, "engine": "com"})
    assert is_error is False, payload
    assert rendered["slides"] == [1]


def test_a_custom_show_takes_slide_numbers_and_names_each_slide_number(tmp_path):
    deck = _deck(tmp_path / "deck.pptx")
    payload, is_error = _call("ppt_deck", "manage_custom_show", {
        "file_path": deck, "action": "create", "name": "Short", "slide_numbers": [4, 2]})
    assert is_error is False, payload
    created = payload["changed"]["slides"]
    assert [(s["index"], s["slide_number"]) for s in created] == [(3, 4), (1, 2)]
    payload, is_error = _call("ppt_deck", "manage_custom_show", {
        "file_path": deck, "action": "list"})
    assert is_error is False, payload
    listed = payload["shows"][0]["slides"]
    assert [(s["index"], s["slide_number"]) for s in listed] == [(3, 4), (1, 2)]


def test_help_lists_slide_numbers_after_slides_on_every_slide_list_operation():
    plan = build(Settings(), public_ops())
    found = [name for name, spec in public_ops().items() if "slides" in spec.args.model_fields]
    assert sorted(found) == sorted(SLIDE_LIST_OPS), "an operation that takes slides was missed"
    for op in SLIDE_LIST_OPS:
        properties = help_payload(plan, op)["args_schema"]["properties"]
        names = list(properties)
        assert names[names.index("slides") + 1] == "slide_numbers", op
        assert properties["slide_numbers"] == copy.SLIDE_NUMBERS_PARAM, op
        assert "counted from 1" in properties["slide_numbers"]["description"]
        assert "slide_numbers" in properties["slides"]["description"], op
