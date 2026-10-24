"""Animations stored in a compatibility wrapper.

PowerPoint wraps a slide's p:timing tree in a slide-level mc:AlternateContent
(an mc:Choice requiring p14, and an mc:Fallback copy) when the animations use
its 2010 features, and plays the mc:Choice branch. The animation operations
looked only for a direct p:timing: list_animations answered that such a slide
has no animations, clear_animations and the prune after delete_shape left the
wrapped tree as it was (a deleted shape's effect kept pointing at it), and
add_entrance_animation wrote a second, direct timing tree beside the wrapped
one. Now reads use the branch PowerPoint plays, prunes and clears change every
branch, and adding to a wrapped tree is refused with the next step.

Every deck here is made up by code.
"""

from __future__ import annotations

import copy

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.errors import UnsupportedStructure
from kitchensink4ppt.core.package import NSMAP, PptxPackage, qn
from kitchensink4ppt.ops import animations as an
from kitchensink4ppt.ops import shapes as sh

MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
AC, CHOICE, FALLBACK = (f"{{{MC}}}{n}" for n in ("AlternateContent", "Choice", "Fallback"))


def _shape_ids(path):
    prs = Presentation(str(path))
    return [s.shape_id for s in prs.slides[0].shapes]


@pytest.fixture()
def wrapped(tmp_path):
    """One slide, two text boxes, each with an entrance effect, and the
    timing tree moved into mc:AlternateContent the way PowerPoint stores it."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for n in range(2):
        box = slide.shapes.add_textbox(Inches(1), Inches(1 + 2 * n), Inches(4), Inches(1))
        box.text_frame.text = f"Made-up line {n + 1}"
    path = tmp_path / "wrapped.pptx"
    prs.save(str(path))
    a, b = _shape_ids(path)
    pkg = PptxPackage(path)
    an.add_entrance_animation(pkg, 0, a, "fade", "click")
    an.add_entrance_animation(pkg, 0, b, "appear", "click")
    part = pkg.slide_parts()[0]
    root = pkg.root(part)
    timing = root.find(qn("p:timing"))
    wrapper = etree.Element(AC, nsmap={"mc": MC})
    choice = etree.SubElement(wrapper, CHOICE, nsmap={"p14": NSMAP["p14"]})
    choice.set("Requires", "p14")
    fallback = etree.SubElement(wrapper, FALLBACK)
    timing.addprevious(wrapper)
    choice.append(copy.deepcopy(timing))
    fallback.append(timing)
    pkg.mark_dirty(part)
    pkg.save(do_backup=False)
    return path, a, b


def _trees(path):
    root = PptxPackage(path).root(PptxPackage(path).slide_parts()[0])
    direct = root.findall(qn("p:timing"))
    wrapped = [t for w in root.findall(AC) for br in w for t in br.findall(qn("p:timing"))]
    return direct, wrapped


def _targets(timing):
    return sorted({int(t.get("spid")) for t in timing.iter(qn("p:spTgt"))})


def test_list_reads_the_wrapped_tree(wrapped):
    path, a, b = wrapped
    out = an.list_animations(PptxPackage(path), 0)
    assert out["has_timing"] is True
    assert [e["shape_id"] for e in out["effects"]] == [a, b]


def test_adding_to_a_wrapped_tree_is_refused_and_nothing_changes(wrapped):
    path, a, _b = wrapped
    before = path.read_bytes()
    pkg = PptxPackage(path)
    with pytest.raises(UnsupportedStructure) as refused:
        an.add_entrance_animation(pkg, 0, a, "appear", "click")
    assert "clear_animations" in str(refused.value)
    assert path.read_bytes() == before


def test_clear_removes_the_wrapper(wrapped):
    path, _a, _b = wrapped
    pkg = PptxPackage(path)
    out = an.clear_animations(pkg, 0)
    assert out["timing_removed"] is True and out["effects_removed"] == 2
    pkg.save(do_backup=False)
    assert _trees(path) == ([], [])
    assert PptxPackage(path).root(PptxPackage(path).slide_parts()[0]).find(AC) is None


def test_delete_shape_prunes_every_branch(wrapped):
    path, a, b = wrapped
    pkg = PptxPackage(path)
    out = sh.delete_shape(pkg, 0, a)
    assert out["timing_report"]["effects_removed"] == 1
    pkg.save(do_backup=False)
    direct, trees = _trees(path)
    assert direct == [] and len(trees) == 2
    assert [_targets(t) for t in trees] == [[b], [b]]


def test_clear_one_shape_prunes_every_branch(wrapped):
    path, a, b = wrapped
    pkg = PptxPackage(path)
    out = an.clear_animations(pkg, 0, shape=b)
    assert out["effects_removed"] == 1
    pkg.save(do_backup=False)
    _direct, trees = _trees(path)
    assert [_targets(t) for t in trees] == [[a], [a]]
