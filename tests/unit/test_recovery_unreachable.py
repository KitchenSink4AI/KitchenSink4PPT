"""A save journal entry for a deck on a share that does not answer.

A save writes its swap down before it starts, and an entry that a stopped
server left stays until its deck is reached again, for up to a week. A share
that does not answer (a VPN that is down, a server that is off) can take 20
seconds to say no to any file-system call. Such an entry must never hold up a
change to another deck, and must never hold up the start of the server: its
recovery waits in the background until the share answers.

The share here is simulated: every file-system call that names it waits and
then fails, as an unreachable share does, and is recorded. Nothing here
reaches a real network.
"""

from __future__ import annotations

import asyncio
import json
import os
import sys
import threading
import time
from pathlib import Path

import pytest
from pptx import Presentation

from kitchensink4ppt import server
from kitchensink4ppt.core import safesave, swapwait
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.plan import build

pytestmark = pytest.mark.skipif(sys.platform != "win32", reason="ReplaceFileW")

HOST = "share-that-does-not-answer.invalid"
NO_ANSWER_SECONDS = 3.0


@pytest.fixture(autouse=True)
def own_journal(tmp_path, monkeypatch):
    monkeypatch.setenv("KS4P_SAVE_JOURNAL_DIR", str(tmp_path / "journal"))
    safesave.take_notices()
    yield
    safesave.take_notices()


@pytest.fixture
def silent_share(monkeypatch):
    """Every call below that names HOST waits NO_ANSWER_SECONDS and fails, as
    a share that does not answer does; each such call is recorded."""
    touched: list[tuple[str, str]] = []
    lock = threading.Lock()

    def fail(path):
        raise OSError(53, "The network path was not found", path)

    def wrap(name, real, answer):
        def call(*args, **kwargs):
            if args and HOST.lower() in os.fspath(args[0]).lower():
                with lock:
                    touched.append((name, os.fspath(args[0])))
                time.sleep(NO_ANSWER_SECONDS)
                return answer(os.fspath(args[0]))
            return real(*args, **kwargs)
        return call

    gone = lambda path: False  # noqa: E731
    for module, name, answer in ((os, "stat", fail), (os, "lstat", fail),
                                 (os, "listdir", fail), (os, "scandir", fail),
                                 # realpath gives the path back once it gives up
                                 (os.path, "realpath", lambda path: path),
                                 (os.path, "isdir", gone), (os.path, "exists", gone),
                                 (os.path, "lexists", gone), (os.path, "samefile", fail)):
        monkeypatch.setattr(module, name, wrap(name, getattr(module, name), answer))
    return touched


def _deck(path: Path, notes: str) -> Path:
    prs = Presentation()
    for title in ("Alpha", "Bravo"):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.notes_slide.notes_text_frame.text = notes
    prs.save(path)
    return path


def _notes(path) -> str | None:
    if not os.path.lexists(path):
        return None
    return Presentation(str(path)).slides[0].notes_slide.notes_text_frame.text


def _call(op, args):
    plan = build(Settings())
    return asyncio.run(dispatch(plan, plan.specs[op].group, {"op": op, "args": args}))


def _write_entry(name: str, record: dict) -> Path:
    folder = safesave._journal_dir()
    folder.mkdir(parents=True, exist_ok=True)
    entry = folder / f"{name}.json"
    entry.write_text(json.dumps(record), encoding="utf-8")
    return entry


def _stopped_server(**fields) -> dict:
    """An entry whose server is gone (a pid no process has)."""
    return {"pid": 4000000, "token": "4000000-stopped", "pid_created": 1.0,
            "host": safesave._local_host(), "time": time.time(), "staged": [], **fields}


def _share_entry(name: str, share: str = "share") -> Path:
    deck = f"\\\\{HOST}\\{share}\\Deck.pptx"
    return _write_entry(name, _stopped_server(
        deck=deck, aside=deck + ".x", new=deck + ".ppt-mcp-tmp"))


def _left_aside(tmp_path: Path) -> Path:
    """A local deck that a save stopped between its two moves left aside."""
    deck = _deck(tmp_path / "Local.pptx", "v1")
    aside = tmp_path / swapwait.aside_name(deck.name, "0123456789ab")
    os.replace(deck, aside)
    new = _deck(tmp_path / "Local.pptx.ppt-mcp-tmp", "v2")
    _write_entry("zz-local", _stopped_server(deck=str(deck), aside=str(aside), new=str(new)))
    return deck


def _wait_for_background(seconds: float = 60.0) -> None:
    worker = getattr(safesave, "_LATER", None)
    if isinstance(worker, threading.Thread):
        worker.join(seconds)


def test_a_change_to_one_deck_never_reaches_for_another_decks_share(tmp_path, silent_share):
    _share_entry("00-share")
    deck = _deck(tmp_path / "Local.pptx", "v0")
    started = time.perf_counter()
    payload, refused = _call("set_notes", {"file_path": str(deck), "slide": 0, "text": "v1",
                                           "live": "off"})
    took = time.perf_counter() - started
    assert not refused, payload
    assert silent_share == [], "a save of a local deck asked the other deck's share"
    assert took < NO_ANSWER_SECONDS, f"the save took {took:.1f} s"
    assert _notes(deck) == "v1"


def test_the_start_puts_a_local_deck_back_without_waiting_for_a_share(tmp_path, silent_share):
    _share_entry("00-share")
    deck = _left_aside(tmp_path)
    started = time.perf_counter()
    restored = server._recover_at_start(Settings())
    took = time.perf_counter() - started
    assert restored == [str(deck)]
    assert _notes(deck) == "v1"
    assert took < 1.0, f"the start waited {took:.1f} s for a share that does not answer"
    _wait_for_background()


def test_many_entries_for_shares_cost_the_start_nothing(tmp_path, silent_share):
    for i in range(5):
        _share_entry(f"0{i}-share", share=f"share{i}")
    deck = _left_aside(tmp_path)
    started = time.perf_counter()
    server._recover_at_start(Settings())
    took = time.perf_counter() - started
    assert _notes(deck) == "v1", "a local deck after five share entries was not put back"
    assert took < 1.0, f"the start took {took:.1f} s"
    _wait_for_background()


def test_a_share_that_does_not_answer_is_tried_in_the_background_and_marked(tmp_path,
                                                                            silent_share):
    entry = _share_entry("00-share")
    server._recover_at_start(Settings())
    _wait_for_background()
    record = json.loads(entry.read_text(encoding="utf-8"))
    assert "folder_gone_since" in record, "the background run did not mark the entry"
    assert silent_share, "the share was never tried"
