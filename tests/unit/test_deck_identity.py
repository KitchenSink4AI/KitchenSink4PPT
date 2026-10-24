"""A save keeps the deck's identity: its Mark of the Web, permissions,
attributes and creation time.

An edit put the new version in place with os.replace, which gives the path a
NEW file. One ordinary edit then dropped what Windows keeps on the file itself
rather than in its bytes: the Mark of the Web (the Zone.Identifier stream, so a
downloaded deck opened in PowerPoint without Protected View from then on), the
deck's explicit permissions (the new file takes whatever its folder hands
down), its attributes and its creation time. PowerPoint's own save keeps all of
them. safesave.promote now uses ReplaceFileW, which keeps them, for the save,
the backup restore and the overwrite of an existing destination. Where the
file system has no ReplaceFile, a plain replace puts the Mark of the Web, the
creation time and the attributes back by hand.
"""

from __future__ import annotations

import ctypes
import os
import subprocess
import sys
from pathlib import Path

import pytest
from pptx import Presentation

from kitchensink4ppt import server
from kitchensink4ppt.core import safesave
from kitchensink4ppt.ops import backups
from kitchensink4ppt.surface import envelope

pytestmark = pytest.mark.skipif(sys.platform != "win32", reason="Windows file identity")

ZONE = "[ZoneTransfer]\r\nZoneId=3\r\nHostUrl=https://example.invalid/deck.pptx\r\n"
NO_WINDOW = 0x08000000
GENERIC_READ = 0x80000000
SHARE_READ, SHARE_WRITE, SHARE_DELETE = 0x1, 0x2, 0x4


def _icacls(*args):
    return subprocess.run(
        ["icacls", *map(str, args)], capture_output=True, text=True,
        stdin=subprocess.DEVNULL, creationflags=NO_WINDOW,
    )


def _icacls_out(*args) -> str:
    return _icacls(*args).stdout


def _my_sid() -> str:
    out = subprocess.run(
        ["whoami", "/user", "/fo", "csv", "/nh"], capture_output=True, text=True,
        stdin=subprocess.DEVNULL, creationflags=NO_WINDOW,
    ).stdout
    return out.strip().split(",")[-1].strip('"')


def _aces(path) -> list[str]:
    """The deck's permission entries, one per line, without the file name."""
    lines = [ln.strip().replace(str(path), "").strip() for ln in _icacls_out(path).splitlines()
             if ":(" in ln]
    return sorted(lines)


def _title(path) -> str:
    return Presentation(str(path)).core_properties.title


def _edit(path, title="edited"):
    result = server.set_document_properties(str(path), {"title": title})
    assert result["ok"] is True, result
    assert _title(path) == title
    return result


def _zone(path) -> str | None:
    try:
        with open(f"{path}:Zone.Identifier", encoding="utf-8", newline="") as fh:
            return fh.read()
    except OSError:
        return None


def _set_zone(path) -> None:
    with open(f"{path}:Zone.Identifier", "w", encoding="utf-8", newline="") as fh:
        fh.write(ZONE)


def _set_creation_time(path, filetime: int = 132223104000000000) -> None:
    """Give the deck an old creation time (2020-01-01), so that a new file's
    creation time cannot coincide with it."""
    from ctypes import wintypes

    k32 = ctypes.WinDLL("kernel32", use_last_error=True)
    k32.CreateFileW.restype = wintypes.HANDLE
    k32.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD,
                                ctypes.c_void_p, wintypes.DWORD, wintypes.DWORD,
                                wintypes.HANDLE]
    handle = k32.CreateFileW(str(path), 0x100, 0x7, None, 3, 0x80, None)
    assert handle and handle != wintypes.HANDLE(-1).value
    try:
        stamp = wintypes.FILETIME(filetime & 0xFFFFFFFF, filetime >> 32)
        assert k32.SetFileTime(handle, ctypes.byref(stamp), None, None)
    finally:
        k32.CloseHandle(handle)


def _leftovers(folder) -> list[str]:
    return sorted(n for n in os.listdir(folder) if n.startswith(".ks4p-replaced-"))


@pytest.fixture
def deck(make_deck):
    path = make_deck("identity.pptx", extra_slides=0)
    yield path
    _icacls(path, "/reset")


# ------------------------------------------------------------ what survives


def test_the_mark_of_the_web_survives_an_edit(deck):
    _set_zone(deck)
    _edit(deck)
    assert "ZoneId=3" in (_zone(deck) or "")


def test_the_creation_time_survives_an_edit(deck):
    _set_creation_time(deck)
    before = os.stat(deck).st_birthtime
    _edit(deck)
    assert os.stat(deck).st_birthtime == before


def test_another_named_stream_survives_an_edit(deck):
    """The Mark of the Web is one stream among those a file can carry; the
    others (a sync client's, an indexer's) go with it."""
    with open(f"{deck}:other", "w", encoding="utf-8") as fh:
        fh.write("keep me")
    _edit(deck)
    with open(f"{deck}:other", encoding="utf-8") as fh:
        assert fh.read() == "keep me"


def test_a_deck_with_a_korean_name_keeps_its_identity(make_deck):
    path = make_deck("한글 발표 자료.pptx", extra_slides=0)
    try:
        _set_zone(path)
        _set_creation_time(path)
        stamp = os.stat(path).st_birthtime
        _edit(path)
        assert "ZoneId=3" in (_zone(path) or "")
        assert os.stat(path).st_birthtime == stamp
        assert _leftovers(path.parent) == []
    finally:
        _icacls(path, "/reset")


def test_a_restrictive_set_of_permissions_survives_an_edit(deck):
    """Owner only: the inherited entries are removed and one explicit entry
    is left. After an edit the deck must still have exactly that, not the
    entries its folder hands down."""
    assert _icacls(deck, "/inheritance:r").returncode == 0
    assert _icacls(deck, "/grant:r", f"*{_my_sid()}:(F)").returncode == 0
    before = _aces(deck)
    # Nothing inherited, and the user's own entry is there. An elevated
    # account's new files also carry explicit entries for Administrators,
    # SYSTEM and OWNER RIGHTS; the edit must keep those as well, so they stay
    # in the comparison below.
    assert not any("(I)" in ace for ace in before), before
    assert any(ace.endswith(":(F)") and not ace.startswith(("BUILTIN", "NT AUTHORITY", "OWNER RIGHTS"))
               for ace in before), before
    _edit(deck)
    assert _aces(deck) == before


@pytest.mark.parametrize("attribute", [0x2000, 0x1000, 0x80000, 0x100000],
                         ids=["not_content_indexed", "offline", "pinned", "unpinned"])
def test_attributes_survive_an_edit(deck, attribute):
    if not ctypes.windll.kernel32.SetFileAttributesW(str(deck), 0x20 | attribute):
        pytest.skip("this file system does not take that attribute")
    _edit(deck)
    assert ctypes.windll.kernel32.GetFileAttributesW(str(deck)) & attribute


def test_the_prev_slot_still_holds_the_old_content_and_nothing_is_left_beside_it(deck):
    _edit(deck, "one")
    _edit(deck, "two")
    assert _title(safesave.slot_dir(deck) / safesave.PREV_SLOT) == "one"
    assert _leftovers(deck.parent) == []


def test_an_edit_beside_a_reader_that_shares_delete(deck):
    """Another program may have the deck open while sharing delete. The save
    goes through, as it did with os.replace."""
    k32 = ctypes.WinDLL("kernel32", use_last_error=True)
    k32.CreateFileW.restype = ctypes.c_void_p
    k32.CreateFileW.argtypes = [ctypes.c_wchar_p, ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p,
                                ctypes.c_uint32, ctypes.c_uint32, ctypes.c_void_p]
    handle = k32.CreateFileW(str(deck), GENERIC_READ, SHARE_READ | SHARE_WRITE | SHARE_DELETE,
                             None, 3, 0x80, None)
    assert handle and handle != ctypes.c_void_p(-1).value
    try:
        _edit(deck)
    finally:
        k32.CloseHandle(ctypes.c_void_p(handle))


def test_a_reader_that_does_not_share_delete_refuses_the_save_and_changes_nothing(deck):
    before = deck.read_bytes()
    with open(deck, "rb"):  # the C runtime's default: shares read and write, not delete
        try:
            server.set_document_properties(str(deck), {"title": "must not publish"})
        except BaseException as exc:  # noqa: BLE001
            out = envelope.refusal(exc)
        else:
            raise AssertionError("the save went ahead beside a reader that blocks it")
    assert out["error"]["code"] == "DOCUMENT_LOCKED", out
    assert deck.read_bytes() == before
    assert _leftovers(deck.parent) == []


# ---------------------------------- readers beside a save that swaps the deck


def test_readers_never_see_the_deck_missing_while_saves_swap_it(deck):
    """ReplaceFileW moves the old file aside for an instant before the new one
    takes its place, so the path is missing for that instant. A read that
    lands in it waits it out; it is never told the deck is not there."""
    import threading
    import time

    from kitchensink4ppt.core import errors, package

    stop = threading.Event()
    outcomes: dict[str, int] = {}
    lock = threading.Lock()

    def tally(name):
        with lock:
            outcomes[name] = outcomes.get(name, 0) + 1

    def reader():
        while not stop.is_set():
            try:
                package.PptxPackage(deck)
                tally("ok")
            except (errors.DocumentBusy, errors.DocumentLocked):
                tally("busy")  # a save was still running; typed and retryable
            except BaseException as exc:  # noqa: BLE001
                tally(type(exc).__name__)
            time.sleep(0.002)

    readers = [threading.Thread(target=reader)]
    for thread in readers:
        thread.start()
    swaps = 0
    end = time.monotonic() + 4.0
    try:
        while time.monotonic() < end:
            try:
                _edit(deck, f"title {swaps}")
            except errors.DocumentLocked:
                continue  # a reader had the file open just then; the save is refused, nothing changed
            swaps += 1
    finally:
        stop.set()
        for thread in readers:
            thread.join()
    assert swaps >= 5, swaps
    assert outcomes.get("ok", 0) > 0
    assert set(outcomes) <= {"ok", "busy"}, outcomes


# ------------------------------------------- the other writes onto a deck


def test_a_restore_keeps_the_identity_too(deck):
    _edit(deck, "one")
    _set_zone(deck)
    _set_creation_time(deck)
    stamp = os.stat(deck).st_birthtime
    backups.restore_backup(str(deck), "prev")
    assert _title(deck) != "one"
    assert "ZoneId=3" in (_zone(deck) or "")
    assert os.stat(deck).st_birthtime == stamp
    assert _leftovers(deck.parent) == []


def test_overwriting_a_copy_keeps_the_destinations_identity(deck, make_deck):
    source = make_deck("source.pptx", seed=3, extra_slides=0)
    _set_zone(deck)
    _set_creation_time(deck)
    stamp = os.stat(deck).st_birthtime
    result = server.copy_presentation(str(source), str(deck), overwrite=True)
    assert result["overwrote_existing"] is True
    assert deck.read_bytes() == source.read_bytes()
    assert "ZoneId=3" in (_zone(deck) or "")
    assert os.stat(deck).st_birthtime == stamp
    assert _leftovers(deck.parent) == []


def test_a_new_file_is_placed_without_ReplaceFile(tmp_path, monkeypatch):
    """No deck is there to keep the identity of: a plain replace."""
    def not_for_a_new_file():
        raise AssertionError("ReplaceFileW is for a deck that exists")

    monkeypatch.setattr(safesave, "_replace_file", not_for_a_new_file)
    src, dst = tmp_path / "new.tmp", tmp_path / "new.pptx"
    src.write_bytes(b"new")
    safesave.promote(src, dst)
    assert dst.read_bytes() == b"new" and not src.exists()


# ----------------------------------------- where ReplaceFile is not available


@pytest.mark.parametrize("error", [50, 87, 120, 1], ids=lambda e: f"error_{e}")
def test_without_ReplaceFile_the_save_still_goes_through_with_the_identity_put_back(
        deck, monkeypatch, error):
    """A file system that answers 'not supported': a plain replace, then the
    Mark of the Web, creation time and attributes put back by hand."""
    _set_zone(deck)
    _set_creation_time(deck)
    assert ctypes.windll.kernel32.SetFileAttributesW(str(deck), 0x20 | 0x2000)
    stamp = os.stat(deck).st_birthtime
    seen = []

    def refuse(ldst, lsrc, lbak, flags, a, b):
        seen.append(flags)
        return False

    monkeypatch.setattr(safesave, "_REPLACE_FILE",
                        (refuse, lambda *a: False, lambda: error, lambda e: "not supported"))
    _edit(deck)
    assert seen == [0], "ReplaceFileW is called once, with no flags, and then given up on"
    assert "ZoneId=3" in (_zone(deck) or "")
    assert os.stat(deck).st_birthtime == stamp
    assert ctypes.windll.kernel32.GetFileAttributesW(str(deck)) & 0x2000
    assert _leftovers(deck.parent) == []


# ------------------------------------------------------- failures leave the deck


def test_a_failed_swap_that_moved_the_original_aside_puts_it_back(tmp_path, monkeypatch):
    """ReplaceFileW's ERROR_UNABLE_TO_MOVE_REPLACEMENT_2: the original was
    renamed to the backup name; promote moves it back before refusing."""
    dst = tmp_path / "deck.pptx"
    dst.write_bytes(b"original")
    src = tmp_path / "deck.pptx.ppt-mcp-tmp"
    src.write_bytes(b"new")
    state = {}

    def fake_replace(ldst, lsrc, lbak, flags, a, b):
        os.replace(ldst, lbak)
        state["bak"] = lbak
        return False

    def fake_move(a, b, flags):
        os.replace(a, b)
        return True

    monkeypatch.setattr(safesave, "_REPLACE_FILE",
                        (fake_replace, fake_move, lambda: 1177, lambda e: "could not move"))
    with pytest.raises(OSError):
        safesave.promote(str(src), str(dst))
    assert dst.read_bytes() == b"original"
    assert not os.path.exists(state["bak"])
    assert src.read_bytes() == b"new"


def test_a_held_deck_is_a_sharing_refusal_after_the_retries(tmp_path, monkeypatch):
    """ReplaceFileW's ERROR_UNABLE_TO_REMOVE_REPLACED: nothing changed, and the
    save is refused as DOCUMENT_LOCKED like any other sharing violation."""
    dst = tmp_path / "deck.pptx"
    dst.write_bytes(b"original")
    src = tmp_path / "deck.pptx.ppt-mcp-tmp"
    src.write_bytes(b"new")
    tries = []

    def held(*args):
        tries.append(1)
        return False

    monkeypatch.setattr(safesave, "_RETRY_DELAYS", (0.0, 0.0))
    monkeypatch.setattr(safesave, "_REPLACE_FILE",
                        (held, lambda *a: False, lambda: 1175, lambda e: "could not be removed"))
    with pytest.raises(PermissionError) as caught:
        safesave.promote(str(src), str(dst))
    assert caught.value.winerror == 1175
    assert len(tries) == 3  # the first try and one per retry delay
    refusal = safesave.os_write_refusal(dst, caught.value)
    assert refusal is not None and "another process is holding the file" in str(refusal)
    assert dst.read_bytes() == b"original" and src.read_bytes() == b"new"


# ------------------------------- a folder where owners may not set permissions


def _owner_rights_folder(tmp_path) -> Path:
    """A folder that grants Modify to the user and to OWNER RIGHTS, inherited:
    people may write files there but not change their permissions (a common
    file-share hardening). OWNER RIGHTS takes away the owner's implicit right
    to write the permissions, so ReplaceFileW with flags 0, which writes the
    merged permissions onto the new file, is denied there."""
    folder = tmp_path / "share"
    folder.mkdir()
    done = _icacls(folder, "/inheritance:r", "/grant:r", f"*{_my_sid()}:(OI)(CI)M",
                   "*S-1-3-4:(OI)(CI)M")
    assert done.returncode == 0, done.stdout + done.stderr
    return folder


def _deck_in(folder: Path, name: str = "deck.pptx") -> Path:
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "T"
    path = folder / name
    prs.save(path)
    return path


def test_a_folder_where_owners_may_not_change_permissions_still_saves(tmp_path):
    folder = _owner_rights_folder(tmp_path)
    try:
        deck = _deck_in(folder)
        _set_zone(deck)
        _set_creation_time(deck)
        stamp = os.stat(deck).st_birthtime
        result = _edit(deck)
        assert result["warnings"] == [], "the deck's permissions all came from the folder"
        assert "ZoneId=3" in (_zone(deck) or "")
        assert os.stat(deck).st_birthtime == stamp
        assert _leftovers(folder) == []
        restored = backups.restore_backup(str(deck), "prev")
        assert _title(deck) != "edited" and "warnings" not in restored
    finally:
        _icacls(folder, "/reset", "/t", "/q")


def test_a_denied_swap_is_made_once_more_ignoring_permission_errors(tmp_path, monkeypatch):
    """Flags 0 first; on ERROR_ACCESS_DENIED one more call at once with
    REPLACEFILE_IGNORE_ACL_ERRORS (0x4), and no pause before it."""
    dst = tmp_path / "deck.pptx"
    dst.write_bytes(b"original")
    src = tmp_path / "deck.pptx.ppt-mcp-tmp"
    src.write_bytes(b"new")
    seen = []

    def denied_without_the_flag(ldst, lsrc, lbak, flags, a, b):
        seen.append(flags)
        if flags & 0x4:
            os.replace(lsrc, ldst)
            return True
        return False

    monkeypatch.setattr(safesave, "_RETRY_DELAYS", (30.0, 30.0))  # a pause would hang the test
    monkeypatch.setattr(safesave, "_REPLACE_FILE",
                        (denied_without_the_flag, lambda *a: False, lambda: 5, lambda e: "denied"))
    assert safesave.promote(str(src), str(dst)) == []
    assert seen == [0, 0x4]
    assert dst.read_bytes() == b"new" and not src.exists()


def test_a_deck_whose_own_permissions_could_not_be_kept_is_named_in_a_warning(tmp_path, monkeypatch):
    dst = tmp_path / "deck.pptx"
    dst.write_bytes(b"original")
    src = tmp_path / "deck.pptx.ppt-mcp-tmp"
    src.write_bytes(b"new")
    texts = iter(["D:P(A;;FA;;;OW)", "D:AI(A;ID;FA;;;SY)"])
    monkeypatch.setattr(safesave, "_dacl_text", lambda path: next(texts))

    def denied_without_the_flag(ldst, lsrc, lbak, flags, a, b):
        if flags & 0x4:
            os.replace(lsrc, ldst)
            return True
        return False

    monkeypatch.setattr(safesave, "_REPLACE_FILE",
                        (denied_without_the_flag, lambda *a: False, lambda: 5, lambda e: "denied"))
    warnings = safesave.promote(str(src), str(dst))
    assert len(warnings) == 1 and "deck.pptx" in warnings[0] and "permissions" in warnings[0]
