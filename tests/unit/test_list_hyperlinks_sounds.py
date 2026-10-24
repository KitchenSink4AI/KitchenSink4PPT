"""list_hyperlinks lists a click that only plays a sound.

A shape whose click action is "None" with "Play sound" (Action Settings in
PowerPoint) has an hlinkClick with action ppaction://noaction and an a:snd
child. list_hyperlinks dropped it as an inert link, while its op line says
it lists every click action, and PowerPoint's own Hyperlinks collection
counts it. It is now listed as kind "sound" with the sound's name, and a
link or jump that also plays a sound carries the sound's name. Media
playback controls (the play button of a video or audio shape) are still not
listed. Every deck here is made up by code.
"""

from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.package import PptxPackage, qn
from kitchensink4ppt.ops import links

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
RT_AUDIO = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/audio"


def _deck(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for i, name in enumerate(("Chime box", "Jump box", "Quiet box")):
        box = slide.shapes.add_textbox(Inches(1), Inches(1 + 1.5 * i), Inches(4), Inches(1))
        box.name = name
        box.text_frame.text = name + " words"
    path = tmp_path / "made_up_sounds.pptx"
    prs.save(str(path))
    pkg = PptxPackage(path)
    part = pkg.slide_parts()[0]
    pkg.add_part_with_content_type("ppt/media/made_up_chime.wav", b"RIFF made up", "audio/wav")
    rid = pkg.add_relationship(part, RT_AUDIO, "../media/made_up_chime.wav")
    actions = {
        "Chime box": f'<a:hlinkClick xmlns:a="{A}" xmlns:r="{R}" r:id="" action="ppaction://noaction" '
                     f'highlightClick="1"><a:snd r:embed="{rid}" name="chime.wav"/></a:hlinkClick>',
        "Jump box": f'<a:hlinkClick xmlns:a="{A}" xmlns:r="{R}" r:id="" '
                    'action="ppaction://hlinkshowjump?jump=nextslide">'
                    f'<a:snd r:embed="{rid}" name="chime.wav"/></a:hlinkClick>',
        "Quiet box": f'<a:hlinkClick xmlns:a="{A}" xmlns:r="{R}" r:id="" action="ppaction://noaction"/>',
    }
    for cnvpr in pkg.root(part).iter(qn("p:cNvPr")):
        xml = actions.get(cnvpr.get("name"))
        if xml is not None:
            cnvpr.insert(0, etree.fromstring(xml))
    pkg.mark_dirty(part)
    pkg.save(do_backup=False)
    return path


def test_a_click_that_only_plays_a_sound_is_listed(tmp_path):
    out = links.list_hyperlinks(PptxPackage(_deck(tmp_path)))
    by_name = {h["shape_name"]: h for h in out["hyperlinks"]}
    chime = by_name.get("Chime box")
    assert chime is not None, out["hyperlinks"]
    assert (chime["kind"], chime["sound"], chime["where"], chime["trigger"]) == (
        "sound", "chime.wav", "shape", "click"), chime
    assert chime["broken"] is False


def test_a_jump_that_plays_a_sound_names_the_sound(tmp_path):
    out = links.list_hyperlinks(PptxPackage(_deck(tmp_path)))
    jump = next(h for h in out["hyperlinks"] if h["shape_name"] == "Jump box")
    assert (jump["kind"], jump["jump"], jump["sound"]) == ("show_jump", "nextslide", "chime.wav"), jump


def test_no_action_without_a_sound_is_still_not_listed(tmp_path):
    out = links.list_hyperlinks(PptxPackage(_deck(tmp_path)))
    assert "Quiet box" not in {h["shape_name"] for h in out["hyperlinks"]}
    assert out["count"] == 2
