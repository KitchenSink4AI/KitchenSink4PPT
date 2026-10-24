"""Slides are named the way people count, and every answer says which slide.

A weaker model asked "what do the speaker notes on slide 2 say?" sent
get_notes slide=2. Slide arguments count from 0, so that read the third
slide, which has no notes, and the answer
({"slide_index": 2, "has_notes": false, ...}) named neither the slide's
title nor its number. The model told the user, confidently, that slide 2
has no notes. Wrong and unreported, the worst failure class.

The fixes held here: ``slide_number`` (counting from 1) is accepted wherever
``slide`` is; every answer about one slide leads with its slide_number,
slide_count and slide_title; an answer to a 0-based address says in plain
words which slide that was; and the rule is stated where arguments are
chosen. ``slide`` itself still counts from 0, as the 2.0 docs say.
"""

from __future__ import annotations

import asyncio
import shutil
from types import SimpleNamespace

from pptx import Presentation

from kitchensink4ppt import server
from kitchensink4ppt.core import errors, slide_echo
from kitchensink4ppt.registry import public_ops
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface import copy, envelope
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.help import help_payload
from kitchensink4ppt.surface.plan import build


def _deck(path):
    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[0])
    first.shapes.title.text = "Quarterly Review"
    second = prs.slides.add_slide(prs.slide_layouts[1])
    second.shapes.title.text = "Revenue"
    second.placeholders[1].text = "Q3 revenue grew 12 percent"
    second.notes_slide.notes_text_frame.text = "Mention the Lisbon contract here."
    third = prs.slides.add_slide(prs.slide_layouts[5])
    third.shapes.title.text = "Budget"
    prs.save(path)
    return path


def _call(tool: str, op: str, args: dict):
    plan = build(Settings(), public_ops())
    return asyncio.run(dispatch(plan, tool, {"op": op, "args": args}))


def test_the_p02_read_names_the_slide_it_found(tmp_path):
    deck = str(_deck(tmp_path / "deck.pptx"))
    payload, is_error = _call("ppt_read", "get_notes", {"file_path": deck, "slide": 2})
    assert is_error is False
    assert list(payload)[:4] == ["slide_note", "slide_number", "slide_count", "slide_title"]
    assert payload["slide_number"] == 3
    assert payload["slide_count"] == 3
    assert payload["slide_title"] == "Budget"
    assert payload["slide_note"] == (
        'You sent slide=2, which counts from 0, so this is the third slide, '
        'not the second (slide_number 3 of 3), titled "Budget".'
    )
    assert payload["slide_index"] == 2  # slide still counts from 0


def test_slide_number_counts_from_one_and_reaches_the_slide_meant(tmp_path):
    deck = str(_deck(tmp_path / "deck.pptx"))
    payload, is_error = _call(
        "ppt_read", "get_notes", {"file_path": deck, "slide_number": 2}
    )
    assert is_error is False, payload
    assert payload["text"] == "Mention the Lisbon contract here."
    assert payload["slide_number"] == 2
    assert payload["slide_title"] == "Revenue"
    assert "slide_note" not in payload
    assert "normalized_from" not in payload

    # Flat, outside args, the way weak models often send it.
    plan = build(Settings(), public_ops())
    payload, is_error = asyncio.run(dispatch(plan, "ppt_read", {
        "op": "get_notes", "file_path": deck, "slide_number": 2,
    }))
    assert is_error is False, payload
    assert payload["slide_title"] == "Revenue"


def test_other_names_for_one_slide_are_read_one_way_only(tmp_path):
    deck = str(_deck(tmp_path / "deck.pptx"))
    for args in (
        {"slide_index": 1},
        {"slide_id": 257},
        {"slide": "1"},
        {"slide": '{"slide_id": 257}'},
        {"slide": 1, "slide_number": 2},
    ):
        payload, is_error = _call(
            "ppt_read", "get_notes", {"file_path": deck, **args}
        )
        assert is_error is False, (args, payload)
        assert payload["slide_title"] == "Revenue", args

    for args, words in (
        ({"slide": 2, "slide_number": 2}, "name different slides"),
        ({"slide_number": 0}, "counts from 1"),
        ({"slide_number": "two"}, "counts from 1"),
        ({"slide_id": 257, "slide": 1}, "each name a slide"),
    ):
        payload, is_error = _call(
            "ppt_read", "get_notes", {"file_path": deck, **args}
        )
        assert is_error is True, args
        assert payload["error"]["code"] == "BAD_ARGS"
        assert words in payload["error"]["message"], payload


def test_writes_and_refusals_after_the_slide_was_found_name_it_too(tmp_path):
    deck = _deck(tmp_path / "deck.pptx")
    payload, is_error = _call("ppt_edit", "set_notes", {
        "file_path": str(deck), "slide_number": 3, "text": "Check the totals.",
    })
    assert is_error is False, payload
    assert (payload["slide_number"], payload["slide_title"]) == (3, "Budget")
    assert Presentation(str(deck)).slides[2].notes_slide.notes_text_frame.text == (
        "Check the totals."
    )

    payload, is_error = _call("ppt_read", "get_table", {
        "file_path": str(deck), "slide_number": 2,
    })
    assert is_error is True
    assert payload["error"]["code"] == "NOT_FOUND"
    assert payload["error"]["details"]["slide_number"] == 2
    assert payload["error"]["details"]["slide_title"] == "Revenue"

    payload, is_error = _call("ppt_edit", "move_slide", {
        "file_path": str(deck), "slide_number": 3, "to": 0,
    })
    assert is_error is False, payload
    assert payload["slide_title"] == "Budget"
    assert payload["changed"]["new_slide_number"] == 1
    payload, is_error = _call("ppt_edit", "duplicate_slide", {
        "file_path": str(deck), "slide_number": 1,
    })
    assert is_error is False, payload
    assert payload["changed"]["new_slide_number"] == payload["changed"]["index"] + 1
    # The answer names the source of the copy, and slide_count is the deck now.
    assert (payload["source_slide_number"], payload["source_slide_title"]) == (1, "Budget")
    assert payload["slide_count"] == 4
    assert "slide_number" not in payload

    payload, is_error = _call("ppt_edit", "delete_slide", {
        "file_path": str(deck), "slide": 2,
    })
    assert is_error is False, payload
    # After a delete, slide_count is the deck after it. When it described
    # the deck before, a model reported "the presentation now has 4 slides".
    assert payload["deleted_slide_number"] == 3
    assert payload["deleted_slide_title"] == "Quarterly Review"
    assert payload["slide_count"] == 3
    assert "slide_number" not in payload
    assert payload["slide_note"] == (
        "You sent slide=2, which counts from 0, so this was the third slide, "
        'not the second (slide_number 3), titled "Quarterly Review".'
    )


def test_an_out_of_range_slide_gives_both_ranges(tmp_path):
    deck = str(_deck(tmp_path / "deck.pptx"))
    payload, is_error = _call("ppt_read", "get_notes", {"file_path": deck, "slide": 3})
    assert is_error is True
    message = payload["error"]["message"]
    assert "out of range, presentation has 3 slides" in message
    assert "0 to 2" in message and "slide_number 1 to 3" in message


def test_every_slide_operation_echoes_the_slide_it_resolved(tmp_path, monkeypatch):
    """Every operation that takes slide, run for real in file mode.

    This test is not marked live, so it must never start PowerPoint or attach
    to one, which may be the user's own. An operation that would still reach
    PowerPoint fails the test instead of opening anything."""
    import test_family_op_sweep_gate as sweep
    from kitchensink4ppt.com import bridge

    reached_powerpoint: list[str] = []
    current = {"op": None}

    def no_powerpoint(*_args, **_kwargs):
        reached_powerpoint.append(current["op"])
        raise RuntimeError("file mode: this test never starts PowerPoint")

    monkeypatch.setattr(bridge, "_powerpoint", no_powerpoint)
    # File mode: an operation that renders with PowerPoint finds none, so it
    # refuses after it has found the slide, and that refusal is checked too.
    monkeypatch.setattr(bridge, "powerpoint_installed", lambda: False)

    source = sweep._canary_deck(tmp_path / "canary-source.pptx")
    count = len(Presentation(str(source)).slides)
    plan = build(Settings(call_timeout_s=120.0))
    missing, ran = [], 0
    for name, spec in sorted(plan.specs.items()):
        if "slide" not in spec.args.model_fields or name == "live_scroll_to":
            continue
        args = sweep._args_for(name, tmp_path, source)
        selector = args.get("slide")
        if not slide_echo.single(selector):
            continue
        if "live" in spec.args.model_fields:
            args["live"] = "off"
        current["op"] = name
        payload, is_error = asyncio.run(
            dispatch(plan, spec.group, {"op": name, "args": args})
        )
        echo = payload["error"].get("details", {}) if is_error else payload
        if is_error and "slide_number" not in echo:
            continue  # refused before the slide was resolved
        ran += 1
        style = None if is_error else slide_echo.STYLES.get(name)
        number = {
            "deleted": "deleted_slide_number", "copied": "source_slide_number",
            "source_deck": "source_slide_number",
        }.get(style, "slide_number")
        total = {"deleted": count - 1, "copied": count + 1}.get(style, count)
        size = "source_slide_count" if style == "source_deck" else "slide_count"
        if echo.get(number) != selector + 1 or echo.get(size) != total:
            missing.append(f"{name}: {str(payload)[:160]}")
    assert reached_powerpoint == []
    assert missing == []
    assert ran >= 50


def test_help_and_refusals_offer_slide_number_wherever_slide_is_taken():
    plan = build(Settings(), public_ops())
    schema = help_payload(plan, "get_notes")["args_schema"]
    names = list(schema["properties"])
    assert names[names.index("slide") + 1] == "slide_number"
    assert schema["properties"]["slide"]["description"] == copy.SLIDE_PARAM
    assert schema["properties"]["slide_number"]["minimum"] == 1
    # set_footer's own slide_number (show the number field) is left alone.
    footer = help_payload(plan, "set_footer")["args_schema"]["properties"]
    assert footer["slide_number"].get("description") != copy.SLIDE_NUMBER_PARAM["description"]

    payload, is_error = _call("ppt_read", "get_notes", {"file_path": "x.pptx"})
    assert is_error is True
    fields = payload["error"]["details"]["valid_fields"]
    assert fields[fields.index("slide") + 1] == "slide_number"

    for tool in plan.tools.values():
        if tool.name != "ppt_help":
            assert copy.SLIDE_RULE in tool.input_schema["properties"]["args"]["description"]


def test_a_batch_edit_takes_slide_number_and_reports_it(tmp_path):
    deck = _deck(tmp_path / "deck.pptx")
    info = server.get_slide_info(str(deck), 1, live="off")
    title_id = next(s["id"] for s in info["shapes"] if s.get("placeholder_type") == "title")
    payload, is_error = _call("ppt_edit", "apply_edits", {
        "file_path": str(deck),
        "edits": [{"op": "set_shape", "slide_number": 2, "shape": title_id, "dx": 0.1}],
    })
    assert is_error is False, payload
    assert payload["changed"]["0"]["slide_number"] == 2


def test_the_live_route_names_the_open_deck_slide():
    title = SimpleNamespace(TextFrame=SimpleNamespace(TextRange=SimpleNamespace(Text="Revenue\r")))
    slides = [
        SimpleNamespace(SlideID=256 + i, SlideIndex=i + 1,
                        Shapes=SimpleNamespace(HasTitle=-1, Title=title))
        for i in range(4)
    ]
    pres = SimpleNamespace(Slides=SimpleNamespace(Count=4, Item=lambda n: slides[n - 1]))
    from kitchensink4ppt.com import live_ops

    probe, token = slide_echo.arm(1, "slide")
    try:
        live_ops._resolve_slide_live(pres, 1)
    finally:
        slide_echo.disarm(token)
    assert probe.fields()["slide_number"] == 2
    assert probe.fields()["slide_title"] == "Revenue"

    # A file attempt that found the slide before PowerPoint's lock sent the
    # call live must not describe the file copy.
    probe, token = slide_echo.arm(1, "slide")
    try:
        def file_call():
            slide_echo.report(1, 1, 9, "stale file title")
            raise errors.DocumentLocked("open in PowerPoint")

        server._route_live(
            "auto", file_call, lambda: live_ops._resolve_slide_live(pres, 1)
        )
    finally:
        slide_echo.disarm(token)
    assert probe.fields()["slide_count"] == 4
    assert probe.fields()["slide_title"] == "Revenue"


def test_the_echo_is_only_ever_a_description(tmp_path):
    """No selector form changes which slide slide means, and a title that
    cannot be read never fails the call."""
    assert slide_echo.ordinal(2) == "second"
    assert slide_echo.ordinal(13) == "13th"
    assert slide_echo.ordinal(22) == "22nd"
    probe = slide_echo.Probe(0, "slide")

    def broken():
        raise RuntimeError("no title")

    probe.offer(0, 0, 1, broken)
    assert probe.fields()["slide_number"] == 1
    assert "slide_title" not in probe.fields()
    assert slide_echo.clean_title("  A\vlong\ttitle \n") == "A long title"
    assert len(slide_echo.clean_title("x" * 500)) == slide_echo.TITLE_CHARS
    assert envelope.success({"a": 1}) == {"a": 1, "ok": True}
    copy_path = tmp_path / "copy.pptx"
    shutil.copy2(_deck(tmp_path / "deck.pptx"), copy_path)
    payload, _ = _call("ppt_read", "get_slide_info", {"file_path": str(copy_path), "slide": 0})
    assert payload["index"] == 0 and payload["slide_number"] == 1


def _bytes(path) -> bytes:
    with open(path, "rb") as fh:
        return fh.read()


def test_a_null_slide_name_is_read_as_not_sent(tmp_path):
    """An earlier version refused slide_number: null beside slide
    ("slide_number counts from 1") on every slide operation, though ppt_help
    lists slide_number as an optional argument and every other optional
    argument accepts null. Clients that fill unused fields with null are
    common."""
    deck = str(_deck(tmp_path / "deck.pptx"))
    for extra in ({"slide_number": None}, {"slide_index": None}, {"slide_id": None},
                  {"slide_number": None, "slide_index": None}):
        payload, is_error = _call(
            "ppt_read", "get_notes", {"file_path": deck, "slide": 1, **extra}
        )
        assert is_error is False, (extra, payload)
        assert payload["slide_title"] == "Revenue", extra
    # Null on its own names no slide, so the slide is still missing.
    payload, is_error = _call("ppt_read", "get_notes", {"file_path": deck, "slide_number": None})
    assert is_error is True
    assert payload["error"]["code"] == "BAD_ARGS"
    # A batch edit reads a null slide_number the same way.
    info = server.get_slide_info(deck, 1, live="off")
    title_id = next(s["id"] for s in info["shapes"] if s.get("placeholder_type") == "title")
    payload, is_error = _call("ppt_edit", "apply_edits", {"file_path": deck, "edits": [
        {"op": "set_shape", "slide": 1, "slide_number": None, "shape": title_id, "dx": 0.1}]})
    assert is_error is False, payload
    assert payload["changed"]["0"]["slide_number"] == 2


def test_a_batch_edit_refuses_an_anchor_and_a_slide_that_disagree(tmp_path):
    """An edit with a view anchor on slide 2 and slide_number 3 used to edit
    slide 2 silently. A top-level call refuses two slide names that
    disagree; an edit now does too, and agreeing names still work."""
    deck = str(_deck(tmp_path / "deck.pptx"))
    view, _ = _call("ppt_read", "get_presentation_view", {"file_path": deck, "detail": "full"})
    section = view["view"].split("## Slide 2", 1)[1].split("## Slide 3", 1)[0]
    anchor = section.split("[a:", 1)[1].split("]", 1)[0]
    before = _bytes(deck)
    for name, value in (("slide_number", 3), ("slide", 2), ("slide", 0)):
        payload, is_error = _call("ppt_edit", "apply_edits", {"file_path": deck, "edits": [
            {"op": "set_text", "anchor": anchor, name: value, "text": "which slide?"}]})
        assert is_error is True, (name, value, payload)
        assert "slide_number 2" in payload["error"]["message"], payload
        assert _bytes(deck) == before
    for name, value in (("slide_number", 2), ("slide", 1)):
        payload, is_error = _call("ppt_edit", "apply_edits", {"file_path": deck, "edits": [
            {"op": "set_text", "anchor": anchor, name: value, "text": f"by {name}"}]})
        assert is_error is False, (name, payload)
        assert payload["changed"]["0"]["slide_number"] == 2


def test_a_new_slide_position_refusal_states_the_deck_as_it_is(tmp_path):
    """insert_slide position 6 on a 5-slide deck was refused with "the deck
    now has 6 slides (valid: 0..5)": the check ran after the new slide was
    built, and a model relayed the count of a slide never made."""
    deck = str(_deck(tmp_path / "deck.pptx"))  # 3 slides
    other = str(_deck(tmp_path / "other.pptx"))
    before = _bytes(deck)
    for op, args in (
        ("insert_slide", {"layout": 1, "position": 4, "live": "off"}),
        ("duplicate_slide", {"slide": 0, "position": 4}),
    ):
        payload, is_error = _call("ppt_edit", op, {"file_path": deck, **args})
        assert is_error is True, (op, payload)
        message = payload["error"]["message"]
        assert "has 3 slides" in message and "0 to 3" in message, message
        assert "now has" not in message and "will have" not in message
        assert _bytes(deck) == before
    payload, is_error = _call("ppt_deck", "copy_slide_between", {
        "file_path": deck, "source_path": other, "slide": 0, "position": 9})
    assert is_error is True, payload
    assert "the destination has 3 slides" in payload["error"]["message"]
    assert _bytes(deck) == before
    # The last valid position still adds the slide last.
    payload, is_error = _call(
        "ppt_edit", "insert_slide", {"file_path": deck, "layout": 1, "position": 3, "live": "off"}
    )
    assert is_error is False, payload
    assert payload["changed"]["new_slide_number"] == 4


def test_the_help_schema_lets_the_documented_slide_number_call_through(tmp_path):
    """The help says to send slide_number instead of slide, so the schema
    must not list slide as required: a client that checks calls against the
    schema would block the alias. A required slide is "slide or
    slide_number"; the server still refuses a call that names neither, or
    two that disagree."""
    import jsonschema
    import pytest

    plan = build(Settings(), public_ops())
    schema = help_payload(plan, "get_notes")["args_schema"]
    deck = str(_deck(tmp_path / "deck.pptx"))

    alias = {"file_path": deck, "slide_number": 2}
    jsonschema.validate(alias, schema)
    payload, is_error = _call("ppt_read", "get_notes", alias)
    assert is_error is False, payload
    assert payload["text"] == "Mention the Lisbon contract here."

    missing = {"file_path": deck}
    with pytest.raises(jsonschema.ValidationError):
        jsonschema.validate(missing, schema)
    payload, is_error = _call("ppt_read", "get_notes", missing)
    assert is_error is True and payload["error"]["code"] == "BAD_ARGS"

    conflict = {"file_path": deck, "slide": 0, "slide_number": 2}
    jsonschema.validate(conflict, schema)  # values are the server's to compare
    payload, is_error = _call("ppt_read", "get_notes", conflict)
    assert is_error is True and payload["error"]["code"] == "BAD_ARGS"
    assert "name different slides" in payload["error"]["message"]

    for name, spec in public_ops().items():
        field = spec.args.model_fields.get("slide")
        if field is None:
            continue
        detail = help_payload(plan, name)
        rules = detail["args_schema"]
        jsonschema.validate(detail["example"]["args"], rules)
        if field.is_required():
            assert "slide" not in rules.get("required", []), name
            assert rules["anyOf"] == [
                {"required": ["slide"]},
                {"required": ["slide_number"],
                 "properties": {"slide_number": {"type": "integer"}}},
            ], name
        else:  # an optional slide stays optional
            assert "anyOf" not in rules and "allOf" not in rules, name


def test_the_help_schema_admits_the_null_slide_number_the_server_accepts(tmp_path):
    """The server reads a null slide_number as not sent, so a client that
    checks calls against ppt_help's schema must not be blocked on it. Null on
    its own names no slide: the server refuses it, and so does the schema
    wherever the slide is required."""
    import jsonschema
    import pytest

    plan = build(Settings(), public_ops())
    deck = str(_deck(tmp_path / "deck.pptx"))
    call = {"file_path": deck, "slide": 1, "slide_number": None}
    jsonschema.validate(call, help_payload(plan, "get_notes")["args_schema"])
    payload, is_error = _call("ppt_read", "get_notes", call)
    assert is_error is False, payload
    assert payload["slide_title"] == "Revenue"

    checked = 0
    for name, spec in public_ops().items():
        field = spec.args.model_fields.get("slide")
        if field is None:
            continue
        detail = help_payload(plan, name)
        rules = detail["args_schema"]
        args = dict(detail["example"]["args"])
        args.setdefault("slide", 0)
        jsonschema.validate({**args, "slide_number": None}, rules)
        if field.is_required():
            alone = {key: value for key, value in args.items() if key != "slide"}
            with pytest.raises(jsonschema.ValidationError):
                jsonschema.validate({**alone, "slide_number": None}, rules)
            jsonschema.validate({**alone, "slide_number": 2}, rules)
        checked += 1
    assert checked >= 70


def test_every_argument_that_counts_slides_says_how_in_help():
    """Every argument that counts slides says how, not only slide: to_slide,
    to, position, order, slides, from_slide, to_slides and scope also count
    from 0. A weaker model sent the user's slide 3 as to_slide 3 (the fourth
    slide) when the help said nothing."""
    plan = build(Settings(), public_ops())
    checked = 0
    for op, notes in copy.ARG_NOTES.items():
        properties = help_payload(plan, op)["args_schema"]["properties"]
        for name, text in notes.items():
            assert name in properties, (op, name)
            assert properties[name]["description"] == text
            assert "counted from" in text, (op, name)
            checked += 1
    assert checked >= 20
    to_slide = help_payload(plan, "set_hyperlink")["args_schema"]["properties"]["to_slide"]
    assert "counted from 0" in to_slide["description"]
    assert "target_slide_number" in to_slide["description"]


def test_a_jump_link_names_its_target_the_way_people_count(tmp_path):
    deck = str(_deck(tmp_path / "deck.pptx"))
    info = server.get_slide_info(deck, 0, live="off")
    title_id = next(
        s["id"] for s in info["shapes"] if s.get("placeholder_type") in ("title", "ctrTitle")
    )
    payload, is_error = _call("ppt_edit", "set_hyperlink", {
        "file_path": deck, "slide_number": 1, "target": title_id, "to_slide": 2})
    assert is_error is False, payload
    changed = payload["changed"]
    assert changed["target_slide"]["index"] == 2
    assert changed["target_slide_number"] == 3
    assert changed["target_slide_title"] == "Budget"


def test_a_single_slide_scope_says_which_slide_it_read(tmp_path):
    """get_text(scope=1) now leads with the slide it read, as get_notes(slide=1)
    does, and slide_count is the deck in every answer; the number of slides
    a scope selected is slides_in_scope."""
    deck = str(_deck(tmp_path / "deck.pptx"))
    payload, is_error = _call("ppt_read", "get_text", {"file_path": deck, "scope": 1})
    assert is_error is False, payload
    assert payload["slide_number"] == 2
    assert payload["slide_title"] == "Revenue"
    assert payload["slide_count"] == 3
    assert payload["slides_in_scope"] == 1
    assert payload["slide_note"].startswith("You sent scope=1, which counts from 0")
    assert payload["slides"][0]["index"] == 1

    payload, _ = _call("ppt_read", "get_text", {"file_path": deck, "scope": [0, 2]})
    assert "slide_number" not in payload  # a list names no single slide
    assert (payload["slide_count"], payload["slides_in_scope"]) == (3, 2)

    payload, _ = _call("ppt_read", "get_text", {"file_path": deck})
    assert payload["slide_count"] == 3
    assert "slides_in_scope" not in payload and "slide_number" not in payload

    slide_id = payload["slides"][2]["slide_id"]
    payload, is_error = _call("ppt_read", "get_presentation_view", {
        "file_path": deck, "scope": {"slide_id": slide_id}})
    assert is_error is False, payload
    assert (payload["slide_number"], payload["slide_title"]) == (3, "Budget")
    assert (payload["slide_count"], payload["slides_in_scope"]) == (3, 1)
    assert "slide_note" not in payload  # a slide id is not a 0-based count

    payload, is_error = _call("ppt_read", "find_text", {
        "file_path": deck, "query": "revenue", "scope": 1})
    assert is_error is False, payload
    assert payload["slide_number"] == 2
