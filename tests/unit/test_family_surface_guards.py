"""Regression tests for the PowerPoint family surface.

Each test's docstring gives the behaviour it pins and why. None of them
touches PowerPoint: COM is replaced by a refusal wherever a code path could
reach it, except the one COM-gated test (disconnect during a COM call),
which skips unless KS4_VERIFY_PPT_COM=1 is set in a window where no
PowerPoint is running.
"""

from __future__ import annotations

import ast
import asyncio
import hashlib
import json
import os
import queue
import re
import shutil
import subprocess
import sys
import threading
import time
import zipfile
from pathlib import Path

import pytest

from kitchensink4ppt import server
from kitchensink4ppt.core import budget, safesave
from kitchensink4ppt.core.errors import OpError
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface import dispatch as surface_dispatch
from kitchensink4ppt.surface.dispatch import dispatch, run_bounded
from kitchensink4ppt.surface.plan import build

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "src" / "kitchensink4ppt"


# ----------------------------------------------------------------- helpers


def _deck(path: Path, slides: int = 4, body_chars: int = 200, table: bool = True) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Slide {i}"
        word = f"S{i:03d}word "
        slide.placeholders[1].text_frame.text = (word * (body_chars // len(word) + 1))[:body_chars]
    if table:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        tbl = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
        for r in range(2):
            for c in range(2):
                tbl.cell(r, c).text = f"cell{r}{c}"
    prs.save(path)
    return path


def _sha(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _with_slots(tmp_path: Path) -> tuple[Path, Path, Path]:
    """A deck that has been mutated once, so prev and anchor slots exist."""
    deck = _deck(tmp_path / "deck.pptx")
    server.set_slide_hidden(str(deck), 0, True)
    slot = tmp_path / ".ks4p-backups" / "deck.pptx"
    prev, anchor = slot / "prev.pptx", slot / "anchor.pptx"
    assert prev.exists() and anchor.exists()
    return deck, prev, anchor


def _table_id(deck: Path) -> int:
    info = server.get_slide_info(str(deck), 4, live="off")
    return next(s["id"] for s in info["shapes"] if "Table" in (s.get("name") or ""))


def _rewrite(src: Path, dst: Path, fn) -> Path:
    zin = zipfile.ZipFile(src)
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = fn(item.filename, zin.read(item.filename))
            if data is not None:
                zout.writestr(item, data)
    return dst


@pytest.fixture
def no_com(monkeypatch):
    from kitchensink4ppt.com import bridge

    def refuse(*_a, **_k):
        raise RuntimeError("COM disabled in this test")

    monkeypatch.setattr(bridge, "com_validate_opens_clean", refuse)
    return bridge


# ---------------------------------------------------------------- validate on a broken payload


@pytest.mark.parametrize("damage", ["dangling_rel", "truncated_slide", "no_content_types"])
def test_v1_validate_answers_a_broken_payload_with_ok_false(tmp_path, no_com, damage):
    """validate answers a broken payload with the domain verdict
    ok:false / payload_valid:false. The payload branch once read
    `except _CATCHABLE`, a name deleted with the old envelope, so a broken
    deck raised NameError and the surface answered INTERNAL_ERROR."""
    good = _deck(tmp_path / "good.pptx")
    if damage == "dangling_rel":
        fn = lambda n, d: d.replace(b"slideLayout2.xml", b"slideLayout99.xml") if n.endswith("slide1.xml.rels") else d
    elif damage == "truncated_slide":
        fn = lambda n, d: d[:-40] if n == "ppt/slides/slide1.xml" else d
    else:
        fn = lambda n, d: None if n == "[Content_Types].xml" else d
    broken = _rewrite(good, tmp_path / "broken.pptx", fn)
    plan = build(Settings())
    payload, is_error = asyncio.run(dispatch(
        plan, "ppt_deck", {"op": "validate", "args": {"file_path": str(broken)}}
    ))
    assert payload.get("error", {}).get("code") != "INTERNAL_ERROR", payload
    assert payload["ok"] is False
    assert payload.get("payload_valid") is False


# ---------------------------------------------------------------- paging under the output cap


def test_v2_cutting_a_non_paged_list_keeps_the_engine_paging():
    """budget.enforce must not treat the first list it cuts as the engine's
    paged container. Here the engine paged `findings` (1 returned,
    next_offset 1); the wire cap forces a cut of the unrelated `checks`
    list, and rewriting page to returned=<kept checks>,
    next_offset=<kept checks> would make a caller following next_offset
    skip findings."""
    result = {
        "findings": [{"slide_index": 0, "message": "m" * 200}],
        "checks": [{"check": f"c{i}", "caveat": "x" * 900} for i in range(8)],
        "page": {
            "unit": "findings", "total": 70, "offset": 0, "returned": 1,
            "omitted": 69, "truncated": True, "reason": "budget",
            "next_offset": 1, "hint": "69 more findings were not returned; call again with offset=1",
        },
    }
    out = budget.enforce(result, 6000, wrap=lambda r: {**r, "ok": True},
                         supports_offset=True, requested_offset=0)
    page = out["page"]
    assert len(out["findings"]) == 1
    assert page.get("next_offset") == 1, page
    assert page.get("returned") in (1, None) or page.get("field") != "checks", page


def test_v2_check_layout_chain_at_a_small_cap_returns_every_finding(tmp_path, monkeypatch):
    """End to end: with KS4P_MAX_OUTPUT_CHARS=8000, following next_offset
    through check_layout on a 70-slide deck with one overflow finding per
    slide must return all 70 findings (it once returned 12 of 70)."""
    deck = _deck(tmp_path / "d70.pptx", slides=70, body_chars=1400, table=False)
    monkeypatch.setenv("KS4P_MAX_OUTPUT_CHARS", "8000")
    plan = build(Settings(max_output_chars=8000))
    seen, offset = [], None
    for _ in range(200):
        args = {"file_path": str(deck)}
        if offset is not None:
            args["offset"] = offset
        payload, is_error = asyncio.run(dispatch(plan, "ppt_inspect", {"op": "check_layout", "args": args}))
        assert not is_error, payload
        seen += [f["slide_index"] for f in payload["findings"] if f["check"] == "overflow"]
        nxt = payload.get("page", {}).get("next_offset")
        if not nxt:
            break
        offset = nxt
    assert sorted(set(seen)) == list(range(70)), f"missing {sorted(set(range(70)) - set(seen))[:10]}"


# ---------------------------------------------------------------- outputs never land on the deck or its backups


def test_v3_export_table_never_writes_over_the_presentation(tmp_path):
    """export_table never writes over the presentation: with path = the
    deck itself it wrote CSV over it."""
    deck = _deck(tmp_path / "deck.pptx")
    before = _sha(deck)
    try:
        server.export_table(str(deck), 4, _table_id(deck), path=str(deck), format="csv")
    except Exception:
        pass
    assert _sha(deck) == before


@pytest.mark.parametrize("slot_name", ["prev", "anchor"])
def test_v3_export_table_never_writes_into_the_backup_store(tmp_path, slot_name):
    """export_table never writes into the backup store: onto a backup slot
    it overwrote the slot, and through the prev/anchor hardlink both undo
    slots."""
    deck, prev, anchor = _with_slots(tmp_path)
    target = prev if slot_name == "prev" else anchor
    before = (_sha(prev), _sha(anchor))
    try:
        server.export_table(str(deck), 4, _table_id(deck), path=str(target), format="json")
    except Exception:
        pass
    assert (_sha(prev), _sha(anchor)) == before


@pytest.mark.parametrize("slot_name", ["prev", "anchor"])
def test_v3_copy_presentation_never_overwrites_a_backup_slot(tmp_path, slot_name):
    """copy_presentation never overwrites a backup slot: overwrite=True
    onto a slot replaced the undo copy and answered success."""
    deck, prev, anchor = _with_slots(tmp_path)
    target = prev if slot_name == "prev" else anchor
    other = _deck(tmp_path / "other.pptx", slides=2)
    before = _sha(target)
    try:
        server.copy_presentation(str(other), str(target), overwrite=True)
    except Exception:
        pass
    assert _sha(target) == before


@pytest.mark.parametrize("slot_name", ["prev", "anchor"])
def test_v3_a_backup_slot_is_never_an_edit_target(tmp_path, slot_name):
    """A backup slot is never an edit target: set_notes with file_path = a
    backup slot edited the undo copy in place."""
    deck, prev, anchor = _with_slots(tmp_path)
    target = prev if slot_name == "prev" else anchor
    before = _sha(target)
    try:
        server.set_notes(str(target), 0, text="edited the undo", live="off")
    except Exception:
        pass
    assert _sha(target) == before


@pytest.mark.timeout(400)
def test_v3_export_pdf_never_replaces_the_presentation(tmp_path):
    """export_pdf never replaces the presentation: on the LibreOffice
    route, output = the deck itself moved the produced PDF over the
    presentation and answered success."""
    from kitchensink4ppt.ops import export

    if export._find_soffice() is None:
        pytest.skip("LibreOffice not installed")
    deck = _deck(tmp_path / "deck.pptx", slides=2, table=False)
    before = _sha(deck)
    try:
        server.export_pdf(str(deck), output=str(deck), engine="libreoffice")
    except Exception:
        pass
    assert _sha(deck) == before


# ---------------------------------------------------------------- op sweep gate


def test_v4_the_brief_3_8_op_sweep_exists():
    """A test module carries the op sweep: every op once through its tool,
    a MUST_SUCCEED set, reviewed READER/WRITER_OK_REFUSALS, and the
    output-guard cases."""
    texts = [p.read_text(encoding="utf-8") for p in (ROOT / "tests").rglob("test_*.py")
             if p.name != Path(__file__).name]
    assert any("MUST_SUCCEED" in t and ("WRITER_OK_REFUSALS" in t or "READER_OK_REFUSALS" in t) for t in texts)


# ---------------------------------------------------------------- stderr stays short


class _Stdio:
    def __init__(self, env_extra=None, *, drain=True):
        env = dict(os.environ)
        env.update({"PYTHONUNBUFFERED": "1", "KS4P_UPDATE_CHECK": "off"})
        env.update(env_extra or {})
        self.proc = subprocess.Popen(
            [sys.executable, "-X", "utf8", str(ROOT / "scripts" / "_bootstrap.py")],
            stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
            cwd=str(ROOT), env=env,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        self.q: queue.Queue = queue.Queue()
        self.err = bytearray()
        threading.Thread(target=self._out, daemon=True).start()
        if drain:
            threading.Thread(target=self._errpump, daemon=True).start()
        self.id = 0

    def _out(self):
        for line in self.proc.stdout:
            try:
                self.q.put(json.loads(line))
            except Exception:
                pass

    def _errpump(self):
        while True:
            b = self.proc.stderr.read1(65536)
            if not b:
                break
            self.err += b

    def send(self, raw: bytes):
        self.proc.stdin.write(raw + b"\n")
        self.proc.stdin.flush()

    def request(self, method, params, timeout=60):
        self.id += 1
        self.send(json.dumps({"jsonrpc": "2.0", "id": self.id, "method": method, "params": params}).encode())
        end = time.monotonic() + timeout
        while time.monotonic() < end:
            try:
                m = self.q.get(timeout=max(0.1, end - time.monotonic()))
            except queue.Empty:
                break
            if m.get("id") == self.id:
                return m
        return None

    def init(self):
        r = self.request("initialize", {"protocolVersion": "2025-06-18", "capabilities": {},
                                        "clientInfo": {"name": "verify", "version": "0"}}, timeout=90)
        if r:
            self.send(b'{"jsonrpc":"2.0","method":"notifications/initialized"}')
        return r

    def kill(self):
        if self.proc.poll() is None:
            self.proc.kill()
        self.proc.wait(10)


@pytest.mark.timeout(180)
def test_v5_one_malformed_message_writes_one_short_stderr_line():
    """One malformed message writes one short stderr line. The root-LOGGER
    filter does not apply to records propagated from
    `mcp.server.lowlevel.server`, which wrote about 887 bytes over several
    lines for one message whose params is a list; ten of them before
    initialize deadlock a client that pipes stderr without draining it."""
    s = _Stdio()
    try:
        s.send(b'{"jsonrpc":"2.0","id":900,"method":"tools/list","params":[1,2,3]}')
        assert s.init() is not None
        time.sleep(1.0)
        data = bytes(s.err)
    finally:
        s.kill()
    assert len(data) <= 300, data[:400]
    assert data.count(b"\n") <= 1, data[:400]


# ---------------------------------------------------------------- bundled template notice


def test_v6_notice_carries_the_bundled_template_notice_slot():
    """NOTICE.md carries the notice for the bundled template, or a slot for
    it. The core wheel ships templates/default.pptx (python-pptx 1.0.2,
    MIT) and NOTICE.md is one of its license-files; the TEMPLATE_NOTICE
    copy slot is a Python string nothing renders."""
    notice = (ROOT / "NOTICE.md").read_text(encoding="utf-8")
    assert "python-pptx" in notice or "default.pptx" in notice or "TEMPLATE_NOTICE" in notice


# ---------------------------------------------------------------- CI matrix


def test_v7_ci_runs_the_full_matrix_on_demand_and_for_rc_tags():
    """tests.yml must have a path to 3.13/3.14, not run 3.12 only (family
    shape: dispatch `full: true` or a v*rc* tag)."""
    wf = (ROOT / ".github" / "workflows" / "tests.yml").read_text(encoding="utf-8")
    assert "full:" in wf and "workflow_dispatch" in wf
    assert "v*rc*" in wf
    assert '"3.13"' in wf and '"3.14"' in wf
    assert "macos" not in wf.lower()


# ---------------------------------------------------------------- child processes and stdin


def test_v8_child_processes_never_inherit_the_server_stdin(tmp_path, monkeypatch):
    """soffice, pdftoppm and tasklist must be started with stdin=DEVNULL.
    Started without stdin, they inherit the stdio server's JSON-RPC pipe,
    and over stdio soffice then blocks until its 300 s timeout (with
    stdin=DEVNULL the same call returns in about 6 s)."""
    from kitchensink4ppt.com import bridge
    from kitchensink4ppt.ops import export

    seen = []

    def fake_run(cmd, *a, **k):
        seen.append((cmd[0], k.get("stdin", "INHERITED")))
        return subprocess.CompletedProcess(cmd, 0, stdout="", stderr="")

    monkeypatch.setattr(export.subprocess, "run", fake_run)
    monkeypatch.setattr(bridge.subprocess, "run", fake_run)
    export._run_soffice(Path("soffice.exe"), ["--version"], tmp_path)
    bridge.powerpnt_pids()
    assert seen and all(stdin is subprocess.DEVNULL for _cmd, stdin in seen), seen


# ---------------------------------------------------------------- shielded future logging


def test_v9_a_timed_out_call_logs_no_shielded_future_traceback():
    """A timed-out call logs no shielded-future traceback. Under Python 3.14
    (3.12 and 3.13 shield does not log), when a worker reaches its commit
    point after the bound, the TIMEOUT OpError surfaces through
    asyncio.shield and the loop logs 'exception in shielded future' with a
    full traceback (about 3 KB per timeout on stderr; undrained, the server
    stops answering after two)."""
    logged = []

    def work():
        time.sleep(0.6)
        safesave.claim_publish("the test save")
        return {"ok": True}

    async def main():
        loop = asyncio.get_running_loop()
        loop.set_exception_handler(lambda _l, ctx: logged.append(ctx.get("message", "")))
        with pytest.raises(OpError) as caught:
            await run_bounded(work, 0.2)
        assert caught.value.code == "TIMEOUT"
        await asyncio.sleep(1.2)

    asyncio.run(main())
    assert not [m for m in logged if "shielded" in m], logged


# ---------------------------------------------------------------- exit after a timed-out call


_SLOW_LOAD = """
import importlib.abc, importlib.util, sys, time
class F(importlib.abc.MetaPathFinder):
    def find_spec(self, name, path, target=None):
        if name != "kitchensink4ppt.core.package":
            return None
        sys.meta_path.remove(self)
        spec = importlib.util.find_spec(name)
        sys.meta_path.insert(0, self)
        loader = spec.loader
        class L(importlib.abc.Loader):
            def create_module(self, s): return loader.create_module(s)
            def exec_module(self, m):
                loader.exec_module(m)
                orig = m.PptxPackage.__init__
                def slow(self, *a, _o=orig, **k):
                    time.sleep(240)
                    return _o(self, *a, **k)
                m.PptxPackage.__init__ = slow
        spec.loader = L()
        return spec
sys.meta_path.insert(0, F())
"""


@pytest.mark.timeout(420)
def test_v10_exit_after_a_timed_out_call_waits_one_bound_and_says_so_once(tmp_path):
    """main()'s finally and the atexit hook must not both run
    drain_inflight, or the process waits 2 x 60 s and logs the
    'still running' line twice."""
    hook = tmp_path / "hook"
    hook.mkdir()
    (hook / "sitecustomize.py").write_text(_SLOW_LOAD, encoding="utf-8")
    deck = _deck(tmp_path / "deck.pptx", slides=1, table=False)
    s = _Stdio({"PYTHONPATH": str(hook), "KS4P_CALL_TIMEOUT_S": "2"})
    try:
        assert s.init() is not None
        r = s.request("tools/call", {"name": "ppt_read", "arguments": {
            "op": "get_text", "args": {"file_path": str(deck), "live": "off"}}}, timeout=60)
        assert r is not None and r["result"]["isError"]
        t0 = time.monotonic()
        s.proc.stdin.close()
        s.proc.wait(300)
        took = time.monotonic() - t0
        lines = [ln for ln in bytes(s.err).decode("utf-8", "replace").splitlines() if "still running" in ln]
    finally:
        s.kill()
    assert took < 90, took
    assert len(lines) == 1, lines


# ---------------------------------------------------------------- retired pack names


_RETIRED = ("com-live", "assembly-export", "tables-charts", "review-sweeps", "enable_tools", "disable_tools")


def test_v11_no_raised_message_names_a_retired_pack_or_toggle():
    """No raised message may name a retired pack or toggle. com/bridge.py's
    DocumentLocked refusal once told the caller to 'use the com-live
    tools'; the other AST sweep only matches snake_case op-prefixed
    tokens."""
    bad = []
    for path in SRC.rglob("*.py"):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if isinstance(node, ast.Raise) and node.exc is not None:
                for child in ast.walk(node.exc):
                    if isinstance(child, ast.Constant) and isinstance(child.value, str):
                        hits = [w for w in _RETIRED if w in child.value]
                        if hits:
                            bad.append(f"{path.relative_to(ROOT)}:{node.lineno}: {hits}")
    assert not bad, bad


# ---------------------------------------------------------------- wrong-tool message


def test_v12_wrong_tool_message_names_the_tool_that_was_called():
    """WRONG_TOOL_MESSAGE's last literal must be an f-string: otherwise its
    field stays '{{tool}}' and .format leaves '{tool}' in the message."""
    plan = build(Settings())
    payload, is_error = asyncio.run(dispatch(
        plan, "ppt_read", {"op": "set_notes", "args": {}}
    ))
    assert is_error
    message = payload["error"]["message"]
    assert "{tool}" not in message and "ppt_read" in message, message


# ---------------------------------------------------------------- disconnect during a COM call


_SLOW_OPEN = """
import importlib.abc, importlib.util, json, os, sys, time
class F(importlib.abc.MetaPathFinder):
    def find_spec(self, name, path, target=None):
        if name != "kitchensink4ppt.com.bridge":
            return None
        sys.meta_path.remove(self)
        spec = importlib.util.find_spec(name)
        sys.meta_path.insert(0, self)
        loader = spec.loader
        class L(importlib.abc.Loader):
            def create_module(self, s): return loader.create_module(s)
            def exec_module(self, m):
                loader.exec_module(m)
                orig = m.open_presentation
                def slow(*a, _o=orig, _m=m, **k):
                    r = _o(*a, **k)
                    # The PowerPoint processes this server started, as its own
                    # ownership journal records them (pid: creation time).
                    owned = {str(pid): ct for d in list(_m._OWNED_JOURNAL.values())
                             for pid, ct in d.items()}
                    out = os.environ["KS4_V13_OWNED"]
                    with open(out + ".tmp", "w", encoding="utf-8") as f:
                        json.dump({"server_pid": os.getpid(), "owned": owned}, f)
                    os.replace(out + ".tmp", out)
                    time.sleep(400)
                    return r
                m.open_presentation = slow
        spec.loader = L()
        return spec
sys.meta_path.insert(0, F())
"""


def _end_if_same(pid: int, created: int, exe: str = "powerpnt.exe") -> bool:
    """End one process only when a handle to it shows the recorded creation
    time and executable: a recycled pid, or any other process, is left
    alone. Returns whether it was ended."""
    import ctypes
    from ctypes import wintypes

    from kitchensink4ppt.com import bridge

    k32 = bridge._kernel32()
    k32.TerminateProcess.argtypes = [wintypes.HANDLE, wintypes.UINT]
    handle = k32.OpenProcess(0x0001 | 0x1000, False, int(pid))
    if not handle:
        return False
    try:
        if bridge._creation_time(pid, handle=handle) != created:
            return False
        image = ctypes.create_unicode_buffer(1024)
        size = wintypes.DWORD(len(image))
        if not k32.QueryFullProcessImageNameW(handle, 0, image, ctypes.byref(size)):
            return False
        if not image.value.lower().endswith("\\" + exe):
            return False
        return bool(k32.TerminateProcess(handle, 1))
    finally:
        k32.CloseHandle(handle)


@pytest.mark.skipif(sys.platform != "win32", reason="Windows process handles")
def test_v13_cleanup_ends_only_the_process_it_recorded(tmp_path):
    """test_v13's cleanup ends a process only by its pid AND its creation
    time: a process with another creation time (a recycled pid, or one this
    test did not start) is left running."""
    from kitchensink4ppt.com import bridge

    child = subprocess.Popen([sys.executable, "-c", "import time; time.sleep(120)"],
                             stdin=subprocess.DEVNULL,
                             creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    try:
        created = bridge._creation_time(child.pid)
        assert created
        exe = Path(sys.executable).name.lower()
        assert _end_if_same(child.pid, created + 1, exe) is False
        assert _end_if_same(child.pid, created, "powerpnt.exe") is False
        assert child.poll() is None
        assert _end_if_same(child.pid, created, exe) is True
        assert child.wait(30) is not None
    finally:
        if child.poll() is None:
            child.kill()
            child.wait(10)


@pytest.mark.timeout(600)
@pytest.mark.skipif(os.environ.get("KS4_VERIFY_PPT_COM") != "1",
                    reason="needs a window with no PowerPoint running: set KS4_VERIFY_PPT_COM=1")
def test_v13_disconnect_during_a_com_call_leaves_no_powerpoint_this_server_started(tmp_path):
    """A client disconnecting during a COM call leaves no PowerPoint this
    server started. A validate call is inside a hidden PowerPoint the
    server started when the client disconnects; after the bounded exit wait
    the daemon workers die, and the bridge's Quit must still happen, or
    that POWERPNT.EXE /AUTOMATION survives the server.

    Only the PowerPoint this server started is watched and, if it survives,
    ended: the server's own ownership journal names it (a hook in the server
    writes it out), and it is matched by pid AND creation time. A PowerPoint
    that appeared during the run for any other reason (the user's, another
    session's) is never touched. The parent chain cannot name it: COM
    starts PowerPoint from its own service process, not from the server."""
    from kitchensink4ppt.com import bridge

    before = bridge.powerpnt_pids()
    if before is None or before:
        pytest.skip("a PowerPoint is already running; this test only runs on a clean machine")
    hook = tmp_path / "hook"
    hook.mkdir()
    (hook / "sitecustomize.py").write_text(_SLOW_OPEN, encoding="utf-8")
    owned_file = tmp_path / "owned.json"
    deck = _deck(tmp_path / "deck.pptx", slides=1, table=False)
    s = _Stdio({"PYTHONPATH": str(hook), "KS4_V13_OWNED": str(owned_file)})
    ours: dict[int, int] = {}
    left: dict[int, int] = {}
    try:
        assert s.init() is not None
        s.id += 1
        s.send(json.dumps({"jsonrpc": "2.0", "id": s.id, "method": "tools/call", "params": {
            "name": "ppt_deck", "arguments": {"op": "validate", "args": {"file_path": str(deck)}}}}).encode())
        end = time.monotonic() + 120
        while time.monotonic() < end and not owned_file.exists():
            time.sleep(1)
        assert owned_file.exists(), "the server never opened the deck in a PowerPoint"
        # Only this test's server has KS4_V13_OWNED, so the file is its own
        # journal (record["server_pid"] is the server process itself, which a
        # venv's launcher starts as its child).
        record = json.loads(owned_file.read_text(encoding="utf-8"))
        ours = {int(pid): int(ct) for pid, ct in record["owned"].items()}
        if not ours:
            pytest.skip("another PowerPoint started first and the server attached to it; "
                        "this test needs the server to start its own")
        assert all(bridge._creation_time(pid) == ct for pid, ct in ours.items()), ours
        s.proc.stdin.close()
        s.proc.wait(300)
        time.sleep(20)
        left = {pid: ct for pid, ct in ours.items() if bridge._creation_time(pid) == ct}
    finally:
        s.kill()
        for pid, ct in ours.items():
            _end_if_same(pid, ct)
    assert not left, f"PowerPoint started by this server survived its exit: {sorted(left)}"


# ------------------------------------------------- guards no other test pinned


def test_m05_a_wrong_tool_call_with_valid_args_never_runs_the_op(tmp_path):
    """Removing the wrong-tool check in surface/dispatch.resolve must be
    caught. Other surface tests missed it because their wrong-tool calls
    also fail argument validation."""
    deck = _deck(tmp_path / "deck.pptx", slides=2, table=False)
    before = _sha(deck)
    plan = build(Settings())
    payload, is_error = asyncio.run(dispatch(plan, "ppt_inspect", {
        "op": "set_slide_hidden", "args": {"file_path": str(deck), "slide": 0, "hidden": True}}))
    assert is_error and payload["error"]["code"] == "BAD_ARGS"
    assert payload["error"]["details"].get("tool") == "ppt_edit"
    assert _sha(deck) == before


def test_m07_an_expired_deadline_refuses_the_commit_without_a_cancel():
    """try_commit must check the wall clock, not only the cancel flag.
    Other tests missed this because they all cancel before committing."""
    from kitchensink4ppt.core.deadline import Deadline

    d = Deadline(0.05)
    time.sleep(0.15)
    assert d.try_commit() is False
    assert d.committed is False


def test_m17_help_answers_obey_the_output_cap():
    """ppt_help must go through budget.enforce, so its answers obey the
    output cap."""
    from kitchensink4ppt.surface.tools import make_tool

    plan = build(Settings(max_output_chars=2000))
    tool = make_tool(plan.tools["ppt_help"], plan)
    for args in ({}, {"op": "apply_edits"}, {"op": "generate_diagram"}):
        result = asyncio.run(tool.run(args))
        text = result.content[0].text
        assert len(text) <= 2000, (args, len(text))


def test_m18_a_wire_cut_of_an_unpaged_list_continues_from_the_requested_offset():
    """A wire cut must not ignore requested_offset and restart at 0: a
    caller following next_offset would loop over the same records."""
    result = {"items": [{"n": i, "pad": "p" * 300} for i in range(40)]}
    out = budget.enforce(result, 4000, wrap=lambda r: {**r, "ok": True},
                         supports_offset=True, requested_offset=10)
    page = out["page"]
    kept = len(out["items"])
    assert 0 < kept < 40
    assert page["offset"] == 10
    assert page["next_offset"] == 10 + kept


# ---------------------------------------------------------------- star nudge wiring


def test_v14_the_released_star_nudge_is_still_wired_into_startup():
    """main() calls core.star_nudge.announce_once() after a successful
    start. If the call is dropped, the module and its unit tests stay green
    while the feature is dead, so the wiring itself is pinned here."""
    src = (SRC / "server.py").read_text(encoding="utf-8")
    tree = ast.parse(src)
    main = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == "main")
    assert "announce_once" in ast.get_source_segment(src, main)


def test_m19_a_writing_op_in_a_read_only_tool_is_a_build_error():
    """check_registry must compare the read-only/writing group, so a
    writing op in a read-only tool is a build error."""
    from kitchensink4ppt.registry import OpSpec, StrictArgs
    from kitchensink4ppt.surface.plan import PlanError, check_registry

    spec = OpSpec("get_text", "ppt_read", "Read text.", StrictArgs, lambda **k: {},
                  True, False, False, False)
    with pytest.raises(PlanError):
        check_registry({"get_text": spec})
