"""Shrinking a group can shrink its text and its lines with it.

PowerPoint scales a group's members through the group transform, but font
sizes (a:rPr sz) and line widths (a:ln w) are absolute: a figure placed at a
third of its size kept full-size labels, which wrapped, and full-weight
lines. set_shape takes scale_text and scale_lines beside w and h, and
no_wrap for the text boxes inside. Without them a resize behaves exactly as
before.
"""

from __future__ import annotations

import asyncio

import pytest
from lxml import etree
from pptx import Presentation

from kitchensink4ppt import server  # noqa: F401  (the import registers ops)
from kitchensink4ppt.core.package import PptxPackage, qn
from kitchensink4ppt.ops import shapes as shp
from kitchensink4ppt.registry import public_ops
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.help import help_payload
from kitchensink4ppt.surface.plan import build


def _call(tool: str, op: str, args: dict):
    plan = build(Settings(), public_ops())
    return asyncio.run(dispatch(plan, tool, {"op": op, "args": args}))


def _result(payload: dict) -> dict:
    inner = payload.get("changed")
    return inner if isinstance(inner, dict) else payload


def _ok(tool: str, op: str, args: dict) -> dict:
    payload, is_error = _call(tool, op, args)
    assert is_error is False, payload
    return payload


@pytest.fixture
def figure(tmp_path) -> dict:
    """A group "Figure": a box with an 18 pt label and a 2 pt outline, a
    text box with a 14 pt run and a run that inherits its size, an arrow
    with a 3 pt line, and a box drawn with the theme's line."""
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    deck = str(tmp_path / "scale.pptx")
    prs.save(deck)
    box = _result(_ok("ppt_graphics", "insert_shape", {
        "file_path": deck, "slide": 0, "shape_type": "rect", "name": "Box",
        "x": 1, "y": 1, "w": 4, "h": 2, "text": "Label",
        "text_style": {"size": 18}, "line": {"color": "#000000", "width": 2}}))["shape_id"]
    note = _result(_ok("ppt_edit", "insert_textbox", {
        "file_path": deck, "slide": 0, "text": "Fourteen", "x": 1, "y": 3.5,
        "w": 3, "h": 0.6, "size_pt": 14}))
    note_id = note.get("shape_id") or note.get("id")
    _ok("ppt_graphics", "set_shape", {
        "file_path": deck, "slide": 0, "shape": note_id, "name": "Note"})
    themed = _result(_ok("ppt_graphics", "insert_shape", {
        "file_path": deck, "slide": 0, "shape_type": "ellipse", "name": "Themed",
        "x": 5.5, "y": 1, "w": 1, "h": 1}))["shape_id"]
    arrow = _result(_ok("ppt_graphics", "insert_connector", {
        "file_path": deck, "slide": 0, "start": [1, 5], "end": [5, 5],
        "line": {"color": "#333333", "width": 3}}))
    arrow_id = arrow.get("shape_id") or arrow.get("connector_id")
    group = _result(_ok("ppt_graphics", "group_shapes", {
        "file_path": deck, "slide": 0, "ids": [box, note_id, themed, arrow_id],
        "name": "Figure"}))["group_id"]
    # A second run in the text box with no size of its own.
    pkg = PptxPackage(deck)
    part = pkg.slide_parts()[0]
    elem, _chain = shp._find_shape(pkg, part, note_id)
    para = elem.find(f"{qn('p:txBody')}/{qn('a:p')}")
    run = etree.SubElement(para, qn("a:r"))
    etree.SubElement(run, qn("a:rPr")).set("lang", "en-US")
    etree.SubElement(run, qn("a:t")).text = " plain"
    end = para.find(qn("a:endParaRPr"))
    if end is not None:
        para.remove(end)
        para.append(end)
    pkg.mark_dirty(part)
    pkg.save()
    return {"deck": deck, "box": box, "note": note_id, "themed": themed,
            "arrow": arrow_id, "group": group}


def _elem(deck, ref):
    pkg = PptxPackage(deck)
    return shp._find_shape(pkg, pkg.slide_parts()[0], ref)[0]


def _sizes(deck, ref) -> list:
    elem = _elem(deck, ref)
    return [r.get("sz") for r in elem.iter(qn("a:rPr"))]


def _line_w(deck, ref):
    ln = _elem(deck, ref).find(f"{qn('p:spPr')}/{qn('a:ln')}")
    return None if ln is None else ln.get("w")


def test_a_plain_resize_keeps_text_and_lines_as_they_were(figure):
    deck = figure["deck"]
    before = (_sizes(deck, figure["box"]), _sizes(deck, figure["note"]),
              _line_w(deck, figure["arrow"]))
    payload = _ok("ppt_graphics", "set_shape", {
        "file_path": deck, "slide": 0, "shape": "Figure", "w": 2.5})
    assert (_sizes(deck, figure["box"]), _sizes(deck, figure["note"]),
            _line_w(deck, figure["arrow"])) == before
    assert "text_scaled" not in _result(payload)


def test_scale_text_shrinks_every_size_by_the_resize(figure):
    deck = figure["deck"]
    pkg = PptxPackage(deck)
    group = shp._find_shape(pkg, pkg.slide_parts()[0], figure["group"])[0]
    width = int(group.find(f"{qn('p:grpSpPr')}/{qn('a:xfrm')}/{qn('a:ext')}").get("cx"))
    payload = _ok("ppt_graphics", "set_shape", {
        "file_path": deck, "slide": 0, "shape": "Figure",
        "w": width / 914400 / 2, "h": None, "scale_text": True})
    result = _result(payload)
    assert result["text_scaled"]["factor"] == 0.5
    assert set(_sizes(deck, figure["box"])) == {"900"}
    note_sizes = _sizes(deck, figure["note"])
    # The 14 pt run, the run that inherited 18 pt and the end mark.
    assert note_sizes[0] == "700"
    assert note_sizes[1] == "900"
    body = _elem(deck, figure["note"]).find(f"{qn('p:txBody')}/{qn('a:bodyPr')}")
    assert body.get("lIns") == str(91440 // 2) and body.get("tIns") == str(45720 // 2)
    assert _line_w(deck, figure["arrow"]) == str(3 * 12700)


def test_scale_lines_thins_explicit_and_theme_lines(figure):
    deck = figure["deck"]
    _ok("ppt_graphics", "set_shape", {
        "file_path": deck, "slide": 0, "shape": "Figure", "w": 3.25,
        "scale_lines": True})
    pkg = PptxPackage(deck)
    group = shp._find_shape(pkg, pkg.slide_parts()[0], figure["group"])[0]
    xfrm = group.find(f"{qn('p:grpSpPr')}/{qn('a:xfrm')}")
    factor = int(xfrm.find(qn("a:ext")).get("cx")) / int(xfrm.find(qn("a:chExt")).get("cx"))
    assert int(_line_w(deck, figure["arrow"])) == round(3 * 12700 * factor)
    assert int(_line_w(deck, figure["box"])) == round(2 * 12700 * factor)
    themed = _line_w(deck, figure["themed"])
    assert themed is not None and 0 < int(themed) < 12700 * 2
    assert set(_sizes(deck, figure["box"])) == {"1800"}


def test_the_smaller_ratio_is_the_factor_when_only_one_side_changes(figure):
    deck = figure["deck"]
    pkg = PptxPackage(deck)
    group = shp._find_shape(pkg, pkg.slide_parts()[0], figure["group"])[0]
    ext = group.find(f"{qn('p:grpSpPr')}/{qn('a:xfrm')}/{qn('a:ext')}")
    height = int(ext.get("cy")) / 914400
    payload = _ok("ppt_graphics", "set_shape", {
        "file_path": deck, "slide": 0, "shape": "Figure",
        "h": height / 4, "scale_text": True})
    assert _result(payload)["text_scaled"]["factor"] == 0.25
    assert set(_sizes(deck, figure["box"])) == {"450"}


def test_no_wrap_stops_the_text_boxes_wrapping(figure):
    deck = figure["deck"]
    payload = _ok("ppt_graphics", "set_shape", {
        "file_path": deck, "slide": 0, "shape": "Figure", "w": 3,
        "no_wrap": True})
    note = _elem(deck, figure["note"]).find(f"{qn('p:txBody')}/{qn('a:bodyPr')}")
    assert note.get("wrap") == "none"
    box = _elem(deck, figure["box"]).find(f"{qn('p:txBody')}/{qn('a:bodyPr')}")
    assert box.get("wrap") != "none"
    assert _result(payload)["unwrapped"] == [figure["note"]]


def test_scaling_without_a_new_size_is_refused(figure):
    payload, is_error = _call("ppt_graphics", "set_shape", {
        "file_path": figure["deck"], "slide": 0, "shape": "Figure",
        "x": 0.5, "scale_text": True})
    assert is_error is True
    assert "w or h" in payload["error"]["message"]


@pytest.mark.parametrize("size", [{"w": 0}, {"h": 0}, {"h": 0.0}, {"w": 1e-9}])
@pytest.mark.parametrize("switch", ["scale_text", "scale_lines"])
def test_a_zero_size_never_scales_text_or_lines_to_nothing(figure, size, switch):
    """w=0 or h=0 with scale_text was accepted; the factor was 0.0 and
    every run was written at the 1 pt floor (scale_lines zeroed the lines). A
    model that sends h=0 meaning "keep the height" must not lose the text
    sizes: the call is refused, names how to keep a side, and the deck is
    unchanged."""
    deck = figure["deck"]
    with open(deck, "rb") as f:
        before = f.read()
    payload, is_error = _call("ppt_graphics", "set_shape", {
        "file_path": deck, "slide": 0, "shape": "Figure", switch: True, **size})
    assert is_error is True, payload
    message = payload["error"]["message"]
    assert "leave" in message and "out" in message, message
    with open(deck, "rb") as f:
        assert f.read() == before
    assert set(_sizes(deck, figure["box"])) == {"1800"}


def test_a_zero_size_without_scaling_is_still_a_size(figure):
    """A straight line is a shape with no width or no height, so w=0 or h=0
    on its own stays a plain resize, as before."""
    deck = figure["deck"]
    _ok("ppt_graphics", "set_shape", {
        "file_path": deck, "slide": 0, "shape": figure["arrow"], "h": 0})
    ext = _elem(deck, figure["arrow"]).find(
        f"{qn('p:spPr')}/{qn('a:xfrm')}/{qn('a:ext')}")
    assert ext.get("cy") == "0"
    payload, is_error = _call("ppt_graphics", "set_shape", {
        "file_path": deck, "slide": 0, "shape": figure["arrow"], "w": -1})
    assert is_error is True and "negative" in payload["error"]["message"]


def test_a_single_text_box_scales_its_own_text(figure):
    deck = figure["deck"]
    _ok("ppt_graphics", "set_shape", {
        "file_path": deck, "slide": 0, "shape": "Note", "w": 1.5,
        "scale_text": True})
    assert _sizes(deck, figure["note"])[0] == "700"


def test_help_describes_the_scaling_switches():
    plan = build(Settings(), public_ops())
    props = help_payload(plan, "set_shape")["args_schema"]["properties"]
    for name in ("scale_text", "scale_lines", "no_wrap"):
        assert name in props and props[name].get("description")
