"""Family-surface operation sweep and output-guard gate.

Every registered operation receives a schema-valid, semantically meaningful
example against a real canary presentation. Each operation works on its own
copy so one destructive call cannot make a later example meaningless.
"""

from __future__ import annotations

import asyncio
import base64
import hashlib
import io
import os
import shutil
import struct
import wave
from pathlib import Path

import pytest

from kitchensink4ppt import server
from kitchensink4ppt.com import bridge
from kitchensink4ppt.core import safesave
from kitchensink4ppt.registry import public_ops
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.help import example
from kitchensink4ppt.surface.plan import build


CANARY = "K9QCANARY"
MUST_SUCCEED = frozenset({
    "create_presentation", "get_presentation_view", "get_presentation_info",
    "get_server_info", "get_workflows", "get_export_engines", "get_text",
})

# Explicit per-operation environmental/state allowlists. Argument failures
# (BAD_ARGS/BAD_PARAMS) and INTERNAL_ERROR are intentionally absent: either
# means the sweep did not reach a valid operation body.
_STATE_REFUSALS = frozenset({
    "NOT_FOUND", "UNSUPPORTED_CONTENT", "APP_NOT_RUNNING", "APP_BUSY",
    "APP_BLOCKED", "DOCUMENT_LOCKED", "CONFLICT", "HAZARD_REFUSED",
    "VALIDATION_FAILED", "STALE_ANCHOR", "CALC_UNAVAILABLE",
})
READER_OK_REFUSALS = {
    name: _STATE_REFUSALS
    for name, spec in public_ops().items()
    if spec.read_only
}
WRITER_OK_REFUSALS = {
    name: _STATE_REFUSALS
    for name, spec in public_ops().items()
    if not spec.read_only
}

# These operations require an installed desktop application rather than only
# the package under test. Hosted runners intentionally do not install
# PowerPoint or LibreOffice. Accept only the operation-specific dependency
# refusal; any other BAD_PARAMS still proves that the example is defective.
_DEPENDENCY_REFUSALS = {
    "export_handout": ("BAD_PARAMS", "needs PowerPoint COM"),
    "export_pdf": ("BAD_PARAMS", "soffice was not found"),
    "export_slide_image": ("BAD_PARAMS", "PowerPoint COM is not available"),
    "export_shape_image": ("BAD_PARAMS", "needs PowerPoint"),
    "live_save": ("BAD_PARAMS", "pywin32 is not available"),
    "live_scroll_to": ("BAD_PARAMS", "pywin32 is not available"),
}


def _canary_deck(path: Path) -> Path:
    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = f"{CANARY} title"
    slide.placeholders[1].text = f"{CANARY} body"
    slide.shapes.add_textbox(Inches(1), Inches(4), Inches(2), Inches(0.5)).text = "third"

    table_slide = prs.slides.add_slide(prs.slide_layouts[5])
    table_slide.shapes.title.text = "Table"
    table = table_slide.shapes.add_table(
        2, 2, Inches(1), Inches(2), Inches(4), Inches(1)
    ).table
    for r, row in enumerate((("Name", "Qty"), (CANARY, "1"))):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    chart_slide = prs.slides.add_slide(prs.slide_layouts[5])
    chart_slide.shapes.title.text = "Chart"
    data = ChartData()
    data.categories = ["A", "B"]
    data.add_series("Qty", [1, 2])
    chart_slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1), Inches(2), Inches(5), Inches(3), data,
    )

    image_slide = prs.slides.add_slide(prs.slide_layouts[5])
    image_slide.shapes.title.text = "Image"
    png = path.with_name("canary.png")
    png.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    ))
    image_slide.shapes.add_picture(str(png), Inches(1), Inches(2), Inches(1), Inches(1))
    prs.save(path)
    png2 = path.with_name("canary-replacement.png")
    png2.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZK0sAAAAASUVORK5CYII="
    ))
    wav_path = path.with_name("sample.wav")
    buffer = io.BytesIO()
    with wave.open(buffer, "wb") as wav:
        wav.setnchannels(1)
        wav.setsampwidth(2)
        wav.setframerate(8000)
        wav.writeframes(b"\x00\x00" * 200)
    wav_path.write_bytes(buffer.getvalue())
    mp4_path = path.with_name("sample.mp4")
    mp4_path.write_bytes(struct.pack(">I", 24) + b"ftypisom" + b"\x00" * 16)
    server.insert_audio(str(path), 0, str(wav_path), 6.0, 1.0)
    return path


def _shape_ids(deck: Path, slide: int) -> list[int]:
    info = server.get_slide_info(str(deck), slide, live="off")
    return [int(item["id"]) for item in info["shapes"]]


def _typed_shape_id(deck: Path, slide: int, needle: str) -> int:
    info = server.get_slide_info(str(deck), slide, live="off")
    for item in info["shapes"]:
        haystack = f"{item.get('type') or ''} {item.get('name') or ''}".lower()
        if needle.lower() in haystack:
            return int(item["id"])
    return int(info["shapes"][-1]["id"])


_TABLE_OPS = {
    "set_table_cells", "merge_cells", "unmerge_cells", "insert_table_rows",
    "delete_table_rows", "insert_table_cols", "delete_table_cols",
    "format_table_cells", "set_column_widths", "set_row_heights",
    "apply_table_style", "get_table", "export_table",
}
_CHART_OPS = {"update_chart_data", "format_chart", "get_chart_data"}
_IMAGE_OPS = {"replace_image", "set_image"}


def _args_for(name: str, work: Path, source: Path) -> dict:
    spec = public_ops()[name]
    args = example(spec)
    deck = work / f"{name}.pptx"
    if name != "create_presentation":
        shutil.copy2(source, deck)

    if "file_path" in args:
        args["file_path"] = str(deck)
    if "path" in args and name == "create_presentation":
        args["path"] = str(work / "created.pptx")
    if name == "compare_decks":
        other = work / "compare-source.pptx"
        shutil.copy2(source, other)
        args.update(file_path_a=str(deck), file_path_b=str(other))
    if "source_path" in args:
        other = work / f"{name}-source.pptx"
        shutil.copy2(source, other)
        args["source_path"] = str(other)
    if "sources" in args:
        other = work / f"{name}-source.pptx"
        shutil.copy2(source, other)
        args["sources"] = [str(other)]

    png = source.with_name("canary.png")
    png2 = source.with_name("canary-replacement.png")
    csv = work / "canary.csv"
    csv.write_text("Name,Qty\nCanary,1\n", encoding="utf-8")
    for key in ("image", "old_image", "new_image", "poster"):
        if key in args and args[key] is not None:
            args[key] = str(png)
    if name == "replace_image_everywhere":
        args["old_image"], args["new_image"] = str(png), str(png2)
    if name == "import_table":
        args["source"] = str(csv)
    for key in ("video", "audio"):
        if key in args:
            args[key] = str(source.with_name("sample.mp4" if key == "video" else "sample.wav"))

    slide = 1 if name in _TABLE_OPS else 2 if name in _CHART_OPS else 3 if name in _IMAGE_OPS else 0
    if "slide" in args and name not in {"delete_slide", "duplicate_slide", "move_slide", "set_slide_hidden"}:
        args["slide"] = slide
    ids = _shape_ids(deck, slide) if deck.exists() else []
    text_ids = _shape_ids(deck, 0) if deck.exists() else []
    if "shape" in args and name not in _IMAGE_OPS:
        args["shape"] = text_ids[1] if len(text_ids) > 1 else text_ids[0]
    if name in _IMAGE_OPS:
        args["shape"] = _typed_shape_id(deck, 3, "picture")
    if name == "copy_shape_between":
        args["shape"] = text_ids[2]  # the text box: placeholders do not travel
    if name in {"get_media_playback", "set_media_playback"}:
        args["slide"] = 0
        args["shape"] = _typed_shape_id(deck, 0, "audio")
    if "table" in args:
        args["table"] = _typed_shape_id(deck, 1, "table")
    if "chart" in args:
        args["chart"] = _typed_shape_id(deck, 2, "chart")
    if "from_shape" in args:
        args["from_shape"] = text_ids[0]
    if "to_shapes" in args:
        args["to_shapes"] = [text_ids[1]]
    if "ids" in args:
        count = 3 if name == "distribute_shapes" else 2
        args["ids"] = text_ids[:count]
    if name == "set_reading_order":
        args["order"] = text_ids
    if name == "reorder_slides":
        args["order"] = [3, 2, 1, 0]
    if name == "apply_edits":
        args["edits"] = [{"op": "set_shape", "slide": 0, "shape": text_ids[0], "dx": 0.1}]

    file_outputs = {
        "copy_presentation": ("dest_path", "copy.pptx"),
        "export_pdf": ("output", "export.pdf"),
        "export_handout": ("output", "handout.pdf"),
        "export_table": ("path", "table.csv"),
        "export_shape_image": ("output", "shape.png"),
    }
    dir_outputs = {
        "create_snapshot": "dest_dir",
        "split_deck": "output_dir",
        "export_slide_image": "output_dir",
    }
    if name in file_outputs:
        key, filename = file_outputs[name]
        args[key] = str(work / f"{name}-{filename}")
    if name in dir_outputs:
        args[dir_outputs[name]] = str(work / f"{name}-output")
    if name == "export_pdf":
        args["engine"] = "libreoffice"
    if name == "export_slide_image":
        args["engine"] = "com"
    return args


@pytest.fixture(scope="module")
def sweep(tmp_path_factory):
    running = bridge.powerpnt_pids()
    if running is None:
        pytest.skip("the PowerPoint process table could not be read safely")
    if running:
        pytest.skip("the operation sweep requires a clean PowerPoint window")
    work = tmp_path_factory.mktemp("ppt-op-sweep")
    source = _canary_deck(work / "canary-source.pptx")
    plan = build(Settings(call_timeout_s=900.0))
    results = {}
    for name, spec in sorted(plan.specs.items()):
        args = _args_for(name, work, source)
        payload, is_error = asyncio.run(
            dispatch(plan, spec.group, {"op": name, "args": args})
        )
        results[name] = (payload, is_error, args)
    yield plan, work, source, results
    left = bridge.powerpnt_pids()
    assert left == set(), f"the operation sweep left PowerPoint running: {left}"


@pytest.mark.live
@pytest.mark.skipif(
    os.environ.get("KS4_SWEEP_PPT_COM") != "1",
    reason="set KS4_SWEEP_PPT_COM=1 in a clean PowerPoint window",
)
def test_every_operation_runs_with_real_arguments(sweep):
    plan, _work, _source, results = sweep
    bad = []
    for name, (payload, is_error, _args) in results.items():
        if not is_error:
            continue
        code = payload.get("error", {}).get("code")
        message = payload.get("error", {}).get("message", "")
        dependency_refusal = _DEPENDENCY_REFUSALS.get(name)
        if (
            dependency_refusal
            and code == dependency_refusal[0]
            and dependency_refusal[1] in message
        ):
            continue
        allowed = (
            READER_OK_REFUSALS if plan.specs[name].read_only
            else WRITER_OK_REFUSALS
        )
        if name in MUST_SUCCEED or code not in allowed[name]:
            bad.append(f"{name}: {code}: {message[:180]}")
    assert bad == [], (
        "the sweep did not reach a valid operation body:\n" + "\n".join(bad)
    )
    assert set(results) == set(plan.specs)


def _digest(path: Path) -> str | None:
    return hashlib.sha256(path.read_bytes()).hexdigest() if path.is_file() else None


def test_dg12_calls_each_output_site_with_source_and_backup_targets(tmp_path):
    source = _canary_deck(tmp_path / "guarded.pptx")
    server.set_slide_hidden(str(source), 0, True)
    slot = safesave.slot_dir(source) / safesave.PREV_SLOT
    assert slot.is_file()
    before = (_digest(source), _digest(slot))
    plan = build(Settings())
    table = _typed_shape_id(source, 1, "table")
    cases = {
        "copy_presentation": ("dest_path", {}),
        "export_pdf": ("output", {"engine": "libreoffice"}),
        "export_handout": ("output", {}),
        "export_table": ("path", {"slide": 1, "table": table, "format": "csv"}),
    }
    for name, (key, extra) in cases.items():
        for target in (source, slot):
            args = {"file_path": str(source), key: str(target), **extra}
            payload, is_error = asyncio.run(dispatch(
                plan, plan.specs[name].group, {"op": name, "args": args}
            ))
            assert is_error, (name, target, payload)
    for name, key in {
        "create_snapshot": "dest_dir",
        "split_deck": "output_dir",
        "export_slide_image": "output_dir",
    }.items():
        args = {"file_path": str(source), key: str(slot.parent)}
        if name == "export_slide_image":
            args["engine"] = "libreoffice"
        payload, is_error = asyncio.run(dispatch(
            plan, plan.specs[name].group, {"op": name, "args": args}
        ))
        assert is_error, (name, payload)
    assert (_digest(source), _digest(slot)) == before


def test_restore_target_inside_backup_store_is_refused_without_change(tmp_path):
    source = _canary_deck(tmp_path / "restore-guarded.pptx")
    server.set_slide_hidden(str(source), 0, True)
    store = tmp_path / safesave.BACKUP_DIR_NAME
    before = {
        path: _digest(path)
        for path in store.rglob("*.pptx")
        if path.is_file()
    }
    target = safesave.slot_dir(source) / safesave.PREV_SLOT
    with pytest.raises(Exception, match="backup folder"):
        server.manage_backups(
            "restore", file_path=str(target), source=str(source)
        )
    after = {
        path: _digest(path)
        for path in store.rglob("*.pptx")
        if path.is_file()
    }
    assert after == before


def test_extensionless_pdf_output_is_settled_before_existing_file_check(tmp_path):
    source = _canary_deck(tmp_path / "pdf-guarded.pptx")
    existing = tmp_path / "report.pdf"
    existing.write_bytes(b"%PDF-1.4 user-owned\n")
    before = _digest(existing)
    with pytest.raises(Exception, match="already exists"):
        server.export_pdf(
            str(source), output=str(tmp_path / "report"), engine="com"
        )
    assert _digest(existing) == before


def test_trailing_dot_pdf_output_is_normalized_before_existing_file_check(tmp_path):
    source = _canary_deck(tmp_path / "pdf-trailing-dot.pptx")
    existing = tmp_path / "report.pdf"
    existing.write_bytes(b"%PDF-1.4 user-owned\n")
    before = _digest(existing)
    with pytest.raises(Exception, match="already exists"):
        server.export_pdf(
            str(source), output=str(tmp_path / "report."), engine="com"
        )
    assert _digest(existing) == before


def test_com_pdf_path_normalizer_treats_trailing_dot_as_extensionless(tmp_path):
    from kitchensink4ppt.com import bridge

    source = _canary_deck(tmp_path / "bridge-trailing-dot.pptx")
    assert bridge._settled_pdf_output(
        source, str(tmp_path / "report."), source.with_suffix(".pdf")
    ) == tmp_path / "report.pdf"
