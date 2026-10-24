"""Regressions for table selector ambiguity and layout guidance."""

from __future__ import annotations

import asyncio

from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt import server  # noqa: F401  (registers public operations)
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.plan import build


PLAN = build(Settings())


def _call(tool: str, op: str, **args):
    return asyncio.run(dispatch(PLAN, tool, {"op": op, "args": args}))


def _four_tables(path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for i in range(4):
        table_shape = slide.shapes.add_table(
            2, 2, Inches(0.3 + 2.4 * i), Inches(1), Inches(2), Inches(1)
        )
        table_shape.table.cell(0, 0).text = f"Table {i}"
    prs.save(path)


def test_bare_integer_returns_ambiguous_location(tmp_path):
    """A bare integer is ambiguous: 2 means index 2 and shape id 2."""
    deck = tmp_path / "tables.pptx"
    _four_tables(deck)

    payload, is_error = _call(
        "ppt_read", "get_table", file_path=str(deck), slide=0, table=2
    )

    assert is_error
    assert payload["error"]["code"] == "AMBIGUOUS_LOCATION", payload
    assert "{'index': 2}" in payload["error"]["message"]
    assert "{'shape_id': 2}" in payload["error"]["message"]


def test_ambiguous_edit_refuses_without_touching_the_deck(tmp_path):
    """Following a bare selector can no longer silently edit the wrong table."""
    deck = tmp_path / "tables.pptx"
    _four_tables(deck)
    before = deck.read_bytes()

    payload, is_error = _call(
        "ppt_tables",
        "format_table_cells",
        file_path=str(deck),
        slide=0,
        table=2,
        fill="FF0000",
    )

    assert is_error
    assert payload["error"]["code"] == "AMBIGUOUS_LOCATION", payload
    assert deck.read_bytes() == before


def test_explicit_index_and_shape_id_choose_the_requested_tables(tmp_path):
    deck = tmp_path / "tables.pptx"
    _four_tables(deck)

    by_index, index_error = _call(
        "ppt_read", "get_table", file_path=str(deck), slide=0,
        table={"index": 2},
    )
    by_shape, shape_error = _call(
        "ppt_read", "get_table", file_path=str(deck), slide=0,
        table={"shape_id": 2},
    )

    assert not index_error, by_index
    assert not shape_error, by_shape
    assert by_index["shape_id"] == 4
    assert by_shape["table_index"] == 0
    assert by_index["shape_id"] != by_shape["shape_id"]


def test_check_layout_table_fix_names_shape_id_explicitly(tmp_path):
    deck = tmp_path / "tables.pptx"
    _four_tables(deck)

    payload, is_error = _call(
        "ppt_inspect", "check_layout", file_path=str(deck), slide=0
    )
    fixes = [
        finding.get("fix", "")
        for finding in payload.get("findings", [])
        if "format_table_cells" in (finding.get("fix") or "")
    ]

    assert not is_error, payload
    assert fixes, "expected table-style contrast findings with a fix call"
    assert all("'table': {'shape_id':" in fix for fix in fixes), fixes
