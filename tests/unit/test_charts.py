"""Charts: the chart builder (the same design as the Word editor's) against
pptx packaging.

Checks the two render-critical invariants: literal caches
(what renders) and the embedded workbook (what Edit Data opens) are ALWAYS
written together and always agree, and every r:id in the chart part resolves
(the dangling-rel repair trigger). python-pptx is the independent oracle for
reading created charts back.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from lxml import etree

from kitchensink4ppt.core.errors import (
    AmbiguousTarget,
    PptMcpError,
    TargetNotFound,
    UnsupportedStructure,
)
from kitchensink4ppt.core.package import PptxPackage, qn
from kitchensink4ppt.ops import chartread, charts as ch

_C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_SML = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_CT = "http://schemas.openxmlformats.org/package/2006/content-types"


def _qc(name):
    return f"{{{_C}}}{name}"


@pytest.fixture()
def deck(make_deck):
    from kitchensink4ppt.ops import slides as sl

    path = make_deck("charts.pptx", extra_slides=0)
    pkg = PptxPackage(path)
    slide = sl.insert_slide(pkg, 0)["index"]
    return pkg, slide


CATS = ["Alpha", "Beta", "Gamma"]
SERIES = [
    {"name": "One", "values": [4, 2.5, 3]},
    {"name": "Two", "values": [1, 2, 3]},
]


def _num_cache_values(ser):
    cache = ser.find(f"{_qc('val')}/{_qc('numRef')}/{_qc('numCache')}")
    return [pt.find(_qc("v")).text for pt in cache.findall(_qc("pt"))]


def _workbook_sheet(pkg, workbook_part):
    with zipfile.ZipFile(io.BytesIO(pkg.raw_part(workbook_part))) as zf:
        names = set(zf.namelist())
        assert {
            "[Content_Types].xml",
            "_rels/.rels",
            "xl/workbook.xml",
            "xl/_rels/workbook.xml.rels",
            "xl/worksheets/sheet1.xml",
        } <= names
        return etree.fromstring(zf.read("xl/worksheets/sheet1.xml"))


def _rewrite_xlsx(blob: bytes, replacements: dict[str, bytes]) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(blob))
    out = io.BytesIO()
    with source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        seen = set()
        for info in source.infolist():
            target.writestr(info, replacements.get(info.filename, source.read(info)))
            seen.add(info.filename)
        for name, data in replacements.items():
            if name not in seen:
                target.writestr(name, data)
    return out.getvalue()


def _numeric_category_chart(pkg, chart_part, categories):
    root = pkg.root(chart_part)
    for cat in root.findall(f".//{_qc('ser')}/{_qc('cat')}"):
        ref = cat.find(_qc("strRef"))
        ref.tag = _qc("numRef")
        cache = ref.find(_qc("strCache"))
        cache.tag = _qc("numCache")
        cache.insert(0, etree.Element(_qc("formatCode")))
        cache[0].text = "0"
        for pt, value in zip(cache.findall(_qc("pt")), categories):
            pt.find(_qc("v")).text = str(value)
    pkg.mark_dirty(chart_part)


class TestCreateChart:
    def test_column_chart_parts_and_caches(self, deck):
        pkg, slide = deck
        r = ch.create_chart(
            pkg, slide, "column", CATS, SERIES, 1, 1, 6, 4, title="T"
        )
        assert pkg.has_part(r["chart_part"])
        assert pkg.has_part(r["embedded_workbook"])
        root = pkg.root(r["chart_part"])
        group = root.find(f"{_qc('chart')}/{_qc('plotArea')}/{_qc('barChart')}")
        assert group.find(_qc("barDir")).get("val") == "col"
        assert group.find(_qc("grouping")).get("val") == "clustered"
        sers = group.findall(_qc("ser"))
        assert len(sers) == 2
        assert _num_cache_values(sers[0]) == ["4", "2.5", "3"]
        cat_cache = sers[0].find(
            f"{_qc('cat')}/{_qc('strRef')}/{_qc('strCache')}"
        )
        assert [p.find(_qc("v")).text for p in cat_cache.findall(_qc("pt"))] == CATS
        # Deliberate spPr absence: series follow the theme accent cycle.
        assert all(s.find(_qc("spPr")) is None for s in sers)
        # externalData is wired and its target exists.
        ext = root.find(_qc("externalData"))
        assert ext is not None
        pkg.save()

    def test_stacked_gets_overlap_100(self, deck):
        pkg, slide = deck
        r = ch.create_chart(pkg, slide, "bar_stacked", CATS, [SERIES[0]], 1, 1, 5, 3)
        group = pkg.root(r["chart_part"]).find(
            f"{_qc('chart')}/{_qc('plotArea')}/{_qc('barChart')}"
        )
        assert group.find(_qc("grouping")).get("val") == "stacked"
        assert group.find(_qc("overlap")).get("val") == "100"
        assert group.find(_qc("barDir")).get("val") == "bar"

    def test_pie_and_line_shapes(self, deck):
        pkg, slide = deck
        r1 = ch.create_chart(pkg, slide, "pie", CATS, [SERIES[0]], 0.5, 0.5, 4, 3)
        r2 = ch.create_chart(pkg, slide, "line", CATS, SERIES, 5, 0.5, 4, 3)
        pie = pkg.root(r1["chart_part"]).find(
            f"{_qc('chart')}/{_qc('plotArea')}/{_qc('pieChart')}"
        )
        assert pie is not None and pie.find(_qc("varyColors")).get("val") == "1"
        line = pkg.root(r2["chart_part"]).find(
            f"{_qc('chart')}/{_qc('plotArea')}/{_qc('lineChart')}"
        )
        assert line is not None
        assert len(line.findall(_qc("ser"))) == 2
        pkg.save()

    def test_workbook_matches_caches(self, deck):
        pkg, slide = deck
        r = ch.create_chart(pkg, slide, "column", CATS, SERIES, 1, 1, 6, 4)
        sheet = _workbook_sheet(pkg, r["embedded_workbook"])
        cells = {
            c.get("r"): c for c in sheet.iter(f"{{{_SML}}}c")
        }
        # A1 blank; B1/C1 series names; A2.. categories; B2.. values.
        assert cells["B1"].find(f"{{{_SML}}}is/{{{_SML}}}t").text == "One"
        assert cells["A2"].find(f"{{{_SML}}}is/{{{_SML}}}t").text == "Alpha"
        assert cells["B3"].find(f"{{{_SML}}}v").text == "2.5"
        assert cells["C4"].find(f"{{{_SML}}}v").text == "3"

    def test_data_refusals_fire_before_mutation(self, deck):
        pkg, slide = deck
        parts_before = set(pkg.part_names())
        with pytest.raises(PptMcpError):
            ch.create_chart(pkg, slide, "column", CATS, [{"name": "x", "values": [1]}], 1, 1, 4, 3)
        with pytest.raises(PptMcpError):
            ch.create_chart(pkg, slide, "pie", CATS, SERIES, 1, 1, 4, 3)
        with pytest.raises(PptMcpError):
            ch.create_chart(pkg, slide, "donut3d", CATS, SERIES, 1, 1, 4, 3)
        with pytest.raises(PptMcpError):
            ch.create_chart(
                pkg, slide, "line", CATS,
                [{"name": "x", "values": [1, float("nan"), 2]}], 1, 1, 4, 3,
            )
        assert set(pkg.part_names()) == parts_before

    def test_pptx_oracle_reads_chart(self, deck):
        pkg, slide = deck
        ch.create_chart(pkg, slide, "column", CATS, SERIES, 1, 1, 6, 4, title="Oracle")
        path = pkg.save()
        from pptx import Presentation

        prs = Presentation(str(path))
        frames = [s for s in prs.slides[slide].shapes if s.has_chart]
        assert len(frames) == 1
        chart = frames[0].chart
        assert list(chart.plots[0].categories) == CATS
        values = [list(s.values) for s in chart.plots[0].series]
        assert values == [[4.0, 2.5, 3.0], [1.0, 2.0, 3.0]]
        assert chart.has_title and chart.chart_title.text_frame.text == "Oracle"


class TestUpdateChartData:
    def test_numeric_categories_roundtrip_through_reader_and_update(self, deck):
        pkg, slide = deck
        years = [2015, 2016, 2017]
        r = ch.create_chart(
            pkg, slide, "line", [str(y) for y in years],
            [{"name": "Cases", "values": [4, 5, 6]}], 1, 1, 6, 4,
        )
        _numeric_category_chart(pkg, r["chart_part"], years)

        read = chartread.get_chart_data(pkg, slide, {"shape_id": r["shape_id"]})
        categories = read["groups"][0]["series"][0]["categories"]
        assert categories == years
        ch.update_chart_data(
            pkg, slide, {"shape_id": r["shape_id"]}, categories,
            [{"name": "Cases", "values": [7, 8, 9]}],
        )

        reread = chartread.get_chart_data(pkg, slide, {"shape_id": r["shape_id"]})
        assert reread["groups"][0]["series"][0]["categories"] == years
        ser = ch._plot_groups(pkg.root(r["chart_part"]))[0].find(_qc("ser"))
        assert ser.find(f"{_qc('cat')}/{_qc('numRef')}/{_qc('numCache')}") is not None
        sheet = _workbook_sheet(pkg, r["embedded_workbook"])
        cells = {c.get("r"): c for c in sheet.iter(f"{{{_SML}}}c")}
        assert cells["A2"].get("t") is None
        assert cells["A2"].find(f"{{{_SML}}}v").text == "2015"

    def test_linked_chart_is_reported_and_update_refuses_before_mutation(self, deck):
        pkg, slide = deck
        r = ch.create_chart(pkg, slide, "column", CATS, [SERIES[0]], 1, 1, 6, 4)
        rels = pkg.root(ch._rels_part_for(r["chart_part"]))
        package_rel = next(rel for rel in rels if rel.get("Type") == ch._PACKAGE_REL)
        package_rel.set("Target", "file:///C:/Users/example/linked_book.xlsx")
        package_rel.set("TargetMode", "External")
        pkg.mark_dirty(ch._rels_part_for(r["chart_part"]))

        read = chartread.get_chart_data(pkg, slide, {"shape_id": r["shape_id"]})
        assert read["data_source"] == "linked"
        assert read["linked_workbook"] == "file:///C:/Users/example/linked_book.xlsx"
        assert read["embedded_workbook"] is None
        chart_before = etree.tostring(pkg.root(r["chart_part"]))
        workbook_before = pkg.raw_part(r["embedded_workbook"])

        with pytest.raises(UnsupportedStructure, match="linked"):
            ch.update_chart_data(
                pkg, slide, {"shape_id": r["shape_id"]}, CATS,
                [{"name": "Changed", "values": [9, 8, 7]}],
            )
        assert etree.tostring(pkg.root(r["chart_part"])) == chart_before
        assert pkg.raw_part(r["embedded_workbook"]) == workbook_before

    def test_rich_embedded_workbook_is_updated_in_place(self, deck):
        pkg, slide = deck
        r = ch.create_chart(pkg, slide, "column", CATS, [SERIES[0]], 1, 1, 6, 4)
        workbook = pkg.raw_part(r["embedded_workbook"])
        with zipfile.ZipFile(io.BytesIO(workbook)) as zf:
            sheet = etree.fromstring(zf.read("xl/worksheets/sheet1.xml"))
            workbook_xml = etree.fromstring(zf.read("xl/workbook.xml"))
            workbook_rels = etree.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
            content_types = etree.fromstring(zf.read("[Content_Types].xml"))
        sheet_data = sheet.find(f"{{{_SML}}}sheetData")
        row2 = next(row for row in sheet_data if row.get("r") == "2")
        formula_cell = etree.SubElement(row2, f"{{{_SML}}}c", r="D2")
        etree.SubElement(formula_cell, f"{{{_SML}}}f").text = "B2*1.1"
        etree.SubElement(formula_cell, f"{{{_SML}}}v").text = "4.4"
        etree.SubElement(sheet, f"{{{_SML}}}legacyDrawing").set(f"{{{_R}}}id", "rId2")
        sheets = workbook_xml.find(f"{{{_SML}}}sheets")
        notes_sheet = etree.SubElement(
            sheets, f"{{{_SML}}}sheet", name="Notes", sheetId="2"
        )
        notes_sheet.set(f"{{{_R}}}id", "rId2")
        etree.SubElement(
            workbook_rels, f"{{{_REL}}}Relationship",
            Id="rId2",
            Type=(
                "http://schemas.openxmlformats.org/officeDocument/2006/"
                "relationships/worksheet"
            ),
            Target="worksheets/sheet2.xml",
        )
        etree.SubElement(
            content_types, f"{{{_CT}}}Override",
            PartName="/xl/worksheets/sheet2.xml",
            ContentType=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml."
                "worksheet+xml"
            ),
        )
        etree.SubElement(
            content_types, f"{{{_CT}}}Override",
            PartName="/xl/comments1.xml",
            ContentType=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml."
                "comments+xml"
            ),
        )
        etree.SubElement(
            content_types, f"{{{_CT}}}Default", Extension="vml",
            ContentType="application/vnd.openxmlformats-officedocument.vmlDrawing",
        )
        notes = (
            b"<worksheet xmlns='http://schemas.openxmlformats.org/"
            b"spreadsheetml/2006/main'><sheetData/></worksheet>"
        )
        comments = (
            b"<comments xmlns='http://schemas.openxmlformats.org/"
            b"spreadsheetml/2006/main'><authors><author>N</author></authors>"
            b"<commentList><comment ref='A1' authorId='0'><text><t>Keep me"
            b"</t></text></comment></commentList></comments>"
        )
        sheet_rels = (
            b"<Relationships xmlns='http://schemas.openxmlformats.org/package/"
            b"2006/relationships'><Relationship Id='rId1' Type='http://schemas."
            b"openxmlformats.org/officeDocument/2006/relationships/comments' "
            b"Target='../comments1.xml'/><Relationship Id='rId2' Type='http://"
            b"schemas.openxmlformats.org/officeDocument/2006/relationships/"
            b"vmlDrawing' Target='../drawings/commentsDrawing1.vml'/></Relationships>"
        )
        vml = b"<xml xmlns:v='urn:schemas-microsoft-com:vml'><v:shape id='_x0000_s1'/></xml>"
        workbook_xml_bytes = etree.tostring(
            workbook_xml, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        content_types_bytes = etree.tostring(
            content_types, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        rich = _rewrite_xlsx(
            workbook,
            {
                "[Content_Types].xml": content_types_bytes,
                "xl/workbook.xml": workbook_xml_bytes,
                "xl/_rels/workbook.xml.rels": etree.tostring(
                    workbook_rels,
                    xml_declaration=True,
                    encoding="UTF-8",
                    standalone=True,
                ),
                "xl/worksheets/sheet1.xml": etree.tostring(
                    sheet, xml_declaration=True, encoding="UTF-8", standalone=True
                ),
                "xl/worksheets/sheet2.xml": notes,
                "xl/worksheets/_rels/sheet1.xml.rels": sheet_rels,
                "xl/comments1.xml": comments,
                "xl/drawings/commentsDrawing1.vml": vml,
            },
        )
        pkg.set_raw_part(r["embedded_workbook"], rich)

        out = ch.update_chart_data(
            pkg, slide, {"shape_id": r["shape_id"]},
            ["North", "South", "West"],
            [{"name": "Updated", "values": [10, 20, 30]}],
        )
        assert out["embedded_workbook"] == "updated_in_place"
        saved = pkg.save()
        reloaded = PptxPackage(saved)
        with zipfile.ZipFile(io.BytesIO(reloaded.raw_part(r["embedded_workbook"]))) as zf:
            assert zf.read("[Content_Types].xml") == content_types_bytes
            assert zf.read("xl/workbook.xml") == workbook_xml_bytes
            assert zf.read("xl/worksheets/sheet2.xml") == notes
            assert zf.read("xl/worksheets/_rels/sheet1.xml.rels") == sheet_rels
            assert zf.read("xl/comments1.xml") == comments
            assert zf.read("xl/drawings/commentsDrawing1.vml") == vml
            updated_sheet = etree.fromstring(zf.read("xl/worksheets/sheet1.xml"))
        cells = {c.get("r"): c for c in updated_sheet.iter(f"{{{_SML}}}c")}
        assert cells["A2"].find(f"{{{_SML}}}is/{{{_SML}}}t").text == "North"
        assert cells["B2"].find(f"{{{_SML}}}v").text == "10"
        assert cells["D2"].find(f"{{{_SML}}}f").text == "B2*1.1"
        assert cells["D2"].find(f"{{{_SML}}}v").text == "4.4"

    def test_growth_refuses_before_overwriting_adjacent_total_row(self, deck):
        """Scenario: A5:B5 belongs to the workbook user."""
        pkg, slide = deck
        r = ch.create_chart(pkg, slide, "column", CATS, [SERIES[0]], 1, 1, 6, 4)
        workbook = pkg.raw_part(r["embedded_workbook"])
        sheet = _workbook_sheet(pkg, r["embedded_workbook"])
        sheet_data = sheet.find(f"{{{_SML}}}sheetData")
        total = etree.SubElement(sheet_data, f"{{{_SML}}}row", r="5")
        label = etree.SubElement(total, f"{{{_SML}}}c", r="A5", t="inlineStr")
        inline = etree.SubElement(label, f"{{{_SML}}}is")
        etree.SubElement(inline, f"{{{_SML}}}t").text = "Total"
        formula = etree.SubElement(total, f"{{{_SML}}}c", r="B5")
        etree.SubElement(formula, f"{{{_SML}}}f").text = "SUM(B2:B4)"
        etree.SubElement(formula, f"{{{_SML}}}v").text = "9.5"
        guarded = _rewrite_xlsx(
            workbook,
            {"xl/worksheets/sheet1.xml": etree.tostring(
                sheet, xml_declaration=True, encoding="UTF-8", standalone=True
            )},
        )
        pkg.set_raw_part(r["embedded_workbook"], guarded)
        chart_before = etree.tostring(pkg.root(r["chart_part"]))

        with pytest.raises(UnsupportedStructure, match=r"Sheet1!A5|Sheet1!B5"):
            ch.update_chart_data(
                pkg, slide, {"shape_id": r["shape_id"]},
                ["Alpha", "Beta", "Gamma", "Delta"],
                [{"name": "One", "values": [4, 2.5, 3, 7]}],
            )

        assert pkg.raw_part(r["embedded_workbook"]) == guarded
        assert etree.tostring(pkg.root(r["chart_part"])) == chart_before

    def test_formula_inside_plotted_range_refuses_before_mutation(self, deck):
        pkg, slide = deck
        r = ch.create_chart(pkg, slide, "column", CATS, [SERIES[0]], 1, 1, 6, 4)
        workbook = pkg.raw_part(r["embedded_workbook"])
        sheet = _workbook_sheet(pkg, r["embedded_workbook"])
        cells = {c.get("r"): c for c in sheet.iter(f"{{{_SML}}}c")}
        b3 = cells["B3"]
        for child in list(b3):
            b3.remove(child)
        etree.SubElement(b3, f"{{{_SML}}}f").text = "B2+1000"
        etree.SubElement(b3, f"{{{_SML}}}v").text = "1004"
        guarded = _rewrite_xlsx(
            workbook,
            {"xl/worksheets/sheet1.xml": etree.tostring(
                sheet, xml_declaration=True, encoding="UTF-8", standalone=True
            )},
        )
        pkg.set_raw_part(r["embedded_workbook"], guarded)
        chart_before = etree.tostring(pkg.root(r["chart_part"]))

        with pytest.raises(UnsupportedStructure, match=r"Sheet1!B3.*formula"):
            ch.update_chart_data(
                pkg, slide, {"shape_id": r["shape_id"]}, CATS,
                [{"name": "Updated", "values": [10, 20, 30]}],
            )

        assert pkg.raw_part(r["embedded_workbook"]) == guarded
        assert etree.tostring(pkg.root(r["chart_part"])) == chart_before

    def test_shrink_clears_dropped_literals_and_preserves_total_formula(self, deck):
        pkg, slide = deck
        r = ch.create_chart(
            pkg, slide, "column", ["A", "B", "C", "D"],
            [{"name": "One", "values": [1, 2, 3, 4]}], 1, 1, 6, 4,
        )
        workbook = pkg.raw_part(r["embedded_workbook"])
        sheet = _workbook_sheet(pkg, r["embedded_workbook"])
        sheet_data = sheet.find(f"{{{_SML}}}sheetData")
        total = etree.SubElement(sheet_data, f"{{{_SML}}}row", r="6")
        label = etree.SubElement(total, f"{{{_SML}}}c", r="A6", t="inlineStr")
        inline = etree.SubElement(label, f"{{{_SML}}}is")
        etree.SubElement(inline, f"{{{_SML}}}t").text = "Total"
        formula = etree.SubElement(total, f"{{{_SML}}}c", r="B6")
        etree.SubElement(formula, f"{{{_SML}}}f").text = "SUM(B2:B5)"
        etree.SubElement(formula, f"{{{_SML}}}v").text = "10"
        pkg.set_raw_part(
            r["embedded_workbook"],
            _rewrite_xlsx(
                workbook,
                {"xl/worksheets/sheet1.xml": etree.tostring(
                    sheet, xml_declaration=True, encoding="UTF-8", standalone=True
                )},
            ),
        )

        out = ch.update_chart_data(
            pkg, slide, {"shape_id": r["shape_id"]}, ["A", "B", "C"],
            [{"name": "One", "values": [10, 20, 30]}],
        )

        assert out["embedded_workbook"] == "updated_in_place"
        updated = _workbook_sheet(pkg, r["embedded_workbook"])
        cells = {c.get("r"): c for c in updated.iter(f"{{{_SML}}}c")}
        assert not any(
            etree.QName(child).localname in ("f", "v", "is")
            for child in cells["A5"]
        )
        assert not any(
            etree.QName(child).localname in ("f", "v", "is")
            for child in cells["B5"]
        )
        assert cells["A6"].find(f"{{{_SML}}}is/{{{_SML}}}t").text == "Total"
        assert cells["B6"].find(f"{{{_SML}}}f").text == "SUM(B2:B5)"
        assert cells["B6"].find(f"{{{_SML}}}v").text == "10"

    def test_unsupported_workbook_range_refuses_before_mutation(self, deck):
        pkg, slide = deck
        r = ch.create_chart(pkg, slide, "column", CATS, [SERIES[0]], 1, 1, 6, 4)
        ser = ch._plot_groups(pkg.root(r["chart_part"]))[0].find(_qc("ser"))
        formula = ser.find(f"{_qc('val')}/{_qc('numRef')}/{_qc('f')}")
        formula.text = "SalesValues"
        chart_before = etree.tostring(pkg.root(r["chart_part"]))
        workbook_before = pkg.raw_part(r["embedded_workbook"])

        with pytest.raises(
            UnsupportedStructure, match="is not one simple single-column range"
        ):
            ch.update_chart_data(
                pkg, slide, {"shape_id": r["shape_id"]}, CATS,
                [{"name": "One", "values": [9, 8, 7]}],
            )
        assert etree.tostring(pkg.root(r["chart_part"])) == chart_before
        assert pkg.raw_part(r["embedded_workbook"]) == workbook_before

    @pytest.mark.parametrize(
        ("chart_type", "categories", "series", "fill_path", "fill_tag"),
        [
            (
                "column", CATS,
                [{"name": "One", "values": [9, 8, 7], "color": "112233"}],
                f"{_qc('spPr')}/{{{_A}}}solidFill",
                "srgbClr",
            ),
            (
                "line", CATS,
                [{"name": "One", "values": [9, 8, 7], "color": "accent2"}],
                f"{_qc('spPr')}/{{{_A}}}ln/{{{_A}}}solidFill",
                "schemeClr",
            ),
            (
                "scatter", None,
                [{"name": "One", "x": [1, 2, 3], "y": [9, 8, 7], "color": "112233"}],
                f"{_qc('marker')}/{_qc('spPr')}/{{{_A}}}solidFill",
                "srgbClr",
            ),
        ],
    )
    def test_update_applies_series_color(
        self, deck, chart_type, categories, series, fill_path, fill_tag
    ):
        pkg, slide = deck
        original = (
            [{"name": "One", "x": [1, 2, 3], "y": [4, 5, 6]}]
            if chart_type == "scatter"
            else [{"name": "One", "values": [4, 5, 6]}]
        )
        r = ch.create_chart(pkg, slide, chart_type, categories, original, 1, 1, 6, 4)

        ch.update_chart_data(
            pkg, slide, {"shape_id": r["shape_id"]}, categories, series
        )

        group = ch._plot_groups(pkg.root(r["chart_part"]))[0]
        ser = group.find(_qc("ser"))
        fill = ser.find(fill_path)
        assert fill is not None
        color = fill.find(
            f"{{{_A}}}{fill_tag}"
        )
        assert color is not None
        assert color.get("val") == ("112233" if fill_tag == "srgbClr" else "accent2")

    def test_update_preserves_unspecified_color_and_rejects_invalid_color(self, deck):
        pkg, slide = deck
        r = ch.create_chart(
            pkg, slide, "column", CATS,
            [{"name": "One", "values": [1, 2, 3], "color": "112233"}],
            1, 1, 6, 4,
        )
        ch.update_chart_data(
            pkg, slide, {"shape_id": r["shape_id"]}, CATS,
            [{"name": "One", "values": [4, 5, 6]}],
        )
        ser = ch._plot_groups(pkg.root(r["chart_part"]))[0].find(_qc("ser"))
        color = ser.find(
            f"{_qc('spPr')}/{{{_A}}}solidFill/{{{_A}}}srgbClr"
        )
        assert color.get("val") == "112233"

        with pytest.raises(PptMcpError):
            ch.update_chart_data(
                pkg, slide, {"shape_id": r["shape_id"]}, CATS,
                [{"name": "One", "values": [7, 8, 9], "color": "not-a-color"}],
            )
        assert _num_cache_values(ser) == ["4", "5", "6"]

    def test_roundtrip_caches_and_workbook(self, deck):
        pkg, slide = deck
        r = ch.create_chart(pkg, slide, "column", CATS, SERIES, 1, 1, 6, 4)
        out = ch.update_chart_data(
            pkg, slide, {"shape_id": r["shape_id"]},
            ["X", "Y", "Z", "W"],
            [
                {"name": "One2", "values": [9, 8, 7, 6]},
                {"name": "Two2", "values": [1, 1, 1, 1]},
            ],
        )
        assert out["points"] == 4
        assert out["embedded_workbook"] == "updated_in_place"
        root = pkg.root(r["chart_part"])
        group = root.find(f"{_qc('chart')}/{_qc('plotArea')}/{_qc('barChart')}")
        sers = group.findall(_qc("ser"))
        assert _num_cache_values(sers[0]) == ["9", "8", "7", "6"]
        name_cache = sers[1].find(f"{_qc('tx')}/{_qc('strRef')}/{_qc('strCache')}")
        assert name_cache.find(f"{_qc('pt')}/{_qc('v')}").text == "Two2"
        # Workbook agrees with the caches after the rewrite.
        sheet = _workbook_sheet(pkg, r["embedded_workbook"])
        cells = {c.get("r"): c for c in sheet.iter(f"{{{_SML}}}c")}
        assert cells["A5"].find(f"{{{_SML}}}is/{{{_SML}}}t").text == "W"
        assert cells["B2"].find(f"{{{_SML}}}v").text == "9"
        pkg.save()

    def test_series_count_change_refuses(self, deck):
        pkg, slide = deck
        r = ch.create_chart(pkg, slide, "column", CATS, SERIES, 1, 1, 6, 4)
        with pytest.raises(UnsupportedStructure) as exc:
            ch.update_chart_data(
                pkg, slide, {"shape_id": r["shape_id"]}, CATS,
                [{"name": "only", "values": [1, 2, 3]}],
            )
        assert "series count" in str(exc.value)

    def test_type_preserved_on_update(self, deck):
        pkg, slide = deck
        r = ch.create_chart(pkg, slide, "bar_stacked", CATS, SERIES, 1, 1, 6, 4)
        ch.update_chart_data(
            pkg, slide, {"shape_id": r["shape_id"]}, CATS, SERIES
        )
        group = pkg.root(r["chart_part"]).find(
            f"{_qc('chart')}/{_qc('plotArea')}/{_qc('barChart')}"
        )
        assert group.find(_qc("grouping")).get("val") == "stacked"
        assert group.find(_qc("overlap")).get("val") == "100"


class TestChartAddressing:
    def test_no_chart_refuses(self, deck):
        pkg, slide = deck
        with pytest.raises(TargetNotFound):
            ch.update_chart_data(pkg, slide, None, CATS, SERIES)

    def test_multi_chart_ambiguity(self, deck):
        pkg, slide = deck
        a = ch.create_chart(pkg, slide, "pie", CATS, [SERIES[0]], 0.5, 0.5, 4, 3)
        b = ch.create_chart(pkg, slide, "line", CATS, SERIES, 5, 0.5, 4, 3)
        with pytest.raises(AmbiguousTarget) as exc:
            ch.update_chart_data(pkg, slide, None, CATS, SERIES)
        assert str(a["shape_id"]) in str(exc.value)
        # Explicit index addressing picks document order.
        out = ch.update_chart_data(pkg, slide, {"index": 1}, CATS, SERIES)
        assert out["shape_id"] == b["shape_id"]

    def test_explicit_selectors_and_ambiguous_legacy_integer(self, deck):
        from kitchensink4ppt.ops import chartread

        pkg, slide = deck
        made = [
            ch.create_chart(
                pkg, slide, "column", CATS,
                [{"name": "S0", "values": [1, 2, 3]}],
                0.5, 0.5, 2.5, 2,
            )
        ]
        collision = made[0]["shape_id"]
        for i in range(1, collision + 1):
            made.append(
                ch.create_chart(
                    pkg, slide, "column", CATS,
                    [{"name": f"S{i}", "values": [i + 1, i + 2, i + 3]}],
                    0.5, 0.5, 2.5, 2,
                )
            )

        read = chartread.get_chart_data(pkg, slide, {"index": 1})
        assert read["shape_id"] == made[1]["shape_id"]
        updated = ch.update_chart_data(
            pkg, slide, {"shape_id": made[0]["shape_id"]}, CATS,
            [{"name": "changed", "values": [7, 8, 9]}],
        )
        assert updated["shape_id"] == made[0]["shape_id"]
        formatted = ch.format_chart(pkg, slide, {"index": 2}, title="Third")
        assert formatted["shape_id"] == made[2]["shape_id"]

        assert chartread.get_chart_data(pkg, slide, 0)["shape_id"] == made[0]["shape_id"]
        assert ch.update_chart_data(
            pkg, slide, 0, CATS,
            [{"name": "legacy", "values": [3, 2, 1]}],
        )["shape_id"] == made[0]["shape_id"]
        assert ch.format_chart(pkg, slide, 0, title="Legacy")["shape_id"] == made[0]["shape_id"]

        # This bare integer means both an index and a different shape id.
        with pytest.raises(AmbiguousTarget, match="legacy integer"):
            chartread.get_chart_data(pkg, slide, collision)
        with pytest.raises(AmbiguousTarget, match="legacy integer"):
            ch.update_chart_data(
                pkg, slide, collision, CATS,
                [{"name": "changed", "values": [7, 8, 9]}],
            )
        with pytest.raises(AmbiguousTarget, match="legacy integer"):
            ch.format_chart(pkg, slide, collision, title="Wrong target")

    def test_bad_selector(self, deck):
        pkg, slide = deck
        ch.create_chart(pkg, slide, "pie", CATS, [SERIES[0]], 1, 1, 4, 3)
        with pytest.raises(TargetNotFound):
            ch.update_chart_data(pkg, slide, 42, CATS, [SERIES[0]])
        with pytest.raises(PptMcpError):
            ch.update_chart_data(pkg, slide, "first", CATS, [SERIES[0]])


class TestClosure:
    def test_chart_rels_resolve_after_save(self, deck):
        """The saved package passes the dangling-rel validation with charts
        in it, and the chart part's own rels point at real parts."""
        pkg, slide = deck
        r = ch.create_chart(pkg, slide, "column", CATS, SERIES, 1, 1, 6, 4)
        path = pkg.save()
        with zipfile.ZipFile(path) as zf:
            names = set(zf.namelist())
            assert r["chart_part"] in names
            assert r["embedded_workbook"] in names
            rels = etree.fromstring(
                zf.read("ppt/charts/_rels/" + r["chart_part"].rsplit("/", 1)[1] + ".rels")
            )
            targets = [rel.get("Target") for rel in rels]
            assert any("embeddings" in t for t in targets)
            ct = zf.read("[Content_Types].xml").decode()
            assert "/" + r["chart_part"] in ct
            assert 'Extension="xlsx"' in ct


_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

SCATTER_SERIES = [
    {"name": "Alpha", "x": [1, 2, 3], "y": [2.5, 1, 4]},
    {"name": "Beta", "x": [0.5, 1.5], "y": [3, 3.5]},  # shorter on purpose
]


class TestScatter:
    """Scatter charts: c:scatterChart with xVal/yVal numRefs,
    marker-only styling (a:ln noFill), smooth=0, two value axes, and an
    embedded workbook laying out an X/Y column pair per series."""

    def test_scatter_structure_and_caches(self, deck):
        pkg, slide = deck
        r = ch.create_chart(
            pkg, slide, "scatter", None, SCATTER_SERIES, 1, 1, 6, 4,
            title="XY",
        )
        assert r["type"] == "scatter"
        assert r["points"] == 5
        root = pkg.root(r["chart_part"])
        plot = root.find(f"{_qc('chart')}/{_qc('plotArea')}")
        group = plot.find(_qc("scatterChart"))
        assert group is not None
        assert group.find(_qc("scatterStyle")).get("val") == "lineMarker"
        sers = group.findall(_qc("ser"))
        assert len(sers) == 2
        for ser, spec in zip(sers, SCATTER_SERIES):
            # Marker-only: the connecting line is suppressed per-series.
            ln = ser.find(f"{_qc('spPr')}/{{{_A_NS}}}ln")
            assert ln is not None and ln.find(f"{{{_A_NS}}}noFill") is not None
            assert ser.find(
                f"{_qc('marker')}/{_qc('symbol')}"
            ).get("val") == "circle"
            xc = ser.find(f"{_qc('xVal')}/{_qc('numRef')}/{_qc('numCache')}")
            yc = ser.find(f"{_qc('yVal')}/{_qc('numRef')}/{_qc('numCache')}")
            assert int(xc.find(_qc("ptCount")).get("val")) == len(spec["x"])
            assert int(yc.find(_qc("ptCount")).get("val")) == len(spec["y"])
            assert [float(p.find(_qc("v")).text)
                    for p in yc.findall(_qc("pt"))] == [
                float(v) for v in spec["y"]
            ]
            assert ser.find(_qc("smooth")).get("val") == "0"
            assert ser.find(_qc("cat")) is None
            assert ser.find(_qc("val")) is None
        # Two VALUE axes, no category axis, cross-wired ids.
        assert plot.find(_qc("catAx")) is None
        vals = plot.findall(_qc("valAx"))
        assert len(vals) == 2
        ids = [ax.find(_qc("axId")).get("val") for ax in vals]
        crosses = [ax.find(_qc("crossAx")).get("val") for ax in vals]
        assert ids == crosses[::-1]
        # Workbook: X/Y column pair per series (A/B then C/D).
        sheet = _workbook_sheet(pkg, r["embedded_workbook"])
        sml = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        refs = [c.get("r") for c in sheet.iter(f"{sml}c")]
        assert {"A1", "B1", "C1", "D1"} <= set(refs)
        assert "A4" in refs and "C4" not in refs  # series lengths differ
        pkg.save()

    def test_scatter_refusals_fire_before_mutation(self, deck):
        pkg, slide = deck
        cases = [
            (["c1"], SCATTER_SERIES, "no categories"),
            (None, [{"name": "A", "x": [1, 2], "y": [1]}], "x values but"),
            (None, [{"name": "A", "values": [1]}], "x/y value lists"),
            (None, [{"name": "A", "x": [], "y": []}], "no data points"),
            (None, [{"name": "A", "x": [1], "y": [1], "axis": "secondary"}],
             "primary-axis only"),
            (None, [{"name": "A", "x": [1], "y": [1], "type": "line"}],
             "do not join combo"),
        ]
        for cats, series, msg in cases:
            with pytest.raises(PptMcpError, match=msg):
                ch.create_chart(
                    pkg, slide, "scatter", cats, series, 1, 1, 4, 3
                )
        # no chart/workbook parts were created by any refused call
        assert not any(
            n.startswith(("ppt/charts/", "ppt/embeddings/"))
            for n in pkg.part_names()
        )

    def test_scatter_update_roundtrip(self, deck):
        pkg, slide = deck
        r = ch.create_chart(
            pkg, slide, "scatter", None, SCATTER_SERIES, 1, 1, 6, 4
        )
        out = ch.update_chart_data(
            pkg, slide, None, None,
            [{"name": "A2", "x": [9, 8, 7, 6], "y": [1, 2, 3, 4]},
             {"name": "B2", "x": [5], "y": [5]}],
        )
        assert out["type"] == "scatter"
        assert out["embedded_workbook"] == "updated_in_place"
        group = pkg.root(r["chart_part"]).find(
            f"{_qc('chart')}/{_qc('plotArea')}/{_qc('scatterChart')}"
        )
        sers = group.findall(_qc("ser"))
        xc = sers[0].find(f"{_qc('xVal')}/{_qc('numRef')}/{_qc('numCache')}")
        assert [p.find(_qc("v")).text for p in xc.findall(_qc("pt"))] == [
            "9", "8", "7", "6"
        ]
        tx0 = sers[0].find(f"{_qc('tx')}/{_qc('strRef')}/{_qc('strCache')}")
        assert tx0.find(_qc("pt")).find(_qc("v")).text == "A2"
        # Update refuses the category data model on a scatter chart.
        with pytest.raises(PptMcpError, match="no categories"):
            ch.update_chart_data(
                pkg, slide, None, ["c"],
                [{"name": "A", "values": [1]}, {"name": "B", "values": [2]}],
            )
        pkg.save()

    def test_scatter_pptx_oracle(self, deck):
        pytest.importorskip("pptx")
        from pptx import Presentation
        from pptx.enum.chart import XL_CHART_TYPE

        pkg, slide = deck
        ch.create_chart(
            pkg, slide, "scatter", None, SCATTER_SERIES, 1, 1, 6, 4
        )
        path = pkg.save()
        prs = Presentation(str(path))
        chart = next(
            s.chart for sl in prs.slides for s in sl.shapes if s.has_chart
        )
        assert chart.chart_type == XL_CHART_TYPE.XY_SCATTER
        series = list(chart.series)
        assert len(series) == 2
        # python-pptx reads scatter series .values as the Y values.
        assert list(series[0].values) == [2.5, 1.0, 4.0]
        assert list(series[1].values) == [3.0, 3.5]


class TestSeriesColors:
    """Per-series color OPT-IN: solidFill lands exactly where asked (bar/pie
    fill, line stroke, scatter marker) and the deliberate-spPr-absence
    default is untouched for series that did not ask."""

    def test_bar_color_optin_and_default_absence(self, deck):
        pkg, slide = deck
        r = ch.create_chart(
            pkg, slide, "column", CATS,
            [{"name": "One", "values": [1, 2, 3], "color": "1F4E79"},
             {"name": "Two", "values": [4, 5, 6]}],
            1, 1, 6, 4,
        )
        sers = pkg.root(r["chart_part"]).findall(
            f"{_qc('chart')}/{_qc('plotArea')}/{_qc('barChart')}/{_qc('ser')}"
        )
        fill = sers[0].find(
            f"{_qc('spPr')}/{{{_A_NS}}}solidFill/{{{_A_NS}}}srgbClr"
        )
        assert fill is not None and fill.get("val") == "1F4E79"
        assert sers[1].find(_qc("spPr")) is None  # default stays deliberate
        # spPr must sit between c:tx and c:cat (schema order).
        kids = [etree.QName(c).localname for c in sers[0]]
        assert kids.index("tx") < kids.index("spPr") < kids.index("cat")
        pkg.save()

    def test_line_color_is_stroke_and_theme_token_works(self, deck):
        pkg, slide = deck
        r = ch.create_chart(
            pkg, slide, "line", CATS,
            [{"name": "One", "values": [1, 2, 3], "color": "accent2"}],
            1, 1, 6, 4,
        )
        ser = pkg.root(r["chart_part"]).find(
            f"{_qc('chart')}/{_qc('plotArea')}/{_qc('lineChart')}/{_qc('ser')}"
        )
        clr = ser.find(
            f"{_qc('spPr')}/{{{_A_NS}}}ln/{{{_A_NS}}}solidFill"
            f"/{{{_A_NS}}}schemeClr"
        )
        assert clr is not None and clr.get("val") == "accent2"
        # Fill itself untouched: only the stroke was asked for.
        assert ser.find(f"{_qc('spPr')}/{{{_A_NS}}}solidFill") is None

    def test_scatter_color_lands_on_marker(self, deck):
        pkg, slide = deck
        r = ch.create_chart(
            pkg, slide, "scatter", None,
            [{"name": "A", "x": [1, 2], "y": [3, 4], "color": "#FF0000"}],
            1, 1, 6, 4,
        )
        ser = pkg.root(r["chart_part"]).find(
            f"{_qc('chart')}/{_qc('plotArea')}/{_qc('scatterChart')}"
            f"/{_qc('ser')}"
        )
        mfill = ser.find(
            f"{_qc('marker')}/{_qc('spPr')}/{{{_A_NS}}}solidFill"
            f"/{{{_A_NS}}}srgbClr"
        )
        assert mfill is not None and mfill.get("val") == "FF0000"
        # The connecting line stays suppressed regardless of color.
        assert ser.find(
            f"{_qc('spPr')}/{{{_A_NS}}}ln/{{{_A_NS}}}noFill"
        ) is not None

    def test_invalid_color_refuses_before_mutation(self, deck):
        pkg, slide = deck
        with pytest.raises(PptMcpError, match="color must be"):
            ch.create_chart(
                pkg, slide, "pie", CATS,
                [{"name": "S", "values": [1, 2, 3], "color": "red"}],
                1, 1, 4, 3,
            )
        # no chart/workbook parts were created by the refused call
        assert not any(
            n.startswith(("ppt/charts/", "ppt/embeddings/"))
            for n in pkg.part_names()
        )
