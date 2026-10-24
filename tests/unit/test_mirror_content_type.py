"""A part copied between decks keeps the content type it had in its own deck.

A copy used to mirror a source Default additively: when the destination already had a
Default for the same extension with another type, the copied part silently
took the destination's type. And a part the source deck gave no content type
was copied without one, which no program can type. Now a differing Default
becomes an Override for the copied part, and an untyped part is refused
unless the destination's own Default covers it. The decks are made up.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.errors import UnsupportedStructure
from kitchensink4ppt.core.package import PptxPackage
from kitchensink4ppt.ops import interdeck

_PNG_DEFAULT = b'<Default Extension="png" ContentType="image/png"/>'
_X_PNG_DEFAULT = b'<Default Extension="png" ContentType="image/x-png"/>'


def _png(color) -> io.BytesIO:
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), color).save(buf, "PNG")
    buf.seek(0)
    return buf


def _deck(path, with_picture: bool, color=(20, 120, 200)):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if with_picture:
        slide.shapes.add_picture(_png(color), Inches(1), Inches(1), Inches(1), Inches(1))
    else:
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "Hello"
    prs.save(path)
    return path


def _content_types(path, fn) -> None:
    with zipfile.ZipFile(path) as zf:
        entries = [(i.filename, zf.read(i.filename)) for i in zf.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in entries:
            zf.writestr(name, fn(data) if name == "[Content_Types].xml" else data)


def _saved_types(path) -> dict:
    """{media part: the type the saved file gives it (Override, else Default)}."""
    with zipfile.ZipFile(path) as zf:
        root = etree.fromstring(zf.read("[Content_Types].xml"))
        names = [n for n in zf.namelist() if n.startswith("ppt/media/")]
    over = {n.get("PartName", "").lstrip("/").lower(): n.get("ContentType")
            for n in root if etree.QName(n).localname == "Override"}
    dflt = {n.get("Extension", "").lower(): n.get("ContentType")
            for n in root if etree.QName(n).localname == "Default"}
    return {n: over.get(n.lower(), dflt.get(n.rsplit(".", 1)[-1].lower())) for n in names}


def _x_png(data: bytes) -> bytes:
    if _PNG_DEFAULT in data:
        return data.replace(_PNG_DEFAULT, _X_PNG_DEFAULT)
    return data.replace(b"</Types>", _X_PNG_DEFAULT + b"</Types>")


def test_a_differing_destination_default_gets_an_override(tmp_path):
    src = _deck(tmp_path / "src.pptx", with_picture=True)
    dst = _deck(tmp_path / "dst.pptx", with_picture=False)
    _content_types(dst, _x_png)
    pkg = PptxPackage(dst)
    interdeck.copy_slide_between(pkg, str(src), 0)
    pkg.save(do_backup=False)
    types = _saved_types(dst)
    assert len(types) == 1 and set(types.values()) == {"image/png"}, types


def test_an_untyped_source_part_is_refused_when_nothing_covers_it(tmp_path):
    src = _deck(tmp_path / "src.pptx", with_picture=True)
    _content_types(src, lambda d: d.replace(_PNG_DEFAULT, b""))
    dst = _deck(tmp_path / "dst.pptx", with_picture=False)
    _content_types(dst, lambda d: d.replace(_PNG_DEFAULT, b""))
    before = dst.read_bytes()
    pkg = PptxPackage(dst)
    with pytest.raises(UnsupportedStructure, match="content type"):
        interdeck.copy_slide_between(pkg, str(src), 0)
    assert dst.read_bytes() == before


def test_an_untyped_source_part_takes_the_destination_default(tmp_path):
    src = _deck(tmp_path / "src.pptx", with_picture=True)
    _content_types(src, lambda d: d.replace(_PNG_DEFAULT, b""))
    dst = _deck(tmp_path / "dst.pptx", with_picture=True, color=(200, 40, 40))  # its own png Default
    pkg = PptxPackage(dst)
    interdeck.copy_slide_between(pkg, str(src), 0)
    pkg.save(do_backup=False)
    types = _saved_types(dst)
    assert len(types) == 2 and set(types.values()) == {"image/png"}, types
