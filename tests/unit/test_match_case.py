"""Letter case in find_text and search_and_replace: match_case, and the note on
matches that differ only in case.

Asked which slides mention "budget", a model sends query="budget". A slide
that says "Budget risks" was missed, and nothing in the answer or the help
said the search was case-sensitive, so the model told the user that slide
does not mention the word.

find_text now takes match_case, as search_and_replace does, with the same
default: true, exact case, so existing callers see no change. Whenever an
exact-case search or replace leaves out matches that differ only in case,
the answer's case_note says so in one sentence: how many, which slides
(slide_number, counted from 1) and how to include them (match_case=false).
A replace never changes text in another case unless match_case=false. The
help of both operations states the default.

Every deck here is made up by code.
"""

from __future__ import annotations

import asyncio

import pytest
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt import server  # noqa: F401  (the import registers ops)
from kitchensink4ppt.com import live, live_ops
from kitchensink4ppt.registry import public_ops
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.help import help_payload
from kitchensink4ppt.surface.plan import build

from test_live_text_offsets import _Coll, _Pres, _Session, _Shape, _Slide


def _plan():
    return build(Settings(), public_ops())


def _call(tool: str, op: str, args: dict):
    return asyncio.run(dispatch(_plan(), tool, {"op": op, "args": args}))


def _ok(op: str, args: dict, tool: str = "ppt_read") -> dict:
    payload, is_error = _call(tool, op, args)
    assert is_error is False, payload
    return payload


def _edit(op: str, args: dict) -> dict:
    """An edit's own answer, from inside the save envelope."""
    return _ok(op, args, tool="ppt_edit")["changed"]


def _deck(path, bodies, notes=None):
    """One slide per body text, each in its own text box; notes optional."""
    prs = Presentation()
    for i, text in enumerate(bodies):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
        box.text_frame.text = text
        if notes and notes[i]:
            slide.notes_slide.notes_text_frame.text = notes[i]
    prs.save(str(path))
    return str(path)


def _texts(path):
    prs = Presentation(path)
    return [
        " | ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        for s in prs.slides
    ]


BODIES = [
    "We kept the budget flat.",
    "Budget risks are listed here.",
    "Nothing about money.",
    "The BUDGET line and the budget note.",
]


@pytest.fixture()
def deck(tmp_path):
    return _deck(tmp_path / "case.pptx", BODIES)


# ------------------------------------------------------------------ find_text


def test_find_text_with_match_case_false_finds_every_case(deck):
    out = _ok("find_text", {"file_path": deck, "query": "budget", "match_case": False})
    found = sorted((m["slide_index"], m["match"]) for m in out["matches"])
    assert found == [(0, "budget"), (1, "Budget"), (3, "BUDGET"), (3, "budget")]
    assert out["count"] == 4
    assert "case_note" not in out


def test_find_text_default_stays_exact_case(deck):
    out = _ok("find_text", {"file_path": deck, "query": "budget"})
    assert sorted(m["slide_index"] for m in out["matches"]) == [0, 3]
    assert all(m["match"] == "budget" for m in out["matches"])
    assert out["count"] == 2


def test_find_text_says_where_the_other_case_matches_are(deck):
    out = _ok("find_text", {"file_path": deck, "query": "budget"})
    note = out["case_note"]
    assert "2 more matches" in note
    assert "slide_numbers 2 and 4" in note
    assert '"Budget"' in note and '"BUDGET"' in note
    assert "match_case=false" in note
    assert note.count(". ") == 0 and note.endswith(".")  # one sentence


def test_find_text_with_no_exact_match_says_so(deck):
    out = _ok("find_text", {"file_path": deck, "query": "BUDGET RISKS"})
    assert out["count"] == 0
    note = out["case_note"]
    assert note.startswith("No match in this exact case")
    assert "1 match in another case" in note
    assert "slide_number 2" in note
    assert "match_case=false" in note


def test_find_text_scope_limits_the_note_too(deck):
    """Slides 2 to 3 in the user's count: the other-case match on slide 4 is
    out of scope and is not mentioned."""
    out = _ok("find_text", {"file_path": deck, "query": "budget", "slide_numbers": [2, 3]})
    assert out["count"] == 0
    assert "slide_number 2" in out["case_note"]
    assert "4" not in out["case_note"]


def test_find_text_no_note_when_case_adds_nothing(tmp_path):
    path = _deck(tmp_path / "plain.pptx", ["only budget here", "and budget again"])
    out = _ok("find_text", {"file_path": path, "query": "budget"})
    assert out["count"] == 2
    assert "case_note" not in out


def test_find_text_note_counts_notes_only_when_notes_are_searched(tmp_path):
    path = _deck(tmp_path / "notes.pptx", ["no match", "budget"], notes=["Budget in notes", None])
    out = _ok("find_text", {"file_path": path, "query": "budget"})
    assert "slide_number 1" in out["case_note"]
    out = _ok("find_text", {"file_path": path, "query": "budget", "include_notes": False})
    assert "case_note" not in out


def test_find_text_regex_follows_match_case(deck):
    out = _ok("find_text", {"file_path": deck, "query": r"bud\w+", "regex": True,
                            "match_case": False})
    assert out["count"] == 4
    out = _ok("find_text", {"file_path": deck, "query": r"bud\w+", "regex": True})
    assert out["count"] == 2
    assert "slide_numbers 2 and 4" in out["case_note"]


# --------------------------------------------------------- search_and_replace


def test_replace_leaves_other_case_alone_and_says_so(deck):
    out = _edit("search_and_replace", {"file_path": deck, "find": "budget", "replace": "spend"})
    assert out["total"] == 2
    assert _texts(deck) == [
        "We kept the spend flat.",
        "Budget risks are listed here.",
        "Nothing about money.",
        "The BUDGET line and the spend note.",
    ]
    note = out["case_note"]
    assert "2 matches in another case" in note
    assert "slide_numbers 2 and 4" in note
    assert "left unchanged" in note
    assert "match_case=false" in note


def test_replace_with_match_case_false_changes_every_case(deck):
    out = _edit("search_and_replace", {"file_path": deck, "find": "budget", "replace": "spend",
                                     "match_case": False})
    assert out["total"] == 4
    assert "case_note" not in out
    assert _texts(deck)[1] == "spend risks are listed here."


def test_replace_with_no_exact_match_changes_nothing_and_says_where(deck):
    before = _texts(deck)
    out = _edit("search_and_replace", {"file_path": deck, "find": "BUDGET RISKS",
                                     "replace": "X"})
    assert out["total"] == 0
    assert _texts(deck) == before
    assert "slide_number 2" in out["case_note"]


def test_batch_replace_carries_the_case_note(deck):
    out = _ok("apply_edits", {"file_path": deck, "edits": [
        {"op": "search_and_replace", "find": "budget", "replace": "spend"}]}, tool="ppt_edit")
    assert any(w.startswith("edit 0: ") and "match_case=false" in w
               and "slide_numbers 2 and 4" in w for w in out["warnings"]), out
    assert _texts(deck)[1] == "Budget risks are listed here."


# --------------------------------------------- the live route (stand-in COM)


@pytest.fixture()
def fake_live_deck(monkeypatch):
    holder = {}

    def run_live(path, tool, body, *, mutating=True):
        return body(_Session(holder["pres"]))

    monkeypatch.setattr(live, "run_live", run_live)
    monkeypatch.setattr(live_ops.safesave, "claim_publish", lambda *a, **k: None)
    monkeypatch.setattr(live_ops, "check_path", lambda p, what: p)

    def make(texts):
        slides = []
        for n, text in enumerate(texts):
            slide = _Slide([_Shape(10 + n, text)])
            slide.SlideIndex = n + 1
            slide.SlideID = 256 + n
            slides.append(slide)
        pres = _Pres([])
        pres.Slides = _Coll(slides)
        holder["pres"] = pres
        return slides

    return make


def test_live_replace_leaves_other_case_alone_and_says_so(fake_live_deck):
    slides = fake_live_deck(BODIES)
    out = live_ops.live_search_and_replace("deck.pptx", "budget", "spend")
    assert out["total"] == 2
    assert slides[1].Shapes.Item(1).TextFrame.TextRange.Text == "Budget risks are listed here."
    assert "slide_numbers 2 and 4" in out["case_note"]
    assert "match_case=false" in out["case_note"]


def test_live_replace_with_match_case_false_has_no_note(fake_live_deck):
    fake_live_deck(BODIES)
    out = live_ops.live_search_and_replace("deck.pptx", "budget", "spend", match_case=False)
    assert out["total"] == 4
    assert "case_note" not in out


# ------------------------------------------------------------------ the help


@pytest.mark.parametrize("op", ["find_text", "search_and_replace"])
def test_help_states_the_default_and_the_way_to_any_case(op):
    props = help_payload(_plan(), op)["args_schema"]["properties"]
    assert props["match_case"]["default"] is True
    desc = props["match_case"]["description"]
    assert "default" in desc.lower()
    assert "match_case=false" in desc


def test_find_text_op_line_says_how_case_is_matched():
    from kitchensink4ppt.surface.copy import OP_LINES

    line = OP_LINES["find_text"]
    assert "match_case=false" in line
