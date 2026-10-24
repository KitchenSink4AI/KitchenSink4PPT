"""The argument names a model reaches for first: shape_id and font_size.

Field report: format_text {slide_number, shape_id: 7, font_size: 16} was
refused (shape required; shape_id and font_size unknown) because the
operation names them shape and size_pt, while get_slide_info answers with
`id` and other answers say shape_id. The refusal listed the valid fields,
so recovery took one call, but the first call was the natural one.

shape_id is now read as another name for shape on every operation that
takes a shape, and font_size for the font-size argument of the text
operations (size_pt, or size on the placeholder operations), under the same
rule as every such name: sent beside the argument with a different value,
the call is refused with a hint rather than guessed at. The answer says
which name was read as which.
"""

from __future__ import annotations

import asyncio
import hashlib
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from kitchensink4ppt import server  # noqa: F401  (the import registers ops)
from kitchensink4ppt.core.package import PptxPackage, qn
from kitchensink4ppt.ops.read import slide_table
from kitchensink4ppt.registry import public_ops
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.help import help_payload
from kitchensink4ppt.surface.plan import build


def _plan():
    return build(Settings(), public_ops())


def _call(tool: str, op: str, args: dict):
    return asyncio.run(dispatch(_plan(), tool, {"op": op, "args": args}))


def _ok(tool: str, op: str, args: dict) -> dict:
    payload, is_error = _call(tool, op, args)
    assert is_error is False, payload
    return payload


def _refused(tool: str, op: str, args: dict) -> dict:
    payload, is_error = _call(tool, op, args)
    assert is_error is True, payload
    return payload["error"]


def _md5(path) -> str:
    return hashlib.md5(Path(path).read_bytes()).hexdigest()


@pytest.fixture
def deck(tmp_path):
    """One slide with two text boxes; returns (path, first id, second id)."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    a = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    a.text_frame.text = "First box"
    a.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    b = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1))
    b.text_frame.text = "Second box"
    path = tmp_path / "aliases.pptx"
    prs.save(str(path))
    return str(path), a.shape_id, b.shape_id


def _run_sizes(path, shape_id) -> list[str]:
    pkg = PptxPackage(path)
    part = slide_table(pkg)[0]["part"]
    for sp in pkg.root(part).iter(qn("p:sp")):
        cnv = sp.find(f"{qn('p:nvSpPr')}/{qn('p:cNvPr')}")
        if int(cnv.get("id")) == shape_id:
            return [r.get("sz") for r in sp.iter(qn("a:rPr"))]
    raise AssertionError(shape_id)


def test_format_text_takes_shape_id_and_font_size(deck):
    path, first, _second = deck
    out = _ok("ppt_edit", "format_text", {"file_path": path, "slide_number": 1,
                                          "shape_id": first, "font_size": 16, "live": "off"})
    assert _run_sizes(path, first) == ["1600"]
    assert out["normalized_from"]["args"] == {"shape_id": first, "font_size": 16}
    hint = out["normalized_from"]["hint"]
    assert "`shape_id` was read as `shape`" in hint
    assert "`font_size` was read as `size_pt`" in hint


def test_two_names_with_two_values_are_refused(deck):
    path, first, second = deck
    before = _md5(path)
    error = _refused("ppt_edit", "format_text", {"file_path": path, "slide_number": 1,
                                                 "shape": first, "shape_id": second,
                                                 "bold": True, "live": "off"})
    assert error["code"] == "BAD_ARGS"
    assert "shape_id and shape" in error["message"] and error.get("hint")
    error = _refused("ppt_edit", "format_text", {"file_path": path, "slide_number": 1,
                                                 "shape": first, "size_pt": 12, "font_size": 16,
                                                 "live": "off"})
    assert "font_size and size_pt" in error["message"]
    assert _md5(path) == before


def test_two_names_with_one_value_are_fine(deck):
    path, first, _second = deck
    _ok("ppt_edit", "format_text", {"file_path": path, "slide_number": 1, "shape": first,
                                    "shape_id": first, "size_pt": 18, "font_size": 18,
                                    "live": "off"})
    assert _run_sizes(path, first) == ["1800"]


@pytest.mark.parametrize("tool,op,extra", [
    ("ppt_graphics", "set_shape", {"dx": 0.1, "live": "off"}),
    ("ppt_graphics", "set_alt_text", {"text": "A box"}),
    ("ppt_edit", "set_bullets", {"style": "none"}),
    ("ppt_edit", "fit_text", {}),
    ("ppt_inspect", "get_autofit_state", {}),
])
def test_sibling_operations_take_shape_id(deck, tool, op, extra):
    path, first, _second = deck
    out = _ok(tool, op, {"file_path": path, "slide_number": 1, "shape_id": first, **extra})
    assert out["normalized_from"]["args"]["shape_id"] == first


def test_insert_textbox_takes_font_size(deck):
    path, _first, _second = deck
    out = _ok("ppt_edit", "insert_textbox", {"file_path": path, "slide_number": 1,
                                             "text": "Label", "x": 1, "y": 5, "w": 2, "h": 0.5,
                                             "font_size": 14, "live": "off"})
    assert out["normalized_from"]["args"] == {"font_size": 14}
    new_id = out["changed"]["shape_id"]
    assert _run_sizes(path, new_id) == ["1400"]


def test_placeholder_size_takes_font_size(deck):
    path, _first, _second = deck
    out = _ok("ppt_design", "set_master_placeholder", {"file_path": path, "ph": "title",
                                                       "font_size": 30})
    assert out["normalized_from"]["args"] == {"font_size": 30}


def test_an_operation_with_its_own_shape_id_keeps_it():
    fields = public_ops()["delete_master_shape"].args.model_fields
    assert "shape_id" in fields and "shape" not in fields
    from kitchensink4ppt.surface import normalize
    out, sent = normalize.normalize_args({"shape_id": 4}, fields, op="delete_master_shape")
    assert out == {"shape_id": 4} and sent == {}


def test_apply_edits_reads_the_same_names(deck):
    path, first, second = deck
    out = _ok("ppt_edit", "apply_edits", {"file_path": path, "edits": [
        {"op": "format_text", "slide_number": 1, "shape_id": first, "font_size": 24},
        {"op": "set_text", "slide_number": 1, "shape_id": second, "text": "Changed"},
    ]})
    assert out["applied"] == 2
    assert _run_sizes(path, first) == ["2400"]
    before = _md5(path)
    error = _refused("ppt_edit", "apply_edits", {"file_path": path, "edits": [
        {"op": "format_text", "slide_number": 1, "shape": first, "shape_id": second,
         "bold": True},
    ]})
    assert "shape_id and shape" in error["message"]
    assert _md5(path) == before


def test_help_names_the_other_names():
    props = help_payload(_plan(), "format_text")["args_schema"]["properties"]
    assert "shape_id" in props["shape"]["description"]
    assert "font_size" in props["size_pt"]["description"]
