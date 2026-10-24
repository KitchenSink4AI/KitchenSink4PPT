"""What recovery does with a save journal entry after the save it wrote down.

A save writes its swap down in a small journal before ReplaceFileW runs, and
the next start of a server, or the deck's next write, finishes what a stopped
save left. Recovery acts only on the deck's own files, never puts an older
backup over a newer one, never makes a renamed or deleted deck reappear under
its old name, keeps to its time at start when a folder does not answer, and
tidies what it leaves behind.

A stop in a save's last instant is a real kill: the victim's ReplaceFileW
goes through and the process is ended before it removes the old file and
before its backups land.
"""

from __future__ import annotations

import asyncio
import ctypes
import json
import os
import subprocess
import sys
import time
from pathlib import Path

import pytest
from pptx import Presentation

from kitchensink4ppt import server
from kitchensink4ppt.core import safesave, swapwait
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.plan import build

pytestmark = pytest.mark.skipif(sys.platform != "win32", reason="ReplaceFileW")

NO_WINDOW = 0x08000000

LATE_VICTIM = r'''
import asyncio, os, sys, time
from kitchensink4ppt import server  # noqa: F401
from kitchensink4ppt.core import safesave
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.plan import build

deck, marker = sys.argv[1], sys.argv[2]
fn, move, last_error, text = safesave._replace_file()

def went_through_then_stopped(ldst, lsrc, lbak, flags, a, b):
    ok = fn(ldst, lsrc, lbak, flags, a, b)
    open(marker, "w").close()
    time.sleep(600)  # killed here: the old file still aside, the backups not landed
    return ok

safesave._REPLACE_FILE = (went_through_then_stopped, move, last_error, text)
plan = build(Settings())
asyncio.run(dispatch(plan, "ppt_edit", {"op": "set_notes", "args": {
    "file_path": deck, "slide": 0, "text": "v2", "live": "off"}}))
'''

SAVE_IN_A_CHILD = (
    "import asyncio, sys\n"
    "from kitchensink4ppt import server\n"
    "from kitchensink4ppt.settings import Settings\n"
    "from kitchensink4ppt.surface.dispatch import dispatch\n"
    "from kitchensink4ppt.surface.plan import build\n"
    "plan = build(Settings())\n"
    "asyncio.run(dispatch(plan, 'ppt_edit', {'op': 'set_notes', 'args': {\n"
    "    'file_path': sys.argv[1], 'slide': 0, 'text': sys.argv[2], 'live': 'off'}}))\n")


def _deck(path: Path, notes: str) -> Path:
    prs = Presentation()
    for title in ("Alpha", "Bravo"):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.notes_slide.notes_text_frame.text = notes
    prs.save(path)
    return path


def _notes(path) -> str:
    return Presentation(str(path)).slides[0].notes_slide.notes_text_frame.text


def _call(op, args):
    plan = build(Settings())
    return asyncio.run(dispatch(plan, plan.specs[op].group, {"op": op, "args": args}))


def _set_notes(deck: Path, text: str):
    return _call("set_notes", {"file_path": str(deck), "slide": 0, "text": text,
                               "live": "off"})


def _save_in_a_child(deck: Path, text: str, journal: Path | None = None) -> None:
    env = dict(os.environ)
    if journal is not None:
        env["KS4P_SAVE_JOURNAL_DIR"] = str(journal)
    subprocess.run([sys.executable, "-X", "utf8", "-c", SAVE_IN_A_CHILD, str(deck), text],
                   check=True, timeout=300, stdin=subprocess.DEVNULL,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                   creationflags=NO_WINDOW, env=env)


def _killed_after_the_swap(tmp_path: Path) -> Path:
    """deck v2 in place, its old file (v1) still aside, prev still v0 and the
    v1 backup staged: the state a stop in a save's last instant leaves."""
    deck = _deck(tmp_path / "Deck.pptx", "v0")
    assert not _set_notes(deck, "v1")[1]
    script = tmp_path / "victim.py"
    script.write_text(LATE_VICTIM, encoding="utf-8")
    marker = tmp_path / "swapped"
    victim = subprocess.Popen([sys.executable, "-X", "utf8", str(script), str(deck), str(marker)],
                              stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                              stderr=subprocess.DEVNULL, creationflags=NO_WINDOW)
    try:
        deadline = time.monotonic() + 90
        while not marker.exists():
            assert victim.poll() is None, "the saving process ended before the swap"
            assert time.monotonic() < deadline, "the saving process never reached the swap"
            time.sleep(0.02)
    finally:
        victim.kill()  # TerminateProcess on this process's own handle
        victim.wait(timeout=60)
    assert _notes(deck) == "v2" and len(swapwait.asides(deck)) == 1
    return deck


def _prev(deck: Path) -> Path:
    return safesave.slot_dir(deck) / safesave.PREV_SLOT


def _entries_for(deck: Path) -> list[tuple[str, dict]]:
    key = safesave.canonical_key(deck)
    return [(entry, record) for entry, record in safesave._journal_entries()
            if isinstance(record.get("deck"), str)
            and safesave.canonical_key(record["deck"]) == key]


def _stopped_server_entry(**fields) -> dict:
    """A journal entry whose server is gone (a pid no process has)."""
    return {"pid": 4000000, "token": "4000000-stopped", "pid_created": 1.0,
            "host": safesave._local_host(), "time": time.time(), "staged": [], **fields}


def _write_entry(record: dict, name: str = "entry") -> Path:
    folder = safesave._journal_dir()
    folder.mkdir(parents=True, exist_ok=True)
    entry = folder / f"{name}.json"
    entry.write_text(json.dumps(record), encoding="utf-8")
    return entry


# ------------------------------------------------- only the deck's own files


@pytest.mark.parametrize("target", ["the_deck", "another_file"])
def test_recovery_never_removes_a_file_that_is_not_the_decks_own(tmp_path, target):
    """An entry names the deck itself, or another file, as the new version
    whose save did not finish. The product never writes such an entry; a
    damaged or hand-made one must still not cost the user a file."""
    deck = _deck(tmp_path / "Deck.pptx", "v1")
    other = _deck(tmp_path / "Important.pptx", "keep me")
    _write_entry(_stopped_server_entry(
        deck=str(deck), aside=str(tmp_path / "x.tmp"),
        new=str(deck if target == "the_deck" else other)))
    server._recover_at_start(Settings())
    assert deck.exists() and _notes(deck) == "v1"
    assert other.exists() and _notes(other) == "keep me"
    assert _entries_for(deck) == []


def test_recovery_never_moves_a_file_named_as_the_old_copy_that_is_not_the_decks_own(tmp_path):
    deck = tmp_path / "Deck.pptx"
    stranger = _deck(tmp_path / "Somebody else's.pptx", "theirs")
    _write_entry(_stopped_server_entry(deck=str(deck), aside=str(stranger),
                                       new=str(deck) + ".ppt-mcp-tmp"))
    server._recover_at_start(Settings())
    assert not deck.exists() and _notes(stranger) == "theirs"


# ------------------------------------------------- a deck saved again since


def test_a_stop_after_the_swap_never_lands_an_older_backup_over_a_newer_prev(tmp_path):
    """After the stop, a server that cannot see the entry (another journal,
    as an older build has none) saves the deck: prev holds v2, the version
    before that change. The next start must not put the staged v1 over it.
    The older copy is kept, and the next answer says where."""
    deck = _killed_after_the_swap(tmp_path)  # deck v2, aside v1, prev v0, v1 staged
    _save_in_a_child(deck, "v3", journal=tmp_path / "another-journal")
    assert _notes(deck) == "v3" and _notes(_prev(deck)) == "v2"
    server._recover_at_start(Settings())
    assert _notes(_prev(deck)) == "v2", "the start put an older backup over prev"
    aside = swapwait.asides(deck)
    assert len(aside) == 1 and _notes(aside[0]) == "v1", "the older copy is kept"
    payload, is_error = _call("get_text", {"file_path": str(deck), "live": "off"})
    assert is_error is False, payload
    assert any(aside[0] in n for n in payload.get("notices", [])), payload.get("notices")
    payload, is_error = _call("get_text", {"file_path": str(deck), "live": "off"})
    assert "notices" not in payload, "each notice is given once"


def test_the_old_copy_goes_once_its_deck_was_saved_again(tmp_path):
    """Another program held the old file at the end of a save, so it stayed
    beside the deck and the entry was kept; then another server saved the
    deck. Once the hold is over, the next change removes the old copy: its
    backups had landed and the deck has moved on, as after any save."""
    deck = _deck(tmp_path / "Deck.pptx", "v0")
    assert not _set_notes(deck, "v1")[1]
    k32 = ctypes.WinDLL("kernel32", use_last_error=True)
    k32.CreateFileW.restype = ctypes.c_void_p
    k32.CreateFileW.argtypes = [ctypes.c_wchar_p, ctypes.c_uint32, ctypes.c_uint32,
                                ctypes.c_void_p, ctypes.c_uint32, ctypes.c_uint32,
                                ctypes.c_void_p]
    k32.CloseHandle.argtypes = [ctypes.c_void_p]
    fn, move, last_error, text = safesave._replace_file()
    held = {}

    def then_hold(ldst, lsrc, lbak, flags, a, b):
        ok = fn(ldst, lsrc, lbak, flags, a, b)
        if ok and "h" not in held:  # read, shares read and write, not delete
            held["h"] = k32.CreateFileW(lbak, 0x80000000, 0x3, None, 3, 0, None)
        return ok

    safesave._REPLACE_FILE = (then_hold, move, last_error, text)
    try:
        assert not _set_notes(deck, "v2")[1]
    finally:
        safesave._REPLACE_FILE = (fn, move, last_error, text)
    assert len(swapwait.asides(deck)) == 1
    _save_in_a_child(deck, "v3")
    k32.CloseHandle(held["h"])
    assert not _set_notes(deck, "v4")[1]
    assert _notes(deck) == "v4" and _notes(_prev(deck)) == "v3"
    assert swapwait.asides(deck) == [], "an old copy stayed beside the deck"
    assert _entries_for(deck) == []


# ------------------------------------------ renamed or deleted after a save


@pytest.mark.parametrize("user_action", ["rename", "delete"])
def test_a_deck_renamed_or_deleted_after_its_save_went_through_is_not_recreated(
        tmp_path, user_action):
    deck = _killed_after_the_swap(tmp_path)
    if user_action == "rename":
        os.replace(deck, tmp_path / "Final.pptx")
    else:
        os.remove(deck)
    assert server._recover_at_start(Settings()) == []
    assert not deck.exists(), "an older copy was made into the deck again"
    aside = swapwait.asides(deck)
    assert len(aside) == 1 and _notes(aside[0]) == "v1"
    # Not even by a later write lock on the old name.
    with safesave.write_lock(deck):
        pass
    assert not deck.exists()
    # The next answer says where the older copy is, once.
    payload, _ = _call("get_server_info", {})
    assert any(aside[0] in n and "not put back" in n for n in payload.get("notices", []))
    # A read of the old name says the same, and that nothing comes back.
    payload, is_error = _call("get_text", {"file_path": str(deck), "live": "off"})
    assert is_error is True and payload["error"]["code"] == "NOT_FOUND", payload
    message = payload["error"]["message"]
    assert aside[0] in message and "not put back" in message
    assert "puts it back" not in message


def test_a_deck_renamed_back_after_a_stop_gets_its_backups_and_loses_the_old_copy(tmp_path):
    deck = _killed_after_the_swap(tmp_path)
    final = tmp_path / "Final.pptx"
    os.replace(deck, final)
    server._recover_at_start(Settings())
    os.replace(final, deck)  # the user renames it back
    assert not _set_notes(deck, "v3")[1]
    assert _notes(deck) == "v3" and _notes(_prev(deck)) == "v2"
    assert swapwait.asides(deck) == [] and _entries_for(deck) == []


def test_a_swap_that_went_through_with_its_old_copy_held_is_not_undone_by_a_delete(tmp_path):
    """The swap went through and another program held the old file, so the
    entry stayed. The user then deletes the deck: the old copy must not come
    back as the deck."""
    deck = _deck(tmp_path / "Deck.pptx", "v0")
    assert not _set_notes(deck, "v1")[1]
    fn, move, last_error, text = safesave._replace_file()
    k32 = ctypes.WinDLL("kernel32", use_last_error=True)
    k32.CreateFileW.restype = ctypes.c_void_p
    k32.CreateFileW.argtypes = [ctypes.c_wchar_p, ctypes.c_uint32, ctypes.c_uint32,
                                ctypes.c_void_p, ctypes.c_uint32, ctypes.c_uint32,
                                ctypes.c_void_p]
    k32.CloseHandle.argtypes = [ctypes.c_void_p]
    held = {}

    def then_hold(ldst, lsrc, lbak, flags, a, b):
        ok = fn(ldst, lsrc, lbak, flags, a, b)
        if ok and "h" not in held:
            held["h"] = k32.CreateFileW(lbak, 0x80000000, 0x3, None, 3, 0, None)
        return ok

    safesave._REPLACE_FILE = (then_hold, move, last_error, text)
    try:
        assert not _set_notes(deck, "v2")[1]
    finally:
        safesave._REPLACE_FILE = (fn, move, last_error, text)
    k32.CloseHandle(held["h"])
    os.remove(deck)
    assert safesave.recover_interrupted_saves() == []
    assert not deck.exists()


# ------------------------------------------------------------------ the start


def test_start_up_recovery_keeps_to_its_time_when_a_folder_does_not_answer(
        tmp_path, monkeypatch):
    """A deck on a share that does not answer: asking whether its folder is
    there can take 20 seconds. The start goes on without it."""
    slow = str(tmp_path / "offline-share")
    _write_entry(_stopped_server_entry(deck=slow + "\\Deck.pptx", aside=slow + "\\x.tmp",
                                       new=slow + "\\Deck.pptx.ppt-mcp-tmp"))
    real_isdir = os.path.isdir

    def no_answer(path):
        if str(path).startswith(slow):
            time.sleep(12)
            return False
        return real_isdir(path)

    monkeypatch.setattr(safesave.os.path, "isdir", no_answer)
    started = time.monotonic()
    safesave.recover_interrupted_saves()
    assert time.monotonic() - started < 6


def test_an_entry_whose_folder_has_been_gone_for_a_week_is_dropped(tmp_path):
    gone = tmp_path / "moved-away"
    old = _write_entry(_stopped_server_entry(
        deck=str(gone / "Old.pptx"), aside=str(gone / "x.tmp"),
        new=str(gone / "Old.pptx.ppt-mcp-tmp"),
        folder_gone_since=time.time() - 8 * 24 * 3600), name="old")
    fresh = _write_entry(_stopped_server_entry(
        deck=str(gone / "New.pptx"), aside=str(gone / "y.tmp"),
        new=str(gone / "New.pptx.ppt-mcp-tmp")), name="fresh")
    safesave.recover_interrupted_saves()
    assert not old.exists(), "an entry whose folder has been gone a week is dropped"
    assert fresh.exists(), "a folder gone only now is kept for later starts"
    assert "folder_gone_since" in json.loads(fresh.read_text(encoding="utf-8"))


# --------------------------------------------------------------- staged links


def test_staged_backup_links_left_over_an_hour_are_swept_at_the_next_write(tmp_path):
    deck = _deck(tmp_path / "Deck.pptx", "v1")
    assert not _set_notes(deck, "v2")[1]
    folder = safesave.slot_dir(deck)
    hour_and_more = int(time.time()) - 2 * 3600
    old = folder / f".slot-{hour_and_more}-{'a' * 32}.tmp"
    young = folder / f".slot-{int(time.time())}-{'b' * 32}.tmp"
    legacy = folder / f".slot-{'c' * 32}.tmp"
    named = folder / f".slot-{hour_and_more}-{'d' * 32}.tmp"
    for link in (old, young, legacy, named):
        link.write_bytes(b"a staged backup")
    os.utime(legacy, (hour_and_more, hour_and_more))
    _write_entry({"deck": str(tmp_path / "Other.pptx"), "aside": "", "new": "",
                  "staged": [[str(named), str(folder / "prev.pptx")]],
                  "pid": os.getppid(), "token": "a running server",
                  "pid_created": safesave._process_create_time(os.getppid()),
                  "host": safesave._local_host(), "time": time.time()})
    with safesave.write_lock(deck):
        pass
    assert not old.exists() and not legacy.exists()
    assert young.exists(), "a link younger than an hour may belong to a save"
    assert named.exists(), "a link a journal entry names is kept"


# ------------------------------------------------------- what the answers say


def test_a_save_without_replacefile_says_when_the_deck_lost_its_own_permissions(
        tmp_path, monkeypatch):
    deck = _deck(tmp_path / "Deck.pptx", "v0")
    me = subprocess.run(["whoami", "/user", "/fo", "csv", "/nh"], capture_output=True,
                        text=True, stdin=subprocess.DEVNULL, creationflags=NO_WINDOW,
                        timeout=60).stdout.strip()
    sid = me.split(",")[-1].strip('"')
    done = subprocess.run(["icacls", str(deck), "/inheritance:r", "/grant:r", f"*{sid}:(F)"],
                          capture_output=True, text=True, stdin=subprocess.DEVNULL,
                          creationflags=NO_WINDOW, timeout=60)
    assert done.returncode == 0, done.stdout + done.stderr
    before = safesave._dacl_text(str(deck))
    fn, move, last_error, text = safesave._replace_file()

    def not_supported(*_args):
        ctypes.set_last_error(50)
        return False

    monkeypatch.setattr(safesave, "_REPLACE_FILE", (not_supported, move, last_error, text))
    payload, is_error = _set_notes(deck, "v1")
    assert is_error is False, payload
    assert safesave._dacl_text(str(deck)) != before
    assert any("permissions" in w for w in payload["warnings"]), payload


def test_the_refusal_for_a_deck_left_aside_names_no_particular_operation(tmp_path):
    exc = safesave.DeckLeftAside(str(tmp_path / "deck.pptx"),
                                 str(tmp_path / "aside.tmp"), "error 1177", 1177)
    refusal = safesave.os_write_refusal(tmp_path / "deck.pptx", exc)
    assert str(refusal).startswith("The change to deck.pptx could not finish")
    assert "save" not in str(refusal).split(".")[0]
