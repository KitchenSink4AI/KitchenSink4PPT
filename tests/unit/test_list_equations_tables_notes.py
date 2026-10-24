"""list_equations lists the equations in table cells and in the speaker notes.

list_equations read equations in shapes' text only, while its help said
"List every native equation". An equation in a table cell or in a slide's
notes was left out, though get_text and find_text show it. PowerPoint keeps
a table that holds an equation in a slide-level compatibility wrapper, and
a notes body that holds one in a wrapper on the notes page; the equation is
read from the branch PowerPoint shows, so it is listed once.

Each entry now says where the equation is, in find_text's words: "slide"
(a shape's text), "table" (with row and col) or "notes". The wrapped forms
below follow PowerPoint's own save of a made-up deck (an equation in a text
box, a table cell and the notes). Every deck here is made up by code.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.package import PptxPackage
import pytest

from kitchensink4ppt.core.errors import UnsupportedStructure
from kitchensink4ppt.ops import equations, notes

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A14 = "http://schemas.microsoft.com/office/drawing/2010/main"
OMML = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _math_paragraph(chars: str) -> etree._Element:
    """A paragraph holding an equation the way PowerPoint writes one into a
    table cell: math straight in the paragraph (a14:m)."""
    xml = (f'<a:p xmlns:a="{A}" xmlns:a14="{A14}"><a14:m><m:oMathPara xmlns:m="{OMML}"><m:oMath>'
           f"<m:r><m:t>{chars}</m:t></m:r></m:oMath></m:oMathPara></a14:m><a:endParaRPr/></a:p>")
    return etree.fromstring(xml)


def _wrap(element: etree._Element, choice_extra) -> etree._Element:
    """Put `element` in an mc:AlternateContent where it stands: the Choice
    (Requires a14) holds a copy that `choice_extra` gives its equation, the
    Fallback holds the element as it was, without one."""
    wrapper = etree.Element(f"{{{MC}}}AlternateContent", nsmap={"mc": MC})
    choice = etree.SubElement(wrapper, f"{{{MC}}}Choice", nsmap={"a14": A14})
    choice.set("Requires", "a14")
    fallback = etree.SubElement(wrapper, f"{{{MC}}}Fallback")
    element.addprevious(wrapper)
    shown = etree.fromstring(etree.tostring(element))
    choice_extra(shown)
    choice.append(shown)
    fallback.append(element)
    return wrapper


def _deck(path: Path, *, wrapped: bool) -> Path:
    """One slide: a text box with an equation (added by the product), a
    1 x 2 table whose second cell holds an equation, and notes whose second
    paragraph holds one. wrapped=True keeps the table and the notes body in
    wrappers, as PowerPoint saves them; wrapped=False leaves them bare."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.name = "Math box"
    box.text_frame.text = "Made-up words"
    frame = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1))
    frame.name = "Math table"
    frame.table.cell(0, 0).text = "Cell words"
    slide.notes_slide.notes_text_frame.text = "Made-up notes"
    prs.save(str(path))
    pkg = PptxPackage(path)
    equations.add_equation_to_shape(pkg, 0, "Math box", "x+1=2")
    pkg.save(do_backup=False)

    def fix(name, data):
        root = etree.fromstring(data)
        if name == "ppt/slides/slide1.xml":
            gf = next(e for e in root.iter(f"{{{P}}}graphicFrame"))

            def cell_math(frame_el):
                tc = list(frame_el.iter(f"{{{A}}}tc"))[1]
                body = tc.find(f"{{{A}}}txBody")
                for p in body.findall(f"{{{A}}}p"):
                    body.remove(p)
                body.append(_math_paragraph("y=2+3"))
            if wrapped:
                _wrap(gf, cell_math)
            else:
                cell_math(gf)
        else:
            body_sp = next(sp for sp in root.iter(f"{{{P}}}sp")
                           if sp.find(f".//{{{P}}}ph") is not None
                           and sp.find(f".//{{{P}}}ph").get("type") == "body")

            def notes_math(sp_el):
                sp_el.find(f"{{{P}}}txBody").append(_math_paragraph("z=4"))
            if wrapped:
                _wrap(body_sp, notes_math)
            else:
                notes_math(body_sp)
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    tmp = path.with_name(path.name + ".rw")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        notes = [n for n in zin.namelist() if n.startswith("ppt/notesSlides/") and n.endswith(".xml")]
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "ppt/slides/slide1.xml" or info.filename in notes:
                data = fix(info.filename, data)
            zout.writestr(info, data)
    tmp.replace(path)
    return path


def _ids(path: Path) -> dict[str, int]:
    """Shape name -> id on slide 1, read from the XML (python-pptx does not
    see a shape inside a wrapper)."""
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("ppt/slides/slide1.xml"))
    return {c.get("name"): int(c.get("id")) for c in root.iter(f"{{{P}}}cNvPr")}


def _listed(path: Path, scope=None) -> list[dict]:
    out = equations.list_equations(PptxPackage(path), scope)
    assert out["equation_count"] == len(out["equations"])
    return out["equations"]


def _check(path: Path, entries: list[dict]) -> None:
    ids = _ids(path)
    by_where = {e["where"]: e for e in entries}
    assert sorted(by_where) == ["notes", "slide", "table"], entries
    assert len(entries) == 3, entries
    shape = by_where["slide"]
    assert (shape["shape_id"], shape["paragraph"], shape["text"]) == (ids["Math box"], 1, "x+1=2")
    cell = by_where["table"]
    assert (cell["row"], cell["col"], cell["paragraph"], cell["text"]) == (0, 1, 0, "y=2+3"), cell
    assert cell["shape_id"] == ids["Math table"], cell
    note = by_where["notes"]
    assert (note["paragraph"], note["text"]) == (1, "z=4"), note
    assert isinstance(note["shape_id"], int)
    for e in entries:
        assert (e["slide_index"], e["slide_id"]) == (0, 256)


def test_list_equations_lists_a_table_cell_and_the_notes(tmp_path):
    path = _deck(tmp_path / "made_up_bare.pptx", wrapped=False)
    _check(path, _listed(path))


def test_list_equations_lists_them_once_where_powerpoint_wrapped_them(tmp_path):
    """The table and the notes body in PowerPoint's wrappers: listed from the
    branch PowerPoint shows, never twice."""
    path = _deck(tmp_path / "made_up_wrapped.pptx", wrapped=True)
    _check(path, _listed(path))


def test_scope_selects_the_slides_whose_notes_are_read(tmp_path):
    path = _deck(tmp_path / "made_up_scope.pptx", wrapped=True)
    assert len(_listed(path, scope=0)) == 3
    prs = Presentation(str(path))
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(path))
    assert _listed(path, scope=1) == []


def test_the_help_names_table_cells_and_notes():
    from kitchensink4ppt import server  # noqa: F401
    from kitchensink4ppt.settings import Settings
    from kitchensink4ppt.surface.help import help_payload
    from kitchensink4ppt.surface.plan import build

    summary = str(help_payload(build(Settings()), "list_equations").get("summary") or "")
    assert "slides" in summary and "tables" in summary and "notes" in summary, summary


def _notes_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        name = next(n for n in z.namelist() if n.startswith("ppt/notesSlides/") and n.endswith(".xml"))
        return z.read(name).decode("utf-8")


def _notes_equations(path: Path) -> list[str]:
    return [e["text"] for e in _listed(path) if e["where"] == "notes"]


@pytest.mark.parametrize("wrapped", [False, True])
def test_a_plain_notes_write_that_would_drop_an_equation_is_refused(tmp_path, wrapped):
    path = _deck(tmp_path / "made_up_notes.pptx", wrapped=wrapped)
    before = path.read_bytes()
    assert _notes_equations(path) == ["z=4"]
    pkg = PptxPackage(path)
    with pytest.raises(UnsupportedStructure) as refused:
        notes.set_notes(pkg, 0, "Replacement notes")
    assert "equation" in str(refused.value)
    assert "flatten=True" in str(refused.value)
    assert path.read_bytes() == before
    assert _notes_equations(path) == ["z=4"]


def test_a_structured_notes_rebuild_that_would_drop_an_equation_is_refused(tmp_path):
    path = _deck(tmp_path / "made_up_rebuild.pptx", wrapped=False)
    pkg = PptxPackage(path)
    with pytest.raises(UnsupportedStructure) as refused:
        notes.set_notes(pkg, 0, paragraphs=[{"runs": [{"text": "Only one paragraph"}]}])
    assert "equation" in str(refused.value)
    assert _notes_equations(path) == ["z=4"]


def test_flatten_true_removes_the_notes_equation_and_says_so(tmp_path):
    path = _deck(tmp_path / "made_up_flat_bare.pptx", wrapped=False)
    pkg = PptxPackage(path)
    out = notes.set_notes(pkg, 0, "Replacement notes", flatten=True)
    pkg.save(do_backup=False)
    assert out["flattened"]["equations"] == 1
    assert any("equation" in d for d in out["flattened"]["described"]), out
    assert _notes_equations(path) == []
    assert "Replacement notes" in _notes_xml(path)


def test_flatten_true_on_a_wrapped_notes_body_leaves_one_copy_and_says_so(tmp_path):
    path = _deck(tmp_path / "made_up_flat_wrapped.pptx", wrapped=True)
    pkg = PptxPackage(path)
    out = notes.set_notes(pkg, 0, "Replacement notes", flatten=True)
    pkg.save(do_backup=False)
    assert any("wrapper" in d for d in out["flattened"]["described"]), out
    xml = _notes_xml(path)
    assert "AlternateContent" not in xml
    assert xml.count("Replacement notes") == 1 and "Made-up notes" not in xml
    assert _notes_equations(path) == []
    # the slide's and the table's equations stay
    assert sorted(e["where"] for e in _listed(path)) == ["slide", "table"]


def test_an_unchanged_plain_notes_write_still_leaves_the_equation_and_the_wrapper(tmp_path):
    path = _deck(tmp_path / "made_up_same.pptx", wrapped=True)
    before = path.read_bytes()
    pkg = PptxPackage(path)
    out = notes.set_notes(pkg, 0, "Made-up notes\nz=4", flatten=True)
    assert out["changed"] is False
    assert path.read_bytes() == before
