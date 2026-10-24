"""export_slide_image says how it counts, where the model reads it.

Field report: export_slide_image {slides: [44]} meant the user's slide 44,
rendered the 45th slide as slide45.png, and answered {"slide": 44, "file":
".../slide45.png"}: two numbering schemes in one answer and nothing in the
tool description to say slides counts from 0 (only ppt_help said so). The
tool description now says it, names slide_numbers (counted from 1) and
width; every image record leads with the slide_number its file is named
after; and an answer to a call that sent slides says how those numbers
were read.
"""

from __future__ import annotations

import asyncio
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt import server  # noqa: F401  (the import registers ops)
from kitchensink4ppt.com import bridge
from kitchensink4ppt.ops import export as ex
from kitchensink4ppt.registry import public_ops
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.plan import build


def _plan():
    return build(Settings(), public_ops())


def _call(args: dict):
    return asyncio.run(dispatch(_plan(), "ppt_deck", {"op": "export_slide_image",
                                                      "args": args}))


@pytest.fixture
def deck(tmp_path):
    prs = Presentation()
    for n in range(1, 5):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = (
            f"Slide {n}")
    path = tmp_path / "render.pptx"
    prs.save(str(path))
    return str(path)


@pytest.fixture
def fake_com(monkeypatch):
    """PowerPoint present but never started: the bridge writes what
    PowerPoint would and answers as it does."""

    def images(path, output_dir=None, slides=None, width=1280, height=None, **kw):
        out = Path(output_dir)
        out.mkdir(parents=True, exist_ok=True)
        wanted = slides if slides is not None else range(4)
        for index in wanted:
            (out / f"slide{index + 1}.png").write_bytes(b"png")
        return {"images": [{"slide": i, "file": str(out / f"slide{i + 1}.png")}
                           for i in wanted], "engine": "powerpoint-com"}

    monkeypatch.setattr(ex, "_com_available", lambda: True)
    monkeypatch.setattr(bridge, "com_export_slide_images", images)


def test_the_tool_description_says_how_slides_count():
    description = _plan().tools["ppt_deck"].description
    line = next(row for row in description.splitlines()
                if row.startswith("- export_slide_image:"))
    for words in ("slides", "from 0", "slide_numbers", "from 1", "width"):
        assert words in line, (words, line)


def test_each_image_leads_with_the_slide_number_its_file_is_named_after(deck, tmp_path, fake_com):
    payload, is_error = _call({"file_path": deck, "output_dir": str(tmp_path / "png"),
                               "slides": [1, 3], "engine": "com"})
    assert is_error is False, payload
    for image in payload["images"]:
        assert next(iter(image)) == "slide_number"
        assert Path(image["file"]).name == f"slide{image['slide_number']}.png"
        assert image["slide"] == image["slide_number"] - 1


def test_an_answer_to_slides_says_how_they_were_read(deck, tmp_path, fake_com):
    payload, is_error = _call({"file_path": deck, "output_dir": str(tmp_path / "png"),
                               "slides": [1], "engine": "com"})
    assert is_error is False, payload
    note = payload["slides_note"]
    assert "counts from 0" in note and "slide_number" in note and "2" in note


def test_slide_numbers_need_no_note(deck, tmp_path, fake_com):
    payload, is_error = _call({"file_path": deck, "output_dir": str(tmp_path / "png"),
                               "slide_numbers": [2], "engine": "com"})
    assert is_error is False, payload
    assert "slides_note" not in payload
    assert payload["images"][0]["slide_number"] == 2
    assert Path(payload["images"][0]["file"]).name == "slide2.png"


def test_every_record_is_numbered_whatever_the_engine_answered():
    out = ex._numbered({"images": [{"slide": 44, "file": "x/slide45.png"},
                                   {"slide": 0, "slide_number": 1, "file": "x/slide1.png"}]})
    assert [list(image)[0] for image in out["images"]] == ["slide_number", "slide_number"]
    assert out["images"][0]["slide_number"] == 45
