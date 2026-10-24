"""find_text and the other scoped reads: one slide, some slides, a range.

Field report: find_text for a symbol returned 22 matches across the deck when
one slide was wanted, because nothing in the tool description said find_text
takes a scope; and get_presentation_view {slides: [28]} was refused, since
its argument is scope. find_text's scope is now in the description, every
scope takes a range ({"start", "end"}, both included; counted from 0 in
scope and from 1 in slide_numbers), and slides is read as another name for
scope on the operations that take one, under the rule for such names: sent
beside scope with a different value, the call is refused with a hint.
"""

from __future__ import annotations

import asyncio

import pytest
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt import server  # noqa: F401  (the import registers ops)
from kitchensink4ppt.core.package import PptxPackage
from kitchensink4ppt.ops.read import slide_table
from kitchensink4ppt.registry import public_ops
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.help import help_payload
from kitchensink4ppt.surface.plan import build


def _plan():
    return build(Settings(), public_ops())


def _call(tool: str, op: str, args: dict):
    return asyncio.run(dispatch(_plan(), tool, {"op": op, "args": args}))


def _ok(op: str, args: dict, tool: str = "ppt_read") -> dict:
    payload, is_error = _call(tool, op, args)
    assert is_error is False, payload
    return payload


def _refused(op: str, args: dict, tool: str = "ppt_read") -> dict:
    payload, is_error = _call(tool, op, args)
    assert is_error is True, payload
    return payload["error"]


@pytest.fixture()
def deck(tmp_path):
    """Five slides, each with a text box saying "Marker on slide N"."""
    prs = Presentation()
    for n in range(1, 6):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        box.text_frame.text = f"Marker on slide {n}"
    path = tmp_path / "scoped.pptx"
    prs.save(str(path))
    return str(path)


def _slides_hit(payload: dict) -> list[int]:
    return sorted({m["slide_index"] + 1 for m in payload["matches"]})


def _ids(deck) -> list[int]:
    return [rec["slide_id"] for rec in slide_table(PptxPackage(deck))]


def test_find_text_by_slide_id(deck):
    sid = _ids(deck)[3]
    out = _ok("find_text", {"file_path": deck, "query": "Marker", "scope": {"slide_id": sid}})
    assert _slides_hit(out) == [4] and out["count"] == 1


def test_find_text_by_slide_numbers(deck):
    out = _ok("find_text", {"file_path": deck, "query": "Marker", "slide_numbers": [2, 5]})
    assert _slides_hit(out) == [2, 5]


def test_find_text_by_a_range_counted_from_zero(deck):
    out = _ok("find_text", {"file_path": deck, "query": "Marker",
                            "scope": {"start": 1, "end": 3}})
    assert _slides_hit(out) == [2, 3, 4]


def test_find_text_by_a_range_of_slide_numbers(deck):
    out = _ok("find_text", {"file_path": deck, "query": "Marker",
                            "slide_numbers": {"start": 2, "end": 4}})
    assert _slides_hit(out) == [2, 3, 4]


def test_a_list_may_mix_slides_and_ranges(deck):
    out = _ok("find_text", {"file_path": deck, "query": "Marker",
                            "scope": [0, {"start": 3, "end": 4}]})
    assert _slides_hit(out) == [1, 4, 5]


def test_a_range_past_the_deck_names_its_range(deck):
    error = _refused("find_text", {"file_path": deck, "query": "Marker",
                                   "scope": {"start": 3, "end": 5}})
    assert "5 slides" in error["message"]
    error = _refused("find_text", {"file_path": deck, "query": "Marker",
                                   "scope": {"start": 3, "end": 1}})
    assert "start" in error["message"] and "end" in error["message"]


def test_a_range_of_slide_numbers_counts_from_one(deck):
    error = _refused("find_text", {"file_path": deck, "query": "Marker",
                                   "slide_numbers": {"start": 0, "end": 2}})
    assert "counted from 1" in error["message"]
    assert '{"start": 1, "end": 3}' in error["message"]


def test_get_presentation_view_takes_slides_as_scope(deck):
    out = _ok("get_presentation_view", {"file_path": deck, "slides": [1]})
    assert "Marker on slide 2" in out["view"]
    assert "Marker on slide 1" not in out["view"]
    assert out["normalized_from"]["args"] == {"slides": [1]}
    assert "scope" in out["normalized_from"]["hint"]


def test_get_presentation_view_takes_slide_numbers(deck):
    out = _ok("get_presentation_view", {"file_path": deck, "slide_numbers": [3]})
    assert "Marker on slide 3" in out["view"] and "Marker on slide 2" not in out["view"]


def test_slides_beside_scope_with_another_value_is_refused(deck):
    same = _ok("get_presentation_view", {"file_path": deck, "slides": [1], "scope": [1]})
    assert "Marker on slide 2" in same["view"]
    error = _refused("get_presentation_view", {"file_path": deck, "slides": [1], "scope": [2]})
    assert error["code"] == "BAD_ARGS"
    assert "slides" in error["message"] and "scope" in error["message"]
    assert error.get("hint")


def test_slides_beside_slide_numbers_names_both(deck):
    error = _refused("get_presentation_view", {"file_path": deck, "slides": [1],
                                               "slide_numbers": [2]})
    assert "slides and slide_numbers" in error["message"]


@pytest.mark.parametrize("op", ["get_text", "find_text", "list_hyperlinks"])
def test_other_scoped_reads_take_slides_too(deck, op):
    args = {"file_path": deck, "slides": [4]}
    if op == "find_text":
        args["query"] = "Marker"
    out = _ok(op, args)
    assert out["normalized_from"]["args"] == {"slides": [4]}


def test_the_tool_description_says_find_text_takes_a_scope():
    description = _plan().tools["ppt_read"].description
    line = next(row for row in description.splitlines() if row.startswith("- find_text:"))
    assert "scope" in line and "slide_numbers" in line


def test_help_describes_the_range():
    props = help_payload(_plan(), "find_text")["args_schema"]["properties"]
    assert "start" in props["scope"]["description"]
    assert "start" in props["slide_numbers"]["description"]
