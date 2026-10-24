"""The corpus the suite reads is generated, and only generated.

tests/make_corpus.py builds every corpus deck from code. These tests pin the
rules around it on a private copy of the folder: a file under a corpus name
that the generator did not write is refused (never read, never overwritten),
generated decks are reused while the generator is unchanged, and rebuilt
when it changes or when a deck no longer matches what was generated.
"""

from __future__ import annotations

import json

import make_corpus
import pytest
from pptx import Presentation


@pytest.fixture()
def corpus(tmp_path, monkeypatch):
    """make_corpus pointed at an empty folder of this test's own."""
    folder = tmp_path / "corpus"
    monkeypatch.setattr(make_corpus, "CORPUS", folder)
    monkeypatch.setattr(make_corpus, "MANIFEST", folder / ".generated.json")
    monkeypatch.setattr(make_corpus, "LOCK", folder / ".generate.lock")
    monkeypatch.setattr(make_corpus, "GENERATORS", {
        "one_deck.pptx": lambda p: make_corpus.build_deck(p, seed=1, extra_slides=0),
        "two_deck.pptx": lambda p: make_corpus.build_deck(p, seed=2, extra_slides=1),
    })
    return folder


def _foreign_deck(path):
    """A valid deck that the generator did not write."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Somebody else's deck"
    prs.save(str(path))
    return path


def test_every_deck_is_built_and_recorded(corpus):
    made = make_corpus.ensure_corpus()
    assert sorted(made) == ["one_deck.pptx", "two_deck.pptx"]
    for name in made:
        assert make_corpus.is_generated(corpus / name)
    manifest = json.loads((corpus / ".generated.json").read_text(encoding="utf-8"))
    assert manifest["generator"] == make_corpus.GENERATOR_VERSION
    assert sorted(manifest["files"]) == ["one_deck.pptx", "two_deck.pptx"]


def test_generated_decks_are_reused_while_the_generator_is_unchanged(corpus):
    make_corpus.ensure_corpus()
    stamp = (corpus / "one_deck.pptx").stat().st_mtime_ns
    assert make_corpus.ensure_corpus() == []
    assert (corpus / "one_deck.pptx").stat().st_mtime_ns == stamp


def test_a_file_the_generator_did_not_write_is_refused_and_left_alone(corpus):
    corpus.mkdir()
    foreign = _foreign_deck(corpus / "two_deck.pptx")
    before = foreign.read_bytes()
    assert not make_corpus.is_generated(foreign)
    with pytest.raises(make_corpus.ForeignDeck, match="two_deck.pptx"):
        make_corpus.ensure_corpus()
    assert foreign.read_bytes() == before, "the refused file was overwritten"
    assert not (corpus / "one_deck.pptx").exists(), "nothing is built after a refusal"


def test_a_file_that_is_not_a_deck_is_refused_too(corpus):
    corpus.mkdir()
    (corpus / "one_deck.pptx").write_bytes(b"not a presentation")
    with pytest.raises(make_corpus.ForeignDeck):
        make_corpus.ensure_corpus()


def test_a_new_generator_version_rebuilds_every_deck(corpus, monkeypatch):
    make_corpus.ensure_corpus()
    monkeypatch.setattr(make_corpus, "GENERATOR_VERSION", make_corpus.GENERATOR_VERSION + 1)
    assert sorted(make_corpus.ensure_corpus()) == ["one_deck.pptx", "two_deck.pptx"]


def test_a_generated_deck_that_changed_is_rebuilt(corpus):
    make_corpus.ensure_corpus()
    target = corpus / "one_deck.pptx"
    make_corpus.build_deck(target, seed=9, extra_slides=3)  # still generated, different bytes
    assert make_corpus.ensure_corpus() == ["one_deck.pptx"]


def test_the_suite_corpus_holds_only_generated_decks():
    """The real corpus folder after conftest ran: every deck the tests read
    was written by the generator."""
    for name in make_corpus.GENERATORS:
        path = make_corpus.CORPUS / name
        if path.exists():
            assert make_corpus.is_generated(path), name
