"""Text inside a paragraph's compatibility wrapper is read, and never rewritten.

PowerPoint stores an inline equation, or a run written for a newer version,
inside an mc:AlternateContent in the paragraph. The text reads walked only
plain runs, so those words were missing from get_text, get_table, get_notes,
extract_text, find_text and the view, while PowerPoint shows them: the
equation when it reads the Choice's namespace, the Fallback's words when it
does not. Reads now include the branch PowerPoint shows. Edits treat the
wrapper as one unit: offsets after it stay right, a replacement or a format
range never cuts into it, and a write that would flatten it says so.

Every deck here is made up by code.
"""

from __future__ import annotations

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.errors import UnsupportedStructure
from kitchensink4ppt.core.package import PptxPackage, qn
from kitchensink4ppt.ops import batch, notes, read, tables, text, view

MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
A14 = "http://schemas.microsoft.com/office/drawing/2010/main"
OMML = "http://schemas.openxmlformats.org/officeDocument/2006/math"
FUTURE = "http://example.com/made-up/future/2033"
AC = f"{{{MC}}}AlternateContent"


def _math(expr, fallback):
    return etree.fromstring(
        f'<mc:AlternateContent xmlns:mc="{MC}"><mc:Choice xmlns:a14="{A14}" Requires="a14">'
        f'<a14:m><m:oMathPara xmlns:m="{OMML}"><m:oMath><m:r><m:t>{expr}</m:t></m:r></m:oMath>'
        f'</m:oMathPara></a14:m></mc:Choice><mc:Fallback><a:r xmlns:a="{A}"><a:t>{fallback}</a:t>'
        "</a:r></mc:Fallback></mc:AlternateContent>"
    )


def _future(choice, fallback):
    return etree.fromstring(
        f'<mc:AlternateContent xmlns:mc="{MC}" xmlns:x33="{FUTURE}"><mc:Choice Requires="x33">'
        f'<a:r xmlns:a="{A}"><a:t>{choice}</a:t></a:r></mc:Choice><mc:Fallback>'
        f'<a:r xmlns:a="{A}"><a:t>{fallback}</a:t></a:r></mc:Fallback></mc:AlternateContent>'
    )


@pytest.fixture()
def deck(tmp_path):
    """Slide 0: box "Speed" holds 'Speed ' + equation v=d/t + ' stays constant';
    box "Future" holds a wrapper PowerPoint shows by its Fallback; a 1x2 table
    whose cell (0, 1) is such a wrapper; notes with an equation."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    speed = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    speed.name = "Speed"
    speed.text_frame.text = "Speed "
    speed.text_frame.add_paragraph().text = "Second line"
    future = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    future.name = "Future"
    future.text_frame.text = "Note: "
    frame = slide.shapes.add_table(1, 2, Inches(1), Inches(4), Inches(6), Inches(0.8))
    frame.table.cell(0, 0).text = "plain"
    slide.notes_slide.notes_text_frame.text = "Notes start "
    path = tmp_path / "wrapped_text.pptx"
    prs.save(str(path))

    pkg = PptxPackage(path)
    part = pkg.slide_parts()[0]
    root = pkg.root(part)
    boxes = {s.find(f".//{qn('p:cNvPr')}").get("name"): s for s in root.iter(qn("p:sp"))}
    p = boxes["Speed"].find(f".//{qn('a:p')}")
    p.append(_math("v=d/t", "[speed equation]"))
    tail = etree.SubElement(p, qn("a:r"))
    etree.SubElement(tail, qn("a:t")).text = " stays constant"
    boxes["Future"].find(f".//{qn('a:p')}").append(_future("choice words", "fallback words in a box"))
    cell_p = list(root.iter(qn("a:tc")))[1].find(f".//{qn('a:p')}")
    for r in cell_p.findall(qn("a:r")):
        cell_p.remove(r)
    cell_p.insert(0, _future("choice cell", "fallback words in a cell"))
    pkg.mark_dirty(part)
    npart = [n for n in pkg.part_names() if n.startswith("ppt/notesSlides/") and n.endswith(".xml")][0]
    for sp in pkg.root(npart).iter(qn("p:sp")):
        ph = sp.find(f".//{qn('p:ph')}")
        if ph is not None and ph.get("type") == "body":
            sp.find(f".//{qn('a:p')}").append(_math("x=2", "[notes equation]"))
    pkg.mark_dirty(npart)
    pkg.save(do_backup=False)
    return path


def _shape_id(path, name):
    pkg = PptxPackage(path)
    for sp in pkg.root(pkg.slide_parts()[0]).iter(qn("p:sp")):
        cnv = sp.find(f".//{qn('p:cNvPr')}")
        if cnv.get("name") == name:
            return int(cnv.get("id"))
    raise AssertionError(name)


def _wrappers(path):
    pkg = PptxPackage(path)
    xml = etree.tostring(pkg.root(pkg.slide_parts()[0]))
    return xml.count(b"</mc:AlternateContent>"), xml.count(b"<m:t>v=d/t</m:t>")


def _run(path, fn, *args, **kwargs):
    pkg = PptxPackage(path)
    out = fn(pkg, *args, **kwargs)
    pkg.save(do_backup=False)
    return out


# ------------------------------------------------------------------- reads


def test_get_text_reads_what_powerpoint_shows(deck):
    out = read.get_text(PptxPackage(deck), include_notes=True)
    body = out["slides"][0]["text"]
    assert "Speed v=d/t stays constant" in body
    assert "Note: fallback words in a box" in body
    assert "choice words" not in body
    assert "fallback words in a cell" in body
    assert "Notes start x=2" in out["slides"][0]["notes"]


def test_find_text_finds_wrapped_words_with_offsets_the_edits_use(deck):
    pkg = PptxPackage(deck)
    hits = read.find_text(pkg, "fallback")["matches"]
    assert len(hits) == 2
    eq = read.find_text(pkg, "v=d/t")["matches"]
    assert len(eq) == 1 and eq[0]["start"] == 6
    after = read.find_text(pkg, "stays")["matches"][0]
    assert after["start"] == len("Speed v=d/t ")


def test_get_table_get_notes_and_the_view_read_it(deck):
    pkg = PptxPackage(deck)
    grid = tables.get_table(pkg, 0, {"index": 0})
    assert [c["text"] for c in grid["cells"]] == ["plain", "fallback words in a cell"]
    got = notes.get_notes(pkg, 0)
    assert got["text"] == "Notes start x=2"
    assert got["paragraphs"][0]["text"] == "Notes start x=2"
    shown = view.get_presentation_view(pkg)["view"]
    assert "Speed v=d/t stays constant" in shown and "fallback words in a cell" in shown


# ------------------------------------------------------------------ writes


def test_a_replacement_after_the_equation_lands_and_keeps_it(deck):
    out = _run(deck, text.search_and_replace, "constant", "steady")
    assert out["total"] == 1
    assert read.get_text(PptxPackage(deck))["slides"][0]["text"].startswith(
        "Speed v=d/t stays steady")
    assert _wrappers(deck) == (3, 1)


def test_a_match_inside_a_wrapper_is_left_and_counted(deck):
    out = _run(deck, text.search_and_replace, "v=d/t", "speed")
    assert out["total"] == 0
    assert out["wrapped_matches_skipped"] == 1
    assert out["matches_found"] == 1 and out["chart_matches_skipped"] == 0
    assert out["slides"][0]["wrapped_matches_skipped"] == 1
    assert any("compatibility wrapper" in w for w in out.get("warnings", []))
    assert _wrappers(deck) == (3, 1)


def test_format_range_from_find_text_hits_the_word_after_the_equation(deck):
    pkg = PptxPackage(deck)
    hit = read.find_text(pkg, "stays")["matches"][0]
    sid = _shape_id(deck, "Speed")
    _run(deck, text.format_text, 0, sid, paragraph=0, start=hit["start"], end=hit["end"], bold=True)
    pkg = PptxPackage(deck)
    for sp in pkg.root(pkg.slide_parts()[0]).iter(qn("p:sp")):
        if sp.find(f".//{qn('p:cNvPr')}").get("name") == "Speed":
            bold = [r.find(qn("a:t")).text for r in sp.iter(qn("a:r"))
                    if r.find(qn("a:rPr")) is not None and r.find(qn("a:rPr")).get("b") == "1"]
            assert bold == ["stays"]
    assert _wrappers(deck) == (3, 1)


def test_a_format_range_into_the_equation_is_refused(deck):
    before = deck.read_bytes()
    sid = _shape_id(deck, "Speed")
    with pytest.raises(UnsupportedStructure) as refused:
        _run(deck, text.format_text, 0, sid, paragraph=0, start=2, end=8, bold=True)
    assert "compatibility wrapper" in str(refused.value)
    assert deck.read_bytes() == before


def test_set_paragraph_on_a_paragraph_with_an_equation_is_refused(deck):
    before = deck.read_bytes()
    sid = _shape_id(deck, "Speed")
    with pytest.raises(Exception) as refused:
        _run(deck, batch.apply_edits, [
            {"op": "set_paragraph", "slide": 0, "shape": sid, "paragraph": 0, "text": "New words"}])
    assert "compatibility wrapper" in str(refused.value)
    assert deck.read_bytes() == before


def test_emptying_the_words_beside_an_equation_keeps_its_paragraph(deck):
    _run(deck, text.search_and_replace, "Speed ", "", remove_empty_paragraphs=True)
    _run(deck, text.search_and_replace, " stays constant", "", remove_empty_paragraphs=True)
    assert _wrappers(deck) == (3, 1)


def test_plain_notes_that_would_flatten_an_equation_are_refused(deck):
    before = deck.read_bytes()
    with pytest.raises(UnsupportedStructure) as refused:
        _run(deck, notes.set_notes, 0, "Other notes")
    assert "equation" in str(refused.value)
    assert deck.read_bytes() == before
    same = _run(deck, notes.set_notes, 0, "Notes start x=2")
    assert same["changed"] is False


def test_a_cell_format_that_would_flatten_a_wrapper_is_refused(deck):
    before = deck.read_bytes()
    with pytest.raises(UnsupportedStructure) as refused:
        _run(deck, tables.set_table_cells, 0, {"index": 0}, [{"row": 0, "col": 1, "bold": True}])
    assert "compatibility wrapper" in str(refused.value)
    assert deck.read_bytes() == before
