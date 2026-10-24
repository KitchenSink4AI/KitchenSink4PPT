"""Reads list and read a shape PowerPoint keeps in a slide-level wrapper.

When PowerPoint saves a text box that holds an equation, it keeps the whole
box in an mc:AlternateContent at the slide level: an mc:Choice that Requires
a14 holds the box with its words and its math, and an mc:Fallback holds a
picture of the box for older readers. Ink and newer charts are kept the same
way, and ungroup_shapes lifts such a group member out whole. Every read
skipped the wrapper, so the box and all of its words were missing from
get_text, find_text, list_equations, get_slide_info, list_elements and the
view. The reads now take the branch PowerPoint shows, at the wrapper's place
in the stacking order. An edit that names such a shape is refused by name
and nothing is changed.

The wrapper below has the form PowerPoint wrote when it saved a made-up deck
of three text boxes with an equation added to the second one. Every deck
here is made up by code.
"""

from __future__ import annotations

import hashlib

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.errors import UnsupportedStructure
from kitchensink4ppt.core.package import PptxPackage, qn
from kitchensink4ppt.ops import equations, read, view
from kitchensink4ppt.ops import shapes as sh
from kitchensink4ppt.ops import text as tx

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A14 = "http://schemas.microsoft.com/office/drawing/2010/main"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
OMML = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NS = f'xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}"'
RT_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"

#: The equation's characters as the reads give them (the math runs' text).
MATH = "\U0001d44e2+\U0001d44f2=\U0001d4502"


def _sup(base: str) -> str:
    return (f"<m:sSup><m:e><m:r><m:t>{base}</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r>"
            "</m:sup></m:sSup>")


#: A paragraph holding an equation as PowerPoint writes it inside the box.
EQ_PARAGRAPH = (
    '<a:p><a:pPr/><mc:AlternateContent><mc:Choice Requires="a14"><a14:m>'
    f'<m:oMathPara xmlns:m="{OMML}"><m:oMath><m:box><m:e>'
    f'{_sup(chr(0x1D44E))}<m:r><m:t>+</m:t></m:r>{_sup(chr(0x1D44F))}<m:r><m:t>=</m:t></m:r>'
    f'{_sup(chr(0x1D450))}</m:e></m:box></m:oMath></m:oMathPara></a14:m></mc:Choice>'
    '<mc:Fallback xmlns=""><a:r><a:rPr i="1"><a:latin typeface="Cambria Math"/></a:rPr>'
    "<a:t>a2+b2=c2</a:t></a:r></mc:Fallback></mc:AlternateContent><a:endParaRPr/></a:p>"
)


def _box(sid: int, name: str, x: int, paragraphs: str, *, ns: str = "") -> str:
    return (
        f'<p:sp{ns}><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr txBox="1"/><p:nvPr/>'
        f'</p:nvSpPr><p:spPr><a:xfrm><a:off x="{x}" y="1828800"/><a:ext cx="1828800" '
        'cy="914400"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/>'
        '</p:spPr><p:txBody><a:bodyPr wrap="none"><a:spAutoFit/></a:bodyPr><a:lstStyle/>'
        f"{paragraphs}</p:txBody></p:sp>"
    )


def _picture_fallback(sid: int, name: str, x: int) -> str:
    """The Fallback PowerPoint writes for an equation box: the same shape
    filled with a picture of it, holding one blank run."""
    return (
        f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr txBox="1"><a:spLocks '
        'noRot="1" noChangeAspect="1" noMove="1" noResize="1" noTextEdit="1"/></p:cNvSpPr>'
        f'<p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="{x}" y="1828800"/><a:ext '
        'cx="1828800" cy="914400"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        '<a:blipFill><a:blip r:embed="rIdPIC"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>'
        '</p:spPr><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US"><a:noFill/>'
        "</a:rPr><a:t> </a:t></a:r></a:p></p:txBody></p:sp>"
    )


def _wrapper(choice: str, fallback: str, *, requires: str = "a14",
             decl: str = f'xmlns:a14="{A14}"') -> str:
    return (
        f'<mc:AlternateContent xmlns:mc="{MC}" {NS}><mc:Choice {decl} Requires="{requires}">'
        f"{choice}</mc:Choice><mc:Fallback>{fallback}</mc:Fallback></mc:AlternateContent>"
    )


def _plain(sid: int, name: str, x: int, words: str) -> str:
    return _box(sid, name, x, f"<a:p><a:r><a:t>{words}</a:t></a:r></a:p>", ns=f" {NS}")


def _deck(tmp_path, members: list[str], name: str = "wrapped.pptx"):
    """A one-slide deck whose shape tree holds `members` in order, with the
    picture part every Fallback names."""
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / name
    prs.save(str(path))
    pkg = PptxPackage(path)
    part = pkg.slide_parts()[0]
    pkg.add_part_with_content_type("ppt/media/image1.png", b"\x89PNG made up", "image/png")
    rid = pkg.add_relationship(part, RT_IMAGE, "../media/image1.png")
    tree = pkg.root(part).find(f"{qn('p:cSld')}/{qn('p:spTree')}")
    for xml in members:
        tree.append(etree.fromstring(xml.replace("rIdPIC", rid)))
    pkg.mark_dirty(part)
    pkg.save(do_backup=False)
    return path


def _saved_equation_box(sid: int = 3, name: str = "Box Two", x: int = 3200400) -> str:
    """A text box that holds the words "Two" and an equation, kept whole in
    a slide-level wrapper as PowerPoint saves it."""
    choice = _box(sid, name, x, "<a:p><a:r><a:t>Two</a:t></a:r></a:p>" + EQ_PARAGRAPH)
    return _wrapper(choice, _picture_fallback(sid, name, x))


@pytest.fixture()
def saved(tmp_path):
    """Three text boxes, the second one saved by PowerPoint with an equation."""
    return _deck(tmp_path, [
        _plain(2, "Box One", 914400, "One"),
        _saved_equation_box(),
        _plain(4, "Box Three", 5486400, "Three"),
    ])


def _sha(path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ------------------------------------------------------------------ reads


def test_get_text_reads_the_words_and_the_equation_of_the_saved_box(saved):
    out = read.get_text(PptxPackage(saved))
    assert out["slides"][0]["text"] == f"One\nTwo\n{MATH}\nThree"


def test_find_text_finds_a_word_in_the_saved_box(saved):
    out = read.find_text(PptxPackage(saved), "Two")
    assert out["count"] == 1
    match = out["matches"][0]
    assert (match["shape_id"], match["paragraph"], match["where"]) == (3, 0, "slide")


def test_find_text_finds_the_characters_of_the_equation(saved):
    out = read.find_text(PptxPackage(saved), "+\U0001d44f")
    assert out["count"] == 1
    assert (out["matches"][0]["shape_id"], out["matches"][0]["paragraph"]) == (3, 1)


def test_list_equations_lists_the_equation_in_the_saved_box(saved):
    out = equations.list_equations(PptxPackage(saved))
    assert out["equation_count"] == 1
    entry = out["equations"][0]
    assert (entry["shape_id"], entry["paragraph"], entry["text"]) == (3, 1, MATH)


def test_get_slide_info_lists_the_saved_box_at_its_place(saved):
    out = read.get_slide_info(PptxPackage(saved), 0)
    assert out["shape_count"] == 3
    got = [(s["id"], s["name"], s["type"], s["z"], s["has_text"]) for s in out["shapes"]]
    assert got == [
        (2, "Box One", "textbox", 0, True),
        (3, "Box Two", "textbox", 1, True),
        (4, "Box Three", "textbox", 2, True),
    ]
    assert out["shapes"][1]["geometry"]["x"] == 3200400


def test_list_elements_lists_the_saved_box(saved):
    out = read.list_elements(PptxPackage(saved), "shapes")
    assert out["count"] == 3
    assert [(i["id"], i["z"]) for i in out["items"]] == [(2, 0), (3, 1), (4, 2)]


def test_the_view_shows_the_saved_box_with_its_anchor(saved):
    pkg = PptxPackage(saved)
    out = view.get_presentation_view(pkg)
    slide_id = read.slide_table(pkg)[0]["slide_id"]
    anchor = view.anchor_digest(slide_id, 3)[:6]
    assert f"[a:{anchor}] textbox:" in out["view"]
    assert "- Two" in out["view"] and f"- {MATH}" in out["view"]
    assert view.resolve_anchor(pkg, f"a:{anchor}")["shape_id"] == 3


def test_the_slide_statistics_count_the_saved_box(saved):
    from kitchensink4ppt.ops import assembly

    out = assembly.deck_statistics(PptxPackage(saved))
    assert out["shapes_by_type"].get("textbox") == 3
    assert out["words"]["body"] == 3 + 1  # One, Two, Three and the equation's one word


def test_a_box_whose_math_sits_directly_in_the_paragraph_reads_too(tmp_path):
    """PowerPoint can also write the math straight into the paragraph of
    the box the a14 branch holds."""
    math = (f'<a:p><a14:m><m:oMathPara xmlns:m="{OMML}"><m:oMath><m:r><m:t>y=3</m:t></m:r>'
            "</m:oMath></m:oMathPara></a14:m></a:p>")
    choice = _box(3, "Made-up equation", 914400,
                  "<a:p><a:r><a:t>Slope</a:t></a:r></a:p>" + math)
    path = _deck(tmp_path, [_wrapper(choice, _picture_fallback(3, "Made-up equation", 914400))])
    pkg = PptxPackage(path)
    assert read.get_text(pkg)["slides"][0]["text"] == "Slope\ny=3"
    assert read.find_text(pkg, "y=3")["count"] == 1
    assert equations.list_equations(pkg)["equation_count"] == 1


def test_the_reads_show_the_fallback_when_powerpoint_would(tmp_path):
    """A Choice that Requires a namespace PowerPoint does not read is not
    shown; the Fallback is, and the reads follow it."""
    choice = _plain(3, "Made-up future box", 914400, "words for a newer version")
    fallback = _plain(3, "Made-up future box", 914400, "words everyone sees")
    path = _deck(tmp_path, [_wrapper(choice, fallback, requires="fut",
                                     decl='xmlns:fut="urn:made-up:future"')])
    pkg = PptxPackage(path)
    assert read.get_text(pkg)["slides"][0]["text"] == "words everyone sees"
    assert read.find_text(pkg, "newer")["count"] == 0
    assert [s["id"] for s in read.get_slide_info(pkg, 0)["shapes"]] == [3]


def _group(gid: int, members: list[str]) -> str:
    return (
        f'<p:grpSp {NS}><p:nvGrpSpPr><p:cNvPr id="{gid}" name="Made-up group"/><p:cNvGrpSpPr/>'
        '<p:nvPr/></p:nvGrpSpPr><p:grpSpPr><a:xfrm><a:off x="914400" y="914400"/><a:ext '
        'cx="6400800" cy="2743200"/><a:chOff x="914400" y="914400"/><a:chExt cx="6400800" '
        f'cy="2743200"/></a:xfrm></p:grpSpPr>{"".join(members)}</p:grpSp>'
    )


def test_a_wrapped_member_of_a_group_is_listed_in_its_group(tmp_path):
    path = _deck(tmp_path, [_group(10, [_plain(11, "Plain member", 914400, "plain words"),
                                        _saved_equation_box(12, "Equation member", 3200400)])])
    pkg = PptxPackage(path)
    shapes = read.get_slide_info(pkg, 0)["shapes"]
    assert [(s["id"], s.get("group_id"), s["z"]) for s in shapes] == [
        (10, None, 0), (11, 10, 0), (12, 10, 1)]
    assert read.get_text(pkg)["slides"][0]["text"] == f"plain words\nTwo\n{MATH}"


def test_an_equation_member_lifted_out_by_ungroup_is_listed_and_read(tmp_path):
    path = _deck(tmp_path, [_plain(2, "Before", 0, "Before"),
                            _group(10, [_plain(11, "Plain member", 914400, "plain words"),
                                        _saved_equation_box(12, "Equation member", 3200400)]),
                            _plain(20, "After", 0, "After")])
    pkg = PptxPackage(path)
    freed = sh.ungroup_shapes(pkg, 0, 10)["freed_ids"]
    pkg.save(do_backup=False)
    assert sorted(freed) == [11, 12]
    pkg = PptxPackage(path)
    info = read.get_slide_info(pkg, 0)
    assert [(s["id"], s["z"]) for s in info["shapes"]] == [(2, 0), (11, 1), (12, 2), (20, 3)]
    listed = read.list_elements(pkg, "shapes")["items"]
    assert [i["id"] for i in listed] == [2, 11, 12, 20]
    assert read.get_text(pkg)["slides"][0]["text"] == f"Before\nplain words\nTwo\n{MATH}\nAfter"


def test_ink_in_a_wrapper_is_listed_as_ink_with_its_box(tmp_path):
    ink = (
        f'<mc:AlternateContent xmlns:mc="{MC}" xmlns:p14="{P14}" {NS}><mc:Choice Requires="p14">'
        '<p:contentPart p14:bwMode="auto" r:id="rIdPIC"><p14:nvContentPartPr><p14:cNvPr id="5" '
        'name="Made-up ink"/><p14:cNvContentPartPr/><p14:nvPr/></p14:nvContentPartPr><p14:xfrm>'
        '<a:off x="457200" y="457200"/><a:ext cx="914400" cy="457200"/></p14:xfrm></p:contentPart>'
        '</mc:Choice><mc:Fallback><p:pic><p:nvPicPr><p:cNvPr id="5" name="Made-up ink"/>'
        '<p:cNvPicPr/><p:nvPr/></p:nvPicPr><p:blipFill><a:blip r:embed="rIdPIC"/><a:stretch>'
        '<a:fillRect/></a:stretch></p:blipFill><p:spPr><a:xfrm><a:off x="457200" y="457200"/>'
        '<a:ext cx="914400" cy="457200"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/>'
        "</a:prstGeom></p:spPr></p:pic></mc:Fallback></mc:AlternateContent>"
    )
    path = _deck(tmp_path, [_plain(2, "Words", 0, "Words"), ink])
    pkg = PptxPackage(path)
    shapes = read.get_slide_info(pkg, 0)["shapes"]
    assert [(s["id"], s["name"], s["type"], s["z"]) for s in shapes] == [
        (2, "Words", "textbox", 0), (5, "Made-up ink", "ink", 1)]
    assert shapes[1]["geometry"]["x"] == 457200 and shapes[1]["has_text"] is False
    assert read.get_text(pkg)["slides"][0]["text"] == "Words"


# ------------------------------------------------------------------ edits


def test_an_edit_that_names_the_saved_box_is_refused_by_name(saved):
    before = _sha(saved)
    for shape in (3, "Box Two"):
        pkg = PptxPackage(saved)
        rec = read.resolve_slide(pkg, 0)
        with pytest.raises(UnsupportedStructure) as info:
            tx._resolve_shape(pkg, rec, shape)
        message = str(info.value)
        assert "compatibility wrapper" in message and "'Box Two'" in message
        assert "nothing was changed" in message
    assert _sha(saved) == before


def test_an_edit_through_the_view_anchor_of_the_saved_box_is_refused(saved):
    from kitchensink4ppt.core.errors import PptMcpError
    from kitchensink4ppt.ops import batch

    before = _sha(saved)
    pkg = PptxPackage(saved)
    slide_id = read.slide_table(pkg)[0]["slide_id"]
    anchor = view.anchor_digest(slide_id, 3)[:6]
    with pytest.raises(PptMcpError) as info:
        batch.apply_edits(pkg, [{"anchor": f"a:{anchor}", "op": "set_text", "text": "changed"}])
    assert "compatibility wrapper" in str(info.value) and "'Box Two'" in str(info.value)
    assert _sha(saved) == before


def test_the_run_map_reads_math_in_the_paragraph_as_the_reads_do():
    """Edits resolve find_text's offsets through the run map, whose text
    must equal paragraph_text: math written straight into the paragraph is
    one slot there too, which edits go around."""
    from kitchensink4ppt.ops._runmap import build_map

    p = etree.fromstring(
        f'<a:p xmlns:a="{A}" xmlns:a14="{A14}"><a:r><a:t>Slope </a:t></a:r><a14:m>'
        f'<m:oMathPara xmlns:m="{OMML}"><m:oMath><m:r><m:t>y=3</m:t></m:r></m:oMath>'
        "</m:oMathPara></a14:m><a:r><a:t> here</a:t></a:r></a:p>"
    )
    text, segments = build_map(p)
    assert text == read.paragraph_text(p) == "Slope y=3 here"
    assert [s.kind for s in segments] == ["run", "wrapped", "run"]
