"""One deck, two spellings of its booleans, the same answers.

An OOXML boolean is "1"/"true" or "0"/"false". PowerPoint writes the digits;
other producers write the words. Readers that compared only against "1" or
"0" misread the words: a slide hidden with show="false" read as shown, a
shape hidden with hidden="true" read as visible, a text box marked
txBox="true" read as an autoshape, a table's header row flag read as off.

The oracle is mechanical: build a made-up deck whose booleans are digits,
make its twin by respelling every boolean attribute as a word, and require
every read to answer the same for both. A static guard keeps new code from
comparing a boolean attribute against a literal again.
"""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import pytest

from kitchensink4ppt import server
from kitchensink4ppt.core.xmlbool import attr_bool, xml_bool

SRC = Path(__file__).resolve().parents[2] / "src" / "kitchensink4ppt"

#: Attributes the twin respells. Each is an xsd:boolean wherever these parts
#: carry it (slides, layouts, masters and notes; never chart parts).
BOOL_ATTRS = (
    "show", "hidden", "txBox", "b", "i", "firstRow", "lastRow", "firstCol",
    "lastCol", "bandRow", "bandCol", "hMerge", "vMerge", "flipH", "flipV",
    "advClick", "showMasterSp", "showMasterPhAnim", "noGrp", "noRot",
    "noChangeAspect", "noSelect", "noMove", "noResize", "noTextEdit",
    "noChangeArrowheads", "noChangeShapeType", "hasCustomPrompt", "userDrawn",
    "preserve", "dirty", "err", "noProof", "smtClean", "rtlCol", "anchorCtr",
    "eaLnBrk", "latinLnBrk", "hangingPunct", "rtl", "val",
)

#: Only these parts are respelled; "val" is a boolean only on the
#: decorative marker in them.
_PARTS = re.compile(r"ppt/(slides|slideLayouts|slideMasters|notesSlides)/[^/]+\.xml$")
_ATTR = re.compile(r'(\s(?:[A-Za-z0-9]+:)?(' + "|".join(BOOL_ATTRS) + r'))="([01])"')
_DECORATIVE_VAL = re.compile(r'(<adec:decorative [^>]*?val=)"([01])"')


def _respell(xml: str) -> str:
    def word(m):
        name = m.group(2)
        if name == "val":
            return m.group(0)
        return f'{m.group(1)}="{"true" if m.group(3) == "1" else "false"}"'

    xml = _ATTR.sub(word, xml)
    return _DECORATIVE_VAL.sub(
        lambda m: f'{m.group(1)}"{"true" if m.group(2) == "1" else "false"}"', xml)


def _build(path: Path) -> Path:
    """Digits deck: a text box with a bold italic run, a hidden rectangle, a
    decorative oval, a table with header-row and banding flags and a merged
    pair, a flipped connector, a hidden slide, a slide that does not
    advance on click, notes with a bold run, and a layout that hides the
    master's shapes."""
    from lxml import etree
    from pptx import Presentation
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Presentation()
    blank = prs.slide_layouts[6]
    s0 = prs.slides.add_slide(blank)
    tb = s0.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(4), Inches(1))
    tb.name = "Label"
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "Quarterly plan"
    run.font.bold = True
    run.font.italic = True
    run.font.size = Pt(20)
    rect = s0.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5), Inches(0.4),
                               Inches(2), Inches(1))
    rect.name = "Hidden box"
    rect.text_frame.text = "Behind the scenes"
    rect._element.nvSpPr.cNvPr.set("hidden", "1")
    oval = s0.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.5), Inches(0.4),
                               Inches(1), Inches(1))
    oval.name = "Ornament"
    ext_lst = etree.SubElement(oval._element.nvSpPr.cNvPr,
                               "{http://schemas.openxmlformats.org/drawingml/2006/main}extLst")
    ext = etree.SubElement(ext_lst, "{http://schemas.openxmlformats.org/drawingml/2006/main}ext",
                           uri="{C183D7F6-B498-43B3-948B-1728B52AA6E4}")
    etree.SubElement(ext, "{http://schemas.microsoft.com/office/drawing/2017/decorative}decorative",
                     nsmap={"adec": "http://schemas.microsoft.com/office/drawing/2017/decorative"},
                     val="1")
    table = s0.shapes.add_table(3, 3, Inches(0.5), Inches(2), Inches(6), Inches(2)).table
    for r in range(3):
        for c in range(3):
            table.cell(r, c).text = f"cell {r}{c}"
    table.cell(1, 0).merge(table.cell(1, 1))
    con = s0.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(9), Inches(5),
                                  Inches(7), Inches(6))
    con.name = "Back arrow"
    s0.notes_slide.notes_text_frame.text = ""
    nrun = s0.notes_slide.notes_text_frame.paragraphs[0].add_run()
    nrun.text = "Speaker reminder"
    nrun.font.bold = True

    s1 = prs.slides.add_slide(prs.slide_layouts[5])
    s1.shapes.title.text = "Backup material"
    s1._element.set("show", "0")

    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Timed slide"
    s2.placeholders[1].text = "Moves on by itself"
    trans = etree.SubElement(s2._element,
                             "{http://schemas.openxmlformats.org/presentationml/2006/main}transition")
    trans.set("advClick", "0")
    trans.set("advTm", "3000")

    prs.slide_layouts[5]._element.set("showMasterSp", "0")
    prs.save(str(path))
    return path


def _twin(src: Path, dst: Path) -> Path:
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if _PARTS.match(info.filename):
                data = _respell(data.decode("utf-8")).encode("utf-8")
            zout.writestr(info, data)
    return dst


@pytest.fixture()
def twins(tmp_path):
    (tmp_path / "digits").mkdir()
    (tmp_path / "words").mkdir()
    a = _build(tmp_path / "digits" / "plan.pptx")
    b = _twin(a, tmp_path / "words" / "plan.pptx")
    return a, b


def _norm(answer, folder: Path) -> str:
    text = json.dumps(answer, sort_keys=True, default=str)
    for spelled in (str(folder), str(folder).replace("\\", "/"),
                    json.dumps(str(folder))[1:-1]):
        text = text.replace(spelled, "<dir>")
    return text


READS = {
    "presentation_info": lambda p: server.get_presentation_info(p),
    "slide_info_0": lambda p: server.get_slide_info(p, 0, live="off"),
    "slide_info_1": lambda p: server.get_slide_info(p, 1, live="off"),
    "slide_info_2": lambda p: server.get_slide_info(p, 2, live="off"),
    "shapes": lambda p: server.list_elements(p, "shapes"),
    "slides": lambda p: server.list_elements(p, "slides"),
    "accessibility": lambda p: server.audit_accessibility(p),
    "layout_checks": lambda p: server.check_layout(p),
    "table": lambda p: server.get_table(p, 0),
    "transitions": lambda p: server.get_transitions(p),
    "statistics": lambda p: server.deck_statistics(p),
    "master_elements": lambda p: server.list_master_elements(p),
    "notes": lambda p: server.get_notes(p, 0, rich=True),
    "view": lambda p: server.get_presentation_view(p),
}


def test_the_twin_differs_only_in_spelling(twins):
    a, b = twins
    with zipfile.ZipFile(a) as za, zipfile.ZipFile(b) as zb:
        respelled = [n for n in za.namelist() if za.read(n) != zb.read(n)]
        words = sum(zb.read(n).decode("utf-8").count('="true"')
                    + zb.read(n).decode("utf-8").count('="false"') for n in respelled)
    assert any("slides/slide2.xml" in n for n in respelled)
    assert words >= 10


@pytest.mark.parametrize("read", sorted(READS))
def test_every_read_answers_the_same_for_both_spellings(twins, read):
    a, b = twins
    got_a = _norm(READS[read](str(a)), a.parent)
    got_b = _norm(READS[read](str(b)), b.parent)
    assert got_a == got_b


def test_hidden_slide_hidden_shape_and_text_box_in_words(twins):
    _a, b = twins
    assert server.get_slide_info(str(b), 1, live="off")["hidden"] is True
    assert server.get_slide_info(str(b), 0, live="off")["hidden"] is False
    shapes = {s["name"]: s for s in server.get_slide_info(str(b), 0, live="off")["shapes"]}
    assert shapes["Hidden box"]["hidden"] is True
    assert shapes["Label"]["type"] == "textbox"


@pytest.mark.parametrize("value, default, expected", [
    ("1", False, True), ("true", False, True), (" true ", False, True),
    ("0", True, False), ("false", True, False), ("\tfalse\n", True, False),
    (None, False, False), (None, True, True),
    ("TRUE", False, False), ("yes", True, True), ("", True, True), ("2", False, False),
])
def test_xml_bool_spellings(value, default, expected):
    assert xml_bool(value, default) is expected


def test_attr_bool_on_a_missing_element():
    assert attr_bool(None, "hidden") is False
    assert attr_bool(None, "show", True) is True


#: A boolean attribute compared against a literal, or tested for presence as
#: if presence meant true: what this module replaced. flipH/flipV/rot are
#: read through the helper or a number, never by truthiness.
_LITERAL_COMPARE = re.compile(
    r'\.get\(\s*"(?:' + "|".join(a for a in BOOL_ATTRS if a != "val") + r'|mute|fullScrn|'
    r'showWhenStopped|removed|enabled)"[^)]*\)\s*(?:==|!=|in\b|not\s+in\b)')
_TRUTHY_FLIP = re.compile(r'\.get\("(?:flipH|flipV)"\)\s*(?:or|and|\)\s*:|:)')


def test_no_boolean_attribute_is_compared_against_a_literal():
    offenders = []
    for path in sorted(SRC.rglob("*.py")):
        for n, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
            if _LITERAL_COMPARE.search(line) or _TRUTHY_FLIP.search(line):
                offenders.append(f"{path.relative_to(SRC)}:{n}: {line.strip()}")
    assert offenders == []
