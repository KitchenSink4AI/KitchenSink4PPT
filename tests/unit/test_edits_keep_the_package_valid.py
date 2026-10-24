"""Decks PowerPoint opens clean that an edit used to turn into decks it repairs.

Each case below is one where the save-time check could pass and PowerPoint
would then repair the saved file (or show none of the edit). Each is a made-up
deck built by code, edited, saved and judged the way PowerPoint judges a
package:

* a slide relationship spelled with a percent code, a backslash or a
        #fragment: reads answered INTERNAL_ERROR and set_notes wrote a stray
        rels part for a slide that does not exist.
* anonymize_deck gave two comment authors one name and no user id.
* a deck that declares its .rels parts by Override: a new slide's rels
        part had no content type.
* delete_shape on an audio shape left a media node with an empty target.
* a shape id past 4294967295.
* ungroup_shapes left animation targets on the removed group.
* insert_textbox / insert_shape after a trailing p:extLst in the tree.
"""

from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation

from kitchensink4ppt.core.errors import PptMcpError
from kitchensink4ppt.core.package import PptxPackage, qn
from kitchensink4ppt.ops import assembly, notes, read, shapes, slides, text

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"


def _base(tmp_path, name="base.pptx", *, body=True):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Relationship repro"
    if body:
        box = slide.shapes.add_textbox(914400, 1828800, 4572000, 914400)
        box.text_frame.text = "Existing body words"
    path = tmp_path / name
    prs.save(path)
    return path


def _rewrite(path, change):
    """Rewrite members of a package: change(name, data) -> new bytes (or None to drop)."""
    src = zipfile.ZipFile(path)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for item in src.infolist():
            data = change(item.filename, src.read(item.filename))
            if data is not None:
                zf.writestr(item.filename, data)
    src.close()
    Path(path).write_bytes(out.getvalue())


def _members(path):
    with zipfile.ZipFile(path) as zf:
        return {n: zf.read(n) for n in zf.namelist()}


def _assert_package_clean(path):
    """The structural judgement PowerPoint makes: every part has a content
    type, every relationship target exists, and the product's own validate
    agrees."""
    members = _members(path)
    pkg = PptxPackage(path)
    for name in members:
        if name == "[Content_Types].xml":
            continue
        assert pkg.content_type(name) is not None, f"no content type for {name}"
    # the product's own whole-package check, nothing tolerated
    PptxPackage._validate_payload(Path(path).read_bytes(), doc=Path(path).name)


# ---------------------------------------------------------------- slide relationship spellings

SPELLINGS = {
    "percent": "slides/slide%31.xml",
    "backslash": "slides\\slide1.xml",
    "fragment": "slides/slide1.xml#frag",
}


def _odd_target_deck(tmp_path, kind):
    path = _base(tmp_path, f"odd_{kind}.pptx")

    def change(name, data):
        if name == "ppt/_rels/presentation.xml.rels":
            s = data.decode("utf-8")
            assert 'Target="slides/slide1.xml"' in s
            return s.replace('Target="slides/slide1.xml"',
                             f'Target="{SPELLINGS[kind]}"').encode("utf-8")
        return data

    _rewrite(path, change)
    return path


@pytest.mark.parametrize("kind", sorted(SPELLINGS))
def test_cr1_reads_find_the_slide_whatever_the_target_spelling(tmp_path, kind):
    path = _odd_target_deck(tmp_path, kind)
    pkg = PptxPackage(path)
    table = read.slide_table(pkg)
    assert len(table) == 1
    assert table[0]["part"] == "ppt/slides/slide1.xml"
    info = read.get_slide_info(pkg, 0)
    assert info["shape_count"] >= 2


@pytest.mark.parametrize("kind", sorted(SPELLINGS))
def test_cr1_an_edit_saves_a_clean_deck(tmp_path, kind):
    path = _odd_target_deck(tmp_path, kind)
    pkg = PptxPackage(path)
    text.insert_textbox(pkg, 0, "Added words", 1, 4, 3, 1)
    pkg.save(do_backup=False)
    members = _members(path)
    assert not any("%" in n or "\\" in n for n in members)
    pkg = PptxPackage(path)
    assert any("Added words" in (el.text or "") for el in pkg.root("ppt/slides/slide1.xml").iter(qn("a:t")))
    _assert_package_clean(path)


@pytest.mark.parametrize("kind", sorted(SPELLINGS))
def test_cr1_set_notes_writes_notes_for_the_real_slide(tmp_path, kind):
    path = _odd_target_deck(tmp_path, kind)
    pkg = PptxPackage(path)
    notes.set_notes(pkg, 0, "Notes for the repro")
    pkg.save(do_backup=False)
    members = _members(path)
    # no rels part for a slide that does not exist
    assert not any("slide%31" in n or "\\" in n for n in members)
    assert "ppt/slides/_rels/slide1.xml.rels" in members
    pkg = PptxPackage(path)
    assert notes.get_notes(pkg, 0)["text"] == "Notes for the repro"
    _assert_package_clean(path)


def test_cr1_a_target_that_names_no_part_is_refused_by_name(tmp_path):
    """A query-string spelling (PowerPoint repairs this deck) names no part:
    the read says so instead of failing inside the server."""
    path = _base(tmp_path, "query.pptx")

    def change(name, data):
        if name == "ppt/_rels/presentation.xml.rels":
            return data.replace(b'Target="slides/slide1.xml"', b'Target="slides/slide1.xml?x=1"')
        return data

    _rewrite(path, change)
    pkg = PptxPackage(path)
    with pytest.raises(PptMcpError) as refused:
        read.slide_table(pkg)
    assert "slides/slide1.xml?x=1" in str(refused.value)
    assert "nothing was changed" in str(refused.value)


# ---------------------------------------------------------------- shape id range

def test_cr6_a_new_shape_never_gets_an_id_past_the_unsigned_32_bit_range(tmp_path):
    path = _base(tmp_path, "ceiling.pptx")

    def change(name, data):
        if name == "ppt/slides/slide1.xml":
            s = data.decode("utf-8")
            return re.sub(r'(<p:cNvPr id=")\d+(" name="TextBox)', r"\g<1>4294967295\2", s).encode("utf-8")
        return data

    _rewrite(path, change)
    pkg = PptxPackage(path)
    out = text.insert_textbox(pkg, 0, "Added words", 1, 4, 3, 1)
    pkg.save(do_backup=False)
    new_id = out.get("shape_id")
    assert new_id is not None and 1 <= int(new_id) <= 4294967295
    ids = [int(el.get("id")) for el in PptxPackage(path).root("ppt/slides/slide1.xml").iter(qn("p:cNvPr"))]
    assert max(ids) <= 4294967295
    assert len(ids) == len(set(ids))


def test_cr6_the_save_check_refuses_an_id_outside_the_range(tmp_path):
    path = _base(tmp_path, "ceiling2.pptx")
    pkg = PptxPackage(path)
    root = pkg.root("ppt/slides/slide1.xml")
    list(root.iter(qn("p:cNvPr")))[-1].set("id", "4294967296")
    pkg.mark_dirty("ppt/slides/slide1.xml")
    with pytest.raises(PptMcpError) as refused:
        pkg.save(do_backup=False)
    assert "4294967296" in str(refused.value)


# ---------------------------------------------------------------- trailing extLst

def _extlst_deck(tmp_path):
    path = _base(tmp_path, "extlst.pptx")

    def change(name, data):
        if name == "ppt/slides/slide1.xml":
            s = data.decode("utf-8")
            ext = ('<p:extLst xmlns:p="%s"><p:ext uri="{00000000-0000-4000-8000-0000000000AA}"/>'
                   "</p:extLst>" % P)
            assert "</p:spTree>" in s
            return s.replace("</p:spTree>", ext + "</p:spTree>").encode("utf-8")
        return data

    _rewrite(path, change)
    return path


def _tree_tags(path):
    root = PptxPackage(path).root("ppt/slides/slide1.xml")
    sp_tree = root.find(".//" + qn("p:spTree"))
    return [etree.QName(c).localname for c in sp_tree]


def test_cr8_insert_textbox_lands_before_a_trailing_extlst(tmp_path):
    path = _extlst_deck(tmp_path)
    pkg = PptxPackage(path)
    text.insert_textbox(pkg, 0, "Added words", 1, 4, 3, 1)
    pkg.save(do_backup=False)
    tags = _tree_tags(path)
    assert tags[-1] == "extLst", tags
    assert tags.count("extLst") == 1


def test_cr8_insert_shape_lands_before_a_trailing_extlst(tmp_path):
    path = _extlst_deck(tmp_path)
    pkg = PptxPackage(path)
    shapes.insert_shape(pkg, 0, "rect", 1, 4, 2, 1)
    pkg.save(do_backup=False)
    tags = _tree_tags(path)
    assert tags[-1] == "extLst", tags


# ---------------------------------------------------------------- comment authors

P15 = "http://schemas.microsoft.com/office/powerpoint/2012/main"
RT_AUTHORS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/commentAuthors")


def _authors_xml(people):
    """people: [(name, initials, user_id or None)]"""
    rows = []
    for i, (name, initials, uid) in enumerate(people, start=1):
        ext = ""
        if uid is not None:
            ext = ('<p:extLst><p:ext uri="{19B8F6BF-5375-455C-9EA6-DF929625EA0E}">'
                   f'<p15:presenceInfo userId="{uid}" providerId="AD"/></p:ext></p:extLst>')
        rows.append(f'<p:cmAuthor id="{i}" name="{name}" initials="{initials}" '
                    f'lastIdx="1" clrIdx="{i - 1}">{ext}</p:cmAuthor>')
    return ('<p:cmAuthorLst xmlns:p="%s" xmlns:p15="%s">%s</p:cmAuthorLst>'
            % (P, P15, "".join(rows))).encode()


def _author_deck(tmp_path, people, name="authors.pptx"):
    path = tmp_path / name
    slides.create_presentation(path)
    pkg = PptxPackage(path)
    slides.insert_slide(pkg, 6)
    pkg.add_part_with_content_type(
        "ppt/commentAuthors.xml", _authors_xml(people),
        "application/vnd.openxmlformats-officedocument.presentationml.commentAuthors+xml")
    pkg.add_relationship("ppt/presentation.xml", RT_AUTHORS, "commentAuthors.xml")
    pkg.save(do_backup=False)
    return path


def _saved_authors(path):
    root = PptxPackage(path).root("ppt/commentAuthors.xml")
    out = []
    for a in root.findall(qn("p:cmAuthor")):
        presence = next(a.iter("{%s}presenceInfo" % P15), None)
        out.append((a.get("name"), a.get("initials"),
                    None if presence is None else presence.get("userId")))
    return out


def test_cr3_two_authors_with_one_name_stay_two_authors(tmp_path):
    path = _author_deck(tmp_path, [
        ("Rowan Example", "RE", "S::rowan-one-made-up::0001"),
        ("Rowan Example", "RX", "S::rowan-two-made-up::0002"),
    ])
    pkg = PptxPackage(path)
    out = assembly.anonymize_deck(pkg)
    pkg.save(do_backup=False)
    saved = _saved_authors(path)
    assert len(saved) == 2
    # nothing identifying is left, and the two are still told apart
    assert all(uid == "" for _n, _i, uid in saved)
    assert saved[0][0] != saved[1][0], saved
    assert saved[0][1] != saved[1][1], saved
    assert len({(n, u) for n, _i, u in saved}) == 2
    assert "Rowan Example" not in str(saved)
    assert len(set(out["author_aliases"].values())) == 2
    assert "made-up" not in str(out)


def test_cr3_authors_with_different_names_are_unchanged_in_kind(tmp_path):
    path = _author_deck(tmp_path, [
        ("Rowan Example", "RE", "S::rowan-one-made-up::0001"),
        ("Sam Example", "SE", "S::sam-made-up::0003"),
    ])
    pkg = PptxPackage(path)
    assembly.anonymize_deck(pkg)
    pkg.save(do_backup=False)
    saved = _saved_authors(path)
    assert len({n for n, _i, _u in saved}) == 2


def test_cr3_one_name_without_user_ids_keeps_one_alias(tmp_path):
    path = _author_deck(tmp_path, [
        ("Rowan Example", "RE", None),
        ("Rowan Example", "RX", None),
    ])
    pkg = PptxPackage(path)
    out = assembly.anonymize_deck(pkg)
    pkg.save(do_backup=False)
    assert len(set(out["author_aliases"].values())) == 1


def test_cr3_the_save_check_refuses_two_authors_alike(tmp_path):
    path = _author_deck(tmp_path, [
        ("Rowan Example", "RE", "S::rowan-one-made-up::0001"),
        ("Rowan Example", "RX", "S::rowan-two-made-up::0002"),
    ])
    pkg = PptxPackage(path)
    root = pkg.root("ppt/commentAuthors.xml")
    for presence in root.iter("{%s}presenceInfo" % P15):
        presence.set("userId", "")
    pkg.mark_dirty("ppt/commentAuthors.xml")
    with pytest.raises(PptMcpError) as refused:
        pkg.save(do_backup=False)
    assert "comment authors" in str(refused.value)


# ---------------------------------------------------------------- .rels by Override

def _override_declared_deck(tmp_path):
    path = _base(tmp_path, "override_rels.pptx")
    with zipfile.ZipFile(path) as zf:
        rels_parts = [n for n in zf.namelist() if n.endswith(".rels")]

    def change(name, data):
        if name != "[Content_Types].xml":
            return data
        root = etree.fromstring(data)
        for d in list(root):
            if etree.QName(d).localname == "Default" and d.get("Extension") == "rels":
                root.remove(d)
        for n in rels_parts:
            o = etree.SubElement(root, "{%s}Override" % CT)
            o.set("PartName", "/" + n)
            o.set("ContentType", "application/vnd.openxmlformats-package.relationships+xml")
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    _rewrite(path, change)
    return path


def test_cr4_the_fixture_declares_every_rels_part_by_override(tmp_path):
    path = _override_declared_deck(tmp_path)
    ct = _members(path)["[Content_Types].xml"].decode()
    assert 'Extension="rels"' not in ct
    assert 'PartName="/ppt/slides/_rels/slide1.xml.rels"' in ct


@pytest.mark.parametrize("op", ["duplicate", "insert"])
def test_cr4_a_new_slide_rels_part_gets_a_content_type(tmp_path, op):
    path = _override_declared_deck(tmp_path)
    pkg = PptxPackage(path)
    if op == "duplicate":
        slides.duplicate_slide(pkg, 0)
    else:
        slides.insert_slide(pkg, 6)
    pkg.save(do_backup=False)
    members = _members(path)
    assert "ppt/slides/_rels/slide2.xml.rels" in members
    _assert_package_clean(path)


def test_cr4_the_save_check_refuses_a_part_with_no_content_type(tmp_path):
    path = _base(tmp_path, "typeless.pptx")
    pkg = PptxPackage(path)
    pkg.set_raw_part("ppt/extra/unknown.zzz", b"made-up bytes")
    with pytest.raises(PptMcpError) as refused:
        pkg.save(do_backup=False)
    assert "no content type" in str(refused.value)
    assert "unknown.zzz" in str(refused.value)


# ---------------------------------------------------------------- animation targets

def _timing_xml(spid, *, audio_spid=None):
    audio = ""
    if audio_spid is not None:
        audio = ('<p:audio><p:cMediaNode vol="80000" numSld="3" showWhenStopped="0">'
                 '<p:cTn id="11" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/>'
                 '</p:stCondLst><p:endCondLst><p:cond evt="onStopAudio" delay="0"><p:tgtEl><p:sldTgt/>'
                 '</p:tgtEl></p:cond></p:endCondLst></p:cTn>'
                 f'<p:tgtEl><p:spTgt spid="{audio_spid}"/></p:tgtEl></p:cMediaNode></p:audio>')
    effect = (
        '<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
        '<p:childTnLst><p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
        '<p:childTnLst><p:par><p:cTn id="5" presetID="1" presetClass="entr" presetSubtype="0" '
        'fill="hold" grpId="0" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
        '<p:childTnLst><p:set><p:cBhvr><p:cTn id="6" dur="1" fill="hold"><p:stCondLst>'
        '<p:cond delay="0"/></p:stCondLst></p:cTn>'
        f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility'
        '</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
        '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    return (
        f'<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" '
        'nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" '
        f'dur="indefinite" nodeType="mainSeq"><p:childTnLst>{effect}</p:childTnLst></p:cTn>'
        '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>'
        '</p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl>'
        f'</p:cond></p:nextCondLst></p:seq>{audio}</p:childTnLst></p:cTn></p:par></p:tnLst>'
        f'<p:bldLst><p:bldP spid="{spid}" grpId="0" animBg="1"/></p:bldLst></p:timing>')


def _give_timing(path, xml):
    def change(name, data):
        if name == "ppt/slides/slide1.xml":
            s = data.decode("utf-8")
            assert "<p:timing" not in s
            return s.replace("</p:sld>", xml + "</p:sld>").encode("utf-8")
        return data

    _rewrite(path, change)


def _empty_targets(path):
    root = PptxPackage(path).root("ppt/slides/slide1.xml")
    return [t for t in root.iter(qn("p:tgtEl")) if len(t) == 0]


def _targeted_ids(path):
    root = PptxPackage(path).root("ppt/slides/slide1.xml")
    return {t.get("spid") for t in root.iter(qn("p:spTgt"))}


def _shape_ids(path):
    root = PptxPackage(path).root("ppt/slides/slide1.xml")
    return {str(el.get("id")) for el in root.iter(qn("p:cNvPr"))}


def _slide_ids(path):
    root = PptxPackage(path).root("ppt/slides/slide1.xml")
    return [int(e.get("id")) for e in root.iter(qn("p:cNvPr"))]


def test_cr5_deleting_an_audio_shape_leaves_no_empty_media_target(tmp_path):
    path = _base(tmp_path, "audio.pptx")
    shape_id = _slide_ids(path)[-1]
    # a click effect on the shape, and the standalone media node PowerPoint
    # writes for a click-to-play clip, both naming it
    _give_timing(path, _timing_xml(shape_id, audio_spid=shape_id))
    pkg = PptxPackage(path)
    out = shapes.delete_shape(pkg, 0, shape_id)
    pkg.save(do_backup=False)
    assert out["timing_report"]["effects_removed"] == 1
    assert _empty_targets(path) == []
    assert _targeted_ids(path) <= _shape_ids(path)
    assert b"<p:audio" not in _members(path)["ppt/slides/slide1.xml"]
    _assert_package_clean(path)


def test_cr5_a_media_node_for_another_shape_stays(tmp_path):
    path = _base(tmp_path, "audio2.pptx")
    ids = _slide_ids(path)
    keep, drop = ids[1], ids[-1]
    _give_timing(path, _timing_xml(drop, audio_spid=keep))
    pkg = PptxPackage(path)
    shapes.delete_shape(pkg, 0, drop)
    pkg.save(do_backup=False)
    assert b"<p:audio" in _members(path)["ppt/slides/slide1.xml"]
    assert _empty_targets(path) == []


def _grouped_deck(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Animated group"
    a = slide.shapes.add_shape(1, 914400, 1828800, 1828800, 914400)
    b = slide.shapes.add_shape(1, 3657600, 1828800, 1828800, 914400)
    group = slide.shapes.add_group_shape([a, b])
    path = tmp_path / "group.pptx"
    prs.save(path)
    return path, group.shape_id


def test_cr7_ungrouping_an_animated_group_leaves_no_target_on_the_group(tmp_path):
    path, group_id = _grouped_deck(tmp_path)
    _give_timing(path, _timing_xml(group_id))
    pkg = PptxPackage(path)
    out = shapes.ungroup_shapes(pkg, 0, group_id)
    pkg.save(do_backup=False)
    assert out["timing_report"]["effects_removed"] == 1
    assert any("animation" in w for w in out["warnings"])
    assert str(group_id) not in _targeted_ids(path)
    assert _targeted_ids(path) <= _shape_ids(path)
    root = PptxPackage(path).root("ppt/slides/slide1.xml")
    assert not list(root.iter(qn("p:bldP")))
    assert len(list(root.iter(qn("p:sp")))) >= 3  # title plus both members
    _assert_package_clean(path)


def test_cr7_an_unanimated_group_ungroups_as_before(tmp_path):
    path, group_id = _grouped_deck(tmp_path)
    pkg = PptxPackage(path)
    out = shapes.ungroup_shapes(pkg, 0, group_id)
    pkg.save(do_backup=False)
    assert "timing_report" not in out


def test_the_save_check_refuses_a_timing_target_that_names_no_shape(tmp_path):
    path = _base(tmp_path, "dangling_timing.pptx")
    _give_timing(path, _timing_xml(4242))
    # the deck came in with the problem, so touching the slide keeps it
    pkg = PptxPackage(path)
    pkg.root("ppt/slides/slide1.xml")
    pkg.mark_dirty("ppt/slides/slide1.xml")
    pkg.save(do_backup=False)
    # a change that ADDS one is refused
    pkg = PptxPackage(path)
    pkg.root("ppt/slides/slide1.xml").find(".//" + qn("p:spTgt")).set("spid", "4343")
    pkg.mark_dirty("ppt/slides/slide1.xml")
    with pytest.raises(PptMcpError) as refused:
        pkg.save(do_backup=False)
    assert "4343" in str(refused.value)


def test_the_save_check_refuses_an_empty_animation_target(tmp_path):
    path = _base(tmp_path, "empty_tgt.pptx")
    last = _slide_ids(path)[-1]
    _give_timing(path, _timing_xml(last, audio_spid=last))
    pkg = PptxPackage(path)
    audio_tgt = pkg.root("ppt/slides/slide1.xml").find(
        ".//" + qn("p:cMediaNode") + "/" + qn("p:tgtEl"))
    for child in list(audio_tgt):
        audio_tgt.remove(child)
    pkg.mark_dirty("ppt/slides/slide1.xml")
    with pytest.raises(PptMcpError) as refused:
        pkg.save(do_backup=False)
    assert "names nothing" in str(refused.value)


def test_the_save_check_refuses_a_shape_written_after_the_tree_extlst(tmp_path):
    path = _extlst_deck(tmp_path)
    pkg = PptxPackage(path)
    sp_tree = pkg.root("ppt/slides/slide1.xml").find(".//" + qn("p:spTree"))
    sp = etree.SubElement(sp_tree, qn("p:sp"))
    nv = etree.SubElement(sp, qn("p:nvSpPr"))
    c = etree.SubElement(nv, qn("p:cNvPr"))
    c.set("id", "99")
    c.set("name", "Late")
    etree.SubElement(nv, qn("p:cNvSpPr"))
    etree.SubElement(nv, qn("p:nvPr"))
    etree.SubElement(sp, qn("p:spPr"))
    pkg.mark_dirty("ppt/slides/slide1.xml")
    with pytest.raises(PptMcpError) as refused:
        pkg.save(do_backup=False)
    assert "extLst" in str(refused.value)


def test_cr1_other_parts_keep_the_spelling_of_their_targets(tmp_path):
    """Only the presentation's own relationships are read in another spelling:
    a slide's picture target written with a percent code stays as written."""
    path = _base(tmp_path, "keep_spelling.pptx")
    pkg = PptxPackage(path)
    rels = pkg.rels_for("ppt/slides/slide1.xml")
    node = etree.SubElement(rels.getroot(), "{%s}Relationship" % REL)
    node.set("Id", "rId77")
    node.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image")
    node.set("Target", "../media/my%20picture.jpeg")
    pkg.mark_dirty("ppt/slides/_rels/slide1.xml.rels")
    pkg.set_raw_part("ppt/media/my picture.jpeg", b"\x89PNG made-up bytes")
    pkg.save(do_backup=False)
    again = PptxPackage(path)
    targets = [r.get("Target") for r in again.rels_for("ppt/slides/slide1.xml").getroot()]
    assert "../media/my%20picture.jpeg" in targets
