"""validate never hands PowerPoint a file the server itself refused to read.

validate's payload check refuses a package built to exhaust memory, a
package that names one part twice, and a part that declares a document type
(core/zipsafe.py, core/xmlsafe.py), each with "the file was not read".
validate then still asked PowerPoint to open that same file, and PowerPoint
is one instance per machine, so that could be a PowerPoint the user has
open. A refused file now gets the refusal and no PowerPoint step; a file
with any other problem (a reference to a part that is gone, damage
PowerPoint repairs) still gets PowerPoint's verdict.

Every deck here is made up. PowerPoint is never started: the bridge's
opens-clean call is replaced by a recorder.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt import server

MiB = 1024 * 1024


@pytest.fixture()
def handed(monkeypatch):
    """Every path validate hands to PowerPoint's opens-clean check."""
    from kitchensink4ppt.com import bridge

    paths: list[str] = []

    def recorder(path, *_a, **_k):
        paths.append(str(path))
        return {"opens_clean": True}

    monkeypatch.setattr(bridge, "com_validate_opens_clean", recorder)
    return paths


def _deck(path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Made-up words"
    prs.save(str(path))
    return path


def _rewrite(path: Path, fn) -> Path:
    tmp = path.with_name(path.name + ".rw")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = fn(info.filename, zin.read(info.filename))
            if data is not None:
                zout.writestr(info.filename, data)
    tmp.replace(path)
    return path


def _bomb(path: Path) -> Path:
    """A deck plus one part of 300 MiB of zeros, about 300 KB in the file."""
    _deck(path)
    with zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED, allowZip64=True) as zf:
        info = zipfile.ZipInfo("ppt/made_up_zeros.xml", (2026, 1, 1, 0, 0, 0))
        info.compress_type = zipfile.ZIP_DEFLATED
        with zf.open(info, "w", force_zip64=True) as fh:
            for _ in range(300):
                fh.write(bytes(MiB))
    return path


def _dtd(path: Path) -> Path:
    """A deck whose slide part declares a document type."""
    _deck(path)

    def fn(name, data):
        if name == "ppt/slides/slide1.xml":
            head = re.match(rb"(<\?xml[^>]*\?>)", data)
            return head.group(1) + b'<!DOCTYPE x [<!ENTITY e "made-up">]>' + data[head.end():]
        return data
    return _rewrite(path, fn)


def _twice(path: Path) -> Path:
    """A deck whose ZIP directory names one part twice, in two spellings."""
    _deck(path)
    with zipfile.ZipFile(path) as z:
        slide = z.read("ppt/slides/slide1.xml")
    with zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED) as z:
        z.writestr("ppt/slides/Slide1.xml", slide)
    return path


def _dangling(path: Path) -> Path:
    """A deck with an ordinary problem: a slide relationship to a picture that is gone."""
    _deck(path)

    def fn(name, data):
        if name == "ppt/slides/_rels/slide1.xml.rels":
            return data.replace(
                b"</Relationships>",
                b'<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/'
                b'officeDocument/2006/relationships/image" Target="../media/made_up_gone.png"/>'
                b"</Relationships>")
        return data
    return _rewrite(path, fn)


@pytest.mark.parametrize("kind,make,words", [
    ("bomb", _bomb, "exhaust memory"),
    ("dtd", _dtd, "document type declaration"),
    ("twice", _twice, "name the same part"),
], ids=["bomb", "dtd", "twice"])
def test_validate_does_not_hand_a_refused_file_to_powerpoint(tmp_path, handed, kind, make, words):
    deck = make(tmp_path / f"made_up_{kind}.pptx")
    before = deck.read_bytes()
    out = server.validate(str(deck))
    assert out["payload_valid"] is False, out
    assert words in out["payload_error"], out
    assert handed == [], f"validate refused the {kind} file and still handed it to PowerPoint: {handed}"
    assert out["powerpoint"] is None, out
    note = out.get("powerpoint_note") or ""
    assert "PowerPoint was not asked to open" in note, note
    assert out["ok"] is False
    assert deck.read_bytes() == before


def test_validate_still_asks_powerpoint_about_an_ordinary_problem(tmp_path, handed):
    """A file the server read but found a problem in still gets PowerPoint's
    verdict, which is the authoritative one."""
    deck = _dangling(tmp_path / "made_up_dangling.pptx")
    out = server.validate(str(deck))
    assert out["payload_valid"] is False, out
    assert handed == [str(deck)] or [Path(p) for p in handed] == [deck], handed
    assert out["powerpoint"] == {"opens_clean": True}


def test_validate_still_asks_powerpoint_about_a_sound_file(tmp_path, handed):
    deck = _deck(tmp_path / "made_up_sound.pptx")
    out = server.validate(str(deck))
    assert out["payload_valid"] is True, out
    assert len(handed) == 1
    assert out["ok"] is True
