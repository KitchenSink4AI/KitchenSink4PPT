"""A table PowerPoint keeps in a compatibility wrapper is refused by name, never called missing.

PowerPoint keeps a table whose cell holds an equation in a slide-level
mc:AlternateContent, as it keeps an equation text box and ink. The reads list
and read that table, and an edit that names such a shape is refused by name
with nothing changed. get_table, set_table_cells and merge_cells instead
answered "slide index 0 has no tables", which contradicted the reads and sent
the caller back to a listing that shows the table. They now refuse by shape
id and name, and a slide with a wrapped table and a plain one no longer lets
a selector-less call pick the plain one.

Every deck here is made up by code.
"""

from __future__ import annotations

import copy
import hashlib

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.errors import (
    AmbiguousTarget,
    TargetNotFound,
    UnsupportedStructure,
)
from kitchensink4ppt.core.package import PptxPackage
from kitchensink4ppt.ops import read, tables

MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A14 = "http://schemas.microsoft.com/office/drawing/2010/main"


def _deck(tmp_path, *, plain_too=False):
    """A slide with a 1 x 2 table kept in a slide-level wrapper, and, when asked, a plain table."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    frame = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(6), Inches(1))
    frame.name = "Made-up wrapped table"
    frame.table.cell(0, 0).text = "Left words"
    frame.table.cell(0, 1).text = "Right words"
    elem = frame._element
    tree = elem.getparent()
    ac = etree.Element(f"{{{MC}}}AlternateContent", nsmap={"mc": MC})
    choice = etree.SubElement(ac, f"{{{MC}}}Choice", nsmap={"a14": A14})
    choice.set("Requires", "a14")
    fallback = etree.SubElement(ac, f"{{{MC}}}Fallback")
    at = list(tree).index(elem)
    tree.remove(elem)
    choice.append(elem)
    fallback.append(copy.deepcopy(elem))
    tree.insert(at, ac)
    wrapped_id = frame.shape_id
    if plain_too:
        plain = slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(3), Inches(1))
        plain.name = "Made-up plain table"
        plain.table.cell(0, 0).text = "Plain words"
    path = tmp_path / "wrapped_table.pptx"
    prs.save(str(path))
    return path, wrapped_id


def _md5(path):
    return hashlib.md5(path.read_bytes()).hexdigest()


def test_the_reads_list_the_wrapped_table(tmp_path):
    path, sid = _deck(tmp_path)
    pkg = PptxPackage(path)
    info = read.get_slide_info(pkg, 0)
    assert [s["id"] for s in info["shapes"] if s["type"] == "table"] == [sid]


@pytest.mark.parametrize("selector", [None, {"shape_id": "ID"}, {"index": 0}, "ID"])
def test_get_table_refuses_a_wrapped_table_by_name(tmp_path, selector):
    path, sid = _deck(tmp_path)
    if selector == "ID":
        selector = sid
    elif isinstance(selector, dict) and selector.get("shape_id") == "ID":
        selector = {"shape_id": sid}
    with pytest.raises(UnsupportedStructure) as err:
        tables.get_table(PptxPackage(path), 0, selector)
    msg = str(err.value)
    assert "no tables" not in msg
    assert str(sid) in msg and "Made-up wrapped table" in msg
    assert "compatibility wrapper" in msg


def test_set_table_cells_refuses_a_wrapped_table_and_changes_nothing(tmp_path):
    path, sid = _deck(tmp_path)
    before = _md5(path)
    pkg = PptxPackage(path)
    with pytest.raises(UnsupportedStructure) as err:
        tables.set_table_cells(pkg, 0, {"shape_id": sid}, [{"row": 0, "col": 0, "text": "New words"}])
    assert "no tables" not in str(err.value)
    assert "Made-up wrapped table" in str(err.value)
    assert _md5(path) == before


def test_merge_cells_refuses_a_wrapped_table_and_changes_nothing(tmp_path):
    path, sid = _deck(tmp_path)
    before = _md5(path)
    with pytest.raises(UnsupportedStructure) as err:
        tables.merge_cells(PptxPackage(path), 0, {"shape_id": sid}, 0, 0, 0, 1)
    assert "no tables" not in str(err.value)
    assert _md5(path) == before


def test_a_slide_with_no_table_at_all_still_says_so(tmp_path):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "none.pptx"
    prs.save(str(path))
    with pytest.raises(TargetNotFound, match="has no tables"):
        tables.get_table(PptxPackage(path), 0, None)


def test_a_plain_table_beside_a_wrapped_one_is_still_addressed_by_id(tmp_path):
    path, sid = _deck(tmp_path, plain_too=True)
    pkg = PptxPackage(path)
    plain_id = [t["shape_id"] for t in tables._tables_on_slide(pkg, pkg.slide_parts()[0])][0]
    out = tables.get_table(pkg, 0, {"shape_id": plain_id})
    assert "Plain words" in str(out)
    with pytest.raises(UnsupportedStructure):
        tables.get_table(pkg, 0, {"shape_id": sid})


def test_a_call_without_a_selector_is_ambiguous_when_a_wrapped_table_shares_the_slide(tmp_path):
    path, sid = _deck(tmp_path, plain_too=True)
    with pytest.raises(AmbiguousTarget) as err:
        tables.get_table(PptxPackage(path), 0, None)
    assert str(sid) in str(err.value)


def _deck_with_three_plain_tables(tmp_path):
    """The wrapped table has shape id 2; the plain table at index 2 has shape
    id 5, so the bare integer 2 reads as two different tables."""
    path, sid = _deck(tmp_path, plain_too=True)
    prs = Presentation(str(path))
    slide = prs.slides[0]
    for n in (1, 2):
        extra = slide.shapes.add_table(1, 1, Inches(1), Inches(4 + n), Inches(3), Inches(1))
        extra.name = f"Made-up plain table {n + 1}"
    prs.save(str(path))
    pkg = PptxPackage(path)
    plain = tables._tables_on_slide(pkg, pkg.slide_parts()[0])
    assert plain[2]["shape_id"] != sid and sid == 2
    return path, sid, plain[2]["shape_id"]


def test_a_bare_integer_that_is_a_wrapped_shape_id_and_a_plain_index_is_refused(tmp_path):
    path, sid, plain_id = _deck_with_three_plain_tables(tmp_path)
    before = _md5(path)
    pkg = PptxPackage(path)
    with pytest.raises(AmbiguousTarget) as err:
        tables.set_table_cells(pkg, 0, sid, [{"row": 0, "col": 0, "text": "Wrong target"}])
    msg = str(err.value)
    assert str(plain_id) in msg and "Made-up wrapped table" in msg
    assert "{'index': 2}" in msg and "{'shape_id': 2}" in msg
    with pytest.raises(AmbiguousTarget):
        tables.get_table(pkg, 0, sid)
    assert _md5(path) == before


def test_the_explicit_selectors_still_work_beside_a_wrapped_table(tmp_path):
    path, sid, plain_id = _deck_with_three_plain_tables(tmp_path)
    pkg = PptxPackage(path)
    assert tables.resolve_table(pkg, 0, {"index": 2})["shape_id"] == plain_id
    with pytest.raises(UnsupportedStructure):
        tables.resolve_table(pkg, 0, {"shape_id": sid})
