"""A server stopped in the middle of a save never loses the deck.

ReplaceFileW moves the deck aside and then moves the new version in. A server
stopped between those two moves (a crash, an app closed while it saves, a
client that ends its servers) left the deck's path empty, with the deck only
under a backup name that said nothing to a person, and nothing ever put it
back. Now the aside name says what it holds, each swap is written down in a
journal first, and the deck is put back at the next start of a server and at
the deck's next write. A read never moves it: it says where the deck is.

The kill here is real: a server process saving the deck is ended by its
exact process handle at the instant between the two moves. ReplaceFileW
cannot be paused there, so the victim's ReplaceFileW does the first move
itself and then waits to be killed, which is the state a kill at that
instant leaves.
"""

from __future__ import annotations

import asyncio
import inspect
import json
import os
import subprocess
import sys
import time
from pathlib import Path

import pytest
from pptx import Presentation

from kitchensink4ppt import server
from kitchensink4ppt.core import safesave, swapwait
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.plan import build

pytestmark = pytest.mark.skipif(sys.platform != "win32", reason="ReplaceFileW")

NO_WINDOW = 0x08000000

VICTIM = r'''
import asyncio, os, sys, time
from kitchensink4ppt import server  # noqa: F401
from kitchensink4ppt.core import safesave
from kitchensink4ppt.settings import Settings
from kitchensink4ppt.surface.dispatch import dispatch
from kitchensink4ppt.surface.plan import build

deck, marker = sys.argv[1], sys.argv[2]
fn, move, last_error, text = safesave._replace_file()

def stopped_between_the_moves(ldst, lsrc, lbak, flags, a, b):
    os.replace(ldst, lbak)  # ReplaceFileW's first move: the deck aside
    open(marker, "w").close()
    time.sleep(600)  # killed here, before the new version moves in

safesave._REPLACE_FILE = (stopped_between_the_moves, move, last_error, text)
plan = build(Settings())
asyncio.run(dispatch(plan, "ppt_edit", {"op": "set_notes", "args": {
    "file_path": deck, "slide": 0, "text": "never saved", "live": "off"}}))
'''


@pytest.fixture(autouse=True)
def own_journal(tmp_path, monkeypatch):
    """Each test reads only the swaps it wrote down (the victim process
    inherits the folder through the environment)."""
    monkeypatch.setenv("KS4P_SAVE_JOURNAL_DIR", str(tmp_path / "journal"))


def _deck(path: Path, notes: str) -> Path:
    prs = Presentation()
    for title in ("Alpha", "Bravo"):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.notes_slide.notes_text_frame.text = notes
    prs.save(path)
    return path


def _notes(path) -> str:
    return Presentation(str(path)).slides[0].notes_slide.notes_text_frame.text


def _call(op, args, settings=None):
    plan = build(settings or Settings())
    return asyncio.run(dispatch(plan, plan.specs[op].group, {"op": op, "args": args}))


def _entries_for(deck: Path) -> list[dict]:
    key = safesave.canonical_key(deck)
    return [record for _entry, record in safesave._journal_entries()
            if isinstance(record.get("deck"), str)
            and safesave.canonical_key(record["deck"]) == key]


@pytest.fixture
def killed_mid_save(tmp_path):
    """A deck whose save was stopped between ReplaceFileW's two moves by a
    real kill of the saving process."""
    deck = _deck(tmp_path / "Quarterly review.pptx", "as it was")
    script = tmp_path / "victim.py"
    script.write_text(VICTIM, encoding="utf-8")
    marker = tmp_path / "between-the-moves"
    victim = subprocess.Popen(
        [sys.executable, "-X", "utf8", str(script), str(deck), str(marker)],
        stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        creationflags=NO_WINDOW)
    try:
        deadline = time.monotonic() + 90
        while not marker.exists():
            assert victim.poll() is None, "the saving process ended before the swap"
            assert time.monotonic() < deadline, "the saving process never reached the swap"
            time.sleep(0.02)
    finally:
        victim.kill()  # TerminateProcess on this process's own handle
        victim.wait(timeout=60)
    return deck


def test_a_kill_between_the_moves_leaves_the_deck_aside_under_a_name_that_says_so(killed_mid_save):
    deck = killed_mid_save
    assert not deck.exists(), "the state a kill at that instant leaves"
    aside = swapwait.asides(deck)
    assert len(aside) == 1
    name = os.path.basename(aside[0])
    assert name.startswith(safesave._swap_prefix(deck.name))
    assert "Quarterly review.pptx (as it was before a save)" in name
    assert _notes(aside[0]) == "as it was"
    assert len(_entries_for(deck)) == 1, "the swap was written down before it started"


def test_the_next_start_of_a_server_puts_the_deck_back(killed_mid_save):
    deck = killed_mid_save
    assert server._recover_at_start(Settings()) == [str(deck)]
    assert _notes(deck) == "as it was"
    assert swapwait.asides(deck) == []
    assert not (deck.parent / (deck.name + ".ppt-mcp-tmp")).exists(), (
        "the version whose save did not finish is not left beside the deck")
    assert _entries_for(deck) == []
    payload, is_error = _call("get_text", {"file_path": str(deck), "live": "off"})
    assert is_error is False, payload


def test_the_next_change_puts_the_deck_back_first(killed_mid_save):
    deck = killed_mid_save
    payload, is_error = _call("set_notes", {"file_path": str(deck), "slide": 0,
                                            "text": "next", "live": "off"})
    assert is_error is False, payload
    assert _notes(deck) == "next"
    assert _notes(safesave.slot_dir(deck) / safesave.PREV_SLOT) == "as it was"
    assert swapwait.asides(deck) == [] and _entries_for(deck) == []


def test_a_read_says_where_the_deck_is_and_moves_nothing(killed_mid_save):
    deck = killed_mid_save
    aside = swapwait.asides(deck)[0]
    payload, is_error = _call("get_text", {"file_path": str(deck), "live": "off"})
    assert is_error is True and payload["error"]["code"] == "NOT_FOUND", payload
    assert aside in payload["error"]["message"]
    assert "puts it back" in payload["error"]["message"]
    assert os.path.exists(aside) and not deck.exists(), "a read changes nothing"


def test_a_read_only_server_leaves_it_to_a_server_that_may_write(killed_mid_save):
    deck = killed_mid_save
    assert server._recover_at_start(Settings(read_only=True)) == []
    assert not deck.exists() and len(swapwait.asides(deck)) == 1


def test_the_server_start_runs_the_recovery():
    source = inspect.getsource(server.main)
    assert "_recover_at_start(settings)" in source
    assert source.index("_recover_at_start(settings)") < source.index("mcp.run(")


def test_a_deck_that_is_aside_keeps_its_backups_through_a_purge_of_orphans(killed_mid_save):
    deck = killed_mid_save
    slots = safesave.slot_dir(deck)
    kept = sorted(os.listdir(slots))
    assert kept, "the stopped save left its backup folder"
    payload, is_error = _call("manage_backups", {"action": "purge", "scope": "orphans",
                                                 "directory": str(deck.parent),
                                                 "dry_run": False})
    assert is_error is False, payload
    assert payload["count"] == 0, payload
    assert sorted(os.listdir(slots)) == kept


def test_a_new_deck_is_never_made_over_one_that_is_aside(killed_mid_save):
    deck = killed_mid_save
    payload, is_error = _call("create_presentation", {"path": str(deck)})
    assert is_error is True, payload
    assert _notes(deck) == "as it was", "the deck was put back and kept"


def test_a_copy_onto_a_deck_that_is_aside_does_not_write_over_it(killed_mid_save, tmp_path):
    deck = killed_mid_save
    other = _deck(tmp_path / "other.pptx", "other")
    payload, is_error = _call("copy_presentation", {"file_path": str(other),
                                                    "dest_path": str(deck)})
    assert is_error is True and payload["error"]["code"] == "CONFLICT", payload
    assert _notes(deck) == "as it was"


# ----------------------------------------- a swap that finished, then a kill


def _write_entry(record: dict) -> str:
    folder = safesave._journal_dir()
    folder.mkdir(parents=True, exist_ok=True)
    entry = folder / f"{os.urandom(8).hex()}.json"
    entry.write_text(json.dumps(record), encoding="utf-8")
    return str(entry)


def _ours(**fields) -> dict:
    return {"pid": os.getpid(), "token": safesave._OWNER_TOKEN,
            "pid_created": safesave._OWNER_CREATED, "host": safesave._local_host(),
            "time": time.time(), **fields}


def test_a_stray_copy_of_the_old_version_goes_once_the_prev_slot_holds_it(tmp_path):
    """The swap ran to its end and the server was stopped before it removed
    the old file: the deck is in place and the old version is still aside.
    It is the very file the prev slot holds, so it goes at the next write."""
    deck = _deck(tmp_path / "deck.pptx", "v1")
    payload, is_error = _call("set_notes", {"file_path": str(deck), "slide": 0,
                                            "text": "v2", "live": "off"})
    assert not is_error, payload
    prev = safesave.slot_dir(deck) / safesave.PREV_SLOT
    stray = deck.parent / swapwait.aside_name(deck.name, "0123456789ab")
    os.link(prev, stray)
    entry = _write_entry(_ours(deck=str(deck), aside=str(stray),
                               new=str(deck.parent / "deck.pptx.ppt-mcp-tmp")))
    with safesave.write_lock(deck):
        pass
    assert not stray.exists() and not os.path.exists(entry)
    assert _notes(deck) == "v2" and _notes(prev) == "v1"


def test_a_stray_that_is_not_in_the_prev_slot_is_kept(tmp_path):
    """The only copy of a version is kept beside the deck, the next answer
    names it, and its record stays while it is there, so a later run names
    it again; once it is gone, the record goes too."""
    safesave.take_notices()
    deck = _deck(tmp_path / "deck.pptx", "v1")
    stray = deck.parent / swapwait.aside_name(deck.name, "0123456789ab")
    _deck(stray, "only copy")
    entry = _write_entry(_ours(deck=str(deck), aside=str(stray),
                               new=str(deck.parent / "deck.pptx.ppt-mcp-tmp")))
    with safesave.write_lock(deck):
        pass
    assert stray.exists() and _notes(stray) == "only copy"
    assert any(str(stray) in n for n in safesave.take_notices())
    assert os.path.exists(entry)
    os.remove(stray)
    with safesave.write_lock(deck):
        pass
    assert not os.path.exists(entry)


def test_an_entry_of_a_server_that_is_still_running_is_left_alone(tmp_path):
    """Another server's swap that is running right now: its parent process
    stands in for it (alive, with its own creation time)."""
    deck = tmp_path / "deck.pptx"
    aside = deck.parent / swapwait.aside_name(deck.name, "0123456789ab")
    _deck(aside, "in flight")
    parent = os.getppid()
    entry = _write_entry({"deck": str(deck), "aside": str(aside),
                          "new": str(deck) + ".ppt-mcp-tmp", "pid": parent,
                          "token": "another-server",
                          "pid_created": safesave._process_create_time(parent),
                          "host": safesave._local_host(), "time": time.time()})
    try:
        assert safesave.recover_interrupted_saves() == []
        assert aside.exists() and not deck.exists() and os.path.exists(entry)
    finally:
        os.remove(entry)


def test_the_aside_name_fits_windows_limits_for_a_long_korean_name():
    name = "가" * 200 + ".pptx"
    aside = swapwait.aside_name(name, "0123456789ab")
    assert aside.startswith(safesave._swap_prefix(name))
    assert len(aside.encode("utf-16-le")) // 2 <= 255
    assert aside.endswith("(as it was before a save).tmp")


# ------------------------------- a swap that ReplaceFileW did not go through


def _replace_fails(monkeypatch, err: int, move_back_ok: bool) -> None:
    """ReplaceFileW answers ``err``. For 1177 it first does what ReplaceFileW
    has done by then: the deck is at the backup name."""

    def replace(ldst, lsrc, lbak, flags, a, b):
        if err == 1177:
            os.replace(ldst, lbak)
        return False

    def move(a, b, flags):
        if not move_back_ok:
            return False
        os.replace(a, b)
        return True

    monkeypatch.setattr(safesave, "_RETRY_DELAYS", (0.0,))
    monkeypatch.setattr(safesave, "_REPLACE_FILE", (replace, move, lambda: err, lambda e: f"error {e}"))


@pytest.mark.parametrize("err", [1176, 1177])
def test_a_swap_that_changed_nothing_says_so_and_can_be_retried(tmp_path, monkeypatch, err):
    deck = _deck(tmp_path / "deck.pptx", "v1")
    before = deck.read_bytes()
    _replace_fails(monkeypatch, err, move_back_ok=True)
    payload, is_error = _call("set_notes", {"file_path": str(deck), "slide": 0,
                                            "text": "v2", "live": "off"})
    assert is_error is True and payload["error"]["code"] == "DOCUMENT_LOCKED", payload
    assert "NOT modified" in payload["error"]["message"]
    assert "Retry" in payload["error"]["hint"]
    assert deck.read_bytes() == before and swapwait.asides(deck) == []
    assert _entries_for(deck) == []


ROUTES_THAT_SWAP = {
    "set_notes": lambda deck, tmp: {"file_path": str(deck), "slide": 0, "text": "v2", "live": "off"},
    "manage_backups": lambda deck, tmp: {"action": "restore", "file_path": str(deck),
                                         "source": str(_deck(tmp / "snapshot.pptx", "snap"))},
    "copy_presentation": lambda deck, tmp: {"file_path": str(_deck(tmp / "other.pptx", "other")),
                                            "dest_path": str(deck), "overwrite": True},
}


@pytest.mark.parametrize("op", sorted(ROUTES_THAT_SWAP))
def test_a_deck_that_could_not_be_put_back_is_named_and_comes_back_at_the_next_change(
        tmp_path, monkeypatch, op):
    """1177 and the move back fails too: the deck is not at its path. The
    answer says where it is (it used to be "Something inside the server
    failed"), and the next change puts it back before it runs."""
    deck = _deck(tmp_path / "deck.pptx", "v1")
    _replace_fails(monkeypatch, 1177, move_back_ok=False)
    payload, is_error = _call(op, ROUTES_THAT_SWAP[op](deck, tmp_path))
    assert is_error is True and payload["error"]["code"] == "CONFLICT", payload
    aside = swapwait.asides(deck)
    assert len(aside) == 1 and not deck.exists()
    assert aside[0] in payload["error"]["message"]
    assert "deck.pptx" in payload["error"]["message"] and "lost" in payload["error"]["message"]
    assert _notes(aside[0]) == "v1"

    monkeypatch.setattr(safesave, "_REPLACE_FILE", None)  # the real ReplaceFileW again
    payload, is_error = _call("set_notes", {"file_path": str(deck), "slide": 0,
                                            "text": "v3", "live": "off"})
    assert is_error is False, payload
    assert _notes(deck) == "v3" and swapwait.asides(deck) == []
    assert _entries_for(deck) == []
