"""Regressions for the 2026-09-21 defense-deck field report (45 slides
built end to end through the server).

P3  copy_presentation to a folder that does not exist refused with the
    raw WinError text and a NOT_FOUND hint telling the caller to re-read
    the deck, which is not the problem and not the fix.
P5  set_placeholder_text rebuilt the body as bare a:p / a:r, so a text
    box whose formatting lives on paragraph-level pPr/defRPr came back
    with no indents, no bullet definition, no size, no colour and no
    typeface. (The set_shape / apply_edits set_text route was fixed for
    the 2026-09-20 report; the end-to-end test here pins all four
    text-writing routes against the same body.)
P6  a p:txBody with zero a:p passed layer 1 and made PowerPoint refuse
    the whole deck, and layer 2 reported "it failed" without naming the
    slide or the shape: about 40 minutes of manual bisection.
P8  create_table split the box evenly and offered no way to state row
    heights or column widths at creation.
887 the enable_tools note promised a tool list refresh that a worker or
    subagent never gets.
"""

from __future__ import annotations

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.package import PptxPackage, qn
from kitchensink4ppt.ops import shapes as shp
from kitchensink4ppt.ops.read import get_slide_info

# ------------------------------------------------------------------ P5

#: The body the field run actually had: every visual property on a:pPr
#: (indent, spacing, bullet) and a:defRPr (size, colour, typeface), and
#: runs carrying nothing but text. Nothing states an alignment or an
#: anchor, so a write that introduces one has invented it.
_MARL = "285750"
_INDENT = "-285750"


def _paragraph_level_body(body, lines):
    """Rebuild `body` as paragraphs whose formatting lives entirely on
    a:pPr / a:defRPr, with bare runs. `lines` is [(text, level), ...]."""
    for p in list(body.findall(qn("a:p"))):
        body.remove(p)
    bodypr = body.find(qn("a:bodyPr"))
    bodypr.attrib.pop("anchor", None)
    for text, level in lines:
        p = etree.SubElement(body, qn("a:p"))
        ppr = etree.SubElement(p, qn("a:pPr"))
        if level:
            ppr.set("lvl", str(level))
        ppr.set("marL", _MARL)
        ppr.set("indent", _INDENT)
        spc = etree.SubElement(ppr, qn("a:spcAft"))
        pts = etree.SubElement(spc, qn("a:spcPts"))
        pts.set("val", "600")
        bufont = etree.SubElement(ppr, qn("a:buFont"))
        bufont.set("typeface", "Arial")
        buchar = etree.SubElement(ppr, qn("a:buChar"))
        buchar.set("char", "•")
        defrpr = etree.SubElement(ppr, qn("a:defRPr"))
        defrpr.set("sz", "1400")
        fill = etree.SubElement(defrpr, qn("a:solidFill"))
        clr = etree.SubElement(fill, qn("a:srgbClr"))
        clr.set("val", "1F3864")
        latin = etree.SubElement(defrpr, qn("a:latin"))
        latin.set("typeface", "Calibri")
        r = etree.SubElement(p, qn("a:r"))  # bare run: no a:rPr at all
        t = etree.SubElement(r, qn("a:t"))
        t.text = text
    return body


@pytest.fixture()
def field_deck(tmp_path):
    """(pkg, slide, textbox_id, placeholder_id): one slide carrying a free
    text box and a body placeholder, both formatted the field-run way."""
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1])  # title + content
    path = tmp_path / "field.pptx"
    prs.save(str(path))
    pkg = PptxPackage(str(path))
    sid = shp.insert_shape(
        pkg, 0, "rect", 1, 1, 6, 3, text="placeholder"
    )["shape_id"]
    part = get_slide_info(pkg, 0)["part"]
    elem, _chain = shp._find_shape(pkg, part, sid)
    _paragraph_level_body(
        elem.find(qn("p:txBody")),
        [("Alpha one", 0), ("Beta two", 1), ("Gamma three", 0)],
    )
    ph_id = None
    tree = pkg.root(part).find(f"{qn('p:cSld')}/{qn('p:spTree')}")
    for sp in tree.findall(qn("p:sp")):
        ph = sp.find(
            f"{qn('p:nvSpPr')}/{qn('p:nvPr')}/{qn('p:ph')}"
        )
        if ph is None or ph.get("type") not in (None, "body", "obj"):
            continue
        ph_id = int(
            sp.find(f"{qn('p:nvSpPr')}/{qn('p:cNvPr')}").get("id")
        )
        _paragraph_level_body(
            sp.find(qn("p:txBody")),
            [("Alpha one", 0), ("Beta two", 1), ("Gamma three", 0)],
        )
    assert ph_id is not None, "layout 1 gave no body placeholder"
    pkg.mark_dirty(part)
    return pkg, 0, sid, ph_id


def _paragraphs_of(pkg, slide, shape_id):
    part = get_slide_info(pkg, slide)["part"]
    elem, _chain = shp._find_shape(pkg, part, shape_id)
    body = elem.find(qn("p:txBody"))
    return body, body.findall(qn("a:p"))


def _assert_field_formatting_intact(pkg, slide, shape_id, route):
    """Every property the field report watched die, on every paragraph,
    plus the two attributes the old rebuild invented."""
    body, paras = _paragraphs_of(pkg, slide, shape_id)
    bodypr = body.find(qn("a:bodyPr"))
    assert bodypr.get("anchor") is None, (
        f"{route} invented a vertical anchor of {bodypr.get('anchor')!r}"
    )
    assert paras, f"{route} left the body with no paragraphs"
    for i, p in enumerate(paras):
        where = f"{route} paragraph {i}"
        ppr = p.find(qn("a:pPr"))
        assert ppr is not None, f"{where}: a:pPr is gone"
        assert ppr.get("algn") is None, f"{where}: alignment was invented"
        assert ppr.get("marL") == _MARL, f"{where}: marL lost"
        assert ppr.get("indent") == _INDENT, f"{where}: indent lost"
        spc = ppr.find(f"{qn('a:spcAft')}/{qn('a:spcPts')}")
        assert spc is not None and spc.get("val") == "600", (
            f"{where}: spcAft lost"
        )
        assert ppr.find(qn("a:buFont")).get("typeface") == "Arial", (
            f"{where}: buFont lost"
        )
        assert ppr.find(qn("a:buChar")).get("char") == "•", (
            f"{where}: buChar lost"
        )
        defrpr = ppr.find(qn("a:defRPr"))
        assert defrpr is not None, f"{where}: defRPr is gone"
        assert defrpr.get("sz") == "1400", f"{where}: defRPr size lost"
        clr = defrpr.find(f"{qn('a:solidFill')}/{qn('a:srgbClr')}")
        assert clr is not None and clr.get("val") == "1F3864", (
            f"{where}: defRPr colour lost"
        )
        assert defrpr.find(qn("a:latin")).get("typeface") == "Calibri", (
            f"{where}: defRPr typeface lost"
        )
    # the two outline levels survive as themselves
    assert [p.find(qn("a:pPr")).get("lvl") for p in paras[:3]] == [
        None, "1", None
    ], f"{route}: outline levels moved"


def test_set_placeholder_text_keeps_paragraph_level_formatting(field_deck):
    """P5: the route the 2026-09-20 fix did not cover. The rebuild dropped
    every a:pPr child and every run property; a 45-slide deck had to be
    written through python-pptx because of it."""
    from kitchensink4ppt.ops import text as txt

    pkg, slide, _sid, ph_id = field_deck
    out = txt.set_placeholder_text(
        pkg, slide, "body", paragraphs=[
            {"text": "Alpha one edited"},
            {"text": "Beta two edited", "level": 1},
            {"text": "Gamma three edited"},
        ],
    )
    assert out["paragraphs"] == 3
    _assert_field_formatting_intact(
        pkg, slide, ph_id, "set_placeholder_text"
    )
    _body, paras = _paragraphs_of(pkg, slide, ph_id)
    assert [
        p.find(f"{qn('a:r')}/{qn('a:t')}").text for p in paras
    ] == ["Alpha one edited", "Beta two edited", "Gamma three edited"]
    assert "preserved" in out, "the carry stayed silent about what it kept"
    assert set(out["preserved"]) >= {"bullets"}


def test_set_placeholder_text_carries_run_properties_too(field_deck):
    """P5: runs that DO carry rPr keep their size, colour and typeface."""
    from kitchensink4ppt.ops import text as txt

    pkg, slide, _sid, ph_id = field_deck
    part = get_slide_info(pkg, slide)["part"]
    elem, _chain = shp._find_shape(pkg, part, ph_id)
    for r in elem.find(qn("p:txBody")).iter(qn("a:r")):
        rpr = etree.SubElement(r, qn("a:rPr"))
        r.remove(rpr)
        r.insert(0, rpr)
        rpr.set("sz", "1100")
        fill = etree.SubElement(rpr, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", "C00000")
        etree.SubElement(rpr, qn("a:latin")).set("typeface", "Georgia")

    txt.set_placeholder_text(pkg, slide, "body", text="One\n\tTwo\nThree")
    _body, paras = _paragraphs_of(pkg, slide, ph_id)
    for i, p in enumerate(paras):
        rpr = p.find(f"{qn('a:r')}/{qn('a:rPr')}")
        assert rpr is not None, f"paragraph {i}: run properties are gone"
        assert rpr.get("sz") == "1100", f"paragraph {i}: run size lost"
        assert rpr.find(
            f"{qn('a:solidFill')}/{qn('a:srgbClr')}"
        ).get("val") == "C00000", f"paragraph {i}: run colour lost"
        assert rpr.find(qn("a:latin")).get("typeface") == "Georgia"


def test_set_placeholder_text_more_paragraphs_than_before(field_deck):
    """P5: extra paragraphs take the last old one's shape, the way
    pressing Enter in PowerPoint does, instead of arriving unformatted."""
    from kitchensink4ppt.ops import text as txt

    pkg, slide, _sid, ph_id = field_deck
    txt.set_placeholder_text(
        pkg, slide, "body",
        paragraphs=[{"text": f"Line {i}"} for i in range(6)],
    )
    _body, paras = _paragraphs_of(pkg, slide, ph_id)
    assert len(paras) == 6
    for i, p in enumerate(paras):
        ppr = p.find(qn("a:pPr"))
        assert ppr is not None and ppr.get("marL") == _MARL, (
            f"paragraph {i} arrived with no formatting"
        )


def test_set_placeholder_text_stated_level_beats_the_carried_one(field_deck):
    """P5 guard: preservation is not stickiness. A level the caller states
    wins over the level of the paragraph it replaces."""
    from kitchensink4ppt.ops import text as txt

    pkg, slide, _sid, ph_id = field_deck
    txt.set_placeholder_text(
        pkg, slide, "body",
        paragraphs=[
            {"text": "A", "level": 3},
            {"text": "B", "level": 0},
            {"text": "C", "level": 2},
        ],
    )
    _body, paras = _paragraphs_of(pkg, slide, ph_id)
    assert [p.find(qn("a:pPr")).get("lvl") for p in paras] == ["3", "0", "2"]
    # and the rest of the paragraph still rode through
    assert paras[1].find(qn("a:pPr")).get("marL") == _MARL


def test_every_text_writing_route_keeps_the_field_formatting(field_deck):
    """P5 end to end: the original repro, run against all four routes that
    write text into a shape. Any one of them failing is the field bug."""
    from kitchensink4ppt.ops import batch, text as txt

    pkg, slide, sid, ph_id = field_deck

    batch.apply_edits(pkg, [{
        "op": "set_text", "slide": slide, "shape": sid,
        "text": "Alpha one\nBeta two\nGamma three",
    }])
    _assert_field_formatting_intact(pkg, slide, sid, "apply_edits set_text")

    shp.set_shape(pkg, slide, sid, text="Alpha one\nBeta two\nGamma three")
    _assert_field_formatting_intact(pkg, slide, sid, "set_shape(text=)")

    txt.set_placeholder_text(
        pkg, slide, "body", text="Alpha one\n\tBeta two\nGamma three"
    )
    _assert_field_formatting_intact(
        pkg, slide, ph_id, "set_placeholder_text"
    )

    txt.search_and_replace(pkg, "Alpha", "Omega")
    _assert_field_formatting_intact(pkg, slide, sid, "search_and_replace")
    _assert_field_formatting_intact(
        pkg, slide, ph_id, "search_and_replace"
    )


def test_set_placeholder_text_with_no_paragraphs_leaves_one(field_deck):
    """P6 producer 1: paragraphs=[] used to leave a p:txBody with zero
    a:p, which is the corruption PowerPoint refuses to open."""
    from kitchensink4ppt.ops import text as txt

    pkg, slide, _sid, ph_id = field_deck
    txt.set_placeholder_text(pkg, slide, "body", paragraphs=[])
    _body, paras = _paragraphs_of(pkg, slide, ph_id)
    assert len(paras) == 1, "an empty body is a deck PowerPoint refuses"
    assert paras[0].find(qn("a:r")) is None


# ------------------------------------------------------------------ P6


def _empty_the_first_body(pkg, slide=0):
    """Strip every a:p from the first text body on a slide: the exact
    corruption the field run produced and then had to bisect by hand."""
    part = get_slide_info(pkg, slide)["part"]
    tree = pkg.root(part).find(f"{qn('p:cSld')}/{qn('p:spTree')}")
    for sp in tree.findall(qn("p:sp")):
        body = sp.find(qn("p:txBody"))
        if body is None:
            continue
        for p in list(body.findall(qn("a:p"))):
            body.remove(p)
        pkg.mark_dirty(part)
        return int(sp.find(f"{qn('p:nvSpPr')}/{qn('p:cNvPr')}").get("id"))
    raise AssertionError("no text body on the slide")


def test_save_refuses_a_text_body_with_no_paragraphs(tmp_path):
    """P6 layer 1: payload_valid said true on a deck PowerPoint would not
    open. The content model wants at least one a:p per body."""
    from kitchensink4ppt.core.errors import ValidationFailed

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1])
    path = tmp_path / "empty_body.pptx"
    prs.save(str(path))
    pkg = PptxPackage(str(path))
    shape_id = _empty_the_first_body(pkg)

    with pytest.raises(ValidationFailed) as excinfo:
        pkg.save(str(tmp_path / "out.pptx"))
    message = str(excinfo.value)
    assert "ppt/slides/slide1.xml" in message, message
    assert "a:p" in message or "paragraph" in message, message
    assert str(shape_id) in message, (
        f"the failure does not name the shape: {message}"
    )


def test_empty_table_cell_body_is_refused_too(tmp_path):
    """P6 layer 1: a:txBody in a table cell has the same content model."""
    from kitchensink4ppt.core.errors import ValidationFailed
    from kitchensink4ppt.ops import tables as tb

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "table.pptx"
    prs.save(str(path))
    pkg = PptxPackage(str(path))
    tb.create_table(pkg, 0, 2, 2, 1, 1, 6, 2, data=[["a", "b"], ["c", "d"]])
    part = get_slide_info(pkg, 0)["part"]
    tbl = pkg.root(part).find(f".//{qn('a:tbl')}")
    cell_body = tbl.find(f".//{qn('a:tc')}/{qn('a:txBody')}")
    for p in list(cell_body.findall(qn("a:p"))):
        cell_body.remove(p)
    pkg.mark_dirty(part)

    with pytest.raises(ValidationFailed) as excinfo:
        pkg.save(str(tmp_path / "out.pptx"))
    assert "a:p" in str(excinfo.value) or "paragraph" in str(excinfo.value)


def test_the_whole_corpus_still_validates(tmp_path):
    """P6 guard: _validate_payload runs on every save, so a new assertion
    that refuses a legitimate deck would brick the server. Every corpus
    deck (real or generated) must still round-trip."""
    from pathlib import Path

    corpus = Path(__file__).resolve().parents[1] / "corpus"
    decks = sorted(corpus.glob("*.pptx")) + sorted(corpus.glob("*.potx"))
    assert decks, "no corpus to check"
    for deck in decks:
        pkg = PptxPackage(str(deck))
        out = tmp_path / f"rt_{deck.stem}.pptx"
        pkg.save(str(out))  # raises ValidationFailed if the check is wrong
        assert out.is_file()


def test_full_load_names_the_slide_and_shape_that_failed():
    """P6 layer 2: 'it failed' with no slide and no shape cost about 40
    minutes of manual bisection on a 45-slide deck."""
    from kitchensink4ppt.com import bridge

    class _Shape:
        def __init__(self, name, ok=True):
            self.Name = name
            self._ok = ok

        @property
        def HasTextFrame(self):
            if not self._ok:
                raise OSError("the shape refused to load")
            return False

    class _Shapes:
        def __init__(self, shapes):
            self._shapes = shapes
            self.Count = len(shapes)

        def Item(self, i):
            return self._shapes[i - 1]

    class _Slide:
        def __init__(self, sid, shapes):
            self.SlideID = sid
            self.Shapes = _Shapes(shapes)

    class _Slides:
        def __init__(self, slides):
            self._slides = slides
            self.Count = len(slides)

        def Item(self, i):
            return self._slides[i - 1]

    class _Pres:
        def __init__(self, slides):
            self.Slides = _Slides(slides)

    pres = _Pres([
        _Slide(256, [_Shape("Title 1")]),
        _Slide(257, [_Shape("Body 2"), _Shape("Broken 3", ok=False)]),
    ])
    with pytest.raises(bridge.FullLoadFailed) as excinfo:
        bridge._full_load(pres)
    exc = excinfo.value
    assert exc.slide_index == 2
    assert exc.slide_id == 257
    assert exc.shape_index == 2
    assert exc.shape_name == "Broken 3"
    assert "slide 2" in str(exc) and "Broken 3" in str(exc)


def test_opens_clean_surfaces_the_failing_slide_and_shape(monkeypatch):
    """P6 layer 2: and the identity reaches the validate tool's result."""
    from kitchensink4ppt.com import bridge

    exc = bridge.FullLoadFailed(
        "slide 4 (id 260) shape 7 'Agenda Body': boom",
        slide_index=4, slide_id=260, shape_index=7, shape_name="Agenda Body",
    )
    out = bridge._opens_clean_failure(exc)
    assert out["opens_clean"] is False
    assert out["failed_slide_index"] == 4
    assert out["failed_slide_id"] == 260
    assert out["failed_shape_index"] == 7
    assert out["failed_shape_name"] == "Agenda Body"
    assert "boom" in out["error"]


def test_agenda_body_with_no_entries_still_has_a_paragraph(tmp_path):
    """P6 producer 2: the agenda body emptied itself and appended one
    paragraph per entry, so an empty entry list left zero a:p. Both public
    callers refuse an empty list today; the body builder is the layer that
    must not be able to produce the corruption at all."""
    from kitchensink4ppt.ops import assembly as asm

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1])
    path = tmp_path / "agenda.pptx"
    prs.save(str(path))
    pkg = PptxPackage(str(path))
    body_elem = asm._body_shape(pkg, pkg.slide_parts()[0], tagged=False)
    asm._fill_agenda_body(pkg, 0, body_elem, [], False, 0, [])
    body = body_elem.find(qn("p:txBody"))
    assert body.findall(qn("a:p")), "the agenda body has no paragraphs"
    pkg.save(str(tmp_path / "agenda_out.pptx"))


def test_notes_rebuilt_from_nothing_still_has_a_paragraph(tmp_path):
    """P6 producer 3: ops/notes.py _rebuild_body with an empty spec list."""
    from kitchensink4ppt.ops import notes as nt

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "notes.pptx"
    prs.save(str(path))
    pkg = PptxPackage(str(path))
    nt.set_notes(pkg, 0, "first draft of the note")
    nt.set_notes(pkg, 0, paragraphs=[])
    out = tmp_path / "notes_out.pptx"
    pkg.save(str(out))  # layer 1 refuses a body with no paragraphs
    assert out.is_file()


# ------------------------------------------------------------------ P3


def _copy_presentation(**kwargs):
    from kitchensink4ppt import packs, server  # noqa: F401  (registers tools)

    return packs.tool_objects()["copy_presentation"].fn(**kwargs)


def _deck_at(path):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(path))
    return path


def test_copy_presentation_happy_path(tmp_path):
    """P3: copy_presentation had no behavioural test at all."""
    src = _deck_at(tmp_path / "src.pptx")
    dest = tmp_path / "copy.pptx"
    out = _copy_presentation(file_path=str(src), dest_path=str(dest))
    assert out["ok"] is True
    assert out["overwrote_existing"] is False
    assert dest.is_file()
    assert dest.read_bytes() == src.read_bytes()


def test_copy_presentation_refuses_an_existing_destination(tmp_path):
    src = _deck_at(tmp_path / "src.pptx")
    dest = _deck_at(tmp_path / "dest.pptx")
    out = _copy_presentation(file_path=str(src), dest_path=str(dest))
    assert out["ok"] is False
    assert out["error"]["code"] == "CONFLICT"
    assert "overwrite=True" in out["error"]["message"]


def test_copy_presentation_overwrite_rotates_the_previous_content(tmp_path):
    src = _deck_at(tmp_path / "src.pptx")
    dest = tmp_path / "dest.pptx"
    _deck_at(dest)
    before = dest.read_bytes()
    out = _copy_presentation(
        file_path=str(src), dest_path=str(dest), overwrite=True
    )
    assert out["ok"] is True
    assert out["overwrote_existing"] is True
    assert dest.read_bytes() == src.read_bytes()
    from kitchensink4ppt.core import safesave

    prev = safesave.slot_dir(dest) / safesave.PREV_SLOT
    anchor = safesave.slot_dir(dest) / safesave.ANCHOR_SLOT
    kept = [p for p in (prev, anchor) if p.is_file()]
    assert kept, "the overwritten content was not kept in any backup slot"
    assert any(p.read_bytes() == before for p in kept)


def test_copy_presentation_missing_parent_says_so(tmp_path):
    """P3: the refusal said '[WinError 3] The system cannot find the path
    specified' and hinted at re-reading the deck. Neither is true: the
    deck is fine, the destination folder does not exist."""
    src = _deck_at(tmp_path / "src.pptx")
    dest = tmp_path / "no_such_folder" / "copy.pptx"
    out = _copy_presentation(file_path=str(src), dest_path=str(dest))
    assert out["ok"] is False
    error = out["error"]
    assert "destination folder does not exist" in error["message"]
    assert str(dest.parent) in error["message"]
    assert "get_presentation_view" not in error["hint"], (
        f"the misleading hint is still attached: {error['hint']!r}"
    )
    assert error["hint"], "the refusal gives the caller nothing to do"
    assert not dest.parent.exists(), "a typo in a path created a folder"


def test_a_declared_hint_overrides_the_blanket_one():
    """P3 mechanism: the per-code _HINTS entry is a fallback, not a
    verdict, so a refusal that knows better can say so."""
    from kitchensink4ppt import server
    from kitchensink4ppt.core.errors import TargetNotFound

    plain = TargetNotFound("no shape 9 on slide 1")
    assert "get_presentation_view" in server._refusal(plain)["error"]["hint"]

    exc = TargetNotFound("no shape 9 on slide 1")
    exc.hint = "do this other thing instead"
    assert (
        server._refusal(exc)["error"]["hint"] == "do this other thing instead"
    )


# ------------------------------------------------------------------ P8


def test_create_table_accepts_row_heights_and_col_widths(tmp_path):
    """P8: a 27-row reference table at 10 pt needed per-row heights, and
    the only route was a second call after the fact."""
    from kitchensink4ppt.ops import tables as tb
    from kitchensink4ppt.ops import geometry as g

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "t.pptx"
    prs.save(str(path))
    pkg = PptxPackage(str(path))
    tb.create_table(
        pkg, 0, 3, 2, 1, 1, 6, 3,
        row_heights=[0.5, 1.0, 1.5], col_widths=[2.0, 4.0],
    )
    part = get_slide_info(pkg, 0)["part"]
    tbl = pkg.root(part).find(f".//{qn('a:tbl')}")
    widths = [int(c.get("w")) for c in tbl.iter(qn("a:gridCol"))]
    heights = [int(r.get("h")) for r in tbl.findall(qn("a:tr"))]
    assert widths == [g.in_to_emu(2.0), g.in_to_emu(4.0)]
    assert heights == [g.in_to_emu(h) for h in (0.5, 1.0, 1.5)]


def test_create_table_accepts_an_index_dict(tmp_path):
    from kitchensink4ppt.ops import tables as tb
    from kitchensink4ppt.ops import geometry as g

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "t.pptx"
    prs.save(str(path))
    pkg = PptxPackage(str(path))
    tb.create_table(pkg, 0, 3, 2, 1, 1, 6, 3, row_heights={1: 2.0})
    part = get_slide_info(pkg, 0)["part"]
    tbl = pkg.root(part).find(f".//{qn('a:tbl')}")
    heights = [int(r.get("h")) for r in tbl.findall(qn("a:tr"))]
    assert heights[1] == g.in_to_emu(2.0)
    # rows the dict did not name split what is left of the box, so the
    # table still fits the h it was given
    assert heights[0] == heights[2]
    assert sum(heights) == g.in_to_emu(3)


def test_create_table_refuses_a_wrong_length_list(tmp_path):
    from kitchensink4ppt.core.errors import PptMcpError
    from kitchensink4ppt.ops import tables as tb

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "t.pptx"
    prs.save(str(path))
    pkg = PptxPackage(str(path))
    with pytest.raises(PptMcpError) as excinfo:
        tb.create_table(pkg, 0, 3, 2, 1, 1, 6, 3, row_heights=[1.0, 1.0])
    assert "3 row" in str(excinfo.value)


def test_create_table_refuses_sizes_that_overflow_the_box(tmp_path):
    """The refusal case is a PARTIAL dict that leaves the unstated columns
    nothing. A COMPLETE list is the caller stating the whole grid, so it
    resizes the frame instead of refusing (review finding M5)."""
    from kitchensink4ppt.core.errors import PptMcpError
    from kitchensink4ppt.ops import tables as tb

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "t.pptx"
    prs.save(str(path))
    pkg = PptxPackage(str(path))
    with pytest.raises(PptMcpError) as excinfo:
        tb.create_table(pkg, 0, 2, 3, 1, 1, 6, 3, col_widths={0: 7.0})
    message = str(excinfo.value)
    assert "col_widths" in message and "unstated" in message


def test_create_table_description_stays_in_budget():
    from kitchensink4ppt import packs, server  # noqa: F401  (registers tools)

    tool = packs.tool_objects()["create_table"]
    desc = tool.description or ""
    assert "row_heights" in desc and "col_widths" in desc
    assert round(len(desc) / 4) <= 130, (
        f"create_table description is ~{round(len(desc) / 4)} tokens"
    )


# ----------------------------------------------------------------- #887


_BODY = (
    "If the new tools are not in your tool "
    "list, this client fixed its list when the session or worker started: "
    "do not retry here. What works in every client: ask the user to add "
    "the packs to KS4P_MODE (comma list) in this server's launch "
    "settings, restart the app or session, then start a new worker if "
    "needed. Claude Code only: the orchestrator can instead call "
    "enable_tools in the main session and then start a new worker. If "
    "enable_tools refuses a pack, an administrator locked the tool set: "
    "do not retry."
)

_NOTE = "tools/list_changed was sent. " + _BODY

_NOOP_NOTE = (
    "These packs were already on, so no list change was sent. " + _BODY
)

_WORKER_SENTENCE = (
    "Workers and subagents only see the tools that were on when they "
    "started: start the server with KS4P_MODE set to a comma list of "
    "packs, or, in Claude Code, enable packs in the main session before "
    "starting workers."
)


@pytest.fixture()
def pristine_packs():
    """Restore the process-wide enabled set after a test flips it."""
    from kitchensink4ppt import packs, server  # noqa: F401

    before = dict(packs._ENABLED)
    yield packs
    packs._ENABLED.clear()
    packs._ENABLED.update(before)


def test_enable_tools_note_is_the_887_wording(pristine_packs):
    """#887: the old note promised a tool-list refresh that a worker or
    subagent never gets, so the worker retried the same dead call."""
    packs = pristine_packs
    out = packs.enable(["graphics"])
    assert out["enabled"] == ["graphics"]
    assert out["note"] == _NOTE


def test_the_note_is_emitted_on_a_no_op_re_enable(pristine_packs):
    """#887: the caller who cannot see the tools is exactly the caller
    whose second attempt enables nothing new, and that was the call the
    old note stayed silent on."""
    packs = pristine_packs
    packs.enable(["graphics"])
    again = packs.enable(["graphics"])
    assert again["enabled"] == []
    assert again["already_enabled"] == ["graphics"]
    # N1: _sync fires the visibility hook only when a tool actually
    # flipped, so claiming a notification here was a plain falsehood.
    assert again["note"] == _NOOP_NOTE
    assert "tools/list_changed was sent" not in again["note"]
    # the guidance itself is identical either way
    assert _BODY in again["note"]


def test_the_worker_sentence_is_in_the_instructions_and_the_index():
    """#887: said once where a client actually reads the surface."""
    from kitchensink4ppt import server
    from kitchensink4ppt.ops import workflows as wf

    assert _WORKER_SENTENCE in (server.mcp.instructions or "")
    assert _WORKER_SENTENCE in wf.get_workflows()["note"]


def test_the_locked_refusal_names_the_administrator(monkeypatch):
    """#887 (c): a refusal that does not say a human set the lock reads
    like a bug to retry."""
    from kitchensink4ppt import packs
    from kitchensink4ppt.core.errors import PptMcpError

    monkeypatch.setenv(packs.ENV_PACK_POLICY, "locked")
    with pytest.raises(PptMcpError) as excinfo:
        packs.enable(["graphics"])
    assert excinfo.value.code == "CONFLICT"
    assert (
        "An administrator locked the tool packs for this install."
        in str(excinfo.value)
    )


def test_a_pack_policy_typo_does_not_fail_open(monkeypatch):
    """The one env whose whole purpose is a lock used to shrug off a
    misspelling: KS4P_PACK_POLICY=lockedd served UNLOCKED while its
    sibling KS4P_MODE=fulll refused to start."""
    from kitchensink4ppt import packs
    from kitchensink4ppt.core.errors import PptMcpError

    monkeypatch.setenv(packs.ENV_PACK_POLICY, "lockedd")
    with pytest.raises(PptMcpError) as excinfo:
        packs.resolve_lock()
    message = str(excinfo.value)
    assert "KS4P_PACK_POLICY" in message and "lockedd" in message
    with pytest.raises(PptMcpError):
        packs.apply_startup_mode()


def test_the_accepted_pack_policies_still_work(monkeypatch):
    from kitchensink4ppt import packs

    monkeypatch.setenv(packs.ENV_PACK_POLICY, "locked")
    assert packs.resolve_lock() is True
    monkeypatch.setenv(packs.ENV_PACK_POLICY, "auto")
    assert packs.resolve_lock() is False
    monkeypatch.setenv(packs.ENV_PACK_POLICY, "")
    assert packs.resolve_lock() is False


# ============================================================ ROUND 2
# Findings from the second review of PR #32 (2026-09-22). The test bodies
# for B1 and B2 are the reviewer's, adopted as written where they fit.


def _bare_body():
    body = etree.Element(qn("p:txBody"))
    etree.SubElement(body, qn("a:bodyPr"))
    etree.SubElement(body, qn("a:lstStyle"))
    return body


def _spec(text, level=0, explicit=False):
    return {"text": text, "level": level, "level_explicit": explicit}


def test_placeholder_retext_uses_field_character_properties():
    """B1: a placeholder holding nothing but a slide-number field lost its
    size, colour and East Asian font the moment the text was replaced, and
    reported no loss. a:fld carries a real a:rPr."""
    from kitchensink4ppt.ops.text import _replace_body_paragraphs

    body = _bare_body()
    p = etree.SubElement(body, qn("a:p"))
    fld = etree.SubElement(p, qn("a:fld"))
    rpr = etree.SubElement(fld, qn("a:rPr"))
    rpr.set("sz", "2400")
    etree.SubElement(rpr, qn("a:ea")).set("typeface", "Malgun Gothic")
    etree.SubElement(fld, qn("a:t")).text = "1"

    _replace_body_paragraphs(body, [_spec("Literal")])
    new = body.find(f"{qn('a:p')}/{qn('a:r')}/{qn('a:rPr')}")
    assert new is not None and new.get("sz") == "2400"
    assert new.find(qn("a:ea")).get("typeface") == "Malgun Gothic"


def test_placeholder_retext_uses_break_character_properties():
    """B1: the same for a paragraph whose only property carrier is a:br."""
    from kitchensink4ppt.ops.text import _replace_body_paragraphs

    body = _bare_body()
    p = etree.SubElement(body, qn("a:p"))
    br = etree.SubElement(p, qn("a:br"))
    rpr = etree.SubElement(br, qn("a:rPr"))
    rpr.set("sz", "1600")
    etree.SubElement(rpr, qn("a:latin")).set("typeface", "Consolas")

    _replace_body_paragraphs(body, [_spec("Literal")])
    new = body.find(f"{qn('a:p')}/{qn('a:r')}/{qn('a:rPr')}")
    assert new is not None and new.get("sz") == "1600"
    assert new.find(qn("a:latin")).get("typeface") == "Consolas"


def test_a_field_in_the_replacement_gets_the_properties_too():
    """B1: a:fld and a:br are landing spots as well as sources."""
    from kitchensink4ppt.ops import shapes as _shp

    old = _bare_body()
    p = etree.SubElement(old, qn("a:p"))
    r = etree.SubElement(p, qn("a:r"))
    etree.SubElement(r, qn("a:rPr")).set("sz", "2800")
    etree.SubElement(r, qn("a:t")).text = "old"

    new = _bare_body()
    np_ = etree.SubElement(new, qn("a:p"))
    fld = etree.SubElement(np_, qn("a:fld"))
    etree.SubElement(fld, qn("a:rPr"))
    etree.SubElement(fld, qn("a:t")).text = "3"

    _shp._carry_text_properties(old, new, set())
    assert new.find(
        f"{qn('a:p')}/{qn('a:fld')}/{qn('a:rPr')}"
    ).get("sz") == "2800"


def test_empty_replacement_keeps_the_end_paragraph_style():
    """M2: the paragraph-ending insertion style is what an empty paragraph
    renders at. Taking the first visible run's style instead buried a
    24 pt insertion point under 10 pt body text and called it preserved."""
    from kitchensink4ppt.ops.text import _replace_body_paragraphs

    body = _bare_body()
    p = etree.SubElement(body, qn("a:p"))
    r = etree.SubElement(p, qn("a:r"))
    etree.SubElement(r, qn("a:rPr")).set("sz", "1000")
    etree.SubElement(r, qn("a:t")).text = "visible"
    etree.SubElement(p, qn("a:endParaRPr")).set("sz", "2400")

    _replace_body_paragraphs(body, [_spec("")])
    end = body.find(f"{qn('a:p')}/{qn('a:endParaRPr')}")
    assert end is not None, "the end-paragraph style is gone"
    assert end.get("sz") == "2400", (
        f"empty replacement took the run style, not endParaRPr: "
        f"sz={end.get('sz')}"
    )


def test_a_visible_run_still_takes_the_run_style_not_end_para():
    """M2 guard: the other direction must not flip."""
    from kitchensink4ppt.ops.text import _replace_body_paragraphs

    body = _bare_body()
    p = etree.SubElement(body, qn("a:p"))
    r = etree.SubElement(p, qn("a:r"))
    etree.SubElement(r, qn("a:rPr")).set("sz", "1000")
    etree.SubElement(r, qn("a:t")).text = "visible"
    etree.SubElement(p, qn("a:endParaRPr")).set("sz", "2400")

    _replace_body_paragraphs(body, [_spec("new words")])
    assert body.find(
        f"{qn('a:p')}/{qn('a:r')}/{qn('a:rPr')}"
    ).get("sz") == "1000"


def _two_run_para(body, second_rpr_edit):
    p = etree.SubElement(body, qn("a:p"))
    for i, edit in enumerate((None, second_rpr_edit)):
        r = etree.SubElement(p, qn("a:r"))
        rpr = etree.SubElement(r, qn("a:rPr"))
        rpr.set("sz", "1800")
        etree.SubElement(rpr, qn("a:ea")).set("typeface", "Yu Gothic")
        if edit is not None:
            edit(rpr)
        etree.SubElement(r, qn("a:t")).text = f"run{i}"
    return p


def _collapse_probe(old):
    from kitchensink4ppt.ops import shapes as _shp

    new = _bare_body()
    np_ = etree.SubElement(new, qn("a:p"))
    r = etree.SubElement(np_, qn("a:r"))
    etree.SubElement(r, qn("a:rPr"))
    etree.SubElement(r, qn("a:t")).text = "merged"
    carried, facts = _shp._carry_text_properties(old, new, set())
    return new, carried, facts


def test_east_asian_difference_is_reported_as_a_collapse():
    """M1: two runs differing only in East Asian typeface compared EQUAL
    under the old hand-picked signature, so the collapse onto the first
    run happened with nothing said."""
    old = _bare_body()
    _two_run_para(
        old, lambda rpr: rpr.find(qn("a:ea")).set("typeface", "Malgun Gothic")
    )
    _new, _carried, facts = _collapse_probe(old)
    assert facts.get("runs_collapsed_to_first") is True, (
        "an East Asian typeface difference was collapsed in silence"
    )


def test_rtl_difference_is_reported_as_a_collapse():
    """M1: same for a:rtl, which the old signature never looked at."""
    old = _bare_body()
    _two_run_para(
        old, lambda rpr: etree.SubElement(rpr, qn("a:rtl")).set("val", "1")
    )
    _new, _carried, facts = _collapse_probe(old)
    assert facts.get("runs_collapsed_to_first") is True


def test_identical_runs_are_not_reported_as_a_collapse():
    """M1 guard: the widened signature must not cry wolf on runs that
    really are alike, including a bare a:br between them."""
    old = _bare_body()
    p = _two_run_para(old, None)
    p.insert(1, etree.Element(qn("a:br")))
    _new, _carried, facts = _collapse_probe(old)
    assert "runs_collapsed_to_first" not in facts


def test_disagreeing_hyperlinks_are_dropped_not_spread():
    """M1: a paragraph whose first run links to A and whose second links
    nowhere became one replacement run entirely linked to A, silently."""
    from kitchensink4ppt.core.package import NSMAP

    old = _bare_body()
    p = etree.SubElement(old, qn("a:p"))
    for i in range(2):
        r = etree.SubElement(p, qn("a:r"))
        rpr = etree.SubElement(r, qn("a:rPr"))
        rpr.set("sz", "1800")
        if i == 0:
            link = etree.SubElement(rpr, qn("a:hlinkClick"))
            link.set(f"{{{NSMAP['r']}}}id", "rId7")
        etree.SubElement(r, qn("a:t")).text = f"run{i}"

    new, _carried, facts = _collapse_probe(old)
    merged = new.find(f"{qn('a:p')}/{qn('a:r')}/{qn('a:rPr')}")
    assert merged.find(qn("a:hlinkClick")) is None, (
        "the first run's hyperlink was spread over the whole replacement"
    )
    assert "hyperlinks_dropped" in facts
    # what every run DID agree on still rides through
    assert merged.get("sz") == "1800"


def test_one_shared_hyperlink_still_rides_through():
    """M1 guard: agreement is not a loss. A link every run shared stays."""
    from kitchensink4ppt.core.package import NSMAP

    old = _bare_body()
    p = etree.SubElement(old, qn("a:p"))
    for i in range(2):
        r = etree.SubElement(p, qn("a:r"))
        rpr = etree.SubElement(r, qn("a:rPr"))
        etree.SubElement(rpr, qn("a:hlinkClick")).set(
            f"{{{NSMAP['r']}}}id", "rId7"
        )
        etree.SubElement(r, qn("a:t")).text = f"run{i}"

    new, _carried, facts = _collapse_probe(old)
    merged = new.find(f"{qn('a:p')}/{qn('a:r')}/{qn('a:rPr')}")
    assert merged.find(qn("a:hlinkClick")) is not None
    assert "hyperlinks_dropped" not in facts


def test_a_carried_level_survives_a_later_explicit_one(field_deck):
    """N2: `carried` is body-global, so a later paragraph whose level the
    caller stated used to erase the accurate report from an earlier one."""
    from kitchensink4ppt.ops import text as txt

    pkg, slide, _sid, ph_id = field_deck
    out = txt.set_placeholder_text(
        pkg, slide, "body",
        paragraphs=[
            {"text": "keeps its carried level"},
            {"text": "restyled", "level": 1},
            {"text": "states its own", "level": 4},
        ],
    )
    part = get_slide_info(pkg, slide)["part"]
    elem, _chain = shp._find_shape(pkg, part, ph_id)
    paras = elem.find(qn("p:txBody")).findall(qn("a:p"))
    assert paras[2].find(qn("a:pPr")).get("lvl") == "4"
    assert "level" in out.get("preserved", []), (
        "paragraph 1 carried a level and the report lost it"
    )


# --------------------------------------------------- B2 and M6 (validator)

_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
_C = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def _alternate_content(choice_para, fallback_para):
    """An mc:AlternateContent wrapper with the given paragraphs (None for
    a branch that supplies none)."""
    ac = etree.Element(
        f"{{{_MC}}}AlternateContent", nsmap={"mc": _MC, "p14": _P14}
    )
    choice = etree.SubElement(ac, f"{{{_MC}}}Choice")
    choice.set("Requires", "p14")
    if choice_para is not None:
        choice.append(choice_para)
    fallback = etree.SubElement(ac, f"{{{_MC}}}Fallback")
    if fallback_para is not None:
        fallback.append(fallback_para)
    return ac


def _first_body(pkg):
    part = pkg.slide_parts()[0]
    body = pkg.root(part).find(f".//{qn('p:txBody')}")
    assert body is not None, "no text body on the first slide"
    return part, body


def test_txbody_accepts_mce_wrapped_paragraph(make_deck, tmp_path):
    """B2: Markup Compatibility lets a body's paragraphs arrive through an
    mc:AlternateContent branch. Requiring a DIRECT a:p refused a legal
    deck at save, which is the worst thing this gate can do."""
    import copy

    pkg = PptxPackage(make_deck("mce.pptx"))
    part, body = _first_body(pkg)
    para = body.find(qn("a:p"))
    body.replace(
        para,
        _alternate_content(copy.deepcopy(para), copy.deepcopy(para)),
    )
    pkg.mark_dirty(part)
    out = tmp_path / "out.pptx"
    pkg.save(str(out))  # used to raise ValidationFailed
    assert out.is_file()


def test_an_empty_selectable_branch_is_still_refused(make_deck, tmp_path):
    """B2: a valid Fallback must not hide an empty Choice. PowerPoint
    picks ONE branch, so every branch has to supply a paragraph."""
    import copy

    from kitchensink4ppt.core.errors import ValidationFailed

    pkg = PptxPackage(make_deck("mce_bad.pptx"))
    part, body = _first_body(pkg)
    para = body.find(qn("a:p"))
    body.replace(para, _alternate_content(None, copy.deepcopy(para)))
    pkg.mark_dirty(part)
    with pytest.raises(ValidationFailed) as excinfo:
        pkg.save(str(tmp_path / "out.pptx"))
    assert "AlternateContent" in str(excinfo.value)


def test_a_direct_paragraph_beside_alternate_content_is_enough(
    make_deck, tmp_path
):
    """B2: a body that keeps a real paragraph AND carries an
    AlternateContent for something else is never empty."""
    import copy

    pkg = PptxPackage(make_deck("mce_mixed.pptx"))
    part, body = _first_body(pkg)
    para = body.find(qn("a:p"))
    body.append(_alternate_content(None, copy.deepcopy(para)))
    pkg.mark_dirty(part)
    out = tmp_path / "out.pptx"
    pkg.save(str(out))
    assert out.is_file()


def test_nested_alternate_content_branches_resolve(make_deck, tmp_path):
    """B2: a branch may itself hold an AlternateContent."""
    import copy

    pkg = PptxPackage(make_deck("mce_nested.pptx"))
    part, body = _first_body(pkg)
    para = body.find(qn("a:p"))
    inner = _alternate_content(copy.deepcopy(para), copy.deepcopy(para))
    outer = etree.Element(
        f"{{{_MC}}}AlternateContent", nsmap={"mc": _MC, "p14": _P14}
    )
    choice = etree.SubElement(outer, f"{{{_MC}}}Choice")
    choice.set("Requires", "p14")
    choice.append(inner)
    fallback = etree.SubElement(outer, f"{{{_MC}}}Fallback")
    fallback.append(copy.deepcopy(para))
    body.replace(para, outer)
    pkg.mark_dirty(part)
    out = tmp_path / "out.pptx"
    pkg.save(str(out))
    assert out.is_file()


def test_chart_text_bodies_are_covered(tmp_path):
    """M6: c:rich and c:txPr use the same paragraph-bearing model. Leaving
    them out made the 'every text body' claim false."""
    from kitchensink4ppt.core.errors import ValidationFailed
    from kitchensink4ppt.ops import charts as ch

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "chart.pptx"
    prs.save(str(path))
    pkg = PptxPackage(str(path))
    ch.create_chart(
        pkg, 0, "bar", ["a", "b"], [{"name": "s", "values": [1, 2]}],
        1, 1, 6, 4, title="Numbers",
    )
    pkg.save(str(tmp_path / "chart_ok.pptx"))  # a real chart still saves

    chart_part = next(
        p for p in pkg.part_names() if p.startswith("ppt/charts/chart")
    )
    rich = pkg.root(chart_part).find(f".//{{{_C}}}rich")
    assert rich is not None, "the chart title has no c:rich to empty"
    for p in list(rich.findall(qn("a:p"))):
        rich.remove(p)
    pkg.mark_dirty(chart_part)
    with pytest.raises(ValidationFailed) as excinfo:
        pkg.save(str(tmp_path / "chart_bad.pptx"))
    assert "charts" in str(excinfo.value)


def test_a_real_chart_round_trips(tmp_path):
    """M6 guard: covering charts must not refuse a chart the server built."""
    from kitchensink4ppt.ops import charts as ch

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "chart2.pptx"
    prs.save(str(path))
    pkg = PptxPackage(str(path))
    for kind in ("bar", "line", "pie"):
        ch.create_chart(
            pkg, 0, kind, ["a", "b"], [{"name": kind, "values": [1, 2]}],
            1, 1, 4, 3, title=f"{kind} title",
        )
    out = tmp_path / "charts_ok.pptx"
    pkg.save(str(out))
    assert out.is_file()


# ------------------------------------------------- M3, M4, N3 (the COM walk)


class _FakeFrame:
    def __init__(self, text="words", raises=False):
        self._text = text
        self._raises = raises
        self.HasText = -1

    @property
    def TextRange(self):
        if self._raises:
            raise OSError("the text refused to load")
        return self

    @property
    def Text(self):
        return self._text


class _FakeShape:
    """A PowerPoint shape proxy with only the members a real one would
    answer; anything else raises AttributeError the way getattr sees a
    missing COM member."""

    def __init__(self, name, *, text=None, raises=False, group=None,
                 table=None, fault=None, has_text_frame=-1):
        self.Name = name
        self._fault = fault
        self._has_text_frame = has_text_frame
        self._frame = (
            _FakeFrame(text, raises) if text is not None or raises else None
        )
        self._group = group
        self._table = table

    @property
    def HasTextFrame(self):
        if self._fault is not None:
            raise self._fault
        return self._has_text_frame if self._frame is not None else 0

    @property
    def TextFrame(self):
        return self._frame

    @property
    def HasTable(self):
        return -1 if self._table is not None else 0

    @property
    def Table(self):
        return self._table

    @property
    def Type(self):
        return 6 if self._group is not None else 1

    @property
    def GroupItems(self):
        return _FakeShapes(self._group)


class _FakeShapes:
    def __init__(self, shapes):
        self._shapes = shapes
        self.Count = len(shapes)

    def Item(self, i):
        return self._shapes[i - 1]


class _FakeCell:
    def __init__(self, shape):
        self.Shape = shape


class _FakeTable:
    def __init__(self, grid):
        self._grid = grid
        self.Rows = _FakeShapes([None] * len(grid))
        self.Columns = _FakeShapes([None] * len(grid[0]))

    def Cell(self, r, c):
        return _FakeCell(self._grid[r - 1][c - 1])


class _FakeSlide:
    def __init__(self, sid, shapes):
        self.SlideID = sid
        self.Shapes = _FakeShapes(shapes)


class _FakePres:
    def __init__(self, slides):
        self.Slides = _FakeShapes(slides)


def test_every_text_shape_is_read_not_just_the_first():
    """M3: after the first successful text read the walk stopped reading,
    so a second shape that raises on access was never touched and the deck
    came back clean."""
    from kitchensink4ppt.com import bridge

    pres = _FakePres([
        _FakeSlide(256, [
            _FakeShape("Good 1", text="fine"),
            _FakeShape("Broken 2", raises=True),
        ]),
    ])
    with pytest.raises(bridge.FullLoadFailed) as excinfo:
        bridge._full_load(pres)
    exc = excinfo.value
    assert exc.slide_index == 1 and exc.shape_index == 2
    assert exc.shape_name == "Broken 2"


def test_a_failure_inside_a_group_is_found_and_named():
    """M3: group children were never enumerated at all."""
    from kitchensink4ppt.com import bridge

    inner = _FakeShape("Deep Label", raises=True)
    group = _FakeShape("Diagram", group=[_FakeShape("Ok", text="a"), inner])
    pres = _FakePres([_FakeSlide(300, [_FakeShape("Title", text="t"), group])])
    with pytest.raises(bridge.FullLoadFailed) as excinfo:
        bridge._full_load(pres)
    message = str(excinfo.value)
    assert "Deep Label" in message and "Diagram" in message
    assert excinfo.value.shape_index == 2


def test_a_failure_inside_a_table_cell_is_found_and_named():
    """M3: table cells were never read."""
    from kitchensink4ppt.com import bridge

    grid = [
        [_FakeShape("c00", text="a"), _FakeShape("c01", text="b")],
        [_FakeShape("c10", text="c"), _FakeShape("c11", raises=True)],
    ]
    table = _FakeShape("Grid", table=_FakeTable(grid))
    pres = _FakePres([_FakeSlide(301, [table])])
    with pytest.raises(bridge.FullLoadFailed) as excinfo:
        bridge._full_load(pres)
    message = str(excinfo.value)
    assert "row 2" in message and "column 2" in message


def test_a_clean_walk_reports_what_it_touched():
    from kitchensink4ppt.com import bridge

    group = _FakeShape("G", group=[_FakeShape("g1", text="x")])
    pres = _FakePres([
        _FakeSlide(256, [_FakeShape("A", text="a"), group]),
        _FakeSlide(257, [_FakeShape("B", text="b")]),
    ])
    out = bridge._full_load(pres)
    assert out["slides"] == 2
    assert out["shapes"] == 3          # top level only, as before
    assert out["shapes_walked"] == 4   # including the group member
    assert out["text_reads"] == 3      # the group frame itself has none
    assert "walk_truncated" not in out


def test_an_enormous_deck_says_the_walk_was_truncated(monkeypatch):
    """M3: a cap, and an honest flag when it bites, beats sitting inside
    the bounded operation until it times out and reports nothing."""
    from kitchensink4ppt.com import bridge

    monkeypatch.setattr(bridge, "FULL_LOAD_MAX_SHAPES", 3)
    pres = _FakePres([
        _FakeSlide(256, [_FakeShape(f"S{i}", text="t") for i in range(10)]),
    ])
    out = bridge._full_load(pres)
    assert "walk_truncated" in out
    assert out["shapes_walked"] == 3
    assert "cap 3" in out["walk_truncated"]


def _com_error(hresult):
    """A stand-in for a pywin32 com_error carrying one HRESULT."""

    class _ComError(Exception):
        pass

    exc = _ComError("com fault")
    exc.hresult = hresult
    return exc


def test_a_busy_powerpoint_mid_walk_is_not_a_corrupt_file():
    """M4: a dialog opening halfway through the walk came back as an
    authoritative corruption verdict naming an innocent shape."""
    from kitchensink4ppt.com import bridge
    from kitchensink4ppt.core.errors import PowerPointBusy

    pres = _FakePres([
        _FakeSlide(256, [
            _FakeShape("Fine", text="a"),
            _FakeShape("Innocent", fault=_com_error(
                bridge.RPC_E_CALL_REJECTED)),
        ]),
    ])
    with pytest.raises(PowerPointBusy):
        bridge._full_load(pres)


def test_a_disconnected_powerpoint_mid_walk_is_not_a_corrupt_file():
    """M4: the same for a PowerPoint that went away under the proxy."""
    from kitchensink4ppt.com import bridge
    from kitchensink4ppt.core.errors import PowerPointDisconnected

    pres = _FakePres([
        _FakeSlide(256, [
            _FakeShape("Innocent", fault=_com_error(
                bridge.RPC_E_DISCONNECTED)),
        ]),
    ])
    with pytest.raises(PowerPointDisconnected):
        bridge._full_load(pres)


def test_a_busy_fault_inside_a_group_still_blames_the_application():
    """M4: the classifier has to be consulted at every depth."""
    from kitchensink4ppt.com import bridge
    from kitchensink4ppt.core.errors import PowerPointBusy

    inner = _FakeShape("Deep", fault=_com_error(
        bridge.RPC_E_SERVERCALL_RETRYLATER))
    group = _FakeShape("G", group=[inner])
    pres = _FakePres([_FakeSlide(256, [group])])
    with pytest.raises(PowerPointBusy):
        bridge._full_load(pres)


def test_an_unclassified_fault_is_still_a_file_verdict():
    """M4 guard: classification must not swallow real content failures."""
    from kitchensink4ppt.com import bridge

    pres = _FakePres([
        _FakeSlide(256, [_FakeShape("Broken", fault=OSError("bad record"))]),
    ])
    with pytest.raises(bridge.FullLoadFailed):
        bridge._full_load(pres)


def test_mixed_tristate_is_not_treated_as_yes():
    """N3: `if shp.HasTextFrame` counted msoTriStateMixed (-2) as true, so
    an unusual shape was probed for a TextFrame it does not support and
    became a false failure."""
    from kitchensink4ppt.com import bridge

    assert bridge._mso_true(-1) is True
    assert bridge._mso_true(-2) is False
    assert bridge._mso_true(0) is False
    assert bridge._mso_true(None) is False

    shape = _FakeShape("Odd", text="never read", has_text_frame=-2)
    pres = _FakePres([_FakeSlide(256, [shape])])
    out = bridge._full_load(pres)
    assert out["text_reads"] == 0, "a mixed tri-state was probed as a yes"


# ------------------------------------------------------- M5 (table geometry)


def _table_pkg(tmp_path, name="m5.pptx"):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / name
    prs.save(str(path))
    return PptxPackage(str(path))


def _geometry(pkg):
    from kitchensink4ppt.ops import geometry as g

    part = get_slide_info(pkg, 0)["part"]
    frame = pkg.root(part).find(f".//{qn('p:graphicFrame')}")
    ext = frame.find(f"{qn('p:xfrm')}/{qn('a:ext')}")
    tbl = frame.find(f".//{qn('a:tbl')}")
    widths = [int(c.get("w")) for c in tbl.iter(qn("a:gridCol"))]
    heights = [int(r.get("h")) for r in tbl.findall(qn("a:tr"))]
    return {
        "frame_cx": int(ext.get("cx")),
        "frame_cy": int(ext.get("cy")),
        "widths": widths,
        "heights": heights,
        "in": g.emu_to_in,
    }


def test_a_complete_underfilling_list_resizes_the_frame(tmp_path):
    """M5: a 6-inch frame around a grid declaring 2 inches is geometry
    PowerPoint has to normalize away, so the widths that rendered were not
    the widths the caller asked for."""
    from kitchensink4ppt.ops import tables as tb
    from kitchensink4ppt.ops import geometry as g

    pkg = _table_pkg(tmp_path)
    out = tb.create_table(
        pkg, 0, 2, 2, 1, 1, 6, 3, col_widths=[1.0, 1.0]
    )
    geo = _geometry(pkg)
    assert geo["widths"] == [g.in_to_emu(1.0), g.in_to_emu(1.0)]
    assert geo["frame_cx"] == sum(geo["widths"]), (
        "the frame still claims a width the grid does not fill"
    )
    assert out["w_in"] == 2.0
    assert "box_resized" in out and "width" in out["box_resized"]


def test_a_complete_overfilling_list_resizes_the_frame_too(tmp_path):
    """M5: the same rule in the other direction, instead of a refusal."""
    from kitchensink4ppt.ops import tables as tb
    from kitchensink4ppt.ops import geometry as g

    pkg = _table_pkg(tmp_path)
    out = tb.create_table(
        pkg, 0, 3, 2, 1, 1, 6, 3, row_heights=[2.0, 2.0, 2.0]
    )
    geo = _geometry(pkg)
    assert geo["heights"] == [g.in_to_emu(2.0)] * 3
    assert geo["frame_cy"] == sum(geo["heights"])
    assert out["h_in"] == 6.0
    assert "height" in out["box_resized"]


def test_a_complete_exact_list_does_not_claim_a_resize(tmp_path):
    """M5 guard: a list that already totals the box says nothing."""
    from kitchensink4ppt.ops import tables as tb

    pkg = _table_pkg(tmp_path)
    out = tb.create_table(
        pkg, 0, 2, 2, 1, 1, 6, 3, col_widths=[2.0, 4.0],
        row_heights=[1.5, 1.5],
    )
    geo = _geometry(pkg)
    assert geo["frame_cx"] == sum(geo["widths"])
    assert geo["frame_cy"] == sum(geo["heights"])
    assert "box_resized" not in out


def test_a_single_row_underfill_resizes(tmp_path):
    """M5: the 1-row and 1-column edges the reviewer asked for."""
    from kitchensink4ppt.ops import tables as tb
    from kitchensink4ppt.ops import geometry as g

    pkg = _table_pkg(tmp_path)
    tb.create_table(pkg, 0, 1, 1, 1, 1, 6, 3, row_heights=[0.4],
                    col_widths=[0.9])
    geo = _geometry(pkg)
    assert geo["heights"] == [g.in_to_emu(0.4)]
    assert geo["widths"] == [g.in_to_emu(0.9)]
    assert geo["frame_cy"] == g.in_to_emu(0.4)
    assert geo["frame_cx"] == g.in_to_emu(0.9)


def test_a_partial_dict_still_fits_the_box_it_was_given(tmp_path):
    """M5 guard: a dict names some rows and leaves the rest to the box, so
    the box still rules and the frame does not move."""
    from kitchensink4ppt.ops import tables as tb
    from kitchensink4ppt.ops import geometry as g

    pkg = _table_pkg(tmp_path)
    out = tb.create_table(pkg, 0, 3, 2, 1, 1, 6, 3, row_heights={1: 2.0})
    geo = _geometry(pkg)
    assert sum(geo["heights"]) == g.in_to_emu(3)
    assert geo["frame_cy"] == g.in_to_emu(3)
    assert "box_resized" not in out


def test_a_partial_dict_that_leaves_no_room_still_refuses(tmp_path):
    """M5: the one case a dict cannot resolve."""
    from kitchensink4ppt.core.errors import PptMcpError
    from kitchensink4ppt.ops import tables as tb

    pkg = _table_pkg(tmp_path)
    with pytest.raises(PptMcpError) as excinfo:
        tb.create_table(pkg, 0, 3, 2, 1, 1, 6, 3, row_heights={0: 4.0})
    message = str(excinfo.value)
    assert "row_heights" in message and "unstated" in message


def test_a_sized_table_still_saves(tmp_path):
    """M5: the geometry the tests read has to survive the save gate."""
    from kitchensink4ppt.ops import tables as tb

    pkg = _table_pkg(tmp_path)
    tb.create_table(pkg, 0, 3, 2, 1, 1, 6, 3, col_widths=[1.0, 1.0],
                    data=[["a", "b"], ["c", "d"], ["e", "f"]])
    out = tmp_path / "sized.pptx"
    pkg.save(str(out))
    assert out.is_file()
