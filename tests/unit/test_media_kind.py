"""Video and audio are typed media, not picture.

A video or audio object is a p:pic whose nvPr names its media file.
get_slide_info and list_elements typed it "picture", and list_elements
"images" counted it as an image (its poster frame), while PowerPoint counts
it as media and the live route already says "media". Now the file route says
"media" too; the alt-text audit and the other checks still treat it as a
graphic.

Every deck here is made up by code.
"""

from __future__ import annotations

import struct
import zlib

import pytest
from pptx import Presentation
from pptx.util import Inches

from kitchensink4ppt.core.package import PptxPackage
from kitchensink4ppt.ops import access, read


def _png(path):
    """A 2x2 grey PNG, made by code."""
    raw = b"".join(b"\x00" + b"\x80\x80\x80" * 2 for _ in range(2))

    def chunk(kind, data):
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    path.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0))
                     + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))
    return str(path)


@pytest.fixture()
def deck(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pic = slide.shapes.add_picture(_png(tmp_path / "grey.png"), Inches(1), Inches(1), Inches(1), Inches(1))
    pic.name = "Made-up picture"
    mp4 = tmp_path / "made_up.mp4"
    mp4.write_bytes(struct.pack(">I4s4sI4s4s", 24, b"ftyp", b"isom", 512, b"isom", b"mp41")
                    + struct.pack(">I4s", 8, b"free"))
    movie = slide.shapes.add_movie(str(mp4), Inches(3), Inches(1), Inches(4), Inches(3),
                                   mime_type="video/mp4")
    movie.name = "Made-up video"
    path = tmp_path / "media.pptx"
    prs.save(str(path))
    return path


def _types(path):
    info = read.get_slide_info(PptxPackage(path), 0)
    return {s["name"]: s["type"] for s in info["shapes"]}


def test_a_video_is_typed_media_and_a_picture_stays_picture(deck):
    assert _types(deck) == {"Made-up picture": "picture", "Made-up video": "media"}
    shapes = read.list_elements(PptxPackage(deck), "shapes")["items"]
    assert {s["name"]: s["type"] for s in shapes} == _types(deck)


def test_images_lists_pictures_only(deck):
    images = read.list_elements(PptxPackage(deck), "images")["items"]
    assert [i["name"] for i in images] == ["Made-up picture"]


def test_the_alt_text_audit_still_checks_the_video(deck):
    found = access.audit_accessibility(PptxPackage(deck))
    alt = [f for f in found["findings"] if f["check"] == "alt_text"]
    assert [f["kind"] for f in alt] == ["media"]  # the picture has a descr: its file name
    assert any("'Made-up video'" in f["message"] for f in alt if f["kind"] == "media")
